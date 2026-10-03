#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
TITANGRID / DELTA SELF-HEALING PROTOCOL 888
DIGITAL TRANSFORMER FOR BABY DELTA
NEURO-MONO ATOM SIGNAL MAP ENGINE

STATUS:
    REPORT_ONLY
    INPUT_MODE: READ_ONLY
    OUTPUT_MODE: NO_OVERWRITE_BY_DEFAULT
    EXECUTION_OF_PROJECT_SCRIPTS: NO
    PATCH: NO
    DELETE: NO
    MOVE: NO
    LENDER_USE: NO
    FINAL_USE: NO

PURPOSE:
    This script transforms an input atom register into a controlled Excel workbook with:
        - normalized atom rows
        - truth gates
        - hallucination gates
        - overclaim detection
        - execution blocking
        - risk classification
        - self-healing classifications
        - manifest
        - SHA-256 audit snapshot

INPUT:
    .xlsx, .xlsm, .csv, .tsv, .txt

OUTPUT:
    .xlsx workbook
    .json manifest
    .jsonl audit log
    .sha256.txt digest file

CORE DOCTRINE:
    Evidence precedes intelligence.
    Retrieval precedes generation.
    Audit precedes decision.
    SSOT is sovereign.
    No document = no truth.
    No issuer = no validity.
    No external verification = no lender relevance.
    Folder name is not evidence.
    Hash proves integrity only, not truth.
    Candidate is not approval.
    AI output is not issuer evidence.

Author:
    ChatGPT Control Tower assistant
"""

from __future__ import annotations

import argparse
import csv
import dataclasses
import hashlib
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    import openpyxl
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import FormulaRule
except ImportError as exc:
    raise SystemExit(
        "Missing dependency: openpyxl. Install with: python -m pip install openpyxl"
    ) from exc


# =============================================================================
# 00. CONSTANTS
# =============================================================================

PROTOCOL_ID = "DELTA_SELF_HEALING_PROTOCOL_888"
ENGINE_ID = "TITANGRID_DELTA_NEURO_MONO_TRANSFORMER"
ENGINE_VERSION = "1.0.0"
DEFAULT_STATUS = "OPEN"

FORBIDDEN_TERMS = [
    "FINAL",
    "LENDER_READY",
    "BANKABLE",
    "VERIFIED_FOR_AUDIT",
    "SYSTEM_GREEN",
    "AUDIT_APPROVED",
    "FINANCING_SECURED",
    "SUBMISSION_READY",
    "READY_FOR_SUBMISSION",
    "HARMONIZED_READY",
    "READY",
]

CRITICAL_SIGNALS = [
    "TOKEN",
    "SECRET",
    "DESTRUCTIVE",
    "WATCHDOG",
    "EXECUTIONPOLICY",
    "BYPASS",
    "DELETE",
    "PURGE",
    "SHADOW",
    "VSSADMIN",
    "RUNTIME",
    "DO_NOT_RUN",
    "DAEMON",
    "INSTALLER",
]

HIGH_SIGNALS = [
    "CEDIS",
    "GRID",
    "ESIA",
    "CAPEX",
    "OPEX",
    "LEGAL",
    "KYC",
    "AML",
    "LENDER",
    "BANKABLE",
    "MIA",
    "GRANT",
    "WRONG",
    "VERSION",
    "OFFTAKE",
    "PERMIT",
    "LAND",
    "TITLE",
]

MEDIUM_SIGNALS = [
    "PATH",
    "HASH",
    "MANIFEST",
    "WORKBOOK",
    "EXCEL",
    "SCRIPT",
    "REFACTOR",
    "CANDIDATE",
    "TRANSLATION",
    "REGISTER",
    "MATRIX",
]

LOW_SIGNALS = [
    "LOGO",
    "BRAND",
    "NON-TITAN",
    "HISTORICAL",
    "CONTEXT ONLY",
    "README",
]

DEFAULT_COLUMNS = [
    "ATOM_ID",
    "ATOM",
    "CORE_SIGNAL",
    "PRIMARY_PROBLEM",
    "ROOT_CAUSE",
    "RISK_TYPE",
    "RISK_LEVEL",
    "NEURO_TRIGGER",
    "HEALING_TYPE",
    "OWNER",
    "ALLOWED_ACTIONS",
    "FORBIDDEN_ACTIONS",
    "STATUS",
    "NEXT_SAFE_STEP",
]

OUTPUT_COLUMNS = [
    "ATOM_ID",
    "ATOM",
    "CORE_SIGNAL",
    "PRIMARY_PROBLEM",
    "ROOT_CAUSE",
    "RISK_TYPE",
    "RISK_LEVEL",
    "NEURO_TRIGGER",
    "HEALING_TYPE",
    "OWNER",
    "ALLOWED_ACTIONS",
    "FORBIDDEN_ACTIONS",
    "STATUS",
    "NEXT_SAFE_STEP",
    "SOURCE_CONTEXT",
    "SOURCE_PATH",
    "SOURCE_SHA256",
    "EXTERNAL_VERIFICATION",
    "TRUTH_GATE",
    "HALLUCINATION_GATE",
    "OVERCLAIM_TERMS_DETECTED",
    "EXECUTION_GATE",
    "LENDER_USE_GATE",
    "FINAL_USE_GATE",
    "CANONICAL_DECISION",
    "MONO_SIGNAL_KEY",
    "ZERO_MISTAKE_STATUS",
    "CONTROL_NOTE",
]

HEALING_TYPES = [
    "EVIDENCE_BINDING",
    "OVERCLAIM_REMOVAL",
    "RISK_ISOLATION",
    "SCRIPT_REFACTOR_GATE",
    "FINANCIAL_RECONCILIATION",
    "DOCUMENT_VERSION_CONTROL",
    "PATH_HASH_CONTROL",
    "EVIDENCE_GAP_CLOSURE",
    "SECURITY_CONTROL",
    "BACKLOG_COMPLETION",
    "UNKNOWN_REQUIRES_REVIEW",
]

RISK_LEVELS = ["CRITICAL", "HIGH", "MEDIUM", "LOW", "UNKNOWN"]

STATUS_VALUES = ["OPEN", "PARTIAL", "CONFIRMED", "CONFLICTING", "REJECTED", "DO_NOT_USE", "CLOSED"]

GATE_VALUES = ["PASS", "BLOCK", "WARNING", "GAP", "UNKNOWN"]

BOOL_TEXT = ["YES", "NO", "UNKNOWN"]

CONTROL_RULES = [
    "Evidence precedes intelligence.",
    "Retrieval precedes generation.",
    "Audit precedes decision.",
    "SSOT is sovereign.",
    "No document = no truth.",
    "No issuer = no validity.",
    "No external verification = no lender relevance.",
    "Folder name is not evidence.",
    "Hash proves integrity only, not truth.",
    "Document existence is not issuer validation.",
    "Candidate is not approval.",
    "Register is not external validation.",
    "AI output is not issuer evidence.",
    "No final/lender/public use without human and issuer validation.",
    "No execution, patch, delete, move, overwrite by default.",
]


# =============================================================================
# 01. DATA MODELS
# =============================================================================

@dataclasses.dataclass
class AtomRecord:
    atom_id: str
    atom: str = ""
    core_signal: str = ""
    primary_problem: str = ""
    root_cause: str = ""
    risk_type: str = ""
    risk_level: str = "UNKNOWN"
    neuro_trigger: str = ""
    healing_type: str = "UNKNOWN_REQUIRES_REVIEW"
    owner: str = "Control Tower"
    allowed_actions: str = "copy; classify; review"
    forbidden_actions: str = "execute; patch; delete; move; overwrite; lender_use; final_use"
    status: str = DEFAULT_STATUS
    next_safe_step: str = ""
    source_context: str = ""
    source_path: str = ""
    source_sha256: str = ""
    external_verification: str = "UNKNOWN"
    truth_gate: str = "UNKNOWN"
    hallucination_gate: str = "UNKNOWN"
    overclaim_terms_detected: str = ""
    execution_gate: str = "BLOCK"
    lender_use_gate: str = "BLOCK"
    final_use_gate: str = "BLOCK"
    canonical_decision: str = "REVIEW_REQUIRED"
    mono_signal_key: str = ""
    zero_mistake_status: str = "REVIEW_REQUIRED"
    control_note: str = ""

    def as_row(self) -> Dict[str, Any]:
        return {
            "ATOM_ID": self.atom_id,
            "ATOM": self.atom,
            "CORE_SIGNAL": self.core_signal,
            "PRIMARY_PROBLEM": self.primary_problem,
            "ROOT_CAUSE": self.root_cause,
            "RISK_TYPE": self.risk_type,
            "RISK_LEVEL": self.risk_level,
            "NEURO_TRIGGER": self.neuro_trigger,
            "HEALING_TYPE": self.healing_type,
            "OWNER": self.owner,
            "ALLOWED_ACTIONS": self.allowed_actions,
            "FORBIDDEN_ACTIONS": self.forbidden_actions,
            "STATUS": self.status,
            "NEXT_SAFE_STEP": self.next_safe_step,
            "SOURCE_CONTEXT": self.source_context,
            "SOURCE_PATH": self.source_path,
            "SOURCE_SHA256": self.source_sha256,
            "EXTERNAL_VERIFICATION": self.external_verification,
            "TRUTH_GATE": self.truth_gate,
            "HALLUCINATION_GATE": self.hallucination_gate,
            "OVERCLAIM_TERMS_DETECTED": self.overclaim_terms_detected,
            "EXECUTION_GATE": self.execution_gate,
            "LENDER_USE_GATE": self.lender_use_gate,
            "FINAL_USE_GATE": self.final_use_gate,
            "CANONICAL_DECISION": self.canonical_decision,
            "MONO_SIGNAL_KEY": self.mono_signal_key,
            "ZERO_MISTAKE_STATUS": self.zero_mistake_status,
            "CONTROL_NOTE": self.control_note,
        }


@dataclasses.dataclass
class TransformResult:
    input_path: str
    input_sha256: str
    output_path: str
    output_sha256: str
    manifest_path: str
    audit_log_path: str
    digest_path: str
    row_count: int
    created_at_utc: str
    warnings: List[str]


# =============================================================================
# 02. LOW LEVEL UTILITIES
# =============================================================================

def now_utc_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while True:
            b = f.read(chunk_size)
            if not b:
                break
            h.update(b)
    return h.hexdigest().upper()


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def normalize_header(value: Any) -> str:
    text = normalize_text(value).upper()
    text = text.replace(" ", "_").replace("-", "_").replace("/", "_")
    text = re.sub(r"[^A-Z0-9_]+", "", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return text


def atom_sort_key(atom_id: str) -> Tuple[int, str]:
    m = re.search(r"(\d+)", atom_id or "")
    if not m:
        return (999999999, atom_id)
    return (int(m.group(1)), atom_id)


def ensure_no_overwrite(path: Path, overwrite: bool = False) -> None:
    if path.exists() and not overwrite:
        raise FileExistsError(
            f"Output already exists and overwrite is disabled: {path}"
        )


def safe_sheet_name(name: str) -> str:
    name = re.sub(r"[\[\]\:\*\?\/\\]", "_", name)
    return name[:31] if len(name) > 31 else name


def detect_delimiter(path: Path) -> str:
    sample = path.read_text(encoding="utf-8-sig", errors="replace")[:4096]
    tabs = sample.count("\t")
    commas = sample.count(",")
    semis = sample.count(";")
    if tabs >= max(commas, semis):
        return "\t"
    if semis > commas:
        return ";"
    return ","


def split_action_text(text: str) -> List[str]:
    if not text:
        return []
    parts = re.split(r"[;,|/]+", text)
    return [p.strip().lower() for p in parts if p.strip()]


# =============================================================================
# 03. INPUT READERS
# =============================================================================

def read_input_table(path: Path) -> List[Dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix in [".xlsx", ".xlsm"]:
        return read_xlsx(path)
    if suffix in [".csv", ".tsv", ".txt"]:
        return read_delimited(path)
    # Support double extension .xlsx.xlsx by checking name ending.
    if path.name.lower().endswith(".xlsx.xlsx"):
        return read_xlsx(path)
    raise ValueError(f"Unsupported input format: {path}")


def read_xlsx(path: Path) -> List[Dict[str, Any]]:
    wb = load_workbook(path, data_only=False, read_only=True)
    if not wb.sheetnames:
        return []
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_row_idx = None
    headers: List[str] = []
    for idx, row in enumerate(rows[:20]):
        normalized = [normalize_header(x) for x in row]
        if any(h in normalized for h in ["ATOM_ID", "ATOM", "CORE_SIGNAL", "PRIMARY_PROBLEM"]):
            header_row_idx = idx
            headers = normalized
            break

    if header_row_idx is None:
        header_row_idx = 0
        headers = [normalize_header(x) or f"COL_{i+1}" for i, x in enumerate(rows[0])]

    out: List[Dict[str, Any]] = []
    for row in rows[header_row_idx + 1:]:
        if row is None:
            continue
        values = list(row)
        if not any(normalize_text(v) for v in values):
            continue
        record: Dict[str, Any] = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            record[h] = values[i] if i < len(values) else ""
        out.append(record)
    return out


def read_delimited(path: Path) -> List[Dict[str, Any]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else detect_delimiter(path)
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        out: List[Dict[str, Any]] = []
        for row in reader:
            if not any(normalize_text(v) for v in row.values()):
                continue
            out.append({normalize_header(k): v for k, v in row.items()})
        return out


# =============================================================================
# 04. FIELD MAPPING
# =============================================================================

FIELD_ALIASES: Dict[str, Sequence[str]] = {
    "ATOM_ID": ["ATOM_ID", "ATOMID", "ATOM", "ID"],
    "ATOM": ["ATOM", "STATEMENT", "PROBLEM", "ORIGINAL_PROBLEM"],
    "CORE_SIGNAL": ["CORE_SIGNAL", "SIGNAL", "DETECTED_SIGNAL"],
    "PRIMARY_PROBLEM": ["PRIMARY_PROBLEM", "PROBLEM", "ORIGINAL_PROBLEM"],
    "ROOT_CAUSE": ["ROOT_CAUSE", "CAUSE"],
    "RISK_TYPE": ["RISK_TYPE", "RISK_CLASS", "RISK_FLAG", "RISK_FLAGS"],
    "RISK_LEVEL": ["RISK_LEVEL", "PRIORITY"],
    "NEURO_TRIGGER": ["NEURO_TRIGGER", "TRIGGER"],
    "HEALING_TYPE": ["HEALING_TYPE", "SIGNAL_TYPE"],
    "OWNER": ["OWNER", "RESPONSIBLE", "OWNER_OR_RESPONSIBLE_LAYER"],
    "ALLOWED_ACTIONS": ["ALLOWED_ACTIONS", "ALLOWED_SAFE_ACTIONS"],
    "FORBIDDEN_ACTIONS": ["FORBIDDEN_ACTIONS", "BLOCKED_ACTIONS"],
    "STATUS": ["STATUS", "CURRENT_STATUS"],
    "NEXT_SAFE_STEP": ["NEXT_SAFE_STEP", "NEXT_REQUIRED_ACTION"],
    "SOURCE_CONTEXT": ["SOURCE_CONTEXT", "CONTEXT"],
    "SOURCE_PATH": ["SOURCE_PATH", "PATH", "PATH_IF_KNOWN"],
    "SOURCE_SHA256": ["SOURCE_SHA256", "HASH", "HASH_IF_KNOWN", "SHA256"],
    "EXTERNAL_VERIFICATION": ["EXTERNAL_VERIFICATION", "EXTERNAL_VERIFICATION_STATUS"],
}


def get_value(row: Dict[str, Any], canonical: str) -> str:
    aliases = FIELD_ALIASES.get(canonical, [canonical])
    for alias in aliases:
        if alias in row:
            value = normalize_text(row.get(alias))
            if value:
                return value
    return ""


def derive_atom_id(row: Dict[str, Any], index: int) -> str:
    atom_id = get_value(row, "ATOM_ID")
    if atom_id:
        m = re.search(r"(\d+)", atom_id)
        if m:
            return f"ATOM-{int(m.group(1)):06d}"
        return atom_id.strip().upper()
    return f"ATOM-{index:06d}"


def normalize_status(status: str) -> str:
    s = normalize_header(status)
    if s in STATUS_VALUES:
        return s
    if not s:
        return DEFAULT_STATUS
    if "REJECT" in s:
        return "REJECTED"
    if "DO_NOT" in s:
        return "DO_NOT_USE"
    if "CONFLICT" in s:
        return "CONFLICTING"
    if "CONFIRM" in s:
        return "CONFIRMED"
    if "PARTIAL" in s:
        return "PARTIAL"
    return DEFAULT_STATUS


def normalize_external_verification(value: str) -> str:
    v = normalize_header(value)
    if v in ["YES", "NO", "UNKNOWN"]:
        return v
    if v in ["TRUE", "Y"]:
        return "YES"
    if v in ["FALSE", "N"]:
        return "NO"
    if not v:
        return "UNKNOWN"
    return "UNKNOWN"


# =============================================================================
# 05. SIGNAL DETECTION
# =============================================================================

def detect_forbidden_terms(*texts: str) -> List[str]:
    combined = " ".join(t for t in texts if t).upper()
    found = []
    for term in FORBIDDEN_TERMS:
        pattern = re.escape(term).replace("\\_", r"[_\s-]?")
        if re.search(pattern, combined, flags=re.IGNORECASE):
            found.append(term)
    return sorted(set(found))


def classify_priority(text: str, provided: str = "") -> str:
    provided_norm = normalize_header(provided)
    if provided_norm in RISK_LEVELS:
        return provided_norm

    t = text.upper()
    if any(sig in t for sig in CRITICAL_SIGNALS):
        return "CRITICAL"
    if any(sig in t for sig in HIGH_SIGNALS):
        return "HIGH"
    if any(sig in t for sig in MEDIUM_SIGNALS):
        return "MEDIUM"
    if any(sig in t for sig in LOW_SIGNALS):
        return "LOW"
    return "UNKNOWN"


def classify_healing_type(text: str, provided: str = "") -> str:
    provided_norm = normalize_header(provided)
    if provided_norm in HEALING_TYPES:
        return provided_norm

    t = text.upper()

    if any(x in t for x in ["TOKEN", "SECRET", "WDAC", "SANDBOX", "PID_NAMESPACE", "MNT_NAMESPACE", "NET_NAMESPACE"]):
        return "SECURITY_CONTROL"

    if any(x in t for x in ["DELETE", "DESTRUCTIVE", "DO_NOT_RUN", "WATCHDOG", "DAEMON", "INSTALLER", "EXECUTIONPOLICY", "BYPASS", "ARCHIVE_ONLY", "REJECTED_RUNTIME"]):
        return "RISK_ISOLATION"

    if any(x in t for x in ["SCRIPT", ".PY", ".PS1", ".BAT", "REFACTOR", "CODEX", "HYDRA", "RUNHARVEST", "FORENSIC_ENGINE"]):
        return "SCRIPT_REFACTOR_GATE"

    if any(x in t for x in ["CAPEX", "OPEX", "IRR", "DSCR", "WACC", "EBITDA", "GRANT", "DEBT", "EQUITY", "FINANCIAL"]):
        return "FINANCIAL_RECONCILIATION"

    if any(x in t for x in ["DOCX", "DOCUMENT", "MEMORANDUM", "PACKAGE", "VERSION", "TRANSLATION", "SR", "EN", "TR"]):
        return "DOCUMENT_VERSION_CONTROL"

    if any(x in t for x in ["PATH", "HASH", "SHA", "MANIFEST", "REGISTER", "FILE_INDEX"]):
        return "PATH_HASH_CONTROL"

    if any(x in t for x in ["CEDIS", "CGES", "EPCG", "ESIA", "PERMIT", "KYC", "AML", "OFFTAKE", "LAND", "TITLE", "MUNICIPAL", "MIA", "GOVERNMENT"]):
        return "EVIDENCE_GAP_CLOSURE"

    if any(x in t for x in ["FINAL", "LENDER_READY", "BANKABLE", "SYSTEM_GREEN", "VERIFIED_FOR_AUDIT", "READY"]):
        return "OVERCLAIM_REMOVAL"

    if any(x in t for x in ["INCOMPLETE", "BACKLOG", "MISSING SECTION", "JSONL", "SELF_AUDIT"]):
        return "BACKLOG_COMPLETION"

    return "EVIDENCE_BINDING"


def derive_truth_gate(record: AtomRecord) -> str:
    text = " ".join([
        record.atom,
        record.core_signal,
        record.primary_problem,
        record.root_cause,
        record.risk_type,
        record.next_safe_step,
    ]).upper()

    if record.status in ["REJECTED", "DO_NOT_USE"]:
        return "BLOCK"
    if "MISSING" in text or "GAP" in text or "PENDING" in text or "UNKNOWN" in text:
        return "GAP"
    if record.external_verification == "YES" and record.source_sha256:
        return "PASS"
    if record.source_sha256 and record.external_verification in ["NO", "UNKNOWN"]:
        return "WARNING"
    return "UNKNOWN"


def derive_hallucination_gate(record: AtomRecord) -> str:
    text = " ".join([
        record.atom,
        record.core_signal,
        record.primary_problem,
        record.root_cause,
        record.next_safe_step,
    ]).upper()

    if record.overclaim_terms_detected:
        return "BLOCK"
    if any(x in text for x in ["INVENT", "HALLUCINATION", "NO DOCUMENT", "NO ISSUER", "UNKNOWN", "MISSING"]):
        return "WARNING"
    if record.external_verification == "YES" and record.truth_gate == "PASS":
        return "PASS"
    return "WARNING"


def derive_execution_gate(record: AtomRecord) -> str:
    text = " ".join([
        record.atom,
        record.primary_problem,
        record.risk_type,
        record.forbidden_actions,
        record.status,
    ]).upper()

    if any(x in text for x in ["EXECUTE", "DO_NOT_RUN", "RUNTIME", "DESTRUCTIVE", "WATCHDOG", "BYPASS", "INSTALLER", "DAEMON", "DELETE"]):
        return "BLOCK"
    return "BLOCK"  # Protocol default: block execution.


def derive_canonical_decision(record: AtomRecord) -> str:
    text = " ".join([
        record.atom,
        record.core_signal,
        record.primary_problem,
        record.risk_type,
        record.healing_type,
        record.status,
    ]).upper()

    if record.status in ["REJECTED", "DO_NOT_USE"]:
        return "REJECTED_OR_ARCHIVE_ONLY"
    if "CANONICAL" in text and "CANDIDATE" in text:
        return "CANONICAL_CANDIDATE_REVIEW_REQUIRED"
    if "ADAPTER" in text:
        return "ADAPTER_CANDIDATE_REVIEW_REQUIRED"
    if "REPORT" in text:
        return "REPORTING_ONLY_REVIEW_REQUIRED"
    if "TEST" in text:
        return "TEST_FRAMEWORK_CANDIDATE_REVIEW_REQUIRED"
    return "REVIEW_REQUIRED"


def derive_zero_mistake_status(record: AtomRecord) -> str:
    if record.truth_gate == "PASS" and record.hallucination_gate == "PASS":
        if record.execution_gate == "BLOCK" and record.lender_use_gate == "BLOCK" and record.final_use_gate == "BLOCK":
            return "CONTROLLED_PASS_REPORT_ONLY"
    if record.truth_gate in ["BLOCK", "GAP"] or record.hallucination_gate == "BLOCK":
        return "BLOCK_OR_GAP_REVIEW_REQUIRED"
    return "REVIEW_REQUIRED"


def build_mono_signal_key(record: AtomRecord) -> str:
    base = f"{record.atom_id}|{record.healing_type}|{record.risk_level}|{record.status}|{record.canonical_decision}"
    return hashlib.sha256(base.encode("utf-8")).hexdigest().upper()[:24]


def enrich_record(record: AtomRecord) -> AtomRecord:
    combined = " ".join([
        record.atom,
        record.core_signal,
        record.primary_problem,
        record.root_cause,
        record.risk_type,
        record.neuro_trigger,
        record.healing_type,
        record.next_safe_step,
    ])

    record.risk_level = classify_priority(combined, record.risk_level)
    record.healing_type = classify_healing_type(combined, record.healing_type)

    forbidden = detect_forbidden_terms(combined)
    record.overclaim_terms_detected = "; ".join(forbidden)

    record.execution_gate = derive_execution_gate(record)
    record.lender_use_gate = "BLOCK"
    record.final_use_gate = "BLOCK"
    record.truth_gate = derive_truth_gate(record)
    record.hallucination_gate = derive_hallucination_gate(record)
    record.canonical_decision = derive_canonical_decision(record)
    record.mono_signal_key = build_mono_signal_key(record)
    record.zero_mistake_status = derive_zero_mistake_status(record)

    notes: List[str] = []
    if forbidden:
        notes.append("OVERCLAIM_RISK: forbidden terms detected.")
    if record.source_sha256 and record.external_verification != "YES":
        notes.append("HASH_INTEGRITY_ONLY: hash does not prove external validity.")
    if record.external_verification != "YES":
        notes.append("NO_EXTERNAL_VERIFICATION: no lender relevance.")
    if not record.next_safe_step:
        notes.append("NEXT_SAFE_STEP_MISSING.")
    record.control_note = " ".join(notes) if notes else "REPORT_ONLY_CONTROLLED_ROW."

    return record


def row_to_atom_record(row: Dict[str, Any], index: int, input_path: Path, input_sha: str) -> AtomRecord:
    atom_id = derive_atom_id(row, index)
    atom = get_value(row, "ATOM")

    core_signal = get_value(row, "CORE_SIGNAL")
    primary_problem = get_value(row, "PRIMARY_PROBLEM")
    root_cause = get_value(row, "ROOT_CAUSE")
    risk_type = get_value(row, "RISK_TYPE")
    risk_level = get_value(row, "RISK_LEVEL")
    neuro_trigger = get_value(row, "NEURO_TRIGGER")
    healing_type = get_value(row, "HEALING_TYPE")
    owner = get_value(row, "OWNER") or "Control Tower"
    allowed = get_value(row, "ALLOWED_ACTIONS") or "copy; classify; review"
    forbidden = get_value(row, "FORBIDDEN_ACTIONS") or "execute; patch; delete; move; overwrite; lender_use; final_use"
    status = normalize_status(get_value(row, "STATUS"))
    next_step = get_value(row, "NEXT_SAFE_STEP")

    source_context = get_value(row, "SOURCE_CONTEXT") or f"Imported from {input_path.name}"
    source_path = get_value(row, "SOURCE_PATH") or str(input_path)
    source_sha = get_value(row, "SOURCE_SHA256") or input_sha
    external = normalize_external_verification(get_value(row, "EXTERNAL_VERIFICATION"))

    if not core_signal and atom:
        core_signal = atom[:250]
    if not primary_problem and atom:
        primary_problem = atom
    if not root_cause:
        root_cause = "Source row requires evidence, issuer, hash and review mapping before external use."
    if not risk_type:
        risk_type = "REVIEW_REQUIRED"
    if not neuro_trigger:
        neuro_trigger = "Block unsupported, external, final, lender-ready or overclaim language."
    if not next_step:
        next_step = "Review row, bind evidence, classify risk and update next safe action."

    rec = AtomRecord(
        atom_id=atom_id,
        atom=atom,
        core_signal=core_signal,
        primary_problem=primary_problem,
        root_cause=root_cause,
        risk_type=risk_type,
        risk_level=risk_level or "UNKNOWN",
        neuro_trigger=neuro_trigger,
        healing_type=healing_type or "UNKNOWN_REQUIRES_REVIEW",
        owner=owner,
        allowed_actions=allowed,
        forbidden_actions=forbidden,
        status=status,
        next_safe_step=next_step,
        source_context=source_context,
        source_path=source_path,
        source_sha256=source_sha,
        external_verification=external,
    )
    return enrich_record(rec)


# =============================================================================
# 06. FALLBACK ATOM GENERATOR
# =============================================================================

def build_fallback_atoms(count: int = 416) -> List[AtomRecord]:
    """
    Used only if the user intentionally wants a shell map without an input file.
    Does not invent project facts. Creates review placeholders.
    """
    records: List[AtomRecord] = []
    for i in range(1, count + 1):
        atom_id = f"ATOM-{i:06d}"
        rec = AtomRecord(
            atom_id=atom_id,
            atom=f"{atom_id} placeholder row. Source atom content not provided in input file.",
            core_signal="PLACEHOLDER_REQUIRES_SOURCE_DATA",
            primary_problem="No source row was provided for this atom.",
            root_cause="The transformer cannot infer missing atom content without source evidence.",
            risk_type="SOURCE_DATA_MISSING",
            risk_level="HIGH",
            neuro_trigger="Block generation if source atom text is missing.",
            healing_type="BACKLOG_COMPLETION",
            owner="Control Tower",
            allowed_actions="copy; classify; request source data",
            forbidden_actions="invent; execute; patch; delete; move; overwrite; lender_use; final_use",
            status="OPEN",
            next_safe_step="Provide source atom row and rerun transformer.",
            source_context="Fallback placeholder generated without source input.",
            source_path="",
            source_sha256="",
            external_verification="NO",
        )
        records.append(enrich_record(rec))
    return records


# =============================================================================
# 07. EXCEL WRITER
# =============================================================================

class ExcelStyler:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    subheader_fill = PatternFill("solid", fgColor="D9EAF7")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    block_fill = PatternFill("solid", fgColor="F4CCCC")
    pass_fill = PatternFill("solid", fgColor="D9EAD3")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    @classmethod
    def style_header(cls, ws, row: int = 1) -> None:
        for cell in ws[row]:
            cell.fill = cls.header_fill
            cell.font = cls.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cls.border

    @classmethod
    def style_cells(cls, ws) -> None:
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = cls.border

    @classmethod
    def autosize(cls, ws, max_width: int = 70) -> None:
        for col_idx, col_cells in enumerate(ws.columns, start=1):
            length = 0
            for cell in col_cells:
                text = str(cell.value) if cell.value is not None else ""
                length = max(length, min(len(text), max_width))
            ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(length + 2, max_width))

    @classmethod
    def freeze_filter_table(cls, ws, table_name: str) -> None:
        ws.freeze_panes = "A2"
        max_row = ws.max_row
        max_col = ws.max_column
        if max_row >= 2 and max_col >= 1:
            ref = f"A1:{get_column_letter(max_col)}{max_row}"
            tab = Table(displayName=table_name, ref=ref)
            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            tab.tableStyleInfo = style
            ws.add_table(tab)


def add_validations(ws, header_map: Dict[str, int]) -> None:
    def add_list_validation(col_name: str, values: List[str]) -> None:
        col = header_map.get(col_name)
        if not col:
            return
        formula = '"' + ",".join(values) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}1048576")

    add_list_validation("RISK_LEVEL", RISK_LEVELS)
    add_list_validation("HEALING_TYPE", HEALING_TYPES)
    add_list_validation("STATUS", STATUS_VALUES)
    add_list_validation("EXTERNAL_VERIFICATION", BOOL_TEXT)
    add_list_validation("TRUTH_GATE", GATE_VALUES)
    add_list_validation("HALLUCINATION_GATE", GATE_VALUES)
    add_list_validation("EXECUTION_GATE", GATE_VALUES)
    add_list_validation("LENDER_USE_GATE", GATE_VALUES)
    add_list_validation("FINAL_USE_GATE", GATE_VALUES)


def add_conditional_formatting(ws, header_map: Dict[str, int]) -> None:
    max_row = ws.max_row
    for gate_col in ["TRUTH_GATE", "HALLUCINATION_GATE", "EXECUTION_GATE", "LENDER_USE_GATE", "FINAL_USE_GATE"]:
        col = header_map.get(gate_col)
        if not col:
            continue
        letter = get_column_letter(col)
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{max_row}",
            FormulaRule(formula=[f'${letter}2="BLOCK"'], fill=ExcelStyler.block_fill),
        )
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{max_row}",
            FormulaRule(formula=[f'${letter}2="GAP"'], fill=ExcelStyler.warning_fill),
        )
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{max_row}",
            FormulaRule(formula=[f'${letter}2="PASS"'], fill=ExcelStyler.pass_fill),
        )

    risk_col = header_map.get("RISK_LEVEL")
    if risk_col:
        letter = get_column_letter(risk_col)
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{max_row}",
            FormulaRule(formula=[f'${letter}2="CRITICAL"'], fill=ExcelStyler.block_fill),
        )
        ws.conditional_formatting.add(
            f"{letter}2:{letter}{max_row}",
            FormulaRule(formula=[f'${letter}2="HIGH"'], fill=ExcelStyler.warning_fill),
        )


def write_records_sheet(wb: Workbook, records: List[AtomRecord]) -> None:
    ws = wb.active
    ws.title = "NEURO_MONO_MAP"

    ws.append(OUTPUT_COLUMNS)
    for rec in records:
        row = rec.as_row()
        ws.append([row.get(col, "") for col in OUTPUT_COLUMNS])

    ExcelStyler.style_header(ws, 1)
    ExcelStyler.style_cells(ws)
    ExcelStyler.autosize(ws)
    ExcelStyler.freeze_filter_table(ws, "NEURO_MONO_MAP_TABLE")

    header_map = {cell.value: idx for idx, cell in enumerate(ws[1], start=1)}
    add_validations(ws, header_map)
    add_conditional_formatting(ws, header_map)


def count_by(records: List[AtomRecord], attr: str) -> Dict[str, int]:
    out: Dict[str, int] = {}
    for rec in records:
        value = getattr(rec, attr)
        out[value] = out.get(value, 0) + 1
    return dict(sorted(out.items(), key=lambda kv: (-kv[1], kv[0])))


def write_dashboard_sheet(wb: Workbook, records: List[AtomRecord], input_sha: str) -> None:
    ws = wb.create_sheet("SIGNAL_DASHBOARD")
    rows = [
        ["FIELD", "VALUE"],
        ["PROTOCOL_ID", PROTOCOL_ID],
        ["ENGINE_ID", ENGINE_ID],
        ["ENGINE_VERSION", ENGINE_VERSION],
        ["CREATED_AT_UTC", now_utc_iso()],
        ["ROW_COUNT", len(records)],
        ["INPUT_SHA256", input_sha],
        ["EXECUTION", "NO"],
        ["PATCH", "NO"],
        ["DELETE", "NO"],
        ["MOVE", "NO"],
        ["OVERWRITE", "NO_BY_DEFAULT"],
        ["LENDER_USE", "NO"],
        ["FINAL_USE", "NO"],
    ]

    for row in rows:
        ws.append(row)

    ws.append([])
    ws.append(["RISK_LEVEL", "COUNT"])
    for k, v in count_by(records, "risk_level").items():
        ws.append([k, v])

    ws.append([])
    ws.append(["HEALING_TYPE", "COUNT"])
    for k, v in count_by(records, "healing_type").items():
        ws.append([k, v])

    ws.append([])
    ws.append(["TRUTH_GATE", "COUNT"])
    for k, v in count_by(records, "truth_gate").items():
        ws.append([k, v])

    ExcelStyler.style_header(ws, 1)
    ExcelStyler.style_cells(ws)
    ExcelStyler.autosize(ws)


def write_truth_gates_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("FORMULA_TRUTH_GATES")
    ws.append(["GATE", "RULE", "DECISION"])
    rows = [
        ["TRUTH_GATE", "External verification YES + source hash present", "PASS"],
        ["TRUTH_GATE", "Missing evidence / pending / unknown", "GAP"],
        ["TRUTH_GATE", "Rejected or do-not-use status", "BLOCK"],
        ["HALLUCINATION_GATE", "Forbidden maturity/readiness terms detected", "BLOCK"],
        ["HALLUCINATION_GATE", "Missing source/issuer/evidence language", "WARNING"],
        ["EXECUTION_GATE", "Protocol default", "BLOCK"],
        ["LENDER_USE_GATE", "Protocol default", "BLOCK"],
        ["FINAL_USE_GATE", "Protocol default", "BLOCK"],
        ["ZERO_MISTAKE_STATUS", "PASS truth + PASS hallucination + blocked execution/final/lender", "CONTROLLED_PASS_REPORT_ONLY"],
    ]
    for row in rows:
        ws.append(row)
    ExcelStyler.style_header(ws, 1)
    ExcelStyler.style_cells(ws)
    ExcelStyler.autosize(ws)
    ExcelStyler.freeze_filter_table(ws, "FORMULA_TRUTH_GATES_TABLE")


def write_dictionary_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("CONTROL_DICTIONARY")
    ws.append(["CATEGORY", "VALUE", "MEANING"])

    for h in HEALING_TYPES:
        ws.append(["HEALING_TYPE", h, meaning_for_healing_type(h)])
    for r in RISK_LEVELS:
        ws.append(["RISK_LEVEL", r, meaning_for_risk_level(r)])
    for rule in CONTROL_RULES:
        ws.append(["CONTROL_RULE", rule, "Mandatory interpretation rule"])
    for term in FORBIDDEN_TERMS:
        ws.append(["FORBIDDEN_TERM", term, "Flag as OVERCLAIM_RISK unless backed by explicit external issuer/lender evidence"])

    ExcelStyler.style_header(ws, 1)
    ExcelStyler.style_cells(ws)
    ExcelStyler.autosize(ws)
    ExcelStyler.freeze_filter_table(ws, "CONTROL_DICTIONARY_TABLE")


def meaning_for_healing_type(value: str) -> str:
    return {
        "EVIDENCE_BINDING": "Needs issuer/source/hash/review evidence.",
        "OVERCLAIM_REMOVAL": "Contains or risks maturity/finality/readiness overclaim.",
        "RISK_ISOLATION": "Must stay archive-only/rejected/static-review/destructive-isolate.",
        "SCRIPT_REFACTOR_GATE": "Code may be salvaged only after static scan/refactor gate.",
        "FINANCIAL_RECONCILIATION": "Financial contradiction or missing model support.",
        "DOCUMENT_VERSION_CONTROL": "Wrong version, duplicate, translation, package or review issue.",
        "PATH_HASH_CONTROL": "Needs path/hash/manifest/integrity control.",
        "EVIDENCE_GAP_CLOSURE": "Requires issuer or external evidence.",
        "SECURITY_CONTROL": "Token/security/hardening/sandbox issue.",
        "BACKLOG_COMPLETION": "Missing extraction/export/self-audit work.",
        "UNKNOWN_REQUIRES_REVIEW": "Needs manual classification.",
    }.get(value, "Unknown")


def meaning_for_risk_level(value: str) -> str:
    return {
        "CRITICAL": "Security/runtime/deletion/external misuse danger.",
        "HIGH": "Evidence, lender, institutional or financial contradiction risk.",
        "MEDIUM": "Incomplete control/refactor/path/hash/document review.",
        "LOW": "Brand/context/historical-only signal.",
        "UNKNOWN": "Insufficient classification data.",
    }.get(value, "Unknown")


def write_readme_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("README_PROTOCOL_888")
    rows = [
        ["ITEM", "VALUE"],
        ["Purpose", "Transform atom register into neuro-mono signal map with truth and safety gates."],
        ["Mode", "REPORT_ONLY"],
        ["Input", "Read-only"],
        ["Output", "No-overwrite by default"],
        ["Execution", "NO"],
        ["Patch", "NO"],
        ["Delete", "NO"],
        ["Move", "NO"],
        ["Final use", "NO"],
        ["Lender use", "NO"],
        ["Important", "This workbook is a control artifact. It is not external validation."],
        ["Hash rule", "Hash proves integrity only, not truth."],
        ["Issuer rule", "No issuer = no validity."],
        ["Evidence rule", "No document = no truth."],
    ]
    for row in rows:
        ws.append(row)
    ExcelStyler.style_header(ws, 1)
    ExcelStyler.style_cells(ws)
    ExcelStyler.autosize(ws)


def write_workbook(records: List[AtomRecord], output_path: Path, input_sha: str) -> None:
    wb = Workbook()
    write_records_sheet(wb, records)
    write_dashboard_sheet(wb, records, input_sha)
    write_truth_gates_sheet(wb)
    write_dictionary_sheet(wb)
    write_readme_sheet(wb)

    # Security: no external links created, no macros, no formulas except simple workbook content.
    wb.properties.creator = ENGINE_ID
    wb.properties.title = "TITANGRID DELTA 888 NEURO MONO ATOM SIGNAL MAP"
    wb.properties.subject = "REPORT_ONLY control workbook"
    wb.properties.keywords = "TITANGRID, DELTA, PROTOCOL 888, REPORT_ONLY, NEURO_MONO_MAP"
    wb.save(output_path)


# =============================================================================
# 08. MANIFEST AND AUDIT
# =============================================================================

def build_manifest(
    input_path: Path,
    input_sha: str,
    output_path: Path,
    output_sha: str,
    records: List[AtomRecord],
    warnings: List[str],
) -> Dict[str, Any]:
    return {
        "protocol_id": PROTOCOL_ID,
        "engine_id": ENGINE_ID,
        "engine_version": ENGINE_VERSION,
        "created_at_utc": now_utc_iso(),
        "mode": "REPORT_ONLY",
        "execution": "NO",
        "patch": "NO",
        "delete": "NO",
        "move": "NO",
        "overwrite": "NO_BY_DEFAULT",
        "lender_use": "NO",
        "final_use": "NO",
        "input": {
            "path": str(input_path),
            "sha256": input_sha,
            "exists": input_path.exists(),
            "size_bytes": input_path.stat().st_size if input_path.exists() else None,
        },
        "output": {
            "path": str(output_path),
            "sha256": output_sha,
            "size_bytes": output_path.stat().st_size if output_path.exists() else None,
        },
        "row_count": len(records),
        "risk_counts": count_by(records, "risk_level"),
        "healing_type_counts": count_by(records, "healing_type"),
        "truth_gate_counts": count_by(records, "truth_gate"),
        "hallucination_gate_counts": count_by(records, "hallucination_gate"),
        "warnings": warnings,
        "control_rules": CONTROL_RULES,
    }


def write_json(path: Path, data: Dict[str, Any], overwrite: bool) -> None:
    ensure_no_overwrite(path, overwrite=overwrite)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def write_audit_log(path: Path, events: List[Dict[str, Any]], overwrite: bool) -> None:
    ensure_no_overwrite(path, overwrite=overwrite)
    with path.open("w", encoding="utf-8") as f:
        for event in events:
            f.write(json.dumps(event, ensure_ascii=False) + "\n")


def write_digest(path: Path, digest: str, output_path: Path, overwrite: bool) -> None:
    ensure_no_overwrite(path, overwrite=overwrite)
    path.write_text(f"{digest}  {output_path.name}\n", encoding="utf-8")


# =============================================================================
# 09. TRANSFORMER
# =============================================================================

def transform(
    input_path: Optional[Path],
    output_path: Path,
    overwrite: bool = False,
    fallback_count: int = 0,
) -> TransformResult:
    warnings: List[str] = []
    audit: List[Dict[str, Any]] = []

    ensure_no_overwrite(output_path, overwrite=overwrite)

    created_at = now_utc_iso()
    audit.append({
        "timestamp_utc": created_at,
        "event": "TRANSFORM_START",
        "protocol_id": PROTOCOL_ID,
        "engine_id": ENGINE_ID,
        "output_path": str(output_path),
    })

    if input_path:
        if not input_path.exists():
            raise FileNotFoundError(f"Input file not found: {input_path}")
        input_sha = sha256_file(input_path)
        rows = read_input_table(input_path)
        if not rows:
            warnings.append("Input table has no readable rows.")
        records = [
            row_to_atom_record(row, index=i + 1, input_path=input_path, input_sha=input_sha)
            for i, row in enumerate(rows)
        ]
    else:
        if fallback_count <= 0:
            raise ValueError("Either input_path or fallback_count must be provided.")
        input_sha = ""
        rows = []
        records = build_fallback_atoms(fallback_count)
        warnings.append("Fallback placeholder atoms generated because no input file was provided.")

    records = sorted(records, key=lambda r: atom_sort_key(r.atom_id))

    audit.append({
        "timestamp_utc": now_utc_iso(),
        "event": "RECORDS_NORMALIZED",
        "row_count": len(records),
        "warnings": warnings,
    })

    output_path.parent.mkdir(parents=True, exist_ok=True)
    write_workbook(records, output_path, input_sha)
    output_sha = sha256_file(output_path)

    manifest_path = output_path.with_suffix(output_path.suffix + ".manifest.json")
    audit_path = output_path.with_suffix(output_path.suffix + ".audit.jsonl")
    digest_path = output_path.with_suffix(output_path.suffix + ".sha256.txt")

    manifest = build_manifest(
        input_path=input_path or Path("NO_INPUT_FALLBACK"),
        input_sha=input_sha,
        output_path=output_path,
        output_sha=output_sha,
        records=records,
        warnings=warnings,
    )

    write_json(manifest_path, manifest, overwrite=overwrite)

    audit.append({
        "timestamp_utc": now_utc_iso(),
        "event": "OUTPUT_WRITTEN",
        "output_path": str(output_path),
        "output_sha256": output_sha,
        "manifest_path": str(manifest_path),
        "row_count": len(records),
    })

    write_audit_log(audit_path, audit, overwrite=overwrite)
    write_digest(digest_path, output_sha, output_path, overwrite=overwrite)

    return TransformResult(
        input_path=str(input_path) if input_path else "NO_INPUT_FALLBACK",
        input_sha256=input_sha,
        output_path=str(output_path),
        output_sha256=output_sha,
        manifest_path=str(manifest_path),
        audit_log_path=str(audit_path),
        digest_path=str(digest_path),
        row_count=len(records),
        created_at_utc=created_at,
        warnings=warnings,
    )


# =============================================================================
# 10. COMMAND LINE
# =============================================================================

def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="TITANGRID DELTA 888 Neuro-Mono Atom Signal Map Transformer"
    )
    parser.add_argument(
        "--input",
        "-i",
        type=str,
        default="",
        help="Input Excel/CSV/TSV atom register path. Read-only.",
    )
    parser.add_argument(
        "--output",
        "-o",
        type=str,
        required=True,
        help="Output .xlsx path. No overwrite unless --overwrite is set.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Allow overwriting output files. Default is false. Use only with explicit approval.",
    )
    parser.add_argument(
        "--fallback-count",
        type=int,
        default=0,
        help="Generate placeholder atoms if no input is provided. Does not invent facts.",
    )
    return parser.parse_args(argv)


def print_result(result: TransformResult) -> None:
    print("=======================================================================")
    print("DELTA SELF-HEALING PROTOCOL 888 — DIGITAL TRANSFORMER COMPLETE")
    print("=======================================================================")
    print(f"CREATED_AT_UTC : {result.created_at_utc}")
    print(f"INPUT_PATH     : {result.input_path}")
    print(f"INPUT_SHA256   : {result.input_sha256}")
    print(f"OUTPUT_PATH    : {result.output_path}")
    print(f"OUTPUT_SHA256  : {result.output_sha256}")
    print(f"MANIFEST       : {result.manifest_path}")
    print(f"AUDIT_LOG      : {result.audit_log_path}")
    print(f"DIGEST         : {result.digest_path}")
    print(f"ROW_COUNT      : {result.row_count}")
    print("MODE           : REPORT_ONLY")
    print("EXECUTION      : NO")
    print("PATCH          : NO")
    print("DELETE         : NO")
    print("MOVE           : NO")
    print("LENDER_USE     : NO")
    print("FINAL_USE      : NO")
    if result.warnings:
        print("WARNINGS:")
        for w in result.warnings:
            print(f"  - {w}")
    print("=======================================================================")


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    input_path = Path(args.input).expanduser().resolve() if args.input else None
    output_path = Path(args.output).expanduser().resolve()

    try:
        result = transform(
            input_path=input_path,
            output_path=output_path,
            overwrite=bool(args.overwrite),
            fallback_count=int(args.fallback_count or 0),
        )
        print_result(result)
        return 0
    except Exception as exc:
        print("=======================================================================", file=sys.stderr)
        print("DELTA SELF-HEALING PROTOCOL 888 — DIGITAL TRANSFORMER BLOCKED", file=sys.stderr)
        print("=======================================================================", file=sys.stderr)
        print(f"ERROR: {exc}", file=sys.stderr)
        print("CONTROL_DECISION: NO OUTPUT TRUSTED UNTIL ERROR IS RESOLVED", file=sys.stderr)
        print("EXECUTION: NO | PATCH: NO | DELETE: NO | MOVE: NO", file=sys.stderr)
        print("=======================================================================", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
