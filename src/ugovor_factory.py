import os
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import docx
from docx.shared import Pt

FONT_NAME = "Helvetica"
FONT_BOLD_NAME = "Helvetica-Bold"
DEJAVU_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
DEJAVU_BOLD_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

if os.path.exists(DEJAVU_PATH):
    pdfmetrics.registerFont(TTFont('DejaVuSans', DEJAVU_PATH))
    FONT_NAME = 'DejaVuSans'

if os.path.exists(DEJAVU_BOLD_PATH):
    pdfmetrics.registerFont(TTFont('DejaVuSans-Bold', DEJAVU_BOLD_PATH))
    FONT_BOLD_NAME = 'DejaVuSans-Bold'

def generisi_ugovor_pdf(broj_ugovora, narucilac, izvrsilac, predmet, iznos_eur):
    izlaz_dir = os.path.join(os.getcwd(), "izlaz", "dokumenti")
    os.makedirs(izlaz_dir, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    broj_clean = broj_ugovora.replace('/', '_')
    filepath = os.path.join(izlaz_dir, f"Ugovor_{broj_clean}_{timestamp}.pdf")
    
    doc = SimpleDocTemplate(filepath, pagesize=A4, leftMargin=50, rightMargin=50, topMargin=50, bottomMargin=50)
    styles = getSampleStyleSheet()
    
    h1 = ParagraphStyle('H1', fontName=FONT_BOLD_NAME, fontSize=16, leading=20, alignment=1, textColor=colors.HexColor("#1A365D"), spaceAfter=20)
    body = ParagraphStyle('Body', fontName=FONT_NAME, fontSize=10, leading=15, spaceAfter=10)
    bold = ParagraphStyle('Bold', fontName=FONT_BOLD_NAME, fontSize=10, leading=15)
    
    elements = [
        Paragraph(f"UGOVOR O POSLOVNOJ SARADNJI br. {broj_ugovora}", h1),
        Paragraph(f"Zaključen dana {datetime.now().strftime('%d.%m.%Y.')} godine između:", body),
        Spacer(1, 10),
        Paragraph(f"<b>1. NARUČILAC:</b> {narucilac}", body),
        Paragraph(f"<b>2. IZVRŠILAC:</b> {izvrsilac}", body),
        Spacer(1, 15),
        Paragraph("<b>Član 1. (Predmet ugovora)</b>", bold),
        Paragraph(f"Izvršilac se obavezuje da za potrebe Naručioca izvrši usluge: {predmet}.", body),
        Spacer(1, 10),
        Paragraph("<b>Član 2. (Naknada i plaćanje)</b>", bold),
        Paragraph(f"Naručilac se obavezuje da Izvršiocu isplati ugovoreni iznos od <b>{iznos_eur} EUR</b> nakon realizacije predmeta ugovora.", body),
        Spacer(1, 10),
        Paragraph("<b>Član 3. (Završne odredbe)</b>", bold),
        Paragraph("Ugovor je sačinjen u 2 (dva) istovetna primjerka, po jedan za svaku ugovornu stranu.", body),
        Spacer(1, 40),
    ]
    
    potpisi = [
        [Paragraph("Za Naručioca:<br/><br/>__________________", body), Paragraph("Za Izvršioca:<br/><br/>__________________", body)]
    ]
    t = Table(potpisi, colWidths=[240, 240])
    t.setStyle(TableStyle([('ALIGN', (0,0), (-1,-1), 'CENTER')]))
    elements.append(t)
    
    doc.build(elements)
    return filepath

def generisi_ugovor_docx(broj_ugovora, narucilac, izvrsilac, predmet, iznos_eur):
    izlaz_dir = os.path.join(os.getcwd(), "izlaz", "dokumenti")
    os.makedirs(izlaz_dir, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    broj_clean = broj_ugovora.replace('/', '_')
    filepath = os.path.join(izlaz_dir, f"Ugovor_{broj_clean}_{timestamp}.docx")
    
    doc = docx.Document()
    p = doc.add_paragraph()
    run = p.add_run(f"UGOVOR O POSLOVNOJ SARADNJI br. {broj_ugovora}")
    run.bold = True
    run.font.size = Pt(16)
    p.alignment = 1
    
    doc.add_paragraph(f"Zaključen dana {datetime.now().strftime('%d.%m.%Y.')} godine između:")
    doc.add_paragraph(f"1. NARUČILAC: {narucilac}")
    doc.add_paragraph(f"2. IZVRŠILAC: {izvrsilac}")
    
    p1 = doc.add_paragraph()
    p1.add_run("Član 1. (Predmet ugovora)\n").bold = True
    p1.add_run(f"Izvršilac se obavezuje da za potrebe Naručioca izvrši usluge: {predmet}.")
    
    p2 = doc.add_paragraph()
    p2.add_run("Član 2. (Naknada i plaćanje)\n").bold = True
    p2.add_run(f"Naručilac se obavezuje da Izvršiocu isplati ugovoreni iznos od {iznos_eur} EUR.")
    
    doc.add_paragraph("\nZa Naručioca: ___________________          Za Izvršioca: ___________________")
    
    doc.save(filepath)
    return filepath

if __name__ == "__main__":
    pdf_p = generisi_ugovor_pdf("888/2026", "TITAN GRID DOO", "FREYA SOFTWARE", "Implementacija ugovor_factory.py modula", "1500")
    docx_p = generisi_ugovor_docx("888/2026", "TITAN GRID DOO", "FREYA SOFTWARE", "Implementacija ugovor_factory.py modula", "1500")
    print(f"[OK] Generisan PDF ugovor: {pdf_p}")
    print(f"[OK] Generisan DOCX ugovor: {docx_p}")
