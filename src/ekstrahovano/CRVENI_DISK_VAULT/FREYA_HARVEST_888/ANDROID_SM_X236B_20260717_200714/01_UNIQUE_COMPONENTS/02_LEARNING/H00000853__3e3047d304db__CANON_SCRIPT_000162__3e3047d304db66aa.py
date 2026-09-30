# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import datetime

today = datetime.date.today().strftime("%Y-%m-%d")
out = Path("/storage/emulated/0/Download/TITAN_ECOSYSTEM/TITAN1_SYSTEM/01_ACTIVE")

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)

def p(text="", bold=False, size=10.5, align=None, after=6):
    par = doc.add_paragraph()
    if align:
        par.alignment = align
    par.paragraph_format.space_after = Pt(after)
    r = par.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    return par

p("TEHNIČKI ZAHTJEV ZA PRIKLJUČENJE", True, 18, WD_ALIGN_PARAGRAPH.CENTER, 8)
p("Projekat TITAN 1 – Fabrika transformatorskih kazana (tankova)", True, 14, WD_ALIGN_PARAGRAPH.CENTER, 4)
p(f"Datum: {today}", False, 10.5, WD_ALIGN_PARAGRAPH.CENTER, 20)

p("Prima: Opština Tuzi (za dalje proslijeđivanje CEDIS-u)", True, 12)

p("Poštovani,", False, 10.5, None, 12)

p("Ars Metal Industries D.O.O. Podgorica podnosi formalni tehnički zahtjev za izdavanje tehničkih uslova za priključenje na distributivnu mrežu CEDIS-a za projekat TITAN 1 u industrijskoj zoni Tuzi / KAP.")

p("Predviđena priključna snaga: 12,5 MW (faza realizacije).")
p("Lokacija: Industrijska zona Tuzi / KAP")
p("Investicija: 18,2 miliona EUR")
p("Ovo je prvi projekat u vertikalnoj investicionoj lancu TITAN 1 → TITAN 2 → TITAN 3.")

p("Molimo Opštinu Tuzi da uputi formalni zahtjev CEDIS-u za izdavanje tehničkih uslova priključenja, koji će biti sastavni dio urbanističko-tehničkih uslova.")

p("Spremni smo da dostavimo dodatne tehničke podatke (load profile, jednovremena vršna snaga, broj mjernih mjesta, preliminarne tehničke karakteristike).")

p("Sa poštovanjem,", False, 10.5, None, 14)
p("DANIJELA KESKIN", True, 10.5, None, 0)
p("Director", False, 10.5, None, 0)
p("Ars Metal Industries D.O.O. Podgorica", False, 10.5, None, 0)
p("+382 68 323 339 | d.djurovic@adsmetal.com.tr", False, 10.5, None, 0)

doc.save(out / f"TEHNICKI_ZAHTJEV_CEDIS_TITAN1_{today}.docx")
print(f"✅ Tehnički zahtjev za CEDIS kreiran: TEHNICKI_ZAHTJEV_CEDIS_TITAN1_{today}.docx")
