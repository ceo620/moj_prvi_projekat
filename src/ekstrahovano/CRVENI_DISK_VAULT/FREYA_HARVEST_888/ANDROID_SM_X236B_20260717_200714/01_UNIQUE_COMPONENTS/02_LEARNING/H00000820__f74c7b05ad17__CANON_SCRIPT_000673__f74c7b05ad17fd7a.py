# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

out = Path("02_TECHNICAL_CEDIS_TUZI")
out.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)

def p(text="", bold=False, size=10.5, align=None, after=6):
    par = doc.add_paragraph()
    if align:
        par.alignment = align
    par.paragraph_format.space_after = Pt(after)
    r = par.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    return par

p("TEHNIČKI DOPIS – CEDIS / OPŠTINA TUZI", True, 16, WD_ALIGN_PARAGRAPH.CENTER, 8)
p("Zahtjev za pisano izjašnjenje o statusu elektrodistributivne infrastrukture", True, 12, WD_ALIGN_PARAGRAPH.CENTER, 14)

p("Prima: Opština Tuzi", True, 10.5, None, 0)
p("Za dalje institucionalno upućivanje prema CEDIS-u", False, 10.5, None, 8)

p("Predmet:", True, 10.5, None, 0)
p("Formalni upit CEDIS-u radi dobijanja pisanog izjašnjenja o statusu infrastrukture, raspoloživom kapacitetu i daljim koracima za priključenje projekta TITAN 1", False, 10.5, None, 10)

p("Poštovani,", False, 10.5, None, 8)

p("Ars Metal Industries D.O.O. Podgorica planira realizaciju projekta TITAN 1 – Fabrika transformatorskih kazana (tankova) u industrijskoj zoni Tuzi / KAP.", False, 10.5)

p("Predmet ovog zahtjeva je isključivo projekat TITAN 1 – Fabrika transformatorskih kazana (tankova). Ovaj zahtjev se ne odnosi na druge buduće projekte, faze ili proizvodne programe.", False, 10.5)

p("Za potrebe daljeg planiranja i procjene izvodljivosti projekta, potrebno je pribaviti pisano izjašnjenje CEDIS-a o statusu relevantne elektrodistributivne infrastrukture, raspoloživom kapacitetu i narednim proceduralnim koracima.", False, 10.5)

p("Preliminarni okvir priključne snage koji je predmet provjere iznosi do 12,5 MW, uz faznu realizaciju i naknadnu inženjersku potvrdu kroz odgovarajuću tehničku dokumentaciju.", False, 10.5)

p("Molimo Opštinu Tuzi da uputi formalni upit CEDIS-u radi dobijanja pisanog izjašnjenja o sljedećem:", False, 10.5)

for item in [
    "trenutni status rekonstrukcije ili razvoja relevantne elektrodistributivne infrastrukture za lokaciju Tuzi / KAP,",
    "očekivani rokovi završetka radova ili raspoloživosti kapaciteta,",
    "preliminarna raspoloživost kapaciteta za industrijski priključak,",
    "dalji koraci za izdavanje tehničkih uslova priključenja,",
    "dodatni tehnički podaci koje investitor treba dostaviti u narednoj fazi."
]:
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Inches(0.25)
    par.paragraph_format.space_after = Pt(2)
    r = par.add_run("• " + item)
    r.font.name = "Arial"
    r.font.size = Pt(10.5)

p("Investitor je spreman da, po zahtjevu nadležnih institucija, dostavi dodatne tehničke podatke, uključujući preliminarni load profile, jednovremenu vršnu snagu, broj mjernih mjesta i osnovne tehničko-tehnološke karakteristike pogona.", False, 10.5)

p("Sa poštovanjem,", False, 10.5, None, 14)
p("DANIJELA KESKIN", True, 10.5, None, 0)
p("Director", False, 10.5, None, 0)
p("Ars Metal Industries D.O.O. Podgorica", False, 10.5, None, 0)
p("+382 68 323 339 | d.djurovic@adsmetal.com.tr", False, 10.5, None, 0)

target = out / "TEHNICKI_DOPIS_CEDIS_TUZI_TITAN1_2026-05-10.docx"
doc.save(target)
print("CREATED:", target)
