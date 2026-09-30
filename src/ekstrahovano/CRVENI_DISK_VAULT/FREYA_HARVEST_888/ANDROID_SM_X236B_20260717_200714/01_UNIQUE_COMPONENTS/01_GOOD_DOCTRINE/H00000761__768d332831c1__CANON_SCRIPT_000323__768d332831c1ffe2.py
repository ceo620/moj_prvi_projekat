# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 22_review_decision_importer.py
# PURPOSE: Import manual review queue decisions as non-canonical review log
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_FILE_ERROR = 2
EXIT_SCHEMA_ERROR = 3

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

SHEET = "Review Queue"
REQUIRED = [
    "Review_ID",
    "Item_Type",
    "Source_ID",
    "Priority",
    "Required_Action",
    "Reviewer",
    "Review_Status",
    "Can_Close_Gate",
    "Can_Approve_Evidence",
    "Final_Use_Allowed"
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def norm(value):
    return "" if value is None else str(value).strip()

def headers(ws):
    return {norm(c.value): i for i, c in enumerate(ws[1], start=1) if norm(c.value)}

def emit_jsonl(path, event):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "a", encoding="utf-8") as f:
        f.write(json.dumps(event, ensure_ascii=False) + "\n")

def main():
    parser = argparse.ArgumentParser(description="Import manual review decisions as non-canonical log")
    parser.add_argument("--review-xlsx", required=True)
    parser.add_argument("--out-jsonl", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    if not os.path.exists(args.review_xlsx):
        print(f"❌ Review workbook ne postoji: {args.review_xlsx}")
        return EXIT_FILE_ERROR

    try:
        wb = openpyxl.load_workbook(args.review_xlsx, data_only=True)
    except Exception as exc:
        print(f"❌ Review workbook nije čitljiv: {exc}")
        return EXIT_FILE_ERROR

    if SHEET not in wb.sheetnames:
        print(f"❌ Nedostaje sheet: {SHEET}")
        return EXIT_SCHEMA_ERROR

    ws = wb[SHEET]
    h = headers(ws)
    missing = [c for c in REQUIRED if c not in h]
    if missing:
        print(f"❌ Nedostaju kolone: {missing}")
        return EXIT_SCHEMA_ERROR

    imported = 0
    violations = []

    for r in range(2, ws.max_row + 1):
        row = {key: norm(ws.cell(r, h[key]).value) for key in REQUIRED}
        if not any(row.values()):
            continue

        # hard safety override: even if workbook says YES, canonical permissions stay NO
        if row.get("Can_Close_Gate", "").upper() == "YES":
            violations.append({"row": r, "Review_ID": row["Review_ID"], "violation": "Workbook attempted Can_Close_Gate=YES"})
        if row.get("Can_Approve_Evidence", "").upper() == "YES":
            violations.append({"row": r, "Review_ID": row["Review_ID"], "violation": "Workbook attempted Can_Approve_Evidence=YES"})
        if row.get("Final_Use_Allowed", "").upper() == "YES":
            violations.append({"row": r, "Review_ID": row["Review_ID"], "violation": "Workbook attempted Final_Use_Allowed=YES"})

        event = {
            "timestamp": now_iso(),
            "component": "REVIEW_DECISION_IMPORTER",
            "version": "1.0",
            "status": "REVIEW_LOG_IMPORTED",
            "row": r,
            "review": row,
            "canonical_effect": "NONE",
            "note": "Imported for audit only; does not approve evidence, close gates, write SSoT, or allow final use.",
            **CANON
        }
        emit_jsonl(args.out_jsonl, event)
        imported += 1

    summary = {
        "timestamp": now_iso(),
        "component": "REVIEW_DECISION_IMPORTER",
        "status": "BLOCK" if violations else "REVIEW_REQUIRED",
        "imported_rows": imported,
        "violations": violations,
        "decision": "No canonical status changed.",
        **CANON
    }
    emit_jsonl(args.out_jsonl, summary)

    print(json.dumps(summary, ensure_ascii=False, indent=2) if args.json else f"✅ Imported review rows: {imported}")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
