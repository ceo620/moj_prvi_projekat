# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 15_filefinding_register_builder.py
# PURPOSE: Build or normalize FileFinding Register from a folder scan
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import hashlib
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_BLOCK = 1
EXIT_FILE_ERROR = 2
EXIT_DEPENDENCY_ERROR = 5

SHEET = "FileFinding Register"
DEFAULT_EXTS = {".pdf", ".docx", ".xlsx", ".xls", ".csv", ".txt", ".md", ".json", ".py", ".ps1"}

HEADERS = [
    "FileFinding_ID",
    "Evidence_ID",
    "Signal_ID",
    "File_Path",
    "File_Name",
    "File_Extension",
    "SHA256",
    "Size_Bytes",
    "Modified_UTC",
    "Finding_Type",
    "Weight",
    "Status",
    "Reviewer",
    "Notes",
    "Final_Use_Allowed"
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def ensure_sheet(wb, name, headers):
    if name not in wb.sheetnames:
        ws = wb.create_sheet(name)
        ws.append(headers)
        return ws
    ws = wb[name]
    existing = [cell.value for cell in ws[1]]
    if not any(existing):
        ws.append(headers)
    else:
        for header in headers:
            if header not in existing:
                ws.cell(row=1, column=len(existing) + 1, value=header)
                existing.append(header)
    return ws

def header_map(ws):
    return {str(c.value).strip(): idx for idx, c in enumerate(ws[1], start=1) if c.value}

def existing_sha_set(ws):
    hm = header_map(ws)
    col = hm.get("SHA256")
    if not col:
        return set()
    values = set()
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, col).value
        if v:
            values.add(str(v).strip())
    return values

def next_id(ws):
    return f"FF-{ws.max_row:05d}"

def write_row(ws, data):
    hm = header_map(ws)
    row = ws.max_row + 1
    for key, value in data.items():
        col = hm.get(key)
        if col:
            ws.cell(row=row, column=col, value=value)

def scan_files(folder, recursive):
    folder = Path(folder)
    iterator = folder.rglob("*") if recursive else folder.glob("*")
    for path in iterator:
        if path.is_file() and path.suffix.lower() in DEFAULT_EXTS:
            yield path

def main():
    parser = argparse.ArgumentParser(description="Build FileFinding Register")
    parser.add_argument("--excel", required=True, help="Excel workbook path")
    parser.add_argument("--scan-root", required=True, help="Folder to scan")
    parser.add_argument("--recursive", action="store_true")
    parser.add_argument("--dry-run", action="store_true", help="Do not save workbook")
    args = parser.parse_args()

    if not os.path.exists(args.scan_root):
        print(f"❌ scan-root ne postoji: {args.scan_root}")
        return EXIT_FILE_ERROR

    if os.path.exists(args.excel):
        wb = openpyxl.load_workbook(args.excel)
    else:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    ws = ensure_sheet(wb, SHEET, HEADERS)
    known = existing_sha_set(ws)
    added = 0
    scanned = 0

    for path in scan_files(args.scan_root, args.recursive):
        scanned += 1
        try:
            digest = sha256_file(path)
        except Exception as exc:
            print(f"⚠️ SHA error {path}: {exc}")
            continue

        if digest in known:
            continue

        stat = path.stat()
        data = {
            "FileFinding_ID": next_id(ws),
            "Evidence_ID": "",
            "Signal_ID": "",
            "File_Path": str(path),
            "File_Name": path.name,
            "File_Extension": path.suffix.lower(),
            "SHA256": digest,
            "Size_Bytes": stat.st_size,
            "Modified_UTC": datetime.utcfromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
            "Finding_Type": "UNCLASSIFIED",
            "Weight": 0,
            "Status": "REVIEW_REQUIRED",
            "Reviewer": "",
            "Notes": "Auto-discovered; requires evidence mapping",
            "Final_Use_Allowed": "NO",
        }
        write_row(ws, data)
        known.add(digest)
        added += 1

    if not args.dry_run:
        Path(args.excel).parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.excel)

    print(json.dumps({
        "timestamp": now_iso(),
        "component": "FILEFINDING_REGISTER_BUILDER",
        "status": "PASS",
        "scanned_files": scanned,
        "added_findings": added,
        "excel": args.excel,
        "system_status": "SYSTEM RED",
        "step102": "LOCKED / NOT ACCEPTED",
        "final_use_allowed": "NO"
    }, ensure_ascii=False))

    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
