# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 03_ssot_readonly_exporter.py
# PURPOSE: Export read-only snapshot of Monolith/SSoT candidate
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import hashlib
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_OK = 0
EXIT_SOURCE_MISSING = 1
EXIT_READ_ERROR = 2
EXIT_WRITE_ERROR = 3
EXIT_DEPENDENCY_ERROR = 5

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def load_json(path):
    with open(path, "r", encoding="utf-8-sig") as f:
        return json.load(f)

def flatten_json(obj, prefix=""):
    rows = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            key = f"{prefix}.{k}" if prefix else str(k)
            rows.extend(flatten_json(v, key))
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            key = f"{prefix}[{i}]"
            rows.extend(flatten_json(v, key))
    else:
        rows.append((prefix, obj))
    return rows

def write_excel(snapshot, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Snapshot Summary"
    ws.append(["Field", "Value"])
    for k, v in snapshot["snapshot_metadata"].items():
        ws.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v])

    ws2 = wb.create_sheet("Flattened SSoT")
    ws2.append(["JSON_Path", "Value"])
    for path, value in flatten_json(snapshot["ssot_data"]):
        ws2.append([path, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])

    ws3 = wb.create_sheet("Canonical Locks")
    ws3.append(["Control", "Value"])
    for k, v in CANON.items():
        ws3.append([k, v])

    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def main():
    parser = argparse.ArgumentParser(description="TITAN read-only SSoT snapshot exporter")
    parser.add_argument("--monolith", required=True, help="Path to monolith_registry.json")
    parser.add_argument("--out-json", required=True, help="Output read-only snapshot JSON")
    parser.add_argument("--out-xlsx", help="Optional output XLSX snapshot")
    parser.add_argument("--json", action="store_true", help="Print JSON status")
    args = parser.parse_args()

    if not os.path.exists(args.monolith):
        status = {
            "timestamp": now_iso(),
            "component": "SSOT_READONLY_EXPORTER",
            "status": "BLOCK",
            "message": "Monolith/SSoT source missing",
            "monolith": args.monolith,
            **CANON,
        }
        print(json.dumps(status, ensure_ascii=False) if args.json else status["message"])
        return EXIT_SOURCE_MISSING

    try:
        ssot_data = load_json(args.monolith)
        source_sha256 = sha256_file(args.monolith)
        source_size = os.path.getsize(args.monolith)
    except Exception as exc:
        status = {
            "timestamp": now_iso(),
            "component": "SSOT_READONLY_EXPORTER",
            "status": "BLOCK",
            "message": f"Could not read Monolith/SSoT source: {exc}",
            "monolith": args.monolith,
            **CANON,
        }
        print(json.dumps(status, ensure_ascii=False) if args.json else status["message"])
        return EXIT_READ_ERROR

    snapshot = {
        "snapshot_metadata": {
            "created_at": now_iso(),
            "component": "SSOT_READONLY_EXPORTER",
            "version": "1.0",
            "source_path": args.monolith,
            "source_sha256": source_sha256,
            "source_size_bytes": source_size,
            "export_mode": "READ_ONLY",
            "canonical_status": CANON,
            "note": "This snapshot is an audit/export artifact only. It does not write canonical SSoT and does not allow final use."
        },
        "ssot_data": ssot_data
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")

        if args.out_xlsx:
            write_excel(snapshot, args.out_xlsx)

    except Exception as exc:
        status = {
            "timestamp": now_iso(),
            "component": "SSOT_READONLY_EXPORTER",
            "status": "BLOCK",
            "message": f"Could not write snapshot: {exc}",
            "out_json": args.out_json,
            "out_xlsx": args.out_xlsx,
            **CANON,
        }
        print(json.dumps(status, ensure_ascii=False) if args.json else status["message"])
        return EXIT_WRITE_ERROR

    status = {
        "timestamp": now_iso(),
        "component": "SSOT_READONLY_EXPORTER",
        "status": "PASS",
        "message": "Read-only SSoT snapshot exported",
        "source_sha256": source_sha256,
        "out_json": args.out_json,
        "out_xlsx": args.out_xlsx or "N/A",
        **CANON,
    }

    print(json.dumps(status, ensure_ascii=False) if args.json else "✅ Read-only SSoT snapshot exported")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
