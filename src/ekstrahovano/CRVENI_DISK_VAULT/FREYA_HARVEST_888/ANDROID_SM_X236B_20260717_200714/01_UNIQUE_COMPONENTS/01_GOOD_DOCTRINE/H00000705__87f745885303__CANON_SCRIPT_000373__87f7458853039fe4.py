# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 21_manual_review_queue_builder.py
# PURPOSE: Build manual review queue from evidence/gap candidate reports
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_FILE_ERROR = 2
EXIT_WRITE_ERROR = 3

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def load_json(path):
    if not path or not os.path.exists(path):
        return None
    with open(path, "r", encoding="utf-8-sig") as f:
        return json.load(f)

def append_rows(ws, rows):
    for row in rows:
        ws.append(row)

def main():
    parser = argparse.ArgumentParser(description="Build TITAN manual review queue")
    parser.add_argument("--gap-report", required=True)
    parser.add_argument("--bundle-manifest", help="Optional evidence_bundle_manifest.json")
    parser.add_argument("--out-xlsx", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    gap_report = load_json(args.gap_report)
    if gap_report is None:
        print(f"❌ Gap report ne postoji: {args.gap_report}")
        return EXIT_FILE_ERROR

    bundle_manifest = load_json(args.bundle_manifest) if args.bundle_manifest else None

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Review Queue"
    ws.append([
        "Review_ID",
        "Item_Type",
        "Source_ID",
        "Priority",
        "Reason",
        "Required_Action",
        "Reviewer",
        "Review_Status",
        "Can_Close_Gate",
        "Can_Approve_Evidence",
        "Final_Use_Allowed"
    ])

    idx = 1

    for item in gap_report.get("blocked_gaps", []):
        ws.append([
            f"REV-{idx:05d}",
            "BLOCKED_GAP",
            item.get("Gap_ID", ""),
            "HIGH",
            "; ".join(item.get("Missing", [])) or "Gap requires review",
            "Attach valid FileFinding/Evidence or keep gap open",
            "",
            "PENDING",
            "NO",
            "NO",
            "NO"
        ])
        idx += 1

    for item in gap_report.get("candidates", []):
        ws.append([
            f"REV-{idx:05d}",
            "GAP_CANDIDATE",
            item.get("Gap_ID", ""),
            "CRITICAL",
            "; ".join(item.get("Reasons", [])) or "Candidate requires manual review",
            "Manual reviewer must validate; script cannot close gap",
            "",
            "PENDING",
            "NO",
            "NO",
            "NO"
        ])
        idx += 1

    if bundle_manifest:
        for record in bundle_manifest.get("records", []):
            if not record.get("Hash_Match", False):
                ws.append([
                    f"REV-{idx:05d}",
                    "HASH_MISMATCH_OR_MISSING",
                    record.get("FileFinding_ID", ""),
                    "CRITICAL",
                    "Evidence file hash missing or mismatch",
                    "Verify file path and SHA256",
                    "",
                    "PENDING",
                    "NO",
                    "NO",
                    "NO"
                ])
                idx += 1

    ws2 = wb.create_sheet("Canonical Locks")
    ws2.append(["Control", "Value"])
    for k, v in CANON.items():
        ws2.append([k, v])

    ws3 = wb.create_sheet("Instructions")
    ws3.append(["Instruction_ID", "Instruction"])
    ws3.append(["INS-001", "This workbook is a review queue only."])
    ws3.append(["INS-002", "Do not treat any row as evidence approval."])
    ws3.append(["INS-003", "Do not close STEP102 from this workbook."])
    ws3.append(["INS-004", "SYSTEM RED — STEP102 LOCKED — NO FINAL USE."])

    try:
        Path(args.out_xlsx).parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.out_xlsx)
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    result = {
        "timestamp": now_iso(),
        "component": "MANUAL_REVIEW_QUEUE_BUILDER",
        "status": "REVIEW_REQUIRED",
        "queue_items": idx - 1,
        "out_xlsx": args.out_xlsx,
        **CANON,
    }

    print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else f"✅ Manual review queue: {args.out_xlsx}")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
