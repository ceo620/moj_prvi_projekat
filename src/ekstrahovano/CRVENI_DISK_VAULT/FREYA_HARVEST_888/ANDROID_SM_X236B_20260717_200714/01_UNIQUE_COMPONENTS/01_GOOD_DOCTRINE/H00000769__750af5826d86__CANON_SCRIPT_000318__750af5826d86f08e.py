# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 16_evidence_gap_loader.py
# PURPOSE: Create or normalize Evidence Gap Register with 6 default gaps
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

SHEET = "Evidence Gap Register"
HEADERS = [
    "Gap_ID",
    "Description",
    "Required_Evidence",
    "Linked_FileFinding_ID",
    "Status",
    "Owner",
    "Blocking_STEP",
    "Risk_If_Open",
    "Reviewer",
    "Notes",
    "Final_Use_Allowed"
]

DEFAULT_GAPS = [
    ("GAP-001", "Monolith registry provenance not verified", "FileFinding for monolith_registry.json with SHA256 and source confirmation"),
    ("GAP-002", "Config baseline not independently verified", "FileFinding for config.json with schema validation result"),
    ("GAP-003", "AI signals not fully mapped to FileFinding_ID", "All AI Signal_ID rows linked to FileFinding_ID and Evidence_ID"),
    ("GAP-004", "External API assumptions not verified", "Evidence for LME/Google Maps API source, timestamp and error handling"),
    ("GAP-005", "Forensic registry/log scope not approved", "Read-only scope document and FileFinding evidence for allowed artifacts"),
    ("GAP-006", "Final control tower validation not clean", "control_tower_validation_report.jsonl with no blocking schema errors"),
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def ensure_sheet(wb):
    if SHEET not in wb.sheetnames:
        ws = wb.create_sheet(SHEET)
        ws.append(HEADERS)
        return ws
    ws = wb[SHEET]
    existing = [cell.value for cell in ws[1]]
    if not any(existing):
        ws.append(HEADERS)
    else:
        for header in HEADERS:
            if header not in existing:
                ws.cell(row=1, column=len(existing) + 1, value=header)
                existing.append(header)
    return ws

def header_map(ws):
    return {str(c.value).strip(): idx for idx, c in enumerate(ws[1], start=1) if c.value}

def existing_gaps(ws):
    hm = header_map(ws)
    col = hm.get("Gap_ID")
    out = set()
    if col:
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, col).value
            if v:
                out.add(str(v).strip())
    return out

def append_gap(ws, gap_id, desc, evidence):
    hm = header_map(ws)
    row = ws.max_row + 1
    data = {
        "Gap_ID": gap_id,
        "Description": desc,
        "Required_Evidence": evidence,
        "Linked_FileFinding_ID": "",
        "Status": "OPEN",
        "Owner": "REVIEW_REQUIRED",
        "Blocking_STEP": "STEP102",
        "Risk_If_Open": "STEP102 remains LOCKED; FINAL_USE_ALLOWED remains NO",
        "Reviewer": "",
        "Notes": "Default canonical gap; update only with reviewed evidence",
        "Final_Use_Allowed": "NO",
    }
    for key, value in data.items():
        col = hm.get(key)
        if col:
            ws.cell(row=row, column=col, value=value)

def main():
    parser = argparse.ArgumentParser(description="Load Evidence Gap Register")
    parser.add_argument("--excel", required=True)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    if os.path.exists(args.excel):
        wb = openpyxl.load_workbook(args.excel)
    else:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    ws = ensure_sheet(wb)
    known = existing_gaps(ws)
    added = 0

    for gap_id, desc, evidence in DEFAULT_GAPS:
        if gap_id not in known:
            append_gap(ws, gap_id, desc, evidence)
            added += 1

    if not args.dry_run:
        Path(args.excel).parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.excel)

    print(json.dumps({
        "timestamp": now_iso(),
        "component": "EVIDENCE_GAP_LOADER",
        "status": "PASS",
        "gaps_added": added,
        "required_open_gaps": 6,
        "step102": "LOCKED / NOT ACCEPTED",
        "final_use_allowed": "NO"
    }, ensure_ascii=False))
    return 0

if __name__ == "__main__":
    sys.exit(main())
