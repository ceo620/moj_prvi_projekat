# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 102_final_run_ledger_builder.py
# PURPOSE: Build final run ledger from major report artifacts
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# NOTE: This is NOT STEP102 acceptance. STEP102 remains LOCKED.
# ============================================================

import argparse
import hashlib
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_LOCKED = 1
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

LEDGER_FILES = [
    "pipeline_state.json",
    "master_orchestrator_locked_plan.json",
    "post_orchestration_denial_receipt.json",
    "denial_receipt_consistency_check.json",
    "final_locked_state_certificate.json",
    "no_release_memorandum.json",
    "decision_ledger.json",
    "emergency_stop_ledger.json",
    "prefreeze_denial_precheck.json",
    "release_readiness.json",
    "preproduction_freeze_denied.json",
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def sha256_file(path):
    if not path or not path.exists() or not path.is_file():
        return "MISSING"
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def find_file(root, name):
    root = Path(root)
    direct = root / name
    if direct.exists() and direct.is_file():
        return direct
    matches = list(root.rglob(name)) if root.exists() else []
    return matches[0] if matches else None

def infer_status(path):
    if not path or not path.exists() or path.suffix.lower() != ".json":
        return "MISSING" if path is None else "REVIEW_REQUIRED"
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        return data.get("status", data.get("release_status", data.get("freeze_status", "UNKNOWN")))
    except Exception:
        return "UNREADABLE"

def write_xlsx(payload, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Final Run Ledger"
    headers = ["Ledger_ID", "Artifact", "Path", "Exists", "Status", "SHA256", "Size_Bytes", "Canonical_Effect", "Final_Use_Allowed"]
    ws.append(headers)
    for row in payload["records"]:
        ws.append([row.get(h, "") for h in headers])
    ws2 = wb.create_sheet("Canonical Locks")
    ws2.append(["Control", "Value"])
    for k, v in CANON.items():
        ws2.append([k, v])
    ws3 = wb.create_sheet("Important Note")
    ws3.append(["Note"])
    ws3.append(["102_final_run_ledger_builder.py does not accept STEP102. STEP102 remains LOCKED / NOT ACCEPTED."])
    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def main():
    parser = argparse.ArgumentParser(description="Build final run ledger; does not accept STEP102")
    parser.add_argument("--reports-root", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--out-xlsx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    records = []
    for idx, name in enumerate(LEDGER_FILES, start=1):
        p = find_file(args.reports_root, name)
        exists = p is not None and p.exists()
        records.append({
            "Ledger_ID": f"RUN-{idx:05d}",
            "Artifact": name,
            "Path": str(p) if p else "MISSING",
            "Exists": exists,
            "Status": infer_status(p),
            "SHA256": sha256_file(p),
            "Size_Bytes": p.stat().st_size if exists else 0,
            "Canonical_Effect": "NONE",
            "Final_Use_Allowed": "NO",
        })

    payload = {
        "timestamp": now_iso(),
        "component": "FINAL_RUN_LEDGER_BUILDER",
        "version": "1.0",
        "status": "STEP102_STILL_LOCKED",
        "record_count": len(records),
        "missing_count": sum(1 for r in records if not r["Exists"]),
        "records": records,
        "decision": "Final run ledger is audit-only. It does not accept STEP102 or allow final use.",
        **CANON
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.out_xlsx:
            write_xlsx(payload, args.out_xlsx)
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else "🔒 Final run ledger built; STEP102 still locked")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_LOCKED

if __name__ == "__main__":
    sys.exit(main())
