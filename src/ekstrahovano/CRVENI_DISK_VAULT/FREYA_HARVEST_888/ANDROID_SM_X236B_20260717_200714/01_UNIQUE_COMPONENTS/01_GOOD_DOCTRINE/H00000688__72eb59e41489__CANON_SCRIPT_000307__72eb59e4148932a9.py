# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 123_consolidated_locked_archive_index.py
# PURPOSE: Build consolidated index of locked/no-use archive contents
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import hashlib
import json
import zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_OK = 0
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def sha256_bytes(data):
    return hashlib.sha256(data).hexdigest()

def iter_zips(roots):
    for root in roots:
        p = Path(root)
        if not p.exists():
            continue
        if p.is_file() and p.suffix.lower() == ".zip":
            yield p
        else:
            yield from p.rglob("*.zip")

def collect(roots):
    rows = []
    by_hash = defaultdict(list)

    for archive in iter_zips(roots):
        try:
            with zipfile.ZipFile(archive, "r") as z:
                for info in z.infolist():
                    if info.is_dir():
                        continue
                    try:
                        data = z.read(info.filename)
                        digest = sha256_bytes(data)
                        row = {
                            "Content_ID": f"CNT-{len(rows)+1:05d}",
                            "Archive_Name": archive.name,
                            "Archive_Path": str(archive),
                            "Item_Name": info.filename,
                            "Item_Basename": Path(info.filename).name,
                            "Item_SHA256": digest,
                            "Item_Size_Bytes": info.file_size,
                            "Canonical_Effect": "NONE",
                            "Final_Use_Allowed": "NO",
                        }
                        rows.append(row)
                        by_hash[digest].append(row["Content_ID"])
                    except Exception:
                        continue
        except Exception:
            continue

    duplicate_groups = [
        {"SHA256": digest, "Content_IDs": ids, "Duplicate_Count": len(ids)}
        for digest, ids in by_hash.items()
        if len(ids) > 1
    ]

    return rows, duplicate_groups

def write_xlsx(payload, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Locked Archive Contents"
    headers = ["Content_ID", "Archive_Name", "Archive_Path", "Item_Name", "Item_Basename", "Item_SHA256", "Item_Size_Bytes", "Canonical_Effect", "Final_Use_Allowed"]
    ws.append(headers)
    for row in payload["contents"]:
        ws.append([row.get(h, "") for h in headers])

    ws2 = wb.create_sheet("Duplicate Content")
    ws2.append(["SHA256", "Duplicate_Count", "Content_IDs"])
    for row in payload["duplicate_groups"]:
        ws2.append([row["SHA256"], row["Duplicate_Count"], ", ".join(row["Content_IDs"])])

    ws3 = wb.create_sheet("Canonical Locks")
    ws3.append(["Control", "Value"])
    for k, v in CANON.items():
        ws3.append([k, v])

    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def main():
    parser = argparse.ArgumentParser(description="Build consolidated locked archive index")
    parser.add_argument("--roots", nargs="+", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--out-xlsx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    contents, duplicate_groups = collect(args.roots)

    payload = {
        "timestamp": now_iso(),
        "component": "CONSOLIDATED_LOCKED_ARCHIVE_INDEX",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "content_count": len(contents),
        "duplicate_group_count": len(duplicate_groups),
        "contents": contents,
        "duplicate_groups": duplicate_groups,
        "decision": "Consolidated archive index is read-only inventory. It does not approve final use.",
        **CANON,
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.out_xlsx:
            write_xlsx(payload, args.out_xlsx)
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else f"✅ Consolidated locked archive index built: {len(contents)} items")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    raise SystemExit(main())
