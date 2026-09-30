# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 69_reviewer_attestation_template.py
# PURPOSE: Generate reviewer attestation template with hard NO approval defaults
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

ATTESTATION_HEADERS = [
    "Attestation_ID",
    "Reviewer_Name",
    "Reviewer_Role",
    "Review_Date",
    "Reviewed_Artifact",
    "Artifact_SHA256",
    "Finding",
    "Decision",
    "Evidence_Approved",
    "Gate_Closed",
    "SSOT_Write_Approved",
    "STEP102_Unlock_Approved",
    "Final_Use_Allowed",
    "Reviewer_Notes",
    "Signature_Text"
]

DEFAULT_ROWS = [
    {
        "Attestation_ID": "ATT-00001",
        "Reviewer_Name": "",
        "Reviewer_Role": "",
        "Review_Date": "",
        "Reviewed_Artifact": "NO_RELEASE_CLOSEOUT",
        "Artifact_SHA256": "",
        "Finding": "REVIEW_REQUIRED",
        "Decision": "NO_APPROVAL",
        "Evidence_Approved": "NO",
        "Gate_Closed": "NO",
        "SSOT_Write_Approved": "NO",
        "STEP102_Unlock_Approved": "NO",
        "Final_Use_Allowed": "NO",
        "Reviewer_Notes": "",
        "Signature_Text": ""
    }
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def main():
    parser = argparse.ArgumentParser(description="Build reviewer attestation template")
    parser.add_argument("--out-xlsx", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    payload = {
        "timestamp": now_iso(),
        "component": "REVIEWER_ATTESTATION_TEMPLATE",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "records": DEFAULT_ROWS,
        "decision": "Template defaults to NO approval. It does not approve anything automatically.",
        **CANON
    }

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reviewer Attestation"
        ws.append(ATTESTATION_HEADERS)
        for row in DEFAULT_ROWS:
            ws.append([row.get(h, "") for h in ATTESTATION_HEADERS])

        ws2 = wb.create_sheet("Canonical Locks")
        ws2.append(["Control", "Value"])
        for k, v in CANON.items():
            ws2.append([k, v])

        ws3 = wb.create_sheet("Instructions")
        ws3.append(["Instruction_ID", "Instruction"])
        ws3.append(["INS-001", "This template cannot approve evidence or close gates."])
        ws3.append(["INS-002", "Reviewer notes are advisory until separate governance process accepts them."])
        ws3.append(["INS-003", "SYSTEM RED and STEP102 LOCKED remain unchanged."])

        Path(args.out_xlsx).parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.out_xlsx)

        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else f"✅ Reviewer attestation template: {args.out_xlsx}")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
