# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 19_evidence_bundle_builder.py
# PURPOSE: Build non-approval evidence bundle from FileFinding Register
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import hashlib
import json
import os
import shutil
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_BLOCK = 1
EXIT_FILE_ERROR = 2
EXIT_SCHEMA_ERROR = 3

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

SHEET = "FileFinding Register"
REQUIRED_COLUMNS = ["FileFinding_ID", "Evidence_ID", "Signal_ID", "File_Path", "SHA256", "Weight", "Status"]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def stamp():
    return datetime.now().strftime("%Y%m%d_%H%M%S")

def norm(value):
    return "" if value is None else str(value).strip()

def headers(ws):
    return {norm(c.value): i for i, c in enumerate(ws[1], start=1) if norm(c.value)}

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def safe_name(name):
    keep = []
    for ch in name:
        if ch.isalnum() or ch in "-_. ":
            keep.append(ch)
        else:
            keep.append("_")
    return "".join(keep).strip()[:180] or "unnamed"

def main():
    parser = argparse.ArgumentParser(description="Build TITAN evidence bundle without approval")
    parser.add_argument("--excel", required=True)
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--copy-files", action="store_true", help="Copy source files into evidence bundle")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print(f"❌ Excel ne postoji: {args.excel}")
        return EXIT_FILE_ERROR

    try:
        wb = openpyxl.load_workbook(args.excel, data_only=True)
    except Exception as exc:
        print(f"❌ Excel nije čitljiv: {exc}")
        return EXIT_FILE_ERROR

    if SHEET not in wb.sheetnames:
        print(f"❌ Nedostaje sheet: {SHEET}")
        return EXIT_SCHEMA_ERROR

    ws = wb[SHEET]
    h = headers(ws)
    missing = [c for c in REQUIRED_COLUMNS if c not in h]
    if missing:
        print(f"❌ Nedostaju kolone: {missing}")
        return EXIT_SCHEMA_ERROR

    bundle_id = f"EVIDENCE_BUNDLE_{stamp()}"
    bundle_dir = Path(args.out_dir) / bundle_id
    files_dir = bundle_dir / "files"
    bundle_dir.mkdir(parents=True, exist_ok=True)
    if args.copy_files:
        files_dir.mkdir(parents=True, exist_ok=True)

    manifest = {
        "bundle_id": bundle_id,
        "created_at": now_iso(),
        "component": "EVIDENCE_BUNDLE_BUILDER",
        "version": "1.0",
        "copy_files": args.copy_files,
        "records": [],
        "warnings": [],
        **CANON,
        "note": "Evidence bundle is non-approval artifact. Manual review is still required."
    }

    for r in range(2, ws.max_row + 1):
        row = {k: norm(ws.cell(r, h[k]).value) for k in h.keys()}
        if not any(row.values()):
            continue

        fid = row.get("FileFinding_ID", "")
        file_path = row.get("File_Path", "")
        expected_sha = row.get("SHA256", "").lower()

        record = {
            "row": r,
            "FileFinding_ID": fid,
            "Evidence_ID": row.get("Evidence_ID", ""),
            "Signal_ID": row.get("Signal_ID", ""),
            "File_Path": file_path,
            "Expected_SHA256": expected_sha,
            "Actual_SHA256": "MISSING",
            "Hash_Match": False,
            "Weight": row.get("Weight", ""),
            "Status": row.get("Status", ""),
            "Copied_To": "N/A",
            "Final_Use_Allowed": "NO"
        }

        if file_path and os.path.exists(file_path) and os.path.isfile(file_path):
            actual_sha = sha256_file(file_path).lower()
            record["Actual_SHA256"] = actual_sha
            record["Hash_Match"] = bool(expected_sha and actual_sha == expected_sha)

            if args.copy_files:
                dest_name = f"{safe_name(fid)}__{safe_name(Path(file_path).name)}"
                dest = files_dir / dest_name
                shutil.copy2(file_path, dest)
                record["Copied_To"] = str(dest)
        else:
            manifest["warnings"].append({"row": r, "FileFinding_ID": fid, "warning": "File missing or invalid path"})

        manifest["records"].append(record)

    manifest_path = bundle_dir / "evidence_bundle_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    result = {
        "timestamp": now_iso(),
        "component": "EVIDENCE_BUNDLE_BUILDER",
        "status": "REVIEW_REQUIRED",
        "bundle_dir": str(bundle_dir),
        "manifest": str(manifest_path),
        "records": len(manifest["records"]),
        "warnings": len(manifest["warnings"]),
        **CANON,
    }

    print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else f"✅ Evidence bundle built: {bundle_dir}")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
