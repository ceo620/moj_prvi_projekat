# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 161_review_only_chain_archive_index.py
# PURPOSE: Index review-only chain archives and receipts
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import hashlib
import json
import zipfile
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_OK = 0
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "RISK_BASELINE": 890,
    "P0_GATES": "10/10 BLOCKED",
    "EVIDENCE_GAPS_REMAINING": 6,
    "EVIDENCE_APPROVED": 0,
    "GATES_CLOSED": 0,
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def inspect_zip(path):
    if path.suffix.lower() != ".zip":
        return "N/A", "", 0
    try:
        with zipfile.ZipFile(path, "r") as z:
            bad = z.testzip()
            names = z.namelist()
        return bad is None, bad or "", len(names)
    except Exception as exc:
        return False, str(exc), 0

def classify(path):
    n = path.name.lower()
    if "review_only_chain_archive" in n:
        return "REVIEW_ONLY_CHAIN_ARCHIVE"
    if "review_only_chain_archive_receipt" in n:
        return "REVIEW_ONLY_CHAIN_RECEIPT"
    if "review_only_chain_seal" in n:
        return "REVIEW_ONLY_CHAIN_SEAL"
    if "review_only_super_index" in n:
        return "REVIEW_ONLY_SUPER_INDEX"
    return "REVIEW_ONLY_CHAIN_ARTIFACT"

def collect(roots):
    rows = []
    for root in roots:
        p = Path(root)
        if not p.exists():
            continue
        iterator = [p] if p.is_file() else p.rglob("*")
        for item in iterator:
            if not item.is_file():
                continue
            lname = item.name.lower()
            if "review_only_chain" not in lname and "review_only_super" not in lname:
                continue
            if item.suffix.lower() not in {".zip", ".json", ".md", ".xlsx", ".txt"}:
                continue
            valid_zip, bad_member, count = inspect_zip(item)
            rows.append({
                "Chain_Archive_Index_ID": f"ROCA-{len(rows)+1:05d}",
                "Artifact_Name": item.name,
                "Artifact_Type": classify(item),
                "Path": str(item),
                "SHA256": sha256_file(item),
                "Size_Bytes": item.stat().st_size,
                "Modified_UTC": datetime.utcfromtimestamp(item.stat().st_mtime).isoformat(timespec="seconds"),
                "Valid_Zip": valid_zip,
                "Bad_Member": bad_member,
                "Contained_Item_Count": count,
                "Review_Only": "YES",
                "Operational_Action_Allowed": "NO",
                "Canonical_Effect": "NONE",
                "Final_Use_Allowed": "NO",
            })
    return rows

def index_hash(records):
    text = "\n".join(f"{r['Chain_Archive_Index_ID']}|{r['Path']}|{r['SHA256']}" for r in records)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()

def write_xlsx(payload, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RO Chain Archive Index"
    headers = [
        "Chain_Archive_Index_ID", "Artifact_Name", "Artifact_Type", "Path", "SHA256",
        "Size_Bytes", "Modified_UTC", "Valid_Zip", "Bad_Member", "Contained_Item_Count",
        "Review_Only", "Operational_Action_Allowed", "Canonical_Effect", "Final_Use_Allowed"
    ]
    ws.append(headers)
    for row in payload["records"]:
        ws.append([row.get(h, "") for h in headers])
    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    ws2.append(["Index_Hash", payload["index_hash"]])
    ws2.append(["Record_Count", payload["record_count"]])
    ws3 = wb.create_sheet("Canonical Locks")
    ws3.append(["Control", "Value"])
    for k, v in CANON.items():
        ws3.append([k, v])
    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def main():
    parser = argparse.ArgumentParser(description="Index review-only chain archives")
    parser.add_argument("--roots", nargs="+", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--out-xlsx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    records = collect(args.roots)
    payload = {
        "timestamp": now_iso(),
        "component": "REVIEW_ONLY_CHAIN_ARCHIVE_INDEX",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "record_count": len(records),
        "index_hash": index_hash(records),
        "records": records,
        "decision": "Review-only chain archive index is audit inventory only. It grants no authority.",
        **CANON,
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.out_xlsx:
            write_xlsx(payload, args.out_xlsx)
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else f"✅ Review-only chain archive index built: {len(records)} records")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    raise SystemExit(main())
