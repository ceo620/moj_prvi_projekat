# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 86_archive_reconciliation_builder.py
# PURPOSE: Reconcile closeout, denial, certificate and audit board archives
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import hashlib
import json
import os
import sys
import zipfile
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_OK = 0
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def inspect_zip(path):
    try:
        with zipfile.ZipFile(path, "r") as z:
            bad = z.testzip()
            names = z.namelist()
        return bad is None, bad or "", names
    except Exception as exc:
        return False, str(exc), []

def collect_archives(roots):
    rows = []
    for root in roots:
        p = Path(root)
        if not p.exists():
            continue
        iterator = [p] if p.is_file() else p.rglob("*.zip")
        for item in iterator:
            if item.is_file() and item.suffix.lower() == ".zip":
                valid, bad, names = inspect_zip(item)
                rows.append({
                    "Archive_ID": f"ARCH-{len(rows)+1:05d}",
                    "Archive_Name": item.name,
                    "Path": str(item),
                    "SHA256": sha256_file(item),
                    "Valid_Zip": valid,
                    "Bad_Member": bad,
                    "Item_Count": len(names),
                    "Items": names,
                    "Size_Bytes": item.stat().st_size,
                    "Modified_UTC": datetime.utcfromtimestamp(item.stat().st_mtime).isoformat(timespec="seconds"),
                    "Review_Status": "REVIEW_REQUIRED",
                    "Final_Use_Allowed": "NO"
                })
    return rows

def classify_archive(name):
    n = name.lower()
    if "certificate" in n:
        return "CERTIFICATE_ARCHIVE"
    if "denial" in n:
        return "RELEASE_DENIAL_ARCHIVE"
    if "closeout" in n:
        return "NO_RELEASE_CLOSEOUT"
    if "audit_board" in n:
        return "AUDIT_BOARD_PACK"
    if "audit_pack" in n:
        return "AUDIT_PACK"
    return "OTHER_ARCHIVE"

def write_xlsx(payload, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Archive Reconciliation"
    headers = ["Archive_ID", "Archive_Name", "Archive_Type", "Path", "SHA256", "Valid_Zip", "Bad_Member", "Item_Count", "Size_Bytes", "Modified_UTC", "Review_Status", "Final_Use_Allowed"]
    ws.append(headers)
    for row in payload["archives"]:
        ws.append([row.get(h, "") for h in headers])

    ws2 = wb.create_sheet("Archive Contents")
    ws2.append(["Archive_ID", "Archive_Name", "Item"])
    for row in payload["archives"]:
        for item in row.get("Items", []):
            ws2.append([row["Archive_ID"], row["Archive_Name"], item])

    ws3 = wb.create_sheet("Canonical Locks")
    ws3.append(["Control", "Value"])
    for k, v in CANON.items():
        ws3.append([k, v])

    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def main():
    parser = argparse.ArgumentParser(description="TITAN archive reconciliation builder")
    parser.add_argument("--roots", nargs="+", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--out-xlsx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    archives = collect_archives(args.roots)
    for row in archives:
        row["Archive_Type"] = classify_archive(row["Archive_Name"])

    payload = {
        "timestamp": now_iso(),
        "component": "ARCHIVE_RECONCILIATION_BUILDER",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "archive_count": len(archives),
        "invalid_count": sum(1 for a in archives if not a["Valid_Zip"]),
        "archives": archives,
        "decision": "Archive reconciliation is inventory only. It does not approve release or final use.",
        **CANON
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.out_xlsx:
            write_xlsx(payload, args.out_xlsx)
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else f"✅ Archive reconciliation built: {len(archives)} archives")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
