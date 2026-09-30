# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import hashlib

out = Path("02_TECHNICAL_CEDIS_TUZI")
out.mkdir(parents=True, exist_ok=True)

doc = Document()

sec = doc.sections[0]
sec.top_margin = Inches(0.75)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.95)
sec.right_margin = Inches(0.85)

style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(10.5)

def add_p(text="", bold=False, size=10.5, after=4, align=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    return p

# HEADER
add_p("ARS METAL INDUSTRIES D.O.O. PODGORICA", True, 11, 0)
add_p("Industrijski projekat TITAN 1", False, 10, 10)

# TITLE
add_p("MEMORANDUM", True, 15, 12, WD_ALIGN_PARAGRAPH.CENTER)

# META BLOCK
tbl = doc.add_table(rows=3, cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
tbl.autofit = True

meta = [
    ("Prima:", "Opština Tuzi"),
    ("Kopija:", "CEDIS – Crnogorski elektrodistributivni sistem"),
    ("Predmet:", "Zahtjev za pisano izjašnjenje o statusu elektrodistributivne infrastrukture za projekat TITAN 1")
]

for i, (a,b) in enumerate(meta):
    c1 = tbl.cell(i,0)
    c2 = tbl.cell(i,1)

    p1 = c1.paragraphs[0]
    r1 = p1.add_run(a)
    r1.bold = True
    r1.font.name = "Arial"
    r1.font.size = Pt(10.5)

    p2 = c2.paragraphs[0]
    r2 = p2.add_run(b)
    r2.font.name = "Arial"
    r2.font.size = Pt(10.5)

doc.add_paragraph()

# BODY
add_p("Poštovani,", False, 10.5, 8)

body = [
"Ars Metal Industries D.O.O. Podgorica planira realizaciju projekta TITAN 1 – Fabrika transformatorskih kazana (tankova) u industrijskoj zoni Tuzi / KAP.",

"Predmet ovog zahtjeva odnosi se isključivo na projekat TITAN 1 i ne obuhvata druge buduće projekte, faze ili proizvodne programe.",

"Za potrebe daljeg planiranja i procjene izvodljivosti projekta, potrebno je pribaviti pisano izjašnjenje nadležnih institucija o statusu relevantne elektrodistributivne infrastrukture, raspoloživom kapacitetu i narednim proceduralnim koracima.",

"Preliminarni okvir priključne snage koji je predmet provjere iznosi do 12,5 MW, uz faznu realizaciju i naknadnu tehničku potvrdu kroz odgovarajuću dokumentaciju."
]

for t in body:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(t)
    r.font.name = "Arial"
    r.font.size = Pt(10.5)

add_p("Molimo Opštinu Tuzi da, u okviru svojih nadležnosti, uputi formalni upit prema CEDIS-u radi dobijanja pisanog izjašnjenja o sljedećem:", False, 10.5, 6)

items = [
"trenutnom statusu relevantne elektrodistributivne infrastrukture za lokaciju Tuzi / KAP,",
"očekivanim rokovima raspoloživosti kapaciteta,",
"preliminarnoj mogućnosti industrijskog priključenja,",
"daljim proceduralnim koracima za izdavanje tehničkih uslova,",
"tehničkim podacima koje je potrebno dostaviti u narednoj fazi."
]

for item in items:
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run("• " + item)
    r.font.name = "Arial"
    r.font.size = Pt(10.5)

doc.add_paragraph()

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run(
"Investitor je spreman da, po zahtjevu nadležnih institucija, dostavi dodatne tehničke podatke i raspoloživu projektnu dokumentaciju relevantnu za dalje razmatranje zahtjeva."
)
r.font.name = "Arial"
r.font.size = Pt(10.5)

# SIGNATURE
doc.add_paragraph()

add_p("Sa poštovanjem,", False, 10.5, 12)

add_p("DANIJELA KESKIN", True, 10.5, 0)
add_p("Director", False, 10, 0)
add_p("Ars Metal Industries D.O.O. Podgorica", False, 10, 0)
add_p("+382 68 323 339", False, 10, 0)
add_p("d.djurovic@adsmetal.com.tr", False, 10, 0)

target = out / "TEHNICKI_DOPIS_CEDIS_TUZI_TITAN1_INSTITUTIONAL_v2.docx"
doc.save(target)

h = hashlib.sha256()
with open(target, "rb") as f:
    h.update(f.read())

print("CREATED:", target)
print("SHA256:", h.hexdigest())
