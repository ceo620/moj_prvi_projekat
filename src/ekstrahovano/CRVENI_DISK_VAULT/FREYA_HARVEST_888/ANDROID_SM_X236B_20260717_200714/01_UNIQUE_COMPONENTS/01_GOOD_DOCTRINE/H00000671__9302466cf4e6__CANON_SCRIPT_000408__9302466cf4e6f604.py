# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 129_final_immutable_ledger.py
# PURPOSE: Build final immutable-style ledger of terminal locked artifacts
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import hashlib
import json
import os
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_LOCKED = 1
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "RISK_BASELINE": 890,
    "P0_GATES": "10/10 BLOCKED",
    "EVIDENCE_GAPS_REMAINING": 6,
    "EVIDENCE_APPROVED": 0,
    "GATES_CLOSED": 0,
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

LEDGER_ARTIFACTS = [
    "final_master_seal_record.json",
    "master_seal_consistency_check.json",
    "final_master_index.json",
    "master_index_integrity_validation.json",
    "final_reconciliation_memorandum.json",
    "all_locked_status_register.json",
    "nothing_accepted_memorandum.json",
    "review_only_proof.json",
    "do_not_run_notice.json",
    "system_remains_locked_summary.json",
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def sha256_file(path):
    if not path or not path.exists() or not path.is_file():
        return "MISSING"
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def find_file(root, name):
    root = Path(root)
    direct = root / name
    if direct.exists() and direct.is_file():
        return direct
    matches = list(root.rglob(name)) if root.exists() else []
    return matches[0] if matches else None

def infer_status(path):
    if not path or not path.exists():
        return "MISSING"
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        return data.get("status", data.get("release_status", data.get("freeze_status", "UNKNOWN")))
    except Exception:
        return "UNREADABLE"

def write_xlsx(payload, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Final Immutable Ledger"
    headers = ["Ledger_ID", "Artifact", "Path", "Exists", "Status", "SHA256", "Size_Bytes", "Canonical_Effect", "Final_Use_Allowed"]
    ws.append(headers)
    for row in payload["records"]:
        ws.append([row.get(h, "") for h in headers])

    ws2 = wb.create_sheet("Canonical Locks")
    ws2.append(["Control", "Value"])
    for k, v in CANON.items():
        ws2.append([k, v])

    ws3 = wb.create_sheet("Important")
    ws3.append(["Statement"])
    ws3.append(["Immutable-style ledger is local hash inventory only. It is not legal immutability or approval."])

    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def main():
    parser = argparse.ArgumentParser(description="Build final immutable-style ledger")
    parser.add_argument("--reports-root", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--out-xlsx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    records = []
    concat = ""

    for idx, name in enumerate(LEDGER_ARTIFACTS, start=1):
        p = find_file(args.reports_root, name)
        digest = sha256_file(p)
        exists = p is not None and p.exists()
        rec = {
            "Ledger_ID": f"IML-{idx:05d}",
            "Artifact": name,
            "Path": str(p) if p else "MISSING",
            "Exists": exists,
            "Status": infer_status(p),
            "SHA256": digest,
            "Size_Bytes": p.stat().st_size if exists else 0,
            "Canonical_Effect": "NONE",
            "Final_Use_Allowed": "NO",
        }
        records.append(rec)
        concat += f"{rec['Ledger_ID']}|{rec['Artifact']}|{rec['SHA256']}|{rec['Status']}\n"

    ledger_hash = hashlib.sha256(concat.encode("utf-8")).hexdigest()

    payload = {
        "timestamp": now_iso(),
        "component": "FINAL_IMMUTABLE_LEDGER",
        "version": "1.0",
        "status": "FINAL_IMMUTABLE_LEDGER_LOCKED",
        "ledger_hash": ledger_hash,
        "record_count": len(records),
        "missing_count": sum(1 for r in records if not r["Exists"]),
        "records": records,
        "decision": "Ledger records locked artifact hashes only. It is not approval, release, or final use.",
        **CANON
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.out_xlsx:
            write_xlsx(payload, args.out_xlsx)
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else "🔒 Final immutable-style ledger built")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_LOCKED

if __name__ == "__main__":
    raise SystemExit(main())
