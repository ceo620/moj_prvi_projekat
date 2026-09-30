# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 146_terminal_archive_freeze_index.py
# PURPOSE: Create terminal freeze index for archive-of-archives state
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import hashlib
import json
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_LOCKED = 1
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "RISK_BASELINE": 890,
    "P0_GATES": "10/10 BLOCKED",
    "EVIDENCE_GAPS_REMAINING": 6,
    "EVIDENCE_APPROVED": 0,
    "GATES_CLOSED": 0,
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

EXTS = {".zip", ".json", ".xlsx", ".md", ".txt"}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def collect(roots):
    rows = []
    for root in roots:
        p = Path(root)
        if not p.exists():
            continue
        iterator = [p] if p.is_file() else p.rglob("*")
        for item in iterator:
            if item.is_file() and item.suffix.lower() in EXTS:
                rows.append({
                    "Freeze_Item_ID": f"FRZ-{len(rows)+1:05d}",
                    "Artifact_Name": item.name,
                    "Path": str(item),
                    "SHA256": sha256_file(item),
                    "Size_Bytes": item.stat().st_size,
                    "Modified_UTC": datetime.utcfromtimestamp(item.stat().st_mtime).isoformat(timespec="seconds"),
                    "Freeze_Status": "FROZEN_FOR_REVIEW_ONLY",
                    "Mutable_After_Freeze": "NO_POLICY",
                    "Canonical_Effect": "NONE",
                    "Final_Use_Allowed": "NO",
                })
    return rows

def freeze_hash(records):
    material = "\n".join(f"{r['Freeze_Item_ID']}|{r['Path']}|{r['SHA256']}|{r['Size_Bytes']}" for r in records)
    return hashlib.sha256(material.encode("utf-8")).hexdigest()

def write_xlsx(payload, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Terminal Freeze Index"
    headers = ["Freeze_Item_ID", "Artifact_Name", "Path", "SHA256", "Size_Bytes", "Modified_UTC", "Freeze_Status", "Mutable_After_Freeze", "Canonical_Effect", "Final_Use_Allowed"]
    ws.append(headers)
    for row in payload["records"]:
        ws.append([row.get(h, "") for h in headers])
    ws2 = wb.create_sheet("Freeze Summary")
    ws2.append(["Metric", "Value"])
    ws2.append(["Freeze_Hash", payload["freeze_hash"]])
    ws2.append(["Record_Count", payload["record_count"]])
    ws2.append(["Status", payload["status"]])
    ws3 = wb.create_sheet("Canonical Locks")
    ws3.append(["Control", "Value"])
    for k, v in CANON.items():
        ws3.append([k, v])
    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def main():
    parser = argparse.ArgumentParser(description="Create terminal archive freeze index")
    parser.add_argument("--roots", nargs="+", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--out-xlsx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    records = collect(args.roots)
    payload = {
        "timestamp": now_iso(),
        "component": "TERMINAL_ARCHIVE_FREEZE_INDEX",
        "version": "1.0",
        "status": "TERMINAL_ARCHIVES_FROZEN_LOCKED",
        "freeze_hash": freeze_hash(records),
        "record_count": len(records),
        "records": records,
        "decision": "Terminal freeze index records archive state only. It is not approval, release, handoff, or final use.",
        **CANON,
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.out_xlsx:
            write_xlsx(payload, args.out_xlsx)
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else "🔒 Terminal archive freeze index built")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_LOCKED

if __name__ == "__main__":
    raise SystemExit(main())
