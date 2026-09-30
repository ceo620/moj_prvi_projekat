# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE

# ============================================================
# TITAN_KERNEL: 20_gap_closure_candidate_reporter.py
# PURPOSE: Report closure candidates without closing evidence gaps
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_FILE_ERROR = 2
EXIT_SCHEMA_ERROR = 3

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

FILEFINDING_SHEET = "FileFinding Register"
GAP_SHEET = "Evidence Gap Register"

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def norm(value):
    return "" if value is None else str(value).strip()

def headers(ws):
    return {norm(c.value): i for i, c in enumerate(ws[1], start=1) if norm(c.value)}

def split_ids(value):
    return [x.strip() for x in value.replace(";", ",").split(",") if x.strip()]

def write_json(path, payload):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

def main():
    parser = argparse.ArgumentParser(description="Create non-closing gap candidate report")
    parser.add_argument("--excel", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print(f"❌ Excel ne postoji: {args.excel}")
        return EXIT_FILE_ERROR

    try:
        wb = openpyxl.load_workbook(args.excel, data_only=True)
    except Exception as exc:
        print(f"❌ Excel nije čitljiv: {exc}")
        return EXIT_FILE_ERROR

    missing_sheets = [s for s in [FILEFINDING_SHEET, GAP_SHEET] if s not in wb.sheetnames]
    if missing_sheets:
        print(f"❌ Nedostaju sheetovi: {missing_sheets}")
        return EXIT_SCHEMA_ERROR

    ff_ws = wb[FILEFINDING_SHEET]
    gap_ws = wb[GAP_SHEET]
    ff_h = headers(ff_ws)
    gap_h = headers(gap_ws)

    required_ff = ["FileFinding_ID", "Evidence_ID", "SHA256", "Weight", "Status", "Reviewer"]
    required_gap = ["Gap_ID", "Description", "Required_Evidence", "Linked_FileFinding_ID", "Status", "Blocking_STEP"]

    missing_cols = [c for c in required_ff if c not in ff_h] + [c for c in required_gap if c not in gap_h]
    if missing_cols:
        print(f"❌ Nedostaju kolone: {missing_cols}")
        return EXIT_SCHEMA_ERROR

    findings = {}
    for r in range(2, ff_ws.max_row + 1):
        fid = norm(ff_ws.cell(r, ff_h["FileFinding_ID"]).value)
        if not fid:
            continue
        findings[fid] = {
            "FileFinding_ID": fid,
            "Evidence_ID": norm(ff_ws.cell(r, ff_h["Evidence_ID"]).value),
            "SHA256": norm(ff_ws.cell(r, ff_h["SHA256"]).value),
            "Weight": norm(ff_ws.cell(r, ff_h["Weight"]).value),
            "Status": norm(ff_ws.cell(r, ff_h["Status"]).value),
            "Reviewer": norm(ff_ws.cell(r, ff_h["Reviewer"]).value),
        }

    candidates = []
    blockers = []

    closed_statuses = {"CLOSED", "APPROVED", "RESOLVED", "VALIDATED"}

    for r in range(2, gap_ws.max_row + 1):
        gap_id = norm(gap_ws.cell(r, gap_h["Gap_ID"]).value)
        if not gap_id:
            continue

        linked_raw = norm(gap_ws.cell(r, gap_h["Linked_FileFinding_ID"]).value)
        status = norm(gap_ws.cell(r, gap_h["Status"]).value).upper()
        linked_ids = split_ids(linked_raw)

        candidate = {
            "Gap_ID": gap_id,
            "Row": r,
            "Current_Status": status,
            "Linked_FileFinding_IDs": linked_ids,
            "Candidate_For_Manual_Review": False,
            "Reasons": [],
            "Missing": [],
            "Final_Use_Allowed": "NO",
            "Action": "DO_NOT_CLOSE_AUTOMATICALLY"
        }

        if status in closed_statuses:
            candidate["Reasons"].append("Already marked closed/resolved in sheet; requires audit confirmation")
        if not linked_ids:
            candidate["Missing"].append("Linked_FileFinding_ID")
        else:
            all_ok = True
            for fid in linked_ids:
                f = findings.get(fid)
                if not f:
                    all_ok = False
                    candidate["Missing"].append(f"{fid}: not found in FileFinding Register")
                    continue

                if not f["Evidence_ID"]:
                    all_ok = False
                    candidate["Missing"].append(f"{fid}: Evidence_ID missing")
                if not f["SHA256"]:
                    all_ok = False
                    candidate["Missing"].append(f"{fid}: SHA256 missing")
                if str(f["Weight"]) in {"", "0", "0.0"}:
                    all_ok = False
                    candidate["Missing"].append(f"{fid}: Weight is 0 or missing")
                if not f["Reviewer"]:
                    all_ok = False
                    candidate["Missing"].append(f"{fid}: Reviewer missing")

            if all_ok:
                candidate["Candidate_For_Manual_Review"] = True
                candidate["Reasons"].append("All linked findings have Evidence_ID, SHA256, non-zero weight and reviewer")

        if candidate["Candidate_For_Manual_Review"]:
            candidates.append(candidate)
        else:
            blockers.append(candidate)

    report = {
        "timestamp": now_iso(),
        "component": "GAP_CLOSURE_CANDIDATE_REPORTER",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "candidate_count": len(candidates),
        "blocked_gap_count": len(blockers),
        "candidates": candidates,
        "blocked_gaps": blockers,
        "decision": "No gap is closed by this script. Manual review required.",
        **CANON,
    }

    write_json(args.out_json, report)

    print(json.dumps(report, ensure_ascii=False, indent=2) if args.json else f"✅ Gap candidate report: {args.out_json}")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
