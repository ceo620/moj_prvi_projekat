from __future__ import annotations

# FREYA_DIRECT_REPAIR_RUNTIME_GATE
import os as _freya_os, sys as _freya_sys
if _freya_os.environ.get("HUMAN_GATE_RUNTIME_APPROVED") != "YES":
    print("BLOCKED_BY_FREYA_HUMAN_GATE: runtime not approved")
    _freya_sys.exit(0)
# END_FREYA_DIRECT_REPAIR_RUNTIME_GATE



import concurrent.futures
import hashlib
import json
import os
import platform
import re
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path

import pandas as pd
from docx import Document
from pypdf import PdfReader
from tqdm import tqdm

try:
    import openpyxl
except ImportError:
    openpyxl = None

# =============================================================================
# TITAN V31 - GLOBAL RADAR INSTITUTIONAL
# =============================================================================

SYSTEM_OWNER = "DANIJELA DJUROVIC KESKIN"
RUN_TS = datetime.now()
RUN_ID = RUN_TS.strftime("%Y%m%d_%H%M%S")

DRIVE_LETTERS = ("G", "H", "I", "J", "D", "E", "F")
SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".xlsx", ".txt"})

PDF_PAGE_LIMIT = 10
DOCX_PARAGRAPH_LIMIT = 60
XLSX_SHEET_LIMIT = 5
XLSX_ROW_LIMIT = 30
TXT_CHAR_LIMIT = 20_000
PREVIEW_LIMIT = 300
RELATED_THRESHOLD = 0.60
RELATED_LIMIT = 3
MAX_WORKERS = min(16, max(4, (os.cpu_count() or 4) * 2))

CATEGORY_PREFIXES = {
    "FINANCE": "FIN",
    "FUNDING": "FND",
    "ESG": "ESG",
    "LEGAL": "LEG",
    "TECHNICAL": "TEC",
    "GOVERNANCE": "GOV",
}

# ---------------------------------------------------------------------------
# PATH DETECTION
# ---------------------------------------------------------------------------
def find_global_drive() -> Path | None:
    for letter in DRIVE_LETTERS:
        candidate = Path(f"{letter}:/My Drive")
        if candidate.exists():
            return candidate
    return None


ROOT_DIR = find_global_drive()
DESKTOP = Path.home() / "Desktop"
MASTER_FILE = DESKTOP / f"TITAN_GLOBAL_MASTER_V31_{RUN_ID}.xlsx"
JSON_FILE = DESKTOP / f"TITAN_GLOBAL_MASTER_V31_{RUN_ID}.json"

# ---------------------------------------------------------------------------
# ONTOLOGY
# ---------------------------------------------------------------------------
CATEGORY_RULES = {
    "FINANCE": {
        "dscr": 24,
        "cfads": 22,
        "irr": 18,
        "npv": 18,
        "cash flow": 18,
        "capex": 16,
        "opex": 16,
        "wacc": 18,
        "dsra": 20,
        "sources and uses": 20,
    },
    "FUNDING": {"grant": 24, "wbif": 24, "ipa": 24, "eu fund": 24, "ebrd": 22, "eib": 22},
    "ESG": {"green": 16, "sustainable": 16, "taxonomy": 20, "cbam": 20, "co2": 14, "solar": 14},
    "LEGAL": {
        "statute": 18,
        "contract": 18,
        "agreement": 18,
        "law": 14,
        "articles of association": 24,
        "ppp": 16,
    },
    "TECHNICAL": {"transformer": 14, "tank": 12, "automation": 16, "scada": 16, "mes": 16},
    "GOVERNANCE": {"investment committee": 20, "decision brief": 18, "board": 14},
}

GRANT_KEYWORDS = {"grant": 20, "wbif": 22, "ipa": 22, "eu fund": 22, "green": 10}
LENDER_KEYWORDS = {"ebrd": 20, "eib": 20, "dscr": 18, "cfads": 18, "dsra": 16}
ESG_KEYWORDS = {"esg": 18, "taxonomy": 18, "cbam": 18, "co2": 14, "solar": 14}

# ---------------------------------------------------------------------------
# HELPERS
# ---------------------------------------------------------------------------
def normalize(text: str | None) -> str:
    return re.sub(r"\s+", " ", (text or "")).casefold().strip()


def hash_file(path: Path) -> str:
    digest = hashlib.md5()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def confidence(scores: dict[str, int]) -> float:
    ordered_scores = sorted(scores.values(), reverse=True)
    if not ordered_scores or ordered_scores[0] == 0:
        return 0

    second = ordered_scores[1] if len(ordered_scores) > 1 else 0
    return round(min(100, 55 + (ordered_scores[0] - second) * 1.8), 1)


def similarity(left: str, right: str) -> float:
    return SequenceMatcher(None, left, right).ratio()


def doc_id(index: int, category: str) -> str:
    return f"{CATEGORY_PREFIXES.get(category, 'DOC')}-{index:05d}"


def safe_file_stats(path: Path) -> tuple[int | None, str]:
    try:
        stats = path.stat()
        modified_at = datetime.fromtimestamp(stats.st_mtime).isoformat(timespec="seconds")
        return stats.st_size, modified_at
    except OSError:
        return None, ""


def summarize_keyword_hits(text: str, weights: dict[str, int]) -> int:
    return sum(score for keyword, score in weights.items() if keyword in text)


# ---------------------------------------------------------------------------
# EXTRACTION
# ---------------------------------------------------------------------------
def extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return " ".join(page.extract_text() or "" for page in reader.pages[:PDF_PAGE_LIMIT])


def extract_docx(path: Path) -> str:
    document = Document(str(path))
    return " ".join(paragraph.text for paragraph in document.paragraphs[:DOCX_PARAGRAPH_LIMIT])


def extract_xlsx(path: Path) -> str:
    if not openpyxl:
        raise RuntimeError("openpyxl nije instaliran")

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        fragments: list[str] = []
        for worksheet in workbook.worksheets[:XLSX_SHEET_LIMIT]:
            for row in worksheet.iter_rows(max_row=XLSX_ROW_LIMIT, values_only=True):
                values = [str(value) for value in row if value not in (None, "")]
                if values:
                    fragments.append(" ".join(values))
        return " ".join(fragments)
    finally:
        workbook.close()


def extract_txt(path: Path) -> str:
    with path.open("r", encoding="utf-8", errors="ignore") as handle:
        return handle.read(TXT_CHAR_LIMIT)


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".txt": extract_txt,
}


def extract(path: Path) -> tuple[str, str]:
    extractor = EXTRACTORS.get(path.suffix.lower())
    if not extractor:
        return "", ""

    try:
        return extractor(path), ""
    except Exception as exc:  # noqa: BLE001 - želimo nastaviti skeniranje i evidentirati grešku
        return "", f"{type(exc).__name__}: {exc}"


# ---------------------------------------------------------------------------
# CLASSIFICATION
# ---------------------------------------------------------------------------
def classify(text: str) -> tuple[str, float, dict[str, int]]:
    scores = {
        category: sum(weight for keyword, weight in rules.items() if keyword in text)
        for category, rules in CATEGORY_RULES.items()
    }
    best = max(scores, key=scores.get, default="GENERAL")
    if scores.get(best, 0) == 0:
        best = "GENERAL"
    return best, confidence(scores), scores


# ---------------------------------------------------------------------------
# SCAN
# ---------------------------------------------------------------------------
def scan(index: int, path: Path) -> dict[str, object]:
    file_size, modified_at = safe_file_stats(path)
    raw_content, extraction_error = extract(path)
    content = normalize(raw_content)
    category, conf, _ = classify(content)

    status = "OK"
    if extraction_error:
        status = "WARN"
    elif not content:
        status = "EMPTY"

    try:
        file_hash = hash_file(path)
    except OSError as exc:
        file_hash = ""
        extraction_error = extraction_error or f"{type(exc).__name__}: {exc}"
        status = "WARN"

    return {
        "Document_ID": doc_id(index, category),
        "Name": path.name,
        "Category": category,
        "Confidence": conf,
        "Grant": summarize_keyword_hits(content, GRANT_KEYWORDS),
        "Lender": summarize_keyword_hits(content, LENDER_KEYWORDS),
        "ESG": summarize_keyword_hits(content, ESG_KEYWORDS),
        "Path": str(path),
        "Hash": file_hash,
        "Preview": content[:PREVIEW_LIMIT],
        "Status": status,
        "File_Size_KB": round(file_size / 1024, 2) if file_size is not None else None,
        "Modified_At": modified_at,
        "Error": extraction_error,
    }


# ---------------------------------------------------------------------------
# FILE ITERATOR
# ---------------------------------------------------------------------------
def files(root: Path):
    for current_dir, _, filenames in os.walk(root):
        for filename in filenames:
            if filename.startswith("~$"):
                continue

            path = Path(current_dir) / filename
            if path.suffix.lower() in SUPPORTED_EXTENSIONS:
                yield path


# ---------------------------------------------------------------------------
# POST PROCESSING
# ---------------------------------------------------------------------------
def build_related_documents(frame: pd.DataFrame) -> list[str]:
    if frame.empty:
        return []

    normalized_names = [normalize(Path(name).stem) for name in frame["Name"].fillna("")]
    related_map = {idx: [] for idx in frame.index}

    for left_pos, left_idx in enumerate(frame.index):
        for right_pos in range(left_pos + 1, len(frame.index)):
            right_idx = frame.index[right_pos]
            score = similarity(normalized_names[left_pos], normalized_names[right_pos])
            if score < RELATED_THRESHOLD:
                continue

            left_doc_id = frame.at[left_idx, "Document_ID"]
            right_doc_id = frame.at[right_idx, "Document_ID"]

            if pd.notna(right_doc_id):
                related_map[left_idx].append((score, right_doc_id))
            if pd.notna(left_doc_id):
                related_map[right_idx].append((score, left_doc_id))

    ordered_related = []
    for row_idx in frame.index:
        ranked = sorted(related_map[row_idx], key=lambda item: item[0], reverse=True)
        ordered_related.append(",".join(doc for _, doc in ranked[:RELATED_LIMIT]))
    return ordered_related


def export_json(frame: pd.DataFrame, destination: Path) -> None:
    serializable_frame = frame.astype(object).where(pd.notna(frame), None)
    with destination.open("w", encoding="utf-8") as handle:
        json.dump(serializable_frame.to_dict(orient="records"), handle, ensure_ascii=False, indent=2)


def safe_open_file(path: Path) -> None:
    if platform.system() != "Windows":
        return

    try:
        os.startfile(str(path))
    except OSError:
        pass


# ---------------------------------------------------------------------------
# EXECUTION
# ---------------------------------------------------------------------------
def run() -> pd.DataFrame | None:
    if not ROOT_DIR:
        print("Drive nije pronađen")
        return None

    print(f"TITAN V31 START: {ROOT_DIR}")
    print(f"SYSTEM OWNER: {SYSTEM_OWNER}")

    file_list = list(files(ROOT_DIR))
    if not file_list:
        print("Nema fajlova")
        return None

    results: list[dict[str, object]] = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        iterator = executor.map(scan, range(1, len(file_list) + 1), file_list)
        for result in tqdm(iterator, total=len(file_list), desc="Scanning"):
            results.append(result)

    frame = pd.DataFrame(results).sort_values(["Category", "Name"], na_position="last").reset_index(drop=True)

    if "Hash" in frame.columns:
        hashes = frame["Hash"].fillna("")
        frame["Duplicate"] = hashes.ne("") & hashes.duplicated(keep=False)

    frame["Related"] = build_related_documents(frame)

    frame.to_excel(MASTER_FILE, index=False)
    export_json(frame, JSON_FILE)

    print(f"DONE: {MASTER_FILE}")
    print(f"JSON: {JSON_FILE}")
    safe_open_file(MASTER_FILE)
    return frame


# ---------------------------------------------------------------------------
if __name__ == "__main__":
    run()
