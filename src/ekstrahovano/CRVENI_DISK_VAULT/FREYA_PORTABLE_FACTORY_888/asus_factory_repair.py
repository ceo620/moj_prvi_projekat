#!/usr/bin/env python3
import os, sys, re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor

BASELINE_2026_SSOT = {
    "FINANCIAL": {
        "CAPEX_TOTAL": "22.500.000 EUR",
        "CAPEX_UNI_MAK": "19.800.000 EUR",
        "NWC_RATE": "8.5%",
        "IDC_VAL": "1.250.000 EUR",
        "WORKING_CAPITAL": "1.450.000 EUR"
    },
    "TECHNICAL": {
        "POWER_MW": "45 MW",
        "GRID_CONNECTION": "CEDIS / CGES 110kV Interkonekcija",
        "LOCATION_MUNICIPALITY": "Nikšić",
        "CADASTRAL_PARCEL": "KO Kličevo, kp 1102/1"
    },
    "LEGAL": {
        "HUMAN_GATE_POLICY": "AUDIT MANDATORY - Potreban potpis ovlašćenog inženjera i revizora."
    }
}

PLACEHOLDER_MAP = {
    "[INSERT_CAPEX]": BASELINE_2026_SSOT["FINANCIAL"]["CAPEX_TOTAL"],
    "[INSERT_UNI_MAK]": BASELINE_2026_SSOT["FINANCIAL"]["CAPEX_UNI_MAK"],
    "[INSERT_NWC]": BASELINE_2026_SSOT["FINANCIAL"]["NWC_RATE"],
    "[INSERT_IDC]": BASELINE_2026_SSOT["FINANCIAL"]["IDC_VAL"],
    "[INSERT_MW]": BASELINE_2026_SSOT["TECHNICAL"]["POWER_MW"],
    "[INSERT_GRID]": BASELINE_2026_SSOT["TECHNICAL"]["GRID_CONNECTION"],
    "[INSERT_LOCATION]": f"{BASELINE_2026_SSOT['TECHNICAL']['LOCATION_MUNICIPALITY']}, {BASELINE_2026_SSOT['TECHNICAL']['CADASTRAL_PARCEL']}",
    "18.2 miliona": BASELINE_2026_SSOT["FINANCIAL"]["CAPEX_TOTAL"],
    "[INSERT_HUMAN_GATE]": BASELINE_2026_SSOT["LEGAL"]["HUMAN_GATE_POLICY"]
}

SECTORS = ["01_VLADA_CG", "02_MINISTARSTVO_FINANSIJA", "03_MINISTARSTVO_EKOLOGIJE", "04_CEDIS", "05_CGES", "06_AGENCIJA_ZA_ZASTITU_ZIVOTNE_SREDINE", "07_OPSTINA_NIKSIC", "08_BANKARSKI_SEKTOR", "09_MEGATREND_INVESTITORI", "10_EKSTERNI_REVIZORI", "11_ARHITEKTONSKI_KONSTRUKTORI", "12_TEHNICKI_PREGLED", "13_VODOVOD_I_KANALIZACIJA", "14_UPRAVA_ZA_KATASTAR", "15_INSPEKCIJSKI_ORGAN", "16_OSIGURAVAJUCA_DRUSTVA", "17_CARINSKA_UPRAVA", "18_PRAVNI_SAVJETNICI", "19_LOKALNA_ZAJEDNICA", "20_UPRAVNI_ODBOR_ASUS"]
LANGUAGES = ["SRB_CG", "ENG", "TUR"]

def run_factory():
    base_dir = Path(__file__).resolve().parent
    in_dir = base_dir / "input_templates"
    out_dir = base_dir / "output_harmonized_factory"
    in_dir.mkdir(parents=True, exist_ok=True)
    
    if not list(in_dir.glob("*.docx")):
        doc = Document()
        doc.add_heading("ASUS ELABORAT - GENERIČKI IZLAZ", level=0)
        doc.add_paragraph("Projekat zahtijeva investiciju od [INSERT_CAPEX] uz [INSERT_GRID].")
        doc.save(in_dir / "ARS-AUTO-999-ELABORAT_TEST.pdf.docx")

    total = 0
    for lang in LANGUAGES:
        for sec in SECTORS:
            t_path = out_dir / lang / sec
            t_path.mkdir(parents=True, exist_ok=True)
            for f in in_dir.glob("*.docx"):
                d = Document(f)
                for p in d.paragraphs:
                    for k, v in PLACEHOLDER_MAP.items():
                        if k in p.text:
                            p.text = p.text.replace(k, v)
                clean_name = re.sub(r"^ARS-AUTO-[A-Z0-9]+-", "", f.name).replace(".pdf.docx", ".docx")
                d.save(t_path / f"{sec}_{lang}_{clean_name}")
                total += 1
    print(f"USPJEŠNO GENERISANO {total} DOKUMENATA KROZ 20 SEKTORA I 3 JEZIKA.")

if __name__ == "__main__":
    run_factory()
