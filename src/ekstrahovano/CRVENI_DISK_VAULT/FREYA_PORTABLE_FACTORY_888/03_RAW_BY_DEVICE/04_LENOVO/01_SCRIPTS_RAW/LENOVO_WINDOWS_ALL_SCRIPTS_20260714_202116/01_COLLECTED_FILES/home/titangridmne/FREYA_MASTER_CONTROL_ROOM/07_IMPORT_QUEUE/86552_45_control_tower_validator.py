
import argparse
import json
import os
import sys
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj sa: pip install openpyxl")
    sys.exit(5)

VERSION = "2.0"

EXIT_PASS = 0
EXIT_BLOCK = 1
EXIT_FILE_ERROR = 2
EXIT_SCHEMA_ERROR = 3
EXIT_RUNTIME_ERROR = 4
EXIT_DEPENDENCY_ERROR = 5

REQUIRED_SHEETS = [
    "FileFinding Register",
    "Evidence Gap Register",
    "Control Tower Rules"
]

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102_STATUS": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO"
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def normalize(value):
    if value is None:
        return ""
    return str(value).strip()

def emit_jsonl(report_path, event):
    os.makedirs(os.path.dirname(report_path), exist_ok=True)
    with open(report_path, "a", encoding="utf-8") as f:
        f.write(json.dumps(event, ensure_ascii=False) + "\n")

def make_event(status, rule_id, rule_name, message, severity="INFO", extra=None):
    event = {
        "timestamp": now_iso(),
        "tool": "45_control_tower_validator",
        "version": VERSION,
        "status": status,
        "rule_id": rule_id,
        "rule_name": rule_name,
        "severity": severity,
        "message": message
    }
    if extra:
        event.update(extra)
    return event

def get_headers(ws):
    headers = {}
    for idx, cell in enumerate(ws[1], start=1):
        name = normalize(cell.value)
        if name:
            headers[name] = idx
    return headers

def require_columns(ws, required_columns):
    headers = get_headers(ws)
    missing = [col for col in required_columns if col not in headers]
    return headers, missing

def count_active_evidence_gaps(ws):
    headers = get_headers(ws)
    status_col = headers.get("Status")
    gap_id_col = headers.get("Gap_ID")

    if not status_col:
        return None, "Missing Status column"

    open_count = 0
    total_count = 0

    closed_statuses = {"CLOSED", "APPROVED", "RESOLVED", "VALIDATED"}

    for row_idx in range(2, ws.max_row + 1):
        gap_id = normalize(ws.cell(row=row_idx, column=gap_id_col).value) if gap_id_col else ""
        status = normalize(ws.cell(row=row_idx, column=status_col).value).upper()

        if not gap_id and not status:
            continue

        total_count += 1

        if status not in closed_statuses:
            open_count += 1

    return {"total_gaps": total_count, "open_gaps": open_count}, None

def validate_ai_signal_mapping(ws):
    headers = get_headers(ws)

    filefinding_col = headers.get("FileFinding_ID")
    evidence_col = headers.get("Evidence_ID")
    signal_col = headers.get("Signal_ID")
    weight_col = headers.get("Weight")

    if not filefinding_col:
        return None, "Missing FileFinding_ID column"

    violations = []
    checked = 0

    for row_idx in range(2, ws.max_row + 1):
        signal_id = normalize(ws.cell(row=row_idx, column=signal_col).value) if signal_col else ""
        filefinding_id = normalize(ws.cell(row=row_idx, column=filefinding_col).value)
        evidence_id = normalize(ws.cell(row=row_idx, column=evidence_col).value) if evidence_col else ""
        weight = normalize(ws.cell(row=row_idx, column=weight_col).value) if weight_col else ""

        if not signal_id and not filefinding_id and not evidence_id and not weight:
            continue

        checked += 1
        missing_mapping = not filefinding_id or not evidence_id

        if missing_mapping:
            try:
                numeric_weight = float(weight) if weight else 0
            except ValueError:
                numeric_weight = -1

            if numeric_weight != 0:
                violations.append({
                    "row": row_idx,
                    "Signal_ID": signal_id,
                    "FileFinding_ID": filefinding_id,
                    "Evidence_ID": evidence_id,
                    "Weight": weight,
                    "violation": "AI signal without FileFinding_ID/Evidence_ID must have weight 0"
                })

    return {"checked_rows": checked, "violations": violations}, None

def validate_control_rules(ws):
    headers = get_headers(ws)

    rule_col = headers.get("Rule_ID") or headers.get("Control_ID")
    name_col = headers.get("Rule_Name") or headers.get("Control_Name")
    value_col = headers.get("Required_Value") or headers.get("Value")
    status_col = headers.get("Status")

    rows = []

    for row_idx in range(2, ws.max_row + 1):
        row = {}
        if rule_col:
            row["rule_id"] = normalize(ws.cell(row=row_idx, column=rule_col).value)
        if name_col:
            row["rule_name"] = normalize(ws.cell(row=row_idx, column=name_col).value)
        if value_col:
            row["required_value"] = normalize(ws.cell(row=row_idx, column=value_col).value)
        if status_col:
            row["status"] = normalize(ws.cell(row=row_idx, column=status_col).value)
        if any(row.values()):
            row["row"] = row_idx
            rows.append(row)

    return rows

def validate_workbook(excel_path, report_path):
    events = []
    blocked = False

    if not os.path.exists(excel_path):
        event = make_event(
            "BLOCK",
            "FILE-001",
            "Excel File Exists",
            f"Excel fajl ne postoji: {excel_path}",
            "CRITICAL"
        )
        emit_jsonl(report_path, event)
        return EXIT_FILE_ERROR

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
    except Exception as e:
        event = make_event(
            "BLOCK",
            "FILE-002",
            "Excel File Readable",
            f"Excel fajl nije čitljiv: {e}",
            "CRITICAL"
        )
        emit_jsonl(report_path, event)
        return EXIT_FILE_ERROR

    for sheet_name in REQUIRED_SHEETS:
        if sheet_name not in wb.sheetnames:
            blocked = True
            events.append(make_event(
                "BLOCK",
                "SCHEMA-001",
                "Required Sheet Exists",
                f"Nedostaje sheet: {sheet_name}",
                "CRITICAL"
            ))

    if blocked:
        for event in events:
            emit_jsonl(report_path, event)
        return EXIT_SCHEMA_ERROR

    filefinding_ws = wb["FileFinding Register"]
    gap_ws = wb["Evidence Gap Register"]
    rules_ws = wb["Control Tower Rules"]

    filefinding_required = ["FileFinding_ID", "Evidence_ID", "Signal_ID", "Weight", "Status"]
    gap_required = ["Gap_ID", "Description", "Required_Evidence", "Linked_FileFinding_ID", "Status", "Owner", "Blocking_STEP"]

    _, missing_file_cols = require_columns(filefinding_ws, filefinding_required)
    _, missing_gap_cols = require_columns(gap_ws, gap_required)

    if missing_file_cols:
        blocked = True
        events.append(make_event(
            "BLOCK",
            "SCHEMA-002",
            "FileFinding Required Columns",
            "FileFinding Register nema potrebne kolone.",
            "CRITICAL",
            {"missing_columns": missing_file_cols}
        ))

    if missing_gap_cols:
        blocked = True
        events.append(make_event(
            "BLOCK",
            "SCHEMA-003",
            "Evidence Gap Required Columns",
            "Evidence Gap Register nema potrebne kolone.",
            "CRITICAL",
            {"missing_columns": missing_gap_cols}
        ))

    if blocked:
        for event in events:
            emit_jsonl(report_path, event)
        return EXIT_SCHEMA_ERROR

    events.append(make_event(
        "PASS",
        "RULE-01",
        "SYSTEM RED Hard Lock",
        "SYSTEM RED ostaje hardcoded baseline. AI interpretacija ne može promijeniti status bez manual override ključa.",
        "CRITICAL",
        {"SYSTEM_STATUS": CANON["SYSTEM_STATUS"], "FINAL_USE_ALLOWED": CANON["FINAL_USE_ALLOWED"]}
    ))

    ai_result, ai_error = validate_ai_signal_mapping(filefinding_ws)

    if ai_error:
        blocked = True
        events.append(make_event("BLOCK", "RULE-02", "AI Signal Must Map To FileFinding ID", ai_error, "CRITICAL"))
    else:
        violations = ai_result["violations"]
        if violations:
            blocked = True
            events.append(make_event(
                "BLOCK",
                "RULE-02",
                "AI Signal Must Map To FileFinding ID",
                "Postoje AI/signali bez FileFinding_ID ili Evidence_ID sa težinom različitom od 0.",
                "CRITICAL",
                {"checked_rows": ai_result["checked_rows"], "violations": violations}
            ))
        else:
            events.append(make_event(
                "PASS",
                "RULE-02",
                "AI Signal Must Map To FileFinding ID",
                "Svi signali bez potpunog dokaza imaju težinu 0 ili nema kršenja.",
                "HIGH",
                {"checked_rows": ai_result["checked_rows"]}
            ))

    gap_result, gap_error = count_active_evidence_gaps(gap_ws)

    if gap_error:
        blocked = True
        events.append(make_event("BLOCK", "RULE-03", "STEP102 Evidence Gap Lock", gap_error, "CRITICAL"))
    else:
        open_gaps = gap_result["open_gaps"]
        if open_gaps > 0:
            blocked = True
            events.append(make_event(
                "BLOCK",
                "RULE-03",
                "STEP102 Evidence Gap Lock",
                "Evidence gaps > 0, zato STEP102 ostaje LOCKED / NOT ACCEPTED.",
                "CRITICAL",
                {
                    "total_gaps": gap_result["total_gaps"],
                    "open_gaps": open_gaps,
                    "STEP102": CANON["STEP102_STATUS"],
                    "FINAL_USE_ALLOWED": CANON["FINAL_USE_ALLOWED"]
                }
            ))
        else:
            events.append(make_event(
                "REVIEW_REQUIRED",
                "RULE-03",
                "STEP102 Evidence Gap Lock",
                "Nema otvorenih evidence gaps, ali ovo nije automatsko odobrenje STEP102.",
                "CRITICAL",
                {
                    "total_gaps": gap_result["total_gaps"],
                    "open_gaps": open_gaps,
                    "STEP102": CANON["STEP102_STATUS"],
                    "FINAL_USE_ALLOWED": CANON["FINAL_USE_ALLOWED"]
                }
            ))

    canon_checks = [
        ("CANON-001", "FINAL_USE_ALLOWED", CANON["FINAL_USE_ALLOWED"]),
        ("CANON-002", "SSOT_WRITE_ALLOWED", CANON["SSOT_WRITE_ALLOWED"]),
        ("CANON-003", "EVIDENCE_APPROVAL_ALLOWED", CANON["EVIDENCE_APPROVAL_ALLOWED"]),
        ("CANON-004", "GATE_CLOSURE_ALLOWED", CANON["GATE_CLOSURE_ALLOWED"])
    ]

    for rule_id, key, value in canon_checks:
        events.append(make_event("PASS", rule_id, key, f"{key} ostaje {value}.", "CRITICAL", {key: value}))

    rule_rows = validate_control_rules(rules_ws)
    events.append(make_event(
        "INFO",
        "RULES-TRACE",
        "Control Tower Rules Sheet Read",
        "Control Tower Rules sheet je pročitan za audit trag.",
        "INFO",
        {"rules_rows_found": len(rule_rows)}
    ))

    final_status = "BLOCK" if blocked else "REVIEW_REQUIRED"
    summary = make_event(
        final_status,
        "SUMMARY",
        "Control Tower Validation Summary",
        "SYSTEM RED — STEP102 LOCKED — NO FINAL USE",
        "CRITICAL",
        {
            "SYSTEM_STATUS": CANON["SYSTEM_STATUS"],
            "STEP102": CANON["STEP102_STATUS"],
            "FINAL_USE_ALLOWED": CANON["FINAL_USE_ALLOWED"],
            "blocked": blocked
        }
    )

    for event in events:
        emit_jsonl(report_path, event)

    emit_jsonl(report_path, summary)
    return EXIT_BLOCK if blocked else EXIT_PASS

def main():
    parser = argparse.ArgumentParser(description="TITAN Control Tower Validator v2.0")
    parser.add_argument("--excel", required=True, help="Path to TITAN FileFinding/EvidenceGap Excel file.")
    parser.add_argument(
        "--report",
        default=r"C:\Users\Korisnik\Desktop\GEMINI SKRIPTE\TITAN_KERNEL\REPORTS\control_tower_validation_report.jsonl",
        help="Output JSONL report path."
    )
    parser.add_argument("--clear-report", action="store_true", help="Delete previous report before writing new validation events.")
    args = parser.parse_args()

    print("=====================================================")
    print("🗼 TITAN CONTROL TOWER VALIDATOR v2.0")
    print("=====================================================")

    if args.clear_report and os.path.exists(args.report):
        os.remove(args.report)

    exit_code = validate_workbook(args.excel, args.report)

    if exit_code == EXIT_PASS:
        print("✅ VALIDATION COMPLETE: REVIEW REQUIRED")
    elif exit_code == EXIT_BLOCK:
        print("⛔ VALIDATION COMPLETE: BLOCK")
    else:
        print(f"❌ VALIDATION ERROR: EXIT {exit_code}")

    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    print(f"📄 Report: {args.report}")
    print("=====================================================")
    return exit_code

if __name__ == "__main__":
    sys.exit(main())
