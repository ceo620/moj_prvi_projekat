
#!/usr/bin/env python3
"""
TITAN GRID VDR COMMITTEE PACK GENERATOR
Enterprise Layer v3.4

Enterprise additions over v3.2.3:
- Preflight lender readiness gate
- Role-based delivery groups
- Delivery manifest export (CSV + JSON)
- Audit log export (CSV + JSON)
- Provider abstraction for GDrive + SharePoint/OneDrive placeholder
- Stronger operational controls and reporting
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import os
import re
import smtplib
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

try:
    import yaml
    YAML_AVAILABLE = True
except ImportError:
    YAML_AVAILABLE = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import matplotlib.pyplot as plt
    MPL_AVAILABLE = True
except ImportError:
    MPL_AVAILABLE = False

try:
    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload
    GDRIVE_AVAILABLE = True
except ImportError:
    GDRIVE_AVAILABLE = False


logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)-8s | %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger("TITAN_GRID_VDR_V34")


DEFAULT_CONFIG: Dict[str, Any] = {
    "general": {
        "project_name": "TITAN GRID 1.0",
        "executing_entity": "ARS METAL INDUSTRIES DOO",
        "financing_route": "Corporate financing",
        "classification": "Confidential",
    },
    "colors": {
        "header": "1F4E79",
        "green": "E2F0D9",
        "yellow": "FFF2CC",
        "red": "FCE4D6",
        "blue": "D9EAF7",
        "white_font": "FFFFFF",
    },
    "priority_rules": {
        "critical_ids": ["RPT-002", "SPN-002", "ESG-002", "MKT-006", "PPP-001", "RSK-001"],
        "critical_regex": r"^(FIN|PPP|RSK|LND|RPT|SPN|TEC|ESG|MKT)-001$",
        "high_regex": r"^(FIN|PPP|RSK|LND|TEC|ESG|MKT|SPN)-",
    },
    "round1_rules": {
        "ids": ["RPT-002", "SPN-002", "ESG-002", "MKT-006", "PPP-001", "RSK-001"],
        "regex": r".*-001$",
    },
    "escalation_days": {
        "Level 3 - Immediate Sponsor / IC escalation": 1,
        "Level 2 - Owner close-out within deadline": 3,
        "Level 2 - PMO escalation": 5,
        "Level 1 - Owner follow-up": 7,
        "None": 10,
    },
    "gates": {
        "minimum_overall_readiness_pct": 70.0,
        "minimum_round1_readiness_pct": 80.0,
        "max_round1_red_items": 2,
        "block_delivery_if_gate_fails": True,
    },
    "output": {
        "folder_name": "99_Index_and_Control",
        "excel_filename": "99.86_VDR_Committee_Pack_A4.xlsx",
        "pdf_filename": "99.86_VDR_Committee_Pack_Summary.pdf",
        "docx_filename": "99.86_VDR_Committee_Pack.docx",
        "zip_filename": "99.86_VDR_Committee_Pack_Delivery.zip",
        "manifest_json": "99.87_Delivery_Manifest.json",
        "manifest_csv": "99.87_Delivery_Manifest.csv",
        "audit_json": "99.88_Audit_Log.json",
        "audit_csv": "99.88_Audit_Log.csv",
    },
    "excel": {
        "cover_sheet": "Cover_Page",
        "dashboard_sheet": "Dashboard",
        "submission_sheet": "Lender_Submission_Status",
        "round1_sheet": "Round1_Lender_Checklist",
        "traffic_sheet": "Traffic_Light_Report",
        "escalation_sheet": "Owner_Escalation_Matrix",
        "section_sheet": "Section_Summary",
        "appendix_sheet": "Appendix_Export_Log",
    },
    "pdf": {
        "title": "TITAN GRID VDR Committee Pack - Summary",
        "top_n_critical": 10,
    },
    "email": {
        "enabled": False,
        "smtp_server": "smtp.gmail.com",
        "smtp_port": 587,
        "use_tls": True,
        "use_ssl": False,
        "username": "",
        "password": "",
        "from": "vdr-bot@titan-grid.com",
        "to": ["pmo@titan-grid.com"],
        "subject": "TITAN GRID - VDR Committee Pack Ready",
        "body": (
            "VDR Committee Pack generated.\n\n"
            "Overall Readiness: {OverallReadiness}%\n"
            "Round 1 Readiness: {Round1Readiness}%\n"
            "Assessment: {Assessment}\n"
            "Gate Status: {GateStatus}\n\n"
            "Artifacts:\n{ArtifactLinks}"
        ),
    },
    "delivery": {
        "enabled": False,
        "provider": "gdrive",
        "credentials_json": "",
        "google_drive_folder_id": "",
        "share_role": "reader",
        "notify_users": False,
        "share_emails": [],
        "role_groups": {
            "internal": ["pmo@titan-grid.com"],
            "lenders": [],
            "investors": [],
            "advisors": [],
        },
        "sharepoint": {
            "tenant_id": "",
            "client_id": "",
            "client_secret": "",
            "site_url": "",
            "drive_name": "",
            "target_folder": "",
        },
    },
}


@dataclass
class VDRConfig:
    general: Dict[str, Any] = field(default_factory=dict)
    colors: Dict[str, Any] = field(default_factory=dict)
    priority_rules: Dict[str, Any] = field(default_factory=dict)
    round1_rules: Dict[str, Any] = field(default_factory=dict)
    escalation_days: Dict[str, Any] = field(default_factory=dict)
    gates: Dict[str, Any] = field(default_factory=dict)
    output: Dict[str, Any] = field(default_factory=dict)
    excel: Dict[str, Any] = field(default_factory=dict)
    pdf: Dict[str, Any] = field(default_factory=dict)
    email: Dict[str, Any] = field(default_factory=dict)
    delivery: Dict[str, Any] = field(default_factory=dict)

    def __post_init__(self) -> None:
        for key, value in DEFAULT_CONFIG.items():
            current = getattr(self, key)
            if not current:
                setattr(self, key, value.copy() if isinstance(value, dict) else value)


class TitanGridVDREnterprise:
    REQUIRED_BASE_COLUMNS = [
        "Document ID", "Document Title", "Pillar", "Section", "File Name", "File Path",
        "Status", "Confidentiality", "Owner", "Version", "Last Updated", "Notes",
    ]

    COLUMN_RENAME_MAP = {
        "Document_Name": "Document Title",
        "DocumentName": "Document Title",
        "DocumentID": "Document ID",
        "Document_Title": "Document Title",
        "FilePath": "File Path",
        "File_Path": "File Path",
        "FileName": "File Name",
        "File_Name": "File Name",
        "LastUpdated": "Last Updated",
        "DocumentTitle": "Document Title",
    }

    def __init__(self, args: argparse.Namespace):
        self.args = args
        self.config = self._load_config(args.config)
        self.vdr_root = Path(args.vdr_root).resolve()
        self.index_path = Path(args.index).resolve()
        self.output_folder = Path(args.output_dir).resolve() if args.output_dir else self.vdr_root / self.config.output["folder_name"]
        self.output_folder.mkdir(parents=True, exist_ok=True)

        self.excel_path = self.output_folder / self.config.output["excel_filename"]
        self.pdf_path = self.output_folder / self.config.output["pdf_filename"]
        self.docx_path = self.output_folder / self.config.output["docx_filename"]
        self.zip_path = self.output_folder / self.config.output["zip_filename"]
        self.manifest_json_path = self.output_folder / self.config.output["manifest_json"]
        self.manifest_csv_path = self.output_folder / self.config.output["manifest_csv"]
        self.audit_json_path = self.output_folder / self.config.output["audit_json"]
        self.audit_csv_path = self.output_folder / self.config.output["audit_csv"]

        self.chart_dir = self.output_folder / "_word_charts"
        self.chart_dir.mkdir(parents=True, exist_ok=True)

        self.audit_log: List[Dict[str, Any]] = []
        self.delivery_results: List[Dict[str, Any]] = []

    def _log(self, msg: str, level: str = "info") -> None:
        if self.args.verbose or level.lower() in {"warning", "error"}:
            getattr(logger, level.lower())(msg)

    def _audit(self, event_type: str, status: str, notes: str, path: str = "", meta: Optional[Dict[str, Any]] = None) -> None:
        self.audit_log.append({
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "event_type": event_type,
            "status": status,
            "path": path,
            "notes": notes,
            "meta": meta or {},
        })

    def _load_config(self, path: str) -> VDRConfig:
        cfg = VDRConfig()
        if not path:
            return cfg
        config_path = Path(path)
        if not config_path.exists():
            logger.warning("Config file not found, using defaults: %s", config_path)
            return cfg
        try:
            with open(config_path, "r", encoding="utf-8") as f:
                if config_path.suffix.lower() in {".yaml", ".yml"}:
                    if not YAML_AVAILABLE:
                        raise RuntimeError("PyYAML is not installed.")
                    user = yaml.safe_load(f) or {}
                else:
                    user = json.load(f) or {}
            for section, values in user.items():
                if hasattr(cfg, section) and isinstance(values, dict):
                    getattr(cfg, section).update(values)
            logger.info("Config loaded: %s", config_path)
        except Exception as exc:
            logger.warning("Config load failed (%s). Using defaults.", exc)
        return cfg

    def _validate_paths(self) -> None:
        if not self.vdr_root.exists():
            raise FileNotFoundError(f"VDR root not found: {self.vdr_root}")
        if not self.index_path.exists():
            raise FileNotFoundError(f"Index file not found: {self.index_path}")

    @staticmethod
    def _safe_text(value: Any) -> str:
        if pd.isna(value):
            return ""
        return str(value).strip()

    def _normalize_columns(self, df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        df.columns = [self._safe_text(c) for c in df.columns]
        df.rename(columns=self.COLUMN_RENAME_MAP, inplace=True)
        for col in self.REQUIRED_BASE_COLUMNS:
            if col not in df.columns:
                df[col] = ""
        return df

    def _sheet_is_compatible(self, df: pd.DataFrame) -> bool:
        return {"Document ID", "Document Title", "File Path"}.issubset(set(df.columns))

    def _compute_submission_status(self, row: pd.Series) -> str:
        if not row["Round1Lender"]:
            return "Supporting / internal"
        if row["TrafficLight"] == "GREEN":
            return "Ready for submission"
        if row["TrafficLight"] == "YELLOW":
            return "Conditionally ready - close comments"
        if row["TrafficLight"] == "RED":
            return "Not submittable"
        return "Check"

    def _compute_escalation(self, row: pd.Series) -> str:
        if row["Priority"] == "Critical" and row["TrafficLight"] == "RED":
            return "Level 3 - Immediate Sponsor / IC escalation"
        if row["Priority"] == "Critical" and row["TrafficLight"] == "YELLOW":
            return "Level 2 - Owner close-out within deadline"
        if row["Priority"] == "High" and row["TrafficLight"] == "RED":
            return "Level 2 - PMO escalation"
        if row["Priority"] == "High" and row["TrafficLight"] == "YELLOW":
            return "Level 1 - Owner follow-up"
        return "None"

    def _compute_action(self, row: pd.Series) -> str:
        if not row["Exists"]:
            return "Upload / map file into VDR"
        if row["Status"] == "Draft":
            return "Finalize draft and replace current version"
        if row["Status"] == "Under Review":
            return "Close review comments and issue final"
        return "No immediate action"

    def _analyze(self, df: pd.DataFrame) -> pd.DataFrame:
        cfg = self.config
        df = self._normalize_columns(df)

        df["Exists"] = df["File Path"].apply(
            lambda p: (self.vdr_root / self._safe_text(p)).exists() if self._safe_text(p) else False
        )

        df["TrafficLight"] = "BLUE"
        df.loc[df["Exists"] & (df["Status"] == "Final"), "TrafficLight"] = "GREEN"
        df.loc[df["Exists"] & df["Status"].isin(["Draft", "Under Review"]), "TrafficLight"] = "YELLOW"
        df.loc[~df["Exists"], "TrafficLight"] = "RED"

        df["ReadinessClass"] = df["TrafficLight"].map({
            "GREEN": "READY",
            "YELLOW": "PARTIAL",
            "RED": "MISSING",
            "BLUE": "CHECK",
        }).fillna("CHECK")

        critical_ids = set(cfg.priority_rules["critical_ids"])
        crit_regex = re.compile(cfg.priority_rules["critical_regex"])
        high_regex = re.compile(cfg.priority_rules["high_regex"])
        round1_regex = re.compile(cfg.round1_rules["regex"])
        round1_ids = set(cfg.round1_rules["ids"])

        def priority_for(doc_id: Any) -> str:
            text = self._safe_text(doc_id)
            if crit_regex.match(text) or text in critical_ids:
                return "Critical"
            if high_regex.match(text):
                return "High"
            return "Medium"

        def round1_for(doc_id: Any) -> bool:
            text = self._safe_text(doc_id)
            return bool(round1_regex.match(text) or text in round1_ids)

        df["Priority"] = df["Document ID"].apply(priority_for)
        df["Round1Lender"] = df["Document ID"].apply(round1_for)
        df["SubmissionStatus"] = df.apply(self._compute_submission_status, axis=1)
        df["ActionRequired"] = df.apply(self._compute_action, axis=1)
        df["EscalationLevel"] = df.apply(self._compute_escalation, axis=1)
        df["TargetCloseDate"] = df["EscalationLevel"].apply(
            lambda x: (datetime.now() + timedelta(days=self.config.escalation_days.get(x, self.config.escalation_days["None"]))).strftime("%Y-%m-%d")
        )
        return df

    def process_data(self) -> pd.DataFrame:
        self._validate_paths()
        self._log(f"Loading index: {self.index_path}")
        excel = pd.ExcelFile(self.index_path)
        candidate_sheets = excel.sheet_names if self.args.sheets.lower() == "all" else [s.strip() for s in self.args.sheets.split(",") if s.strip()]
        frames: List[pd.DataFrame] = []
        for sheet in candidate_sheets:
            try:
                raw = pd.read_excel(excel, sheet_name=sheet)
                raw = self._normalize_columns(raw)
                if not self._sheet_is_compatible(raw):
                    msg = f"Skipping incompatible sheet: {sheet}"
                    if self.args.fail_on_missing_columns:
                        raise ValueError(msg)
                    self._log(msg, "warning")
                    self._audit("sheet_skip", "warning", msg)
                    continue
                analyzed = self._analyze(raw)
                analyzed["SourceSheet"] = sheet
                frames.append(analyzed)
                self._log(f"Loaded {sheet}: {len(analyzed)} rows")
            except Exception as exc:
                self._audit("sheet_error", "error", f"{sheet}: {exc}")
                if self.args.fail_on_missing_columns:
                    raise
                self._log(f"Failed sheet {sheet}: {exc}", "error")
        if not frames:
            return pd.DataFrame()
        return pd.concat(frames, ignore_index=True)

    def _build_summary(self, df: pd.DataFrame) -> Dict[str, Any]:
        total = len(df)
        green = int((df["TrafficLight"] == "GREEN").sum())
        yellow = int((df["TrafficLight"] == "YELLOW").sum())
        red = int((df["TrafficLight"] == "RED").sum())
        round1 = df[df["Round1Lender"]]
        round1_total = len(round1)
        round1_green = int((round1["TrafficLight"] == "GREEN").sum())
        round1_yellow = int((round1["TrafficLight"] == "YELLOW").sum())
        round1_red = int((round1["TrafficLight"] == "RED").sum())
        overall = round((green / total) * 100, 2) if total else 0.0
        round1_readiness = round((round1_green / round1_total) * 100, 2) if round1_total else 0.0
        if round1_readiness >= 90 and round1_red == 0:
            assessment = "ROUND 1 LENDER PACK READY"
        elif round1_readiness >= 70:
            assessment = "ROUND 1 PARTIALLY READY - CLOSE RED/YELLOW ITEMS"
        else:
            assessment = "ROUND 1 NOT READY - MATERIAL DEFICIENCIES REMAIN"
        return {
            "TotalDocuments": total,
            "GreenCount": green,
            "YellowCount": yellow,
            "RedCount": red,
            "OverallReadiness": overall,
            "Round1Total": round1_total,
            "Round1Green": round1_green,
            "Round1Yellow": round1_yellow,
            "Round1Red": round1_red,
            "Round1Readiness": round1_readiness,
            "Assessment": assessment,
        }

    def evaluate_gate(self, summary: Dict[str, Any]) -> Dict[str, Any]:
        gates = self.config.gates
        failures: List[str] = []
        if summary["OverallReadiness"] < float(gates["minimum_overall_readiness_pct"]):
            failures.append(f'Overall readiness below threshold ({summary["OverallReadiness"]} < {gates["minimum_overall_readiness_pct"]})')
        if summary["Round1Readiness"] < float(gates["minimum_round1_readiness_pct"]):
            failures.append(f'Round 1 readiness below threshold ({summary["Round1Readiness"]} < {gates["minimum_round1_readiness_pct"]})')
        if summary["Round1Red"] > int(gates["max_round1_red_items"]):
            failures.append(f'Round 1 RED items exceed threshold ({summary["Round1Red"]} > {gates["max_round1_red_items"]})')

        gate_pass = len(failures) == 0
        gate = {
            "gate_pass": gate_pass,
            "gate_status": "PASS" if gate_pass else "FAIL",
            "failures": failures,
            "block_delivery": bool(gates["block_delivery_if_gate_fails"]) and not gate_pass,
        }
        self._audit("preflight_gate", "ok" if gate_pass else "warning", "; ".join(failures) if failures else "Gate passed", meta=gate)
        return gate

    def preflight_check(self, df: pd.DataFrame, summary: Dict[str, Any]) -> Dict[str, Any]:
        checks = {
            "vdr_root_exists": self.vdr_root.exists(),
            "index_exists": self.index_path.exists(),
            "data_not_empty": not df.empty,
            "required_columns_present": all(col in df.columns for col in ["Document ID", "Document Title", "File Path", "TrafficLight"]),
            "artifacts_output_dir_exists": self.output_folder.exists(),
        }
        checks.update(self.evaluate_gate(summary))
        return checks

    def _section_summary_df(self, df: pd.DataFrame) -> pd.DataFrame:
        grouped = (
            df.groupby("Section", dropna=False)
              .agg(
                  TotalDocuments=("Document ID", "count"),
                  Green=("TrafficLight", lambda s: int((s == "GREEN").sum())),
                  Yellow=("TrafficLight", lambda s: int((s == "YELLOW").sum())),
                  Red=("TrafficLight", lambda s: int((s == "RED").sum())),
              ).reset_index()
        )
        grouped["ReadyPct"] = grouped.apply(lambda r: round((r["Green"] / r["TotalDocuments"]) * 100, 2) if r["TotalDocuments"] else 0.0, axis=1)
        grouped.sort_values(by=["Red", "Yellow", "Section"], ascending=[False, False, True], inplace=True)
        return grouped

    def _top_critical_open_df(self, df: pd.DataFrame) -> pd.DataFrame:
        return (
            df[(df["Priority"] == "Critical") & (df["TrafficLight"].isin(["RED", "YELLOW"]))]
            .sort_values(by=["TrafficLight", "Document ID"])
            [["Document ID", "Document Title", "Owner", "TrafficLight", "ActionRequired"]]
            .head(self.config.pdf["top_n_critical"]).copy()
        )

    def _traffic_df(self, df: pd.DataFrame) -> pd.DataFrame:
        cols = ["Document ID", "Document Title", "Pillar", "Section", "Priority", "TrafficLight", "Status", "Exists", "Owner", "ActionRequired", "File Path"]
        return df[cols].sort_values(by=["Priority", "TrafficLight", "Section", "Document ID"]).copy()

    def _round1_df(self, df: pd.DataFrame) -> pd.DataFrame:
        cols = ["Document ID", "Document Title", "Pillar", "Section", "Priority", "TrafficLight", "Status", "Exists", "Confidentiality", "Owner", "Version", "Last Updated", "SubmissionStatus", "ActionRequired", "File Path"]
        return df[df["Round1Lender"]][cols].sort_values(by=["TrafficLight", "Priority", "Document ID"]).copy()

    def _submission_df(self, df: pd.DataFrame) -> pd.DataFrame:
        cols = ["Document ID", "Document Title", "Priority", "TrafficLight", "SubmissionStatus", "Status", "Exists", "Owner", "Confidentiality", "Version", "Last Updated", "ActionRequired"]
        return df[df["Round1Lender"]][cols].sort_values(by=["TrafficLight", "Priority", "Document ID"]).copy()

    def _owner_escalation_df(self, df: pd.DataFrame) -> pd.DataFrame:
        cols = ["Owner", "Document ID", "Document Title", "Priority", "TrafficLight", "Status", "EscalationLevel", "TargetCloseDate", "ActionRequired", "Section", "File Path"]
        return df[df["TrafficLight"].isin(["RED", "YELLOW"])][cols].sort_values(by=["TargetCloseDate", "Owner", "Priority", "Document ID"]).copy()

    def _styles(self) -> Dict[str, Any]:
        colors_cfg = self.config.colors
        thin = Side(style="thin", color="BFBFBF")
        return {
            "header_fill": PatternFill(start_color=colors_cfg["header"], end_color=colors_cfg["header"], fill_type="solid"),
            "header_font": Font(bold=True, color=colors_cfg["white_font"], name="Arial", size=11),
            "title_font": Font(bold=True, size=16, color=colors_cfg["header"]),
            "normal_font": Font(name="Arial", size=10),
            "border": Border(left=thin, right=thin, top=thin, bottom=thin),
            "traffic_fill": {
                "GREEN": PatternFill(start_color=colors_cfg["green"], end_color=colors_cfg["green"], fill_type="solid"),
                "YELLOW": PatternFill(start_color=colors_cfg["yellow"], end_color=colors_cfg["yellow"], fill_type="solid"),
                "RED": PatternFill(start_color=colors_cfg["red"], end_color=colors_cfg["red"], fill_type="solid"),
                "BLUE": PatternFill(start_color=colors_cfg["blue"], end_color=colors_cfg["blue"], fill_type="solid"),
            },
        }

    def _apply_page_setup(self, ws, landscape: bool = True) -> None:
        ws.page_margins = PageMargins(left=0.35, right=0.35, top=0.45, bottom=0.45, header=0.2, footer=0.2)
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE if landscape else ws.ORIENTATION_PORTRAIT
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A3"

    def _auto_width(self, ws, max_width: int = 45) -> None:
        for col_idx, col_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in col_cells:
                max_len = max(max_len, len("" if cell.value is None else str(cell.value)))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), max_width)

    def _apply_conditional_formatting(self, ws, df: pd.DataFrame) -> None:
        start_row, end_row = 3, ws.max_row
        if end_row < start_row:
            return
        columns = list(df.columns)
        if "TrafficLight" in columns:
            traffic_col = get_column_letter(columns.index("TrafficLight") + 1)
            full_range = f"A{start_row}:{get_column_letter(ws.max_column)}{end_row}"
            for status, color_key in [("GREEN", "green"), ("YELLOW", "yellow"), ("RED", "red")]:
                ws.conditional_formatting.add(
                    full_range,
                    FormulaRule(formula=[f'${traffic_col}{start_row}="{status}"'], stopIfTrue=False,
                                fill=PatternFill(start_color=self.config.colors[color_key], end_color=self.config.colors[color_key], fill_type="solid"))
                )
        if "ReadyPct" in columns:
            pct_col = get_column_letter(columns.index("ReadyPct") + 1)
            pct_range = f"{pct_col}{start_row}:{pct_col}{end_row}"
            ws.conditional_formatting.add(pct_range, CellIsRule(operator='lessThan', formula=['70'], fill=PatternFill(start_color=self.config.colors["red"], end_color=self.config.colors["red"], fill_type="solid")))
            ws.conditional_formatting.add(pct_range, CellIsRule(operator='between', formula=['70', '89.99'], fill=PatternFill(start_color=self.config.colors["yellow"], end_color=self.config.colors["yellow"], fill_type="solid")))
            ws.conditional_formatting.add(pct_range, CellIsRule(operator='greaterThanOrEqual', formula=['90'], fill=PatternFill(start_color=self.config.colors["green"], end_color=self.config.colors["green"], fill_type="solid")))

    def _write_dataframe_sheet(self, ws, df: pd.DataFrame, title: str, styles: Dict[str, Any], traffic_col_name: Optional[str] = None, landscape: bool = True) -> None:
        ws["A1"] = title
        ws["A1"].font = styles["title_font"]
        self._apply_page_setup(ws, landscape)
        for idx, col in enumerate(df.columns, start=1):
            cell = ws.cell(row=2, column=idx, value=col)
            cell.fill = styles["header_fill"]
            cell.font = styles["header_font"]
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = styles["border"]
        for row_idx, row in enumerate(df.itertuples(index=False), start=3):
            traffic_value = row[df.columns.get_loc(traffic_col_name)] if traffic_col_name and traffic_col_name in df.columns else None
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = styles["normal_font"]
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = styles["border"]
                if traffic_value in styles["traffic_fill"]:
                    cell.fill = styles["traffic_fill"][traffic_value]
        ws.auto_filter.ref = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
        self._auto_width(ws)
        self._apply_conditional_formatting(ws, df)

    def _create_matplotlib_charts(self, section_summary: pd.DataFrame, summary: Dict[str, Any]) -> List[Path]:
        paths: List[Path] = []
        if not MPL_AVAILABLE:
            return paths

        bar_path = self.chart_dir / "section_summary_bar.png"
        plt.figure(figsize=(10, 4.5))
        x = range(len(section_summary))
        plt.bar(x, section_summary["Green"], label="Green")
        plt.bar(x, section_summary["Yellow"], bottom=section_summary["Green"], label="Yellow")
        plt.bar(x, section_summary["Red"], bottom=section_summary["Green"] + section_summary["Yellow"], label="Red")
        plt.xticks(list(x), section_summary["Section"], rotation=45, ha="right")
        plt.title("Documents by Section and Traffic Light")
        plt.tight_layout()
        plt.legend()
        plt.savefig(bar_path, dpi=180, bbox_inches="tight")
        plt.close()
        paths.append(bar_path)

        pie_path = self.chart_dir / "traffic_mix_pie.png"
        plt.figure(figsize=(5.5, 5.5))
        plt.pie([summary["GreenCount"], summary["YellowCount"], summary["RedCount"]], labels=["GREEN", "YELLOW", "RED"], autopct="%1.1f%%")
        plt.title("Traffic Light Mix")
        plt.tight_layout()
        plt.savefig(pie_path, dpi=180, bbox_inches="tight")
        plt.close()
        paths.append(pie_path)
        return paths

    def _add_dashboard_charts(self, ws, section_summary: pd.DataFrame, summary: Dict[str, Any]) -> None:
        start_row = 20
        for i, header in enumerate(["Section", "Green", "Yellow", "Red"], start=1):
            ws.cell(row=start_row, column=i, value=header)
        for i, row in enumerate(section_summary.itertuples(index=False), start=start_row + 1):
            ws.cell(row=i, column=1, value=row.Section)
            ws.cell(row=i, column=2, value=row.Green)
            ws.cell(row=i, column=3, value=row.Yellow)
            ws.cell(row=i, column=4, value=row.Red)

        bar = BarChart()
        bar.title = "Documents by Section and Traffic Light"
        data = Reference(ws, min_col=2, max_col=4, min_row=start_row, max_row=start_row + len(section_summary))
        cats = Reference(ws, min_col=1, min_row=start_row + 1, max_row=start_row + len(section_summary))
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(cats)
        bar.height = 8
        bar.width = 14
        ws.add_chart(bar, "J20")

        pie_start = 20
        ws.cell(row=pie_start, column=6, value="Traffic")
        ws.cell(row=pie_start, column=7, value="Count")
        for offset, (label, value) in enumerate([("GREEN", summary["GreenCount"]), ("YELLOW", summary["YellowCount"]), ("RED", summary["RedCount"])], start=1):
            ws.cell(row=pie_start + offset, column=6, value=label)
            ws.cell(row=pie_start + offset, column=7, value=value)
        pie = PieChart()
        pie.title = "Traffic Light Mix"
        pdata = Reference(ws, min_col=7, min_row=pie_start, max_row=pie_start + 3)
        labels = Reference(ws, min_col=6, min_row=pie_start + 1, max_row=pie_start + 3)
        pie.add_data(pdata, titles_from_data=True)
        pie.set_categories(labels)
        pie.height = 7
        pie.width = 10
        ws.add_chart(pie, "J3")

    def generate_excel(self, df: pd.DataFrame, summary: Dict[str, Any], gate: Dict[str, Any]) -> str:
        self._log("Creating enterprise Excel workbook...")
        section_summary = self._section_summary_df(df)
        traffic = self._traffic_df(df)
        round1 = self._round1_df(df)
        submission = self._submission_df(df)
        escalation = self._owner_escalation_df(df)
        critical_open = self._top_critical_open_df(df)
        wb = Workbook()
        styles = self._styles()

        ws = wb.active
        ws.title = self.config.excel["cover_sheet"]
        self._apply_page_setup(ws, landscape=False)
        ws.merge_cells("A1:H3")
        c = ws["A1"]
        c.value = f'{self.config.general["project_name"]}\nVDR Committee Pack'
        c.font = Font(bold=True, color=self.config.colors["white_font"], size=22, name="Arial")
        c.fill = styles["header_fill"]
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        cover_rows = [
            ("Date generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Project", self.config.general["project_name"]),
            ("Financing route", self.config.general["financing_route"]),
            ("Executing entity", self.config.general["executing_entity"]),
            ("Overall readiness (%)", summary["OverallReadiness"]),
            ("Round 1 readiness (%)", summary["Round1Readiness"]),
            ("Assessment", summary["Assessment"]),
            ("Gate status", gate["gate_status"]),
        ]
        for i, (label, value) in enumerate(cover_rows, start=8):
            ws.cell(row=i, column=1, value=label).font = Font(bold=True, size=11)
            ws.cell(row=i, column=2, value=value)
        if gate["failures"]:
            ws["A18"] = "Gate failures"
            ws["A18"].font = Font(bold=True)
            for idx, failure in enumerate(gate["failures"], start=19):
                ws.cell(row=idx, column=1, value="- " + failure)
        self._auto_width(ws)

        ws2 = wb.create_sheet(self.config.excel["dashboard_sheet"])
        self._apply_page_setup(ws2, landscape=True)
        ws2["A1"] = "Lender Dashboard"
        ws2["A1"].font = styles["title_font"]
        metric_rows = [
            ("Total indexed documents", summary["TotalDocuments"]),
            ("GREEN documents", summary["GreenCount"]),
            ("YELLOW documents", summary["YellowCount"]),
            ("RED documents", summary["RedCount"]),
            ("Overall readiness (%)", summary["OverallReadiness"]),
            ("Round 1 lender documents", summary["Round1Total"]),
            ("Round 1 readiness (%)", summary["Round1Readiness"]),
            ("Assessment", summary["Assessment"]),
            ("Gate status", gate["gate_status"]),
        ]
        for cell in ("A3", "B3"):
            ws2[cell].fill = styles["header_fill"]
            ws2[cell].font = styles["header_font"]
        ws2["A3"], ws2["B3"] = "Metric", "Value"
        for row_num, (label, value) in enumerate(metric_rows, start=4):
            ws2.cell(row=row_num, column=1, value=label)
            ws2.cell(row=row_num, column=2, value=value)

        for cell in ("D3", "E3", "F3", "G3", "H3"):
            ws2[cell].fill = styles["header_fill"]
            ws2[cell].font = styles["header_font"]
        for i, h in enumerate(["Section", "Total", "Green", "Yellow", "Red"], start=4):
            ws2.cell(row=3, column=i, value=h)
        for idx, row in enumerate(section_summary.itertuples(index=False), start=4):
            ws2.cell(row=idx, column=4, value=row.Section)
            ws2.cell(row=idx, column=5, value=row.TotalDocuments)
            ws2.cell(row=idx, column=6, value=row.Green)
            ws2.cell(row=idx, column=7, value=row.Yellow)
            ws2.cell(row=idx, column=8, value=row.Red)

        for cell in ("J3", "K3", "L3"):
            ws2[cell].fill = styles["header_fill"]
            ws2[cell].font = styles["header_font"]
        ws2["J3"], ws2["K3"], ws2["L3"] = "Top Critical Open Items", "Owner", "Traffic"
        for idx, row in enumerate(critical_open.itertuples(index=False), start=4):
            ws2.cell(row=idx, column=10, value=f'{row[0]} - {row[1]}')
            ws2.cell(row=idx, column=11, value=row[2])
            tcell = ws2.cell(row=idx, column=12, value=row[3])
            if row[3] in styles["traffic_fill"]:
                tcell.fill = styles["traffic_fill"][row[3]]

        self._add_dashboard_charts(ws2, section_summary, summary)
        self._auto_width(ws2)

        self._write_dataframe_sheet(wb.create_sheet(self.config.excel["submission_sheet"]), submission, "Lender Submission Status", styles, "TrafficLight", True)
        self._write_dataframe_sheet(wb.create_sheet(self.config.excel["round1_sheet"]), round1, "Round 1 Lender Checklist", styles, "TrafficLight", True)
        self._write_dataframe_sheet(wb.create_sheet(self.config.excel["traffic_sheet"]), traffic, "Traffic Light Report", styles, "TrafficLight", True)
        self._write_dataframe_sheet(wb.create_sheet(self.config.excel["escalation_sheet"]), escalation, "Owner Escalation Matrix", styles, "TrafficLight", True)
        self._write_dataframe_sheet(wb.create_sheet(self.config.excel["section_sheet"]), section_summary, "Section Summary", styles, None, True)

        appendix = wb.create_sheet(self.config.excel["appendix_sheet"])
        self._apply_page_setup(appendix, landscape=True)
        appendix["A1"] = "Appendix Export Log"
        appendix["A1"].font = styles["title_font"]
        for i, h in enumerate(["Timestamp", "Artifact Type", "Path", "Status", "Notes"], start=1):
            cell = appendix.cell(row=2, column=i, value=h)
            cell.fill = styles["header_fill"]
            cell.font = styles["header_font"]
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = styles["border"]
        self._auto_width(appendix)

        wb.save(self.excel_path)
        self._audit("excel_export", "created", "Enterprise Excel workbook created", str(self.excel_path))
        return str(self.excel_path)

    def _pdf_table(self, data: List[List[Any]], col_widths: Optional[List[int]] = None, traffic_col_index: Optional[int] = None) -> Table:
        tbl = Table(data, colWidths=col_widths)
        style_cmds = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + self.config.colors["header"])),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey]),
        ]
        if traffic_col_index is not None:
            for row_idx in range(1, len(data)):
                val = str(data[row_idx][traffic_col_index])
                if val == "GREEN":
                    style_cmds.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#" + self.config.colors["green"])))
                elif val == "YELLOW":
                    style_cmds.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#" + self.config.colors["yellow"])))
                elif val == "RED":
                    style_cmds.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#" + self.config.colors["red"])))
        tbl.setStyle(TableStyle(style_cmds))
        return tbl

    def _pdf_on_page(self, canv, doc) -> None:
        canv.saveState()
        canv.setFont("Helvetica-Bold", 9)
        canv.drawString(doc.leftMargin, A4[1] - 20, self.config.general["project_name"] + " - Committee Pack")
        canv.setFont("Helvetica", 8)
        canv.drawRightString(A4[0] - doc.rightMargin, A4[1] - 20, datetime.now().strftime("%Y-%m-%d"))
        canv.drawString(doc.leftMargin, 15, self.config.general["classification"])
        canv.drawRightString(A4[0] - doc.rightMargin, 15, f"Page {doc.page}")
        canv.restoreState()

    def generate_pdf(self, df: pd.DataFrame, summary: Dict[str, Any], gate: Dict[str, Any]) -> str:
        if not PDF_AVAILABLE:
            self._log("reportlab not installed - skipping PDF", "warning")
            return ""
        section_summary = self._section_summary_df(df)
        critical_open = self._top_critical_open_df(df)
        round1 = self._round1_df(df)
        doc = SimpleDocTemplate(str(self.pdf_path), pagesize=A4, topMargin=40, bottomMargin=30, leftMargin=30, rightMargin=30)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("TitleStyle", parent=styles["Title"], textColor=colors.HexColor("#" + self.config.colors["header"]), fontSize=18)
        heading = styles["Heading2"]
        normal = styles["BodyText"]

        story: List[Any] = []
        story.append(Paragraph(self.config.pdf["title"], title_style))
        story.append(Spacer(1, 12))
        story.append(Paragraph(f'Project: {self.config.general["project_name"]}', normal))
        story.append(Paragraph(f'Executing entity: {self.config.general["executing_entity"]}', normal))
        story.append(Paragraph(f'Generated on: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}', normal))
        story.append(Spacer(1, 12))
        story.append(Paragraph("Executive Summary", heading))
        story.append(Paragraph(
            f'Overall readiness is {summary["OverallReadiness"]}% and round 1 lender readiness is {summary["Round1Readiness"]}%. '
            f'Current assessment: {summary["Assessment"]}. Gate status: {gate["gate_status"]}.', normal))
        if gate["failures"]:
            for failure in gate["failures"]:
                story.append(Paragraph("- " + failure, normal))
        story.append(Spacer(1, 12))
        kpi_data = [
            ["Metric", "Value"],
            ["Total Documents", summary["TotalDocuments"]],
            ["GREEN", summary["GreenCount"]],
            ["YELLOW", summary["YellowCount"]],
            ["RED", summary["RedCount"]],
            ["Overall Readiness (%)", summary["OverallReadiness"]],
            ["Round 1 Readiness (%)", summary["Round1Readiness"]],
            ["Gate Status", gate["gate_status"]],
            ["Assessment", summary["Assessment"]],
        ]
        story.append(self._pdf_table(kpi_data, [220, 220]))
        story.append(Spacer(1, 16))
        story.append(Paragraph("Section Summary", heading))
        section_data = [["Section", "Total", "Green", "Yellow", "Red", "Ready %"]]
        for row in section_summary.itertuples(index=False):
            section_data.append([row.Section, row.TotalDocuments, row.Green, row.Yellow, row.Red, row.ReadyPct])
        story.append(self._pdf_table(section_data, [190, 50, 50, 50, 50, 60]))
        story.append(Spacer(1, 16))
        story.append(Paragraph("Top Critical Open Items", heading))
        if critical_open.empty:
            story.append(Paragraph("No critical open items.", normal))
        else:
            crit_data = [["Document ID", "Document Title", "Owner", "Traffic", "Action"]]
            for row in critical_open.itertuples(index=False):
                crit_data.append([row[0], row[1], row[2], row[3], row[4]])
            story.append(self._pdf_table(crit_data, [70, 170, 90, 50, 120], 3))
        story.append(Spacer(1, 16))
        story.append(Paragraph("Round 1 Lender Status", heading))
        round1_data = [["Document ID", "Document Title", "Traffic", "Submission Status", "Owner"]]
        for row in round1.head(15).itertuples(index=False):
            round1_data.append([row[0], row[1], row[5], row[12], row[9]])
        story.append(self._pdf_table(round1_data, [70, 170, 50, 130, 100], 2))
        doc.build(story, onFirstPage=self._pdf_on_page, onLaterPages=self._pdf_on_page)
        self._audit("pdf_export", "created", "Enterprise PDF summary created", str(self.pdf_path))
        return str(self.pdf_path)

    def generate_word(self, df: pd.DataFrame, summary: Dict[str, Any], gate: Dict[str, Any]) -> str:
        if not DOCX_AVAILABLE:
            self._log("python-docx not installed - skipping Word export", "warning")
            return ""
        section_summary = self._section_summary_df(df)
        critical_open = self._top_critical_open_df(df)
        submission = self._submission_df(df)
        round1 = self._round1_df(df)
        chart_paths = self._create_matplotlib_charts(section_summary, summary)

        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(self.config.general["project_name"])
        r.bold = True
        r.font.size = Pt(20)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("VDR Committee Pack")
        r.bold = True
        r.font.size = Pt(16)

        doc.add_paragraph(f'Generated on: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
        doc.add_paragraph(f'Executing entity: {self.config.general["executing_entity"]}')
        doc.add_paragraph(f'Financing route: {self.config.general["financing_route"]}')
        doc.add_paragraph(f'Overall readiness: {summary["OverallReadiness"]}%')
        doc.add_paragraph(f'Round 1 readiness: {summary["Round1Readiness"]}%')
        doc.add_paragraph(f'Assessment: {summary["Assessment"]}')
        doc.add_paragraph(f'Gate status: {gate["gate_status"]}')
        if gate["failures"]:
            for failure in gate["failures"]:
                doc.add_paragraph("- " + failure)

        doc.add_page_break()
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(
            f'This committee pack consolidates current VDR readiness for {self.config.general["project_name"]}. '
            f'Overall readiness stands at {summary["OverallReadiness"]}% and round 1 lender readiness at '
            f'{summary["Round1Readiness"]}%. Current assessment: {summary["Assessment"]}.'
        )

        doc.add_heading("Charts", level=1)
        if chart_paths:
            for path in chart_paths:
                if path.exists():
                    doc.add_paragraph(path.stem.replace("_", " ").title())
                    doc.add_picture(str(path), width=Inches(6.5))
        else:
            doc.add_paragraph("Chart image generation skipped because matplotlib is not installed.")

        doc.add_heading("Top Critical Open Items", level=1)
        if critical_open.empty:
            doc.add_paragraph("No critical open items.")
        else:
            crit_table = doc.add_table(rows=1, cols=5)
            crit_table.style = "Table Grid"
            headers = ["Document ID", "Document Title", "Owner", "Traffic", "Action"]
            for i, h in enumerate(headers):
                crit_table.rows[0].cells[i].text = h
            for row in critical_open.itertuples(index=False):
                cells = crit_table.add_row().cells
                cells[0].text, cells[1].text, cells[2].text, cells[3].text, cells[4].text = map(str, row)

        def add_df_section(title: str, frame: pd.DataFrame, max_rows: int = 20) -> None:
            doc.add_page_break()
            doc.add_heading(title, level=1)
            show = frame.head(max_rows)
            table = doc.add_table(rows=1, cols=len(show.columns))
            table.style = "Table Grid"
            for i, h in enumerate(show.columns):
                table.rows[0].cells[i].text = str(h)
            for row in show.itertuples(index=False):
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = "" if pd.isna(value) else str(value)

        add_df_section("Lender Submission Status", submission, 20)
        add_df_section("Round 1 Lender Checklist", round1, 20)

        doc.add_page_break()
        doc.add_heading("Appendix", level=1)
        doc.add_paragraph(f"Source workbook: {self.index_path}")
        doc.add_paragraph(f"Excel export: {self.excel_path}")
        if self.pdf_path.exists():
            doc.add_paragraph(f"PDF export: {self.pdf_path}")
        doc.add_paragraph(f"Chart directory: {self.chart_dir}")
        if self.audit_log:
            t = doc.add_table(rows=1, cols=4)
            t.style = "Table Grid"
            for i, h in enumerate(["Event Type", "Status", "Path", "Notes"]):
                t.rows[0].cells[i].text = h
            for entry in self.audit_log[-20:]:
                cells = t.add_row().cells
                cells[0].text = entry["event_type"]
                cells[1].text = entry["status"]
                cells[2].text = entry["path"]
                cells[3].text = entry["notes"]

        doc.save(self.docx_path)
        self._audit("docx_export", "created", "Enterprise Word committee pack created", str(self.docx_path))
        return str(self.docx_path)

    def update_appendix_sheet(self) -> None:
        if not self.excel_path.exists():
            return
        wb = load_workbook(self.excel_path)
        sheet_name = self.config.excel["appendix_sheet"]
        if sheet_name not in wb.sheetnames:
            wb.create_sheet(sheet_name)
        ws = wb[sheet_name]
        if ws.max_row < 2:
            for i, h in enumerate(["Timestamp", "Artifact Type", "Path", "Status", "Notes"], start=1):
                ws.cell(row=2, column=i, value=h)
        start_row = max(ws.max_row + 1, 3)
        relevant = []
        for entry in self.audit_log:
            if entry["event_type"] in {"excel_export", "pdf_export", "docx_export", "zip_export", "delivery_upload", "delivery_share"}:
                relevant.append(entry)
        for idx, entry in enumerate(relevant, start=start_row):
            ws.cell(row=idx, column=1, value=entry["timestamp"])
            ws.cell(row=idx, column=2, value=entry["event_type"])
            ws.cell(row=idx, column=3, value=entry["path"])
            ws.cell(row=idx, column=4, value=entry["status"])
            ws.cell(row=idx, column=5, value=entry["notes"])
        for col in range(1, 6):
            ws.column_dimensions[get_column_letter(col)].width = 28
        wb.save(self.excel_path)

    def create_zip_pack(self, include_excel: bool, include_pdf: bool, include_docx: bool) -> str:
        if self.zip_path.exists():
            self.zip_path.unlink()
        with zipfile.ZipFile(self.zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            if include_excel and self.excel_path.exists():
                zf.write(self.excel_path, arcname=self.excel_path.name)
            if include_pdf and self.pdf_path.exists():
                zf.write(self.pdf_path, arcname=self.pdf_path.name)
            if include_docx and self.docx_path.exists():
                zf.write(self.docx_path, arcname=self.docx_path.name)
            for chart_file in self.chart_dir.glob("*.png"):
                zf.write(chart_file, arcname=f"charts/{chart_file.name}")
        self._audit("zip_export", "created", "Delivery bundle created", str(self.zip_path))
        return str(self.zip_path)

    def export_audit_logs(self) -> None:
        with open(self.audit_json_path, "w", encoding="utf-8") as f:
            json.dump(self.audit_log, f, indent=2, ensure_ascii=False)
        with open(self.audit_csv_path, "w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=["timestamp", "event_type", "status", "path", "notes", "meta"])
            writer.writeheader()
            for row in self.audit_log:
                writer.writerow(row)
        self._log(f"Audit logs exported: {self.audit_json_path}, {self.audit_csv_path}")

    def export_delivery_manifest(self, gate: Dict[str, Any]) -> None:
        manifest = {
            "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "project_name": self.config.general["project_name"],
            "gate": gate,
            "artifacts": self.delivery_results,
            "output_folder": str(self.output_folder),
        }
        with open(self.manifest_json_path, "w", encoding="utf-8") as f:
            json.dump(manifest, f, indent=2, ensure_ascii=False)
        with open(self.manifest_csv_path, "w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=["artifact", "link", "provider", "shared_to"])
            writer.writeheader()
            for row in self.delivery_results:
                writer.writerow(row)
        self._log(f"Delivery manifest exported: {self.manifest_json_path}, {self.manifest_csv_path}")

    def _build_drive_service(self):
        if not GDRIVE_AVAILABLE:
            raise RuntimeError("Google Drive dependencies are not installed.")
        credentials_json = self.config.delivery.get("credentials_json", "") or self.args.credentials_json
        if not credentials_json:
            raise ValueError("Google Drive credentials_json is not configured.")
        creds_path = Path(credentials_json)
        if not creds_path.exists():
            raise FileNotFoundError(f"credentials.json not found: {creds_path}")
        creds = Credentials.from_service_account_file(
            str(creds_path),
            scopes=["https://www.googleapis.com/auth/drive"],
        )
        return build("drive", "v3", credentials=creds)

    def upload_to_gdrive(self, file_path: str) -> Dict[str, str]:
        service = self._build_drive_service()
        parent_id = self.config.delivery.get("google_drive_folder_id", "")
        if not parent_id:
            raise ValueError("google_drive_folder_id is not configured.")
        target = Path(file_path)
        metadata = {"name": target.name, "parents": [parent_id]}
        media = MediaFileUpload(str(target), resumable=True)
        created = service.files().create(body=metadata, media_body=media, fields="id, webViewLink, webContentLink").execute()
        return {
            "file_id": created.get("id", ""),
            "webViewLink": created.get("webViewLink", ""),
            "webContentLink": created.get("webContentLink", "") or created.get("webViewLink", ""),
        }

    def share_gdrive_file(self, file_id: str, recipients: List[str], role: str, notify_users: bool) -> None:
        if not recipients:
            return
        service = self._build_drive_service()
        for email in recipients:
            service.permissions().create(
                fileId=file_id,
                body={"type": "user", "role": role, "emailAddress": email},
                sendNotificationEmail=notify_users
            ).execute()
            self._audit("delivery_share", "shared", f"Shared file {file_id} with {email}", path=file_id, meta={"recipient": email, "role": role})

    def upload_to_sharepoint(self, file_path: str) -> Dict[str, str]:
        raise NotImplementedError(
            "SharePoint/OneDrive provider scaffold exists, but tenant-specific Graph API implementation "
            "still requires concrete tenant_id, client_id, client_secret, site_url, and target drive/folder logic."
        )

    def _resolve_role_based_recipients(self) -> List[str]:
        groups = self.config.delivery.get("role_groups", {})
        recipients: List[str] = []
        for _, emails in groups.items():
            if isinstance(emails, list):
                recipients.extend(emails)
        recipients.extend(self.config.delivery.get("share_emails", []))
        deduped = sorted(set([e for e in recipients if e]))
        return deduped

    def run_delivery(self, artifact_paths: List[str], gate: Dict[str, Any]) -> List[Dict[str, Any]]:
        if not self.args.run_delivery and not self.config.delivery.get("enabled", False):
            return []
        if gate["block_delivery"]:
            self._audit("delivery_block", "blocked", "Delivery blocked by readiness gate", meta=gate)
            raise RuntimeError("Delivery blocked because readiness gate failed.")

        provider = self.config.delivery.get("provider", "gdrive").lower()
        recipients = self._resolve_role_based_recipients()
        role = self.config.delivery.get("share_role", "reader")
        notify = bool(self.config.delivery.get("notify_users", False))
        results: List[Dict[str, Any]] = []

        for path in artifact_paths:
            if not path:
                continue
            if provider == "gdrive":
                upload_info = self.upload_to_gdrive(path)
                if upload_info.get("file_id"):
                    self._audit("delivery_upload", "uploaded", "Uploaded to Google Drive", path=path, meta=upload_info)
                    self.share_gdrive_file(upload_info["file_id"], recipients, role, notify)
                results.append({
                    "artifact": path,
                    "link": upload_info.get("webViewLink", ""),
                    "provider": "gdrive",
                    "shared_to": ", ".join(recipients),
                })
            elif provider in {"sharepoint", "onedrive"}:
                upload_info = self.upload_to_sharepoint(path)
                results.append({
                    "artifact": path,
                    "link": upload_info.get("webViewLink", ""),
                    "provider": provider,
                    "shared_to": ", ".join(recipients),
                })
            else:
                raise ValueError(f"Unsupported delivery provider: {provider}")

        self.delivery_results = results
        return results

    def send_email(self, attachment: Optional[str], summary: Dict[str, Any], gate: Dict[str, Any], delivery_results: Optional[List[Dict[str, Any]]] = None) -> None:
        cfg = self.config.email
        if not cfg.get("enabled", False):
            self._log("Email disabled in config")
            return

        recipients = cfg.get("to", [])
        if isinstance(recipients, str):
            recipients = [recipients]
        if not recipients:
            raise ValueError("Email recipient list is empty.")

        password = cfg.get("password") or os.getenv("TITAN_GRID_SMTP_PASSWORD", "")
        username = cfg.get("username", "")

        if delivery_results:
            links_block = "\n".join([f'- {Path(item["artifact"]).name}: {item["link"]}' for item in delivery_results])
        elif attachment:
            links_block = f"- Attached artifact: {Path(attachment).name}"
        else:
            links_block = "- No artifact"

        body = cfg["body"].format(
            OverallReadiness=summary["OverallReadiness"],
            Round1Readiness=summary["Round1Readiness"],
            Assessment=summary["Assessment"],
            GateStatus=gate["gate_status"],
            ArtifactLinks=links_block,
        )

        msg = MIMEMultipart()
        msg["From"] = cfg["from"]
        msg["To"] = ", ".join(recipients)
        msg["Subject"] = cfg["subject"]
        msg.attach(MIMEText(body, "plain", "utf-8"))

        if attachment and Path(attachment).exists() and not delivery_results:
            with open(attachment, "rb") as f:
                part = MIMEApplication(f.read(), Name=Path(attachment).name)
                part["Content-Disposition"] = f'attachment; filename="{Path(attachment).name}"'
                msg.attach(part)

        if cfg.get("use_ssl", False):
            with smtplib.SMTP_SSL(cfg["smtp_server"], int(cfg["smtp_port"]), timeout=30) as smtp:
                if username:
                    smtp.login(username, password)
                smtp.send_message(msg)
        else:
            with smtplib.SMTP(cfg["smtp_server"], int(cfg["smtp_port"]), timeout=30) as smtp:
                if cfg.get("use_tls", True):
                    smtp.starttls()
                if username:
                    smtp.login(username, password)
                smtp.send_message(msg)

        self._audit("email_send", "sent", "Email sent successfully", meta={"recipients": recipients})

    def run(self) -> None:
        self._log("=== TITAN GRID ENTERPRISE v3.4 START ===")
        df = self.process_data()
        if df.empty:
            raise ValueError("No compatible data found in workbook.")

        summary = self._build_summary(df)
        gate = self.preflight_check(df, summary)

        excel_output = self.generate_excel(df, summary, gate) if self.args.export_excel else ""
        pdf_output = self.generate_pdf(df, summary, gate) if self.args.export_pdf else ""
        docx_output = self.generate_word(df, summary, gate) if self.args.export_docx else ""

        zip_output = ""
        if self.args.export_zip:
            zip_output = self.create_zip_pack(bool(excel_output), bool(pdf_output), bool(docx_output))

        self.update_appendix_sheet()

        delivery_results: List[Dict[str, Any]] = []
        delivery_artifacts = [p for p in [zip_output, pdf_output, docx_output, excel_output] if p]
        if self.args.run_delivery or self.config.delivery.get("enabled", False):
            delivery_results = self.run_delivery(delivery_artifacts, gate)

        self.export_delivery_manifest(gate)
        self.export_audit_logs()
        self.update_appendix_sheet()

        if self.args.send_email:
            attachment = None if delivery_results else (zip_output or pdf_output or docx_output or excel_output)
            self.send_email(attachment, summary, gate, delivery_results)

        self.export_audit_logs()

        self._log("=== COMPLETED ===")
        print("\n" + "=" * 72)
        print("TITAN GRID ENTERPRISE v3.4 COMPLETE")
        print(f"Gate Status: {gate['gate_status']}")
        if gate["failures"]:
            for failure in gate["failures"]:
                print(f" - {failure}")
        if excel_output:
            print(f"Excel: {excel_output}")
        if pdf_output:
            print(f"PDF  : {pdf_output}")
        if docx_output:
            print(f"DOCX : {docx_output}")
        if zip_output:
            print(f"ZIP  : {zip_output}")
        print(f"Manifest JSON: {self.manifest_json_path}")
        print(f"Manifest CSV : {self.manifest_csv_path}")
        print(f"Audit JSON   : {self.audit_json_path}")
        print(f"Audit CSV    : {self.audit_csv_path}")
        if delivery_results:
            print("DELIVERY LINKS:")
            for item in delivery_results:
                print(f' - {Path(item["artifact"]).name}: {item["link"]}')
        print("=" * 72 + "\n")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="TITAN GRID VDR Enterprise v3.4")
    parser.add_argument("--vdr-root", required=True, help="Root folder of the VDR structure")
    parser.add_argument("--index", required=True, help="Path to MASTER_INDEX workbook")
    parser.add_argument("--config", required=False, default="", help="Optional YAML or JSON config file")
    parser.add_argument("--credentials-json", required=False, default="", help="Optional path to Google service account credentials.json")
    parser.add_argument("--sheets", default="MASTER_INDEX", help='Comma-separated sheet names or "all"')
    parser.add_argument("--output-dir", default="", help="Optional override for output directory")
    parser.add_argument("--export-excel", action="store_true", help="Generate formatted Excel committee pack")
    parser.add_argument("--export-pdf", action="store_true", help="Generate PDF summary pack")
    parser.add_argument("--export-docx", action="store_true", help="Generate Word committee pack")
    parser.add_argument("--export-zip", action="store_true", help="Generate zipped delivery pack")
    parser.add_argument("--run-delivery", action="store_true", help="Upload artifacts to configured delivery provider")
    parser.add_argument("--send-email", action="store_true", help="Send email with generated attachment or links")
    parser.add_argument("--verbose", action="store_true", help="Enable verbose logging")
    parser.add_argument("--fail-on-missing-columns", action="store_true", help="Fail when selected sheets lack required columns")
    return parser


if __name__ == "__main__":
    args = build_parser().parse_args()
    try:
        app = TitanGridVDREnterprise(args)
        app.run()
    except Exception as exc:
        logger.error("Execution failed: %s", exc)
        raise SystemExit(1)
