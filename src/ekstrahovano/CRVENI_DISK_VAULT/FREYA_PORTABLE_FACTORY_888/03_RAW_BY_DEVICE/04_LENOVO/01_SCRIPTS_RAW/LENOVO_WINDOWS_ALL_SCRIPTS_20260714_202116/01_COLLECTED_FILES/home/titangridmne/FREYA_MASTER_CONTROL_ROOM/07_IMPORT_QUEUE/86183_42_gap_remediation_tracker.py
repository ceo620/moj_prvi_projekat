# ============================================================
# TITAN_KERNEL: 42_gap_remediation_tracker.py
# PURPOSE: Build gap remediation tracker from Evidence Gap Register and remediation plan
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_FILE_ERROR = 2
EXIT_SCHEMA_ERROR = 3
EXIT_WRITE_ERROR = 4

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def norm(value):
    return "" if value is None else str(value).strip()

def headers(ws):
    return {norm(c.value): i for i, c in enumerate(ws[1], start=1) if norm(c.value)}

def load_json(path):
    if not path or not os.path.exists(path):
        return None
    with open(path, "r", encoding="utf-8-sig") as f:
        return json.load(f)

def main():
    parser = argparse.ArgumentParser(description="Build gap remediation tracker")
    parser.add_argument("--excel", required=True)
    parser.add_argument("--remediation-plan", help="remediation_plan.json")
    parser.add_argument("--out-xlsx", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print(f"❌ Excel ne postoji: {args.excel}")
        return EXIT_FILE_ERROR

    wb_in = openpyxl.load_workbook(args.excel, data_only=True)
    if "Evidence Gap Register" not in wb_in.sheetnames:
        print("❌ Nedostaje Evidence Gap Register")
        return EXIT_SCHEMA_ERROR

    gap_ws = wb_in["Evidence Gap Register"]
    h = headers(gap_ws)

    required = ["Gap_ID", "Description", "Required_Evidence", "Linked_FileFinding_ID", "Status", "Owner", "Blocking_STEP"]
    missing = [c for c in required if c not in h]
    if missing:
        print(f"❌ Nedostaju kolone: {missing}")
        return EXIT_SCHEMA_ERROR

    remediation = load_json(args.remediation_plan)
    tasks = remediation.get("tasks", []) if remediation else []

    rows = []
    for r in range(2, gap_ws.max_row + 1):
        gap_id = norm(gap_ws.cell(r, h["Gap_ID"]).value)
        if not gap_id:
            continue

        status = norm(gap_ws.cell(r, h["Status"]).value)
        linked = norm(gap_ws.cell(r, h["Linked_FileFinding_ID"]).value)

        related_tasks = []
        for t in tasks:
            text = json.dumps(t, ensure_ascii=False)
            if gap_id in text or "gap" in text.lower():
                related_tasks.append(t.get("Task_ID", ""))

        rows.append({
            "Gap_ID": gap_id,
            "Description": norm(gap_ws.cell(r, h["Description"]).value),
            "Required_Evidence": norm(gap_ws.cell(r, h["Required_Evidence"]).value),
            "Linked_FileFinding_ID": linked,
            "Current_Status": status,
            "Owner": norm(gap_ws.cell(r, h["Owner"]).value),
            "Blocking_STEP": norm(gap_ws.cell(r, h["Blocking_STEP"]).value),
            "Remediation_Tasks": ", ".join([x for x in related_tasks if x]),
            "Next_Action": "Attach valid evidence and submit for manual review" if not linked else "Manual review required",
            "Can_Close_Gap": "NO",
            "Can_Unlock_STEP102": "NO",
            "Final_Use_Allowed": "NO",
        })

    payload = {
        "timestamp": now_iso(),
        "component": "GAP_REMEDIATION_TRACKER",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "gap_count": len(rows),
        "rows": rows,
        "decision": "Tracker does not close gaps. Manual process required.",
        **CANON,
    }

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gap Remediation Tracker"
    headers_out = ["Gap_ID", "Description", "Required_Evidence", "Linked_FileFinding_ID", "Current_Status", "Owner", "Blocking_STEP", "Remediation_Tasks", "Next_Action", "Can_Close_Gap", "Can_Unlock_STEP102", "Final_Use_Allowed"]
    ws.append(headers_out)
    for row in rows:
        ws.append([row.get(hh, "") for hh in headers_out])

    ws2 = wb.create_sheet("Canonical Locks")
    ws2.append(["Control", "Value"])
    for k, v in CANON.items():
        ws2.append([k, v])

    try:
        Path(args.out_xlsx).parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.out_xlsx)
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else "✅ Gap remediation tracker built")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
