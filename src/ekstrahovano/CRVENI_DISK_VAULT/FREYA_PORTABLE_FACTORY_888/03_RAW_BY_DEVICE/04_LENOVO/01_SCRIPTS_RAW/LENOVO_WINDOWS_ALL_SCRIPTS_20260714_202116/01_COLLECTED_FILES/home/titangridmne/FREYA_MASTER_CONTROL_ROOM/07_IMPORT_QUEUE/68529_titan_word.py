import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_titan_document():
    doc = Document()

    # --- STIL I DIZAJN (EU Smart Factory Aesthetic) ---
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # --- HEADER (Zaglavlje) ---
    header = doc.sections[0].header
    htxt = header.paragraphs[0]
    htxt.text = "TITAN GRID | PROJECT CONFIDENTIAL | EU DATA ROOM 2026"
    htxt.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # --- NASLOV DOKUMENTA ---
    title = doc.add_heading('EXECUTIVE SUMMARY: SMART FACTORY TUZI', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # --- TEKST ---
    p = doc.add_paragraph()
    run = p.add_run("Strategic Expansion & Infrastructure Development")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(30, 50, 100) # Deep Blue

    doc.add_paragraph(
        "Ovaj dokument služi kao osnova za Data Room projekta TITAN GRID. "
        "Fokus je na automatizaciji proizvodnje transformatorskih kotlova u Crnoj Gori."
    )

    # --- TABELA (CAPEX/OPEX) ---
    doc.add_heading('FINANCIAL ALLOCATION (CAPEX)', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Shading Accent 1'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Stavka'
    hdr_cells[1].text = 'Iznos (€)'
    hdr_cells[2].text = 'Status'

    data = [
        ('Građevinski radovi - Tuzi', '2,500,000', 'U toku'),
        ('Roboti za zavarivanje', '1,200,000', 'Tender'),
        ('Solarni paneli (ESG)', '450,000', 'Planirano')
    ]

    for item, price, status in data:
        row_cells = table.add_row().cells
        row_cells[0].text = item
        row_cells[1].text = price
        row_cells[2].text = status

    # --- SNIMANJE I OTVARANJE ---
    file_name = "TITAN_Data_Room_Doc.docx"
    doc.save(file_name)
    
    # Automatski otvara Word dokument
    os.startfile(file_name)
    print(f"Word dokument '{file_name}' je uspešno kreiran i otvoren.")

if __name__ == "__main__":
    create_titan_document()