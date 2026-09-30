#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
03_extract_text.py

TITAN LOCAL EVIDENCE RAG ENGINE
Step 03 — Text Extraction Engine
Version: v2.0_AUDIT_LOCKED

Purpose
-------
Read only SAFE_FOR_EXTRACTION records from Step 02 and extract text into
controlled .txt artifacts. Produce an extraction manifest for Step 04.

This script does NOT use AI.
This script does NOT generate answers.
This script does NOT execute user files.
This script never extracts files that Step 02 did not mark as safe.

Core doctrine
-------------
Evidence precedes intelligence.
Retrieval precedes generation.
Audit precedes decision.

Input
-----
01_index/safe_file_index.jsonl

Outputs
-------
02_extracted/texts/*.txt
02_extracted/extraction_manifest.jsonl
02_extracted/extraction_manifest.csv
05_reports/extraction_summary.json
06_logs/extraction_audit.jsonl
06_logs/extraction_errors.jsonl

Supported extraction
--------------------
.txt, .md, .csv, .json, .jsonl, .log, .py, .sql, .yaml, .yml
.docx via python-docx
.xlsx/.xlsm via openpyxl
.pdf via pypdf if installed
.pptx via python-pptx if installed

Install optional libraries
--------------------------
pip install python-docx openpyxl pypdf python-pptx

Recommended command
-------------------
python ".\\08_scripts\\03_extract_text.py"

Incremental command from Step 09
--------------------------------
python ".\\08_scripts\\03_extract_text.py" --input ".\\01_index\\incremental_safe_file_index.jsonl" --overwrite
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import traceback
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable


SCRIPT_NAME = "03_extract_text.py"
SCRIPT_VERSION = "v2.0_AUDIT_LOCKED"

DEFAULT_BASE_DIR = Path(r"C:\Users\Korisnik\Desktop\TITAN_FULL_RAG")

INPUT_SAFE_INDEX = Path("01_index") / "safe_file_index.jsonl"

EXTRACTED_TEXT_DIR = Path("02_extracted") / "texts"
EXTRACTION_MANIFEST_JSONL = Path("02_extracted") / "extraction_manifest.jsonl"
EXTRACTION_MANIFEST_CSV = Path("02_extracted") / "extraction_manifest.csv"

SUMMARY_JSON = Path("05_reports") / "extraction_summary.json"
AUDIT_LOG = Path("06_logs") / "extraction_audit.jsonl"
ERROR_LOG = Path("06_logs") / "extraction_errors.jsonl"

SAFE_FOR_EXTRACTION = "SAFE_FOR_EXTRACTION"

STATUS_EXTRACTED = "EXTRACTED"
STATUS_SKIPPED_NOT_SAFE = "SKIPPED_NOT_SAFE"
STATUS_SKIPPED_EXISTS = "SKIPPED_EXISTS"
STATUS_SKIPPED_EMPTY = "SKIPPED_EMPTY"
STATUS_UNSUPPORTED_EXTENSION = "UNSUPPORTED_EXTENSION"
STATUS_SOURCE_MISSING = "SOURCE_MISSING"
STATUS_EXTRACTION_ERROR = "EXTRACTION_ERROR"

DEFAULT_MAX_CHARS_PER_FILE = 2_000_000

TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".jsonl", ".log", ".py", ".sql", ".yaml", ".yml"}
DOCX_EXTENSIONS = {".docx"}
XLSX_EXTENSIONS = {".xlsx", ".xlsm"}
PDF_EXTENSIONS = {".pdf"}
PPTX_EXTENSIONS = {".pptx"}

SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | DOCX_EXTENSIONS | XLSX_EXTENSIONS | PDF_EXTENSIONS | PPTX_EXTENSIONS


@dataclass(frozen=True)
class ExtractionPaths:
    base_dir: Path
    input_safe_index: Path
    extracted_text_dir: Path
    manifest_jsonl: Path
    manifest_csv: Path
    summary_json: Path
    audit_log: Path
    error_log: Path


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def normalize_ext(value: Any) -> str:
    text = str(value or "").strip().lower()
    if text and not text.startswith("."):
        return "." + text
    return text


def safe_int(value: Any, default: int = 0) -> int:
    try:
        if value is None or value == "":
            return default
        return int(value)
    except Exception:
        return default


def safe_float(value: Any, default: float = 0.0) -> float:
    try:
        if value is None or value == "":
            return default
        return float(value)
    except Exception:
        return default


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", errors="ignore")).hexdigest()


def sha256_file(path: Path) -> str | None:
    try:
        h = hashlib.sha256()
        with path.open("rb") as f:
            for block in iter(lambda: f.read(1024 * 1024), b""):
                h.update(block)
        return h.hexdigest()
    except Exception:
        return None


def append_jsonl(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(payload, ensure_ascii=False) + "\n")


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def write_jsonl(path: Path, records: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for record in records:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")


def write_csv(path: Path, records: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    fieldnames = [
        "extracted_at",
        "extraction_status",
        "extraction_reason",
        "source_path",
        "processed_path",
        "file_name",
        "extension",
        "source_sha256",
        "source_size_bytes",
        "source_modified_time_utc",
        "text_sha256",
        "text_chars",
        "text_words",
        "extractor",
        "source_record_safety_status",
        "chunking_allowed",
        "script",
        "script_version",
    ]

    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for record in records:
            writer.writerow(record)


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"Missing input safe index: {path}")

    records: list[dict[str, Any]] = []

    with path.open("r", encoding="utf-8", errors="ignore") as f:
        for line_no, line in enumerate(f, start=1):
            raw = line.strip()
            if not raw:
                continue

            try:
                item = json.loads(raw)
                if isinstance(item, dict):
                    records.append(item)
                else:
                    records.append({
                        "file_path": None,
                        "safety_status": "INVALID_JSONL",
                        "safety_reason": f"Line {line_no} is not a JSON object",
                    })
            except Exception as exc:
                records.append({
                    "file_path": None,
                    "safety_status": "INVALID_JSONL",
                    "safety_reason": f"Invalid JSONL line {line_no}: {exc}",
                })

    return records


def resolve_paths(base_dir: Path, input_arg: str | None) -> ExtractionPaths:
    return ExtractionPaths(
        base_dir=base_dir,
        input_safe_index=Path(input_arg) if input_arg else base_dir / INPUT_SAFE_INDEX,
        extracted_text_dir=base_dir / EXTRACTED_TEXT_DIR,
        manifest_jsonl=base_dir / EXTRACTION_MANIFEST_JSONL,
        manifest_csv=base_dir / EXTRACTION_MANIFEST_CSV,
        summary_json=base_dir / SUMMARY_JSON,
        audit_log=base_dir / AUDIT_LOG,
        error_log=base_dir / ERROR_LOG,
    )


def safe_output_filename(source_path: str, source_sha256: str | None, index: int) -> str:
    path = Path(source_path)
    stem = path.stem[:80] if path.stem else "unknown"
    suffix = (source_sha256 or sha256_text(source_path))[:16]
    safe_stem = "".join(ch if ch.isalnum() or ch in {"-", "_"} else "_" for ch in stem)
    return f"{index:08d}_{safe_stem}_{suffix}.txt"


def clean_extracted_text(text: str, max_chars: int) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.splitlines())
    text = text.strip()

    if max_chars > 0 and len(text) > max_chars:
        text = text[:max_chars].rstrip() + "\n\n[TRUNCATED_BY_TITAN_EXTRACTOR]"

    return text


def count_words(text: str) -> int:
    return len([x for x in text.split() if x.strip()])


def read_text_file(path: Path) -> str:
    encodings = ["utf-8", "utf-8-sig", "cp1250", "cp1252", "latin-1"]

    last_error: Exception | None = None

    for enc in encodings:
        try:
            return path.read_text(encoding=enc, errors="replace")
        except Exception as exc:
            last_error = exc

    raise RuntimeError(f"Could not read text file with supported encodings: {last_error}")


def extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("Missing dependency for .docx extraction. Run: pip install python-docx") from exc

    doc = Document(str(path))
    parts: list[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        parts.append(f"\n[TABLE {table_index}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def extract_xlsx(path: Path, max_rows_per_sheet: int = 5000) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("Missing dependency for .xlsx/.xlsm extraction. Run: pip install openpyxl") from exc

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    parts: list[str] = []

    for ws in wb.worksheets:
        parts.append(f"\n[SHEET: {ws.title}]")
        row_counter = 0

        for row in ws.iter_rows(values_only=True):
            row_counter += 1
            if row_counter > max_rows_per_sheet:
                parts.append(f"[TRUNCATED_SHEET_ROWS_AFTER_{max_rows_per_sheet}]")
                break

            values = []
            for value in row:
                if value is None:
                    values.append("")
                else:
                    values.append(str(value).replace("\n", " ").strip())

            if any(v for v in values):
                parts.append(" | ".join(values))

    try:
        wb.close()
    except Exception:
        pass

    return "\n".join(parts)


def extract_pdf(path: Path, max_pages: int | None = None) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError("Missing dependency for .pdf extraction. Run: pip install pypdf") from exc

    reader = PdfReader(str(path))
    parts: list[str] = []

    for page_index, page in enumerate(reader.pages, start=1):
        if max_pages is not None and page_index > max_pages:
            parts.append(f"\n[TRUNCATED_PDF_AFTER_{max_pages}_PAGES]")
            break

        try:
            text = page.extract_text() or ""
        except Exception as exc:
            text = f"[PAGE_EXTRACTION_ERROR: {exc}]"

        if text.strip():
            parts.append(f"\n[PDF_PAGE {page_index}]\n{text.strip()}")

    return "\n".join(parts)


def extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ImportError("Missing dependency for .pptx extraction. Run: pip install python-pptx") from exc

    prs = Presentation(str(path))
    parts: list[str] = []

    for slide_index, slide in enumerate(prs.slides, start=1):
        parts.append(f"\n[SLIDE {slide_index}]")
        for shape in slide.shapes:
            text = ""
            if hasattr(shape, "text"):
                text = str(shape.text or "").strip()
            if text:
                parts.append(text)

    return "\n".join(parts)


def extract_by_extension(path: Path, extension: str) -> tuple[str, str]:
    ext = normalize_ext(extension)

    if ext in TEXT_EXTENSIONS:
        return read_text_file(path), "plain_text_reader"

    if ext in DOCX_EXTENSIONS:
        return extract_docx(path), "python_docx"

    if ext in XLSX_EXTENSIONS:
        return extract_xlsx(path), "openpyxl"

    if ext in PDF_EXTENSIONS:
        return extract_pdf(path), "pypdf"

    if ext in PPTX_EXTENSIONS:
        return extract_pptx(path), "python_pptx"

    raise ValueError(f"Unsupported extraction extension: {ext}")


def base_manifest_record(record: dict[str, Any]) -> dict[str, Any]:
    return {
        "extracted_at": utc_now_iso(),
        "script": SCRIPT_NAME,
        "script_version": SCRIPT_VERSION,
        "source_path": record.get("file_path"),
        "processed_path": None,
        "file_name": record.get("file_name"),
        "extension": normalize_ext(record.get("extension")),
        "source_sha256": record.get("sha256"),
        "source_size_bytes": record.get("size_bytes"),
        "source_modified_time_utc": record.get("modified_time_utc"),
        "source_record_safety_status": record.get("safety_status"),
        "text_sha256": None,
        "text_chars": 0,
        "text_words": 0,
        "extractor": None,
        "chunking_allowed": False,
        "extraction_status": None,
        "extraction_reason": None,
        "source_record": {
            "rag_status": record.get("rag_status"),
            "safety_status": record.get("safety_status"),
            "safety_reason": record.get("safety_reason"),
            "risk_level": record.get("risk_level"),
            "hash_status": record.get("hash_status"),
        },
    }


def process_record(
    index: int,
    record: dict[str, Any],
    paths: ExtractionPaths,
    overwrite: bool,
    max_chars: int,
) -> dict[str, Any]:
    manifest = base_manifest_record(record)

    source_path_value = record.get("file_path")
    safety_status = record.get("safety_status")
    extension = normalize_ext(record.get("extension"))

    if safety_status != SAFE_FOR_EXTRACTION:
        manifest["extraction_status"] = STATUS_SKIPPED_NOT_SAFE
        manifest["extraction_reason"] = f"Record is not SAFE_FOR_EXTRACTION: {safety_status}"
        return manifest

    if not source_path_value:
        manifest["extraction_status"] = STATUS_SOURCE_MISSING
        manifest["extraction_reason"] = "Missing source path"
        return manifest

    source_path = Path(str(source_path_value))

    if not source_path.exists():
        manifest["extraction_status"] = STATUS_SOURCE_MISSING
        manifest["extraction_reason"] = "Source file does not exist"
        return manifest

    if extension not in SUPPORTED_EXTENSIONS:
        manifest["extraction_status"] = STATUS_UNSUPPORTED_EXTENSION
        manifest["extraction_reason"] = f"Unsupported extension for extraction: {extension}"
        return manifest

    source_sha = record.get("sha256") or sha256_file(source_path)
    output_name = safe_output_filename(str(source_path), source_sha, index)
    output_path = paths.extracted_text_dir / output_name
    manifest["processed_path"] = str(output_path)

    if output_path.exists() and not overwrite:
        try:
            existing_text = output_path.read_text(encoding="utf-8", errors="ignore")
            manifest["text_sha256"] = sha256_text(existing_text)
            manifest["text_chars"] = len(existing_text)
            manifest["text_words"] = count_words(existing_text)
        except Exception:
            pass

        manifest["extraction_status"] = STATUS_SKIPPED_EXISTS
        manifest["extraction_reason"] = "Processed text already exists. Use --overwrite to regenerate."
        manifest["chunking_allowed"] = True if manifest["text_chars"] else False
        return manifest

    try:
        raw_text, extractor = extract_by_extension(source_path, extension)
        text = clean_extracted_text(raw_text, max_chars=max_chars)

        manifest["extractor"] = extractor

        if not text.strip():
            manifest["extraction_status"] = STATUS_SKIPPED_EMPTY
            manifest["extraction_reason"] = "Extractor returned empty text"
            manifest["chunking_allowed"] = False
            return manifest

        paths.extracted_text_dir.mkdir(parents=True, exist_ok=True)
        output_path.write_text(text, encoding="utf-8")

        manifest["text_sha256"] = sha256_text(text)
        manifest["text_chars"] = len(text)
        manifest["text_words"] = count_words(text)
        manifest["chunking_allowed"] = True
        manifest["extraction_status"] = STATUS_EXTRACTED
        manifest["extraction_reason"] = "Text extracted successfully"
        return manifest

    except Exception as exc:
        manifest["extraction_status"] = STATUS_EXTRACTION_ERROR
        manifest["extraction_reason"] = str(exc)
        manifest["chunking_allowed"] = False

        append_jsonl(paths.error_log, {
            "timestamp": utc_now_iso(),
            "event": "FILE_EXTRACTION_FAILED",
            "script": SCRIPT_NAME,
            "script_version": SCRIPT_VERSION,
            "source_path": str(source_path),
            "extension": extension,
            "error": str(exc),
            "traceback": traceback.format_exc(),
        })

        return manifest


def summarize_manifest(
    manifest: list[dict[str, Any]],
    started_at: str,
    ended_at: str,
    paths: ExtractionPaths,
    input_records: int,
    overwrite: bool,
    max_chars: int,
) -> dict[str, Any]:
    status_counts: dict[str, int] = {}
    extension_counts: dict[str, int] = {}
    extractor_counts: dict[str, int] = {}

    extracted_count = 0
    error_count = 0
    skipped_count = 0
    chunking_allowed_count = 0
    total_chars = 0
    total_words = 0

    for item in manifest:
        status = str(item.get("extraction_status") or "UNKNOWN")
        ext = str(item.get("extension") or "[no extension]")
        extractor = str(item.get("extractor") or "[none]")

        status_counts[status] = status_counts.get(status, 0) + 1
        extension_counts[ext] = extension_counts.get(ext, 0) + 1
        extractor_counts[extractor] = extractor_counts.get(extractor, 0) + 1

        if status == STATUS_EXTRACTED:
            extracted_count += 1

        if status == STATUS_EXTRACTION_ERROR:
            error_count += 1

        if status.startswith("SKIPPED") or status in {STATUS_SOURCE_MISSING, STATUS_UNSUPPORTED_EXTENSION}:
            skipped_count += 1

        if item.get("chunking_allowed"):
            chunking_allowed_count += 1

        total_chars += safe_int(item.get("text_chars"))
        total_words += safe_int(item.get("text_words"))

    return {
        "started_at": started_at,
        "ended_at": ended_at,
        "script": SCRIPT_NAME,
        "script_version": SCRIPT_VERSION,
        "policy": "SAFE_TEXT_EXTRACTION_AUDIT_LOCKED",
        "input_safe_index": str(paths.input_safe_index),
        "input_records": input_records,
        "manifest_records": len(manifest),
        "extracted_count": extracted_count,
        "error_count": error_count,
        "skipped_count": skipped_count,
        "chunking_allowed_count": chunking_allowed_count,
        "total_text_chars": total_chars,
        "total_text_words": total_words,
        "overwrite": overwrite,
        "max_chars_per_file": max_chars,
        "status_counts": dict(sorted(status_counts.items())),
        "extension_counts": dict(sorted(extension_counts.items())),
        "extractor_counts": dict(sorted(extractor_counts.items())),
        "outputs": {
            "extracted_text_dir": str(paths.extracted_text_dir),
            "manifest_jsonl": str(paths.manifest_jsonl),
            "manifest_csv": str(paths.manifest_csv),
            "summary_json": str(paths.summary_json),
        },
        "governance_rule": {
            "extract_only_safe_file_index": True,
            "no_sensitive_skip_extraction": True,
            "no_rejected_extraction": True,
            "no_ai_generation": True,
            "source_file_remains_authoritative": True,
            "manifest_required_for_chunking": True,
        },
    }


def print_summary(summary: dict[str, Any]) -> None:
    print("=" * 100)
    print("TITAN RAG TEXT EXTRACTION v2.0 AUDIT LOCKED")
    print("=" * 100)
    print(f"Started:                  {summary['started_at']}")
    print(f"Ended:                    {summary['ended_at']}")
    print(f"Policy:                   {summary['policy']}")
    print(f"Input records:             {summary['input_records']}")
    print(f"Manifest records:          {summary['manifest_records']}")
    print(f"Extracted count:           {summary['extracted_count']}")
    print(f"Chunking allowed:          {summary['chunking_allowed_count']}")
    print(f"Skipped count:             {summary['skipped_count']}")
    print(f"Error count:               {summary['error_count']}")
    print(f"Total text chars:          {summary['total_text_chars']}")
    print(f"Total text words:          {summary['total_text_words']}")
    print("-" * 100)
    print("Status counts:")
    for key, value in summary["status_counts"].items():
        print(f" - {key}: {value}")
    print("-" * 100)
    for key, value in summary["outputs"].items():
        print(f"{key}: {value}")
    print("=" * 100)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="TITAN RAG Step 03 — safe text extraction engine."
    )

    parser.add_argument(
        "--base-dir",
        default=str(DEFAULT_BASE_DIR),
        help="Base TITAN_FULL_RAG directory.",
    )

    parser.add_argument(
        "--input",
        default=None,
        help="Optional input safe_file_index.jsonl path.",
    )

    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Overwrite existing extracted text artifacts.",
    )

    parser.add_argument(
        "--max-records",
        type=int,
        default=None,
        help="Optional cap for test runs.",
    )

    parser.add_argument(
        "--max-chars-per-file",
        type=int,
        default=DEFAULT_MAX_CHARS_PER_FILE,
        help="Maximum extracted text characters per file.",
    )

    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Read and validate input, but do not extract or write manifest.",
    )

    return parser.parse_args()


def main() -> None:
    args = parse_args()

    base_dir = Path(args.base_dir)
    paths = resolve_paths(base_dir, input_arg=args.input)

    started_at = utc_now_iso()

    try:
        records = read_jsonl(paths.input_safe_index)

        if args.max_records is not None:
            records = records[: int(args.max_records)]

        if args.dry_run:
            safe_count = sum(1 for r in records if r.get("safety_status") == SAFE_FOR_EXTRACTION)
            summary = {
                "started_at": started_at,
                "ended_at": utc_now_iso(),
                "script": SCRIPT_NAME,
                "script_version": SCRIPT_VERSION,
                "policy": "DRY_RUN_NO_EXTRACTION",
                "input_safe_index": str(paths.input_safe_index),
                "input_records": len(records),
                "safe_for_extraction_records": safe_count,
                "planned_output_dir": str(paths.extracted_text_dir),
            }
            print(json.dumps(summary, indent=2, ensure_ascii=False))
            return

        manifest: list[dict[str, Any]] = []

        for index, record in enumerate(records, start=1):
            item = process_record(
                index=index,
                record=record,
                paths=paths,
                overwrite=bool(args.overwrite),
                max_chars=int(args.max_chars_per_file),
            )
            manifest.append(item)

            if index % 250 == 0:
                print(f"[EXTRACTION] processed={index} manifest_records={len(manifest)}")

        ended_at = utc_now_iso()

        summary = summarize_manifest(
            manifest=manifest,
            started_at=started_at,
            ended_at=ended_at,
            paths=paths,
            input_records=len(records),
            overwrite=bool(args.overwrite),
            max_chars=int(args.max_chars_per_file),
        )

        write_jsonl(paths.manifest_jsonl, manifest)
        write_csv(paths.manifest_csv, manifest)
        write_json(paths.summary_json, summary)

        append_jsonl(paths.audit_log, {
            "timestamp": utc_now_iso(),
            "event": "EXTRACTION_COMPLETED",
            "script": SCRIPT_NAME,
            "script_version": SCRIPT_VERSION,
            "input_safe_index": str(paths.input_safe_index),
            "summary": summary,
        })

        print_summary(summary)

    except Exception as exc:
        append_jsonl(paths.error_log, {
            "timestamp": utc_now_iso(),
            "event": "EXTRACTION_FAILED",
            "script": SCRIPT_NAME,
            "script_version": SCRIPT_VERSION,
            "base_dir": str(base_dir),
            "input_safe_index": str(paths.input_safe_index),
            "error": str(exc),
            "traceback": traceback.format_exc(),
        })

        print("=" * 100)
        print("TITAN RAG TEXT EXTRACTION FAILED")
        print("=" * 100)
        print(f"Error: {exc}")
        print(f"Error log: {paths.error_log}")
        print("=" * 100)
        raise


if __name__ == "__main__":
    main()
