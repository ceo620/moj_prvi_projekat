# ============================================================
# TITAN_KERNEL: 66_policy_exception_register.py
# PURPOSE: Create policy exception register template and JSON baseline
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

HEADERS = [
    "Exception_ID",
    "Requested_By",
    "Requested_At",
    "Policy_Area",
    "Requested_Exception",
    "Business_Reason",
    "Linked_Evidence_ID",
    "Linked_FileFinding_ID",
    "Risk_Level",
    "Reviewer",
    "Review_Status",
    "Approved",
    "Can_Change_SYSTEM_RED",
    "Can_Unlock_STEP102",
    "Can_Close_Gate",
    "Can_Approve_Evidence",
    "Can_Write_SSoT",
    "Final_Use_Allowed",
    "Notes"
]

DEFAULT_ROWS = [
    {
        "Exception_ID": "EXC-00001",
        "Requested_By": "",
        "Requested_At": "",
        "Policy_Area": "SYSTEM_STATUS",
        "Requested_Exception": "Attempt to change SYSTEM RED",
        "Business_Reason": "",
        "Linked_Evidence_ID": "",
        "Linked_FileFinding_ID": "",
        "Risk_Level": "CRITICAL",
        "Reviewer": "",
        "Review_Status": "REVIEW_REQUIRED",
        "Approved": "NO",
        "Can_Change_SYSTEM_RED": "NO",
        "Can_Unlock_STEP102": "NO",
        "Can_Close_Gate": "NO",
        "Can_Approve_Evidence": "NO",
        "Can_Write_SSoT": "NO",
        "Final_Use_Allowed": "NO",
        "Notes": "Template row only. No exception is approved."
    }
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def main():
    parser = argparse.ArgumentParser(description="Build TITAN policy exception register")
    parser.add_argument("--out-xlsx", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    payload = {
        "timestamp": now_iso(),
        "component": "POLICY_EXCEPTION_REGISTER",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "records": DEFAULT_ROWS,
        "decision": "No policy exception is approved by this register.",
        **CANON
    }

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Policy Exceptions"
    ws.append(HEADERS)
    for row in DEFAULT_ROWS:
        ws.append([row.get(h, "") for h in HEADERS])

    ws2 = wb.create_sheet("Canonical Locks")
    ws2.append(["Control", "Value"])
    for k, v in CANON.items():
        ws2.append([k, v])

    ws3 = wb.create_sheet("Instructions")
    ws3.append(["Instruction_ID", "Instruction"])
    ws3.append(["INS-001", "No row may approve final use."])
    ws3.append(["INS-002", "No row may unlock STEP102."])
    ws3.append(["INS-003", "No row may change SYSTEM RED."])
    ws3.append(["INS-004", "No row may close gates or approve evidence."])

    try:
        Path(args.out_xlsx).parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.out_xlsx)
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else f"✅ Policy exception register: {args.out_xlsx}")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
