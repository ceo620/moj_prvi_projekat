import os
from docx import Document

# Kompjuter sam pronalazi Desktop
desktop = os.path.join(os.path.join(os.environ['USERPROFILE']), 'Desktop')
folder = os.path.join(desktop, "TITAN_DOKUMENTI")

if not os.path.exists(folder):
    os.makedirs(folder)

print("🚀 TITAN KREĆE... Štampam 309 fajlova na Desktop...")

for i in range(1, 310):
    doc = Document()
    doc.add_heading(f'TITAN GRID - MASTER SCENARIO {i:03d}', 0)
    doc.add_paragraph(f"Scenario ID: TG-{i:03d}")
    doc.save(os.path.join(folder, f"TITAN_{i:03d}.docx"))
    if i % 100 == 0:
        print(f"✅ Završeno {i}...")

print(f"🏁 GOTOVO! Onur, folder TITAN_DOKUMENTI je na tvom Desktopu.")