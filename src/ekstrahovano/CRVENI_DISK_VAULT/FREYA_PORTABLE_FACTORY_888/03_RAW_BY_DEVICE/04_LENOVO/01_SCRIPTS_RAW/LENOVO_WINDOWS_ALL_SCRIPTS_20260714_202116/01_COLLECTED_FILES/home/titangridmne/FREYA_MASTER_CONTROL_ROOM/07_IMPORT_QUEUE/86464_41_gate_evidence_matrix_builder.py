# ============================================================
# TITAN_KERNEL: 41_gate_evidence_matrix_builder.py
# PURPOSE: Build P0 gate to evidence mapping matrix
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_FILE_ERROR = 2
EXIT_SCHEMA_ERROR = 3
EXIT_WRITE_ERROR = 4

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "RISK_BASELINE": 890,
    "P0_GATES": "10/10 BLOCKED",
    "EVIDENCE_GAPS_REMAINING": 6,
    "EVIDENCE_APPROVED": 0,
    "GATES_CLOSED": 0,
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

DEFAULT_GATES = [
    ("P0-01", "Config Validity Gate", "config.json schema and path evidence"),
    ("P0-02", "SSoT Integrity Gate", "monolith_registry.json hash and Lucky_Number evidence"),
    ("P0-03", "Script Integrity Gate", "script manifest and hash verification"),
    ("P0-04", "FileFinding Mapping Gate", "FileFinding_ID and Evidence_ID mapping"),
    ("P0-05", "Evidence Gap Gate", "all gaps linked to valid evidence"),
    ("P0-06", "Control Tower Gate", "Rule 01-03 validator report"),
    ("P0-07", "Review Governance Gate", "manual review queue and decision log"),
    ("P0-08", "Custody/Privacy Gate", "chain-of-custody and sensitive data scan"),
    ("P0-09", "Release Readiness Gate", "NOT_READY until all blockers resolved"),
    ("P0-10", "Preproduction Freeze Gate", "FREEZE_DENIED until active canon permits"),
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def norm(value):
    return "" if value is None else str(value).strip()

def headers(ws):
    return {norm(c.value): i for i, c in enumerate(ws[1], start=1) if norm(c.value)}

def split_ids(value):
    return [x.strip() for x in value.replace(";", ",").split(",") if x.strip()]

def main():
    parser = argparse.ArgumentParser(description="Build P0 gate/evidence matrix")
    parser.add_argument("--excel", required=True, help="FileFinding/EvidenceGap workbook")
    parser.add_argument("--out-xlsx", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print(f"❌ Excel ne postoji: {args.excel}")
        return EXIT_FILE_ERROR

    wb_in = openpyxl.load_workbook(args.excel, data_only=True)
    if "Evidence Gap Register" not in wb_in.sheetnames:
        print("❌ Nedostaje Evidence Gap Register")
        return EXIT_SCHEMA_ERROR

    gap_ws = wb_in["Evidence Gap Register"]
    h = headers(gap_ws)

    required = ["Gap_ID", "Description", "Linked_FileFinding_ID", "Status", "Blocking_STEP"]
    missing = [c for c in required if c not in h]
    if missing:
        print(f"❌ Nedostaju kolone: {missing}")
        return EXIT_SCHEMA_ERROR

    gaps = []
    for r in range(2, gap_ws.max_row + 1):
        gap_id = norm(gap_ws.cell(r, h["Gap_ID"]).value)
        if not gap_id:
            continue
        gaps.append({
            "Gap_ID": gap_id,
            "Description": norm(gap_ws.cell(r, h["Description"]).value) if "Description" in h else "",
            "Linked_FileFinding_ID": norm(gap_ws.cell(r, h["Linked_FileFinding_ID"]).value),
            "Status": norm(gap_ws.cell(r, h["Status"]).value),
            "Blocking_STEP": norm(gap_ws.cell(r, h["Blocking_STEP"]).value),
        })

    matrix = []
    for gate_id, gate_name, required_evidence in DEFAULT_GATES:
        linked_gaps = [g for g in gaps if g.get("Blocking_STEP", "").upper() in {"STEP102", gate_id.upper()}]
        open_gaps = [g for g in linked_gaps if g.get("Status", "").upper() not in {"CLOSED", "APPROVED", "RESOLVED", "VALIDATED"}]
        linked_findings = []
        for g in linked_gaps:
            linked_findings.extend(split_ids(g.get("Linked_FileFinding_ID", "")))

        matrix.append({
            "Gate_ID": gate_id,
            "Gate_Name": gate_name,
            "Required_Evidence": required_evidence,
            "Linked_Gap_Count": len(linked_gaps),
            "Open_Gap_Count": len(open_gaps),
            "Linked_FileFinding_IDs": ", ".join(sorted(set(linked_findings))),
            "Gate_Status": "BLOCKED",
            "Can_Close_Gate": "NO",
            "Final_Use_Allowed": "NO",
        })

    payload = {
        "timestamp": now_iso(),
        "component": "GATE_EVIDENCE_MATRIX_BUILDER",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "matrix": matrix,
        "decision": "P0 gates remain blocked. Matrix is for visibility only.",
        **CANON,
    }

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gate Evidence Matrix"
    headers_out = ["Gate_ID", "Gate_Name", "Required_Evidence", "Linked_Gap_Count", "Open_Gap_Count", "Linked_FileFinding_IDs", "Gate_Status", "Can_Close_Gate", "Final_Use_Allowed"]
    ws.append(headers_out)
    for row in matrix:
        ws.append([row.get(hh, "") for hh in headers_out])

    ws2 = wb.create_sheet("Canonical Locks")
    ws2.append(["Control", "Value"])
    for k, v in CANON.items():
        ws2.append([k, v])

    Path(args.out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out_xlsx)
    Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
    Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else "✅ Gate evidence matrix built")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
