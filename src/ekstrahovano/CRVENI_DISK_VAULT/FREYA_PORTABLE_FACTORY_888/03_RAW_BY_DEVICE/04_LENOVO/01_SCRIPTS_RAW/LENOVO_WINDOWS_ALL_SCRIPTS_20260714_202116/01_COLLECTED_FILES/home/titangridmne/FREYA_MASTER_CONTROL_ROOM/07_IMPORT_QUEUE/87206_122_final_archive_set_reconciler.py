# ============================================================
# TITAN_KERNEL: 122_final_archive_set_reconciler.py
# PURPOSE: Reconcile final closure, read-only proof and non-acceptance archive sets
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import hashlib
import json
import zipfile
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_OK = 0
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def classify_archive(path):
    n = path.name.lower()
    if "nonacceptance" in n:
        return "NONACCEPTANCE_ARCHIVE"
    if "readonly" in n or "read_only" in n:
        return "READONLY_PROOF_ARCHIVE"
    if "final_closure" in n or "system_red_final_closure" in n:
        return "FINAL_CLOSURE_ARCHIVE"
    if "certificate" in n:
        return "CERTIFICATE_ARCHIVE"
    if "denial" in n:
        return "DENIAL_ARCHIVE"
    return "OTHER_LOCKED_ARCHIVE"

def inspect_zip(path):
    try:
        with zipfile.ZipFile(path, "r") as z:
            bad = z.testzip()
            names = z.namelist()
        return bad is None, bad or "", names
    except Exception as exc:
        return False, str(exc), []

def collect(roots):
    records = []
    for root in roots:
        p = Path(root)
        if not p.exists():
            continue
        iterator = [p] if p.is_file() else p.rglob("*.zip")
        for item in iterator:
            if item.is_file() and item.suffix.lower() == ".zip":
                valid, bad, names = inspect_zip(item)
                records.append({
                    "Archive_Set_ID": f"FARC-{len(records)+1:05d}",
                    "Archive_Name": item.name,
                    "Archive_Type": classify_archive(item),
                    "Path": str(item),
                    "SHA256": sha256_file(item),
                    "Valid_Zip": valid,
                    "Bad_Member": bad,
                    "Item_Count": len(names),
                    "Items": names,
                    "Size_Bytes": item.stat().st_size,
                    "Modified_UTC": datetime.utcfromtimestamp(item.stat().st_mtime).isoformat(timespec="seconds"),
                    "Canonical_Effect": "NONE",
                    "Final_Use_Allowed": "NO",
                })
    return records

def write_xlsx(payload, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Final Archive Set"
    headers = ["Archive_Set_ID", "Archive_Name", "Archive_Type", "Path", "SHA256", "Valid_Zip", "Bad_Member", "Item_Count", "Size_Bytes", "Modified_UTC", "Canonical_Effect", "Final_Use_Allowed"]
    ws.append(headers)
    for row in payload["archives"]:
        ws.append([row.get(h, "") for h in headers])

    ws2 = wb.create_sheet("Archive Contents")
    ws2.append(["Archive_Set_ID", "Archive_Name", "Item"])
    for row in payload["archives"]:
        for item in row.get("Items", []):
            ws2.append([row["Archive_Set_ID"], row["Archive_Name"], item])

    ws3 = wb.create_sheet("Canonical Locks")
    ws3.append(["Control", "Value"])
    for k, v in CANON.items():
        ws3.append([k, v])

    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def main():
    parser = argparse.ArgumentParser(description="Reconcile final locked archive sets")
    parser.add_argument("--roots", nargs="+", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--out-xlsx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    archives = collect(args.roots)

    payload = {
        "timestamp": now_iso(),
        "component": "FINAL_ARCHIVE_SET_RECONCILER",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "archive_count": len(archives),
        "invalid_count": sum(1 for a in archives if not a["Valid_Zip"]),
        "archives": archives,
        "decision": "Archive set reconciliation is inventory only. It does not approve or unlock anything.",
        **CANON,
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.out_xlsx:
            write_xlsx(payload, args.out_xlsx)
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else f"✅ Final archive set reconciled: {len(archives)} archives")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    raise SystemExit(main())
