from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

out = Path("03_ANNEXES")
out.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)

def p(text="", bold=False, size=10.5, align=None, after=6):
    par = doc.add_paragraph()
    if align:
        par.alignment = align
    par.paragraph_format.space_after = Pt(after)
    r = par.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(size)
    return par

p("ANEKS 1", True, 18, WD_ALIGN_PARAGRAPH.CENTER, 4)
p("REVIZIJA CAPEX OKVIRA ZA PROJEKAT TITAN 1", True, 13, WD_ALIGN_PARAGRAPH.CENTER, 12)

p("Datum: 2026-05-10")

p("1. PREDMET ANEKSA", True, 11)
p("Ovaj Aneks 1 odnosi se isključivo na projekat TITAN 1 – Fabrika transformatorskih kazana (tankova).")
p("Nosilac projekta je Ars Metal Industries D.O.O. Podgorica.")
p("Lokacija projekta je industrijska zona Tuzi / KAP, Crna Gora.")
p("Ovaj Aneks ne odnosi se na druge buduće projekte, faze ili proizvodne programe koji nijesu predmet aktivnog institucionalnog zahtjeva.")

p("2. SVRHA REVIZIJE", True, 11)
p("Svrha ovog Aneksa je usklađivanje CAPEX okvira projekta TITAN 1 sa ažuriranim internim inženjerskim, tehničkim i troškovnim procjenama nastalim nakon izrade originalnog Elaborata.")
p("Originalni Elaborat izrađen je u ranijoj razvojnoj fazi projekta i sadržao je preliminarne finansijske procjene podložne daljoj reviziji i usklađivanju.")

p("3. AŽURIRANI CAPEX OKVIR", True, 11)
p("Prethodna vrijednost navedena u originalnom Elaboratu: [vrijednost iz originalnog Elaborata]")
p("Ažurirani CAPEX okvir za projekat TITAN 1: 18,2 miliona EUR")
p("Revizija CAPEX okvira zasniva se na preciznijem definisanju proizvodnog procesa, ažuriranim tehničkim parametrima, novim ponudama dobavljača, optimizaciji planirane proizvodne konfiguracije i detaljnijem troškovnom modeliranju specifičnom za proizvodnju transformatorskih kazana.")

p("4. STATUS DOKUMENTA", True, 11)
p("Ovaj Aneks predstavlja prateće objašnjenje uz postojeću projektnu dokumentaciju i koristi se za potrebe institucionalne komunikacije, preliminarnog usklađivanja dokumentacije i daljih procedura vezanih za projekat TITAN 1.")
p("Originalni Elaborat ostaje relevantan u dijelovima koji nijesu predmet ove CAPEX revizije.")
p("Formalna integracija eventualnih ažuriranja u buduću projektnu dokumentaciju biće izvršena u skladu sa važećim procedurama i zahtjevima nadležnih institucija.")

p("5. ZAKLJUČAK", True, 11)
p("Prema trenutno raspoloživim internim procjenama i projektnim parametrima, aktivni CAPEX okvir za projekat TITAN 1 iznosi približno 18,2 miliona EUR.")
p("Ovaj dokument koristi se isključivo za potrebe usklađivanja i pojašnjenja aktivnog investicionog okvira projekta TITAN 1.")

p("Sa poštovanjem,", False, 10.5, None, 14)
p("DANIJELA KESKIN", True, 10.5, None, 0)
p("Director", False, 10.5, None, 0)
p("Ars Metal Industries D.O.O. Podgorica", False, 10.5, None, 0)
p("+382 68 323 339 | d.djurovic@adsmetal.com.tr", False, 10.5, None, 0)

target = out / "ANEKS_1_REVIZIJA_CAPEX_TITAN1_2026-05-10.docx"
doc.save(target)
print("CREATED:", target)
