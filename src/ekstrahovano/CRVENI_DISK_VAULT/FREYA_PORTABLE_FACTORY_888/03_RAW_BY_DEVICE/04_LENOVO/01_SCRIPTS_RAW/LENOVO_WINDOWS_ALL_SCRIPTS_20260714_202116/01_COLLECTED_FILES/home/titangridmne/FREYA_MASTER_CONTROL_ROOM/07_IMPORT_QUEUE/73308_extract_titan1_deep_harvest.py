import os, re, json, glob, zipfile
from openpyxl import load_workbook
from datetime import datetime

# TITAN 1 ACTIVE BASELINE
ALLOWED = [r"18,2", r"18\.2", r"18.2 miliona", r"nova radna mjesta", r"industrijski razvoj", r"jačanje energetskog lanca", r"transformatorskih kazana", r"tankova", r"TITAN 1"]
BLOCKED = [r"43\.5", r"51\.4", r"3 factories", r"TITAN 2", r"TITAN 3", r"38%", r"25% grant", r"Equity IRR", r"Project IRR", r"DSCR", r"lender ready", r"bankable"]

REPORT_DIR = f"/root/Documents/TITAN1_DEEP_HARVEST_{datetime.now().strftime('%Y%m%d_%H%M')}"
os.makedirs(REPORT_DIR, exist_ok=True)

print("🔥 NAJDUBLJI TITAN 1 SIGNAL HARVEST — ACTIVE SCOPE ONLY")
print("   Status: ACTIVE_TITAN1_SCOPE_CANONICAL_WORKING_BASELINE\n")

folder = "/root/Documents/TITAN FINAL"
signals = []

for path in glob.glob(f"{folder}/**/*.*", recursive=True):
    filename = os.path.basename(path).lower()
    try:
        content = ""
        if filename.endswith(('.xlsx')):
            wb = load_workbook(path, data_only=True)
            for sheet in wb.sheetnames:
                for row in wb[sheet].iter_rows(values_only=True):
                    for cell in row:
                        if cell:
                            content += str(cell) + " "
        elif filename.endswith(('.docx', '.pptx')):
            with zipfile.ZipFile(path) as z:
                for name in z.namelist():
                    if name.endswith(('.xml')):
                        content += z.read(name).decode('utf-8', errors='ignore')
        else:
            with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()

        # Traži signale
        for pattern in ALLOWED + BLOCKED:
            matches = re.findall(pattern, content, re.IGNORECASE)
            for match in matches:
                signals.append({
                    "file": os.path.relpath(path, "/root/Documents"),
                    "match": match,
                    "type": "ALLOWED_TITAN1" if pattern in ALLOWED else "BLOCKED_BLEEDTHROUGH"
                })
    except:
        pass

# Sačuvaj
with open(f"{REPORT_DIR}/TITAN1_DEEP_SIGNALS.json", "w") as f:
    json.dump(signals, f, indent=2, ensure_ascii=False)

allowed = len([s for s in signals if s["type"] == "ALLOWED_TITAN1"])
blocked = len([s for s in signals if s["type"] == "BLOCKED_BLEEDTHROUGH"])

print(f"✅ NAJDUBLJI HARVEST ZAVRŠEN!")
print(f"   Report folder: {REPORT_DIR}")
print(f"   Dozvoljeni TITAN 1 signali: {allowed}")
print(f"   Blokirani bleedthrough signali: {blocked}")
print(f"\nGlavna datoteka: {REPORT_DIR}/TITAN1_DEEP_SIGNALS.json")
print("   Spremno za korištenje u ANEKS 1 i SIGNALNI_MEMORANDUM_TITAN1.")
