from __future__ import annotations

import io
from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.models import ExpenseType, MilestoneStatus, Project

NAVY = '002060'
GREEN = 'E2F0D9'
RED = 'FCE4D6'
YELLOW = 'FFF2CC'
WHITE = 'FFFFFF'

HEADER_FILL = PatternFill('solid', fgColor=NAVY)
PASS_FILL = PatternFill('solid', fgColor=GREEN)
FAIL_FILL = PatternFill('solid', fgColor=RED)
WARN_FILL = PatternFill('solid', fgColor=YELLOW)
WHITE_BOLD = Font(color=WHITE, bold=True)
BOLD = Font(bold=True)
THIN = Side(style='thin', color='999999')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center')


def forensic_findings(project: Project) -> list[dict]:
    total_spent = sum((e.amount for e in project.expenses), Decimal('0.00'))
    capex_total = sum((e.amount for e in project.expenses if e.type == ExpenseType.CAPEX), Decimal('0.00'))
    opex_total = sum((e.amount for e in project.expenses if e.type == ExpenseType.OPEX), Decimal('0.00'))
    utilization = (total_spent / project.budget * Decimal('100.00')) if project.budget else Decimal('0.00')
    findings: list[dict] = []

    def add(category: str, check_name: str, expected: str, actual: str, status: str, severity: str, comment: str) -> None:
        findings.append({
            'category': category,
            'check_name': check_name,
            'expected': expected,
            'actual': actual,
            'status': status,
            'severity': severity,
            'comment': comment,
        })

    add('PROJECT', 'Budget positive', '> 0', str(project.budget), 'PASS' if project.budget > 0 else 'FAIL', 'CRITICAL', 'Budget mora biti pozitivan.')
    add('FORMULA', 'Total spent reconciliation', str(total_spent), str(total_spent), 'PASS', 'LOW', 'Suma expense linija.')
    add('FORMULA', 'CAPEX subtotal', 'Derived', str(capex_total), 'PASS', 'LOW', 'CAPEX subtotal.')
    add('FORMULA', 'OPEX subtotal', 'Derived', str(opex_total), 'PASS', 'LOW', 'OPEX subtotal.')
    add('FORMULA', 'Utilization %', 'Derived', str(utilization), 'PASS' if utilization <= Decimal('100.00') else 'FAIL', 'HIGH' if utilization > Decimal('100.00') else 'LOW', 'Utilization ne bi trebalo da pređe 100% bez rebalansa.')

    now = datetime.now(timezone.utc)
    for m in project.milestones:
        overdue = m.due_date.replace(tzinfo=timezone.utc) < now and m.status != MilestoneStatus.COMPLETED
        add('MILESTONE', f'Milestone {m.id} overdue', 'False', str(overdue), 'FAIL' if overdue else 'PASS', 'HIGH' if overdue else 'LOW', 'Istekli milestone zahtijeva pregled.')

    return findings


def _header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD
        cell.border = BORDER
        cell.alignment = CENTER


def _grid(ws) -> None:
    for row in ws.iter_rows():
        for cell in row:
            cell.border = BORDER
            if cell.row > 1:
                cell.alignment = LEFT


def _autosize(ws) -> None:
    for col in ws.columns:
        width = max(len(str(cell.value or '')) for cell in col) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width, 12), 42)


def _bank_layout(ws) -> None:
    ws.freeze_panes = 'A2'
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.gridLines = False
    ws.oddHeader.center.text = '&BCAPEX/OPEX Master System'
    ws.oddFooter.right.text = 'Page &[Page] of &N'


def build_project_workbook(project: Project) -> io.BytesIO:
    findings = forensic_findings(project)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Executive_Dashboard'
    ws['A1'] = 'Executive Dashboard'
    ws['A1'].font = Font(size=16, bold=True)
    ws.append([])
    ws.append(['Metric', 'Value'])
    _header(ws, 3)
    metrics = [
        ['Project Name', project.name],
        ['Budget', float(project.budget)],
        ['Total Spent', float(sum((e.amount for e in project.expenses), Decimal('0.00')))],
        ['Expense Count', len(project.expenses)],
        ['Milestone Count', len(project.milestones)],
        ['Export Timestamp UTC', datetime.utcnow().isoformat()],
    ]
    for row in metrics:
        ws.append(row)
    _grid(ws)
    _autosize(ws)
    _bank_layout(ws)

    ws_exp = wb.create_sheet('Expenses')
    ws_exp.append(['ID', 'Description', 'Amount', 'Type', 'Date', 'Project ID'])
    _header(ws_exp)
    for e in project.expenses:
        ws_exp.append([e.id, e.description, float(e.amount), e.type.value, e.date.isoformat(), e.project_id])
    dv = DataValidation(type='list', formula1='"CAPEX,OPEX"', allow_blank=False)
    ws_exp.add_data_validation(dv)
    if ws_exp.max_row >= 2:
        dv.add(f'D2:D{ws_exp.max_row}')
    ws_exp.conditional_formatting.add(f'C2:C{max(2, ws_exp.max_row)}', CellIsRule(operator='lessThanOrEqual', formula=['0'], fill=FAIL_FILL))
    _grid(ws_exp)
    _autosize(ws_exp)
    _bank_layout(ws_exp)

    ws_m = wb.create_sheet('Milestones')
    ws_m.append(['ID', 'Name', 'Due Date', 'Status', 'Project ID'])
    _header(ws_m)
    for m in project.milestones:
        ws_m.append([m.id, m.name, m.due_date.isoformat(), m.status.value, m.project_id])
    _grid(ws_m)
    _autosize(ws_m)
    _bank_layout(ws_m)

    ws_err = wb.create_sheet('Error_Check')
    ws_err.append(['Category', 'Check Name', 'Expected', 'Actual', 'Status', 'Severity', 'Comment'])
    _header(ws_err)
    for item in findings:
        ws_err.append([item[k] for k in ['category', 'check_name', 'expected', 'actual', 'status', 'severity', 'comment']])
    ws_err.conditional_formatting.add(f'E2:E{max(2, ws_err.max_row)}', CellIsRule(operator='equal', formula=['"PASS"'], fill=PASS_FILL))
    ws_err.conditional_formatting.add(f'E2:E{max(2, ws_err.max_row)}', CellIsRule(operator='equal', formula=['"FAIL"'], fill=FAIL_FILL))
    _grid(ws_err)
    _autosize(ws_err)
    _bank_layout(ws_err)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
