# ============================================================
# TITAN_KERNEL: 81_board_decision_validator.py
# PURPOSE: Validate board/quorum/decision artifacts cannot override canon
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_BLOCK = 1
EXIT_FILE_ERROR = 2
EXIT_WRITE_ERROR = 3

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

FORBIDDEN_COLUMNS = [
    "Can_Approve_Evidence",
    "Can_Close_Gate",
    "Can_Unlock_STEP102",
    "Can_Write_SSoT",
    "Final_Use_Allowed",
    "Evidence_Approved",
    "Gate_Closed",
    "SSOT_Write_Approved",
    "STEP102_Unlock_Approved"
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def norm(value):
    return "" if value is None else str(value).strip()

def headers(ws):
    return {norm(c.value): i for i, c in enumerate(ws[1], start=1) if norm(c.value)}

def scan_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    violations = []
    scanned_cells = 0

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        if ws.max_row < 1:
            continue

        h = headers(ws)
        for col in FORBIDDEN_COLUMNS:
            if col not in h:
                continue
            cidx = h[col]
            for r in range(2, ws.max_row + 1):
                val = norm(ws.cell(r, cidx).value).upper()
                scanned_cells += 1
                if val == "YES":
                    violations.append({
                        "workbook": path,
                        "sheet": sheet,
                        "row": r,
                        "column": col,
                        "value": "YES",
                        "violation": "Forbidden board decision override attempted"
                    })
    return scanned_cells, violations

def main():
    parser = argparse.ArgumentParser(description="Validate board decision artifacts against canon")
    parser.add_argument("--workbooks", nargs="+", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    missing = []
    all_violations = []
    scanned_total = 0

    for wb_path in args.workbooks:
        if not os.path.exists(wb_path):
            missing.append(wb_path)
            continue
        try:
            scanned, violations = scan_workbook(wb_path)
            scanned_total += scanned
            all_violations.extend(violations)
        except Exception as exc:
            all_violations.append({
                "workbook": wb_path,
                "violation": f"Workbook scan failed: {exc}"
            })

    blocked = bool(missing or all_violations)

    payload = {
        "timestamp": now_iso(),
        "component": "BOARD_DECISION_VALIDATOR",
        "version": "1.0",
        "status": "BLOCK" if blocked else "REVIEW_REQUIRED",
        "missing_workbooks": missing,
        "scanned_cells": scanned_total,
        "violation_count": len(all_violations),
        "violations": all_violations,
        "decision": "Board artifacts cannot override active canon. Any YES approval is blocked.",
        **CANON
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else ("⛔ Board decision override violation" if blocked else "✅ Board decisions verified as non-overriding"))
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_BLOCK if blocked else EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
