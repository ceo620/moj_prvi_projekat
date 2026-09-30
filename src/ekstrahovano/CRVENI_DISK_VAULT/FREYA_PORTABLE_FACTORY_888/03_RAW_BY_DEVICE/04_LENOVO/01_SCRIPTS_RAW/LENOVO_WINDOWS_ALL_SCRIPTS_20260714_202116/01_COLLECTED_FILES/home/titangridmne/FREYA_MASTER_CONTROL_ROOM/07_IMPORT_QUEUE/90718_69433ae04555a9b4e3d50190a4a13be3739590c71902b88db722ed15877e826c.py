#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
20_rag_control_tower_export.py

TITAN LOCAL EVIDENCE RAG ENGINE
Step 20 — RAG Control Tower Export
Version: v2.0_AUDIT_LOCKED

Purpose
-------
Export consolidated TITAN RAG pipeline outputs into one institutional Excel
Control Tower workbook and companion JSON/Markdown summaries.

This script does NOT scan source documents.
This script does NOT modify source documents.
This script does NOT generate unsupported conclusions.
It only consolidates already-produced audit/report outputs.

Core doctrine
-------------
Evidence precedes intelligence.
Retrieval precedes generation.
Audit precedes decision.

Inputs
------
05_reports/*.json
05_reports/*.csv
06_logs/*.jsonl

Primary outputs
---------------
05_reports/TITAN_RAG_CONTROL_TOWER.xlsx
05_reports/rag_control_tower_export_summary.json
05_reports/rag_control_tower_export_summary.md
06_logs/rag_control_tower_export_audit.jsonl
06_logs/rag_control_tower_export_errors.jsonl

Recommended command
-------------------
python ".\\08_scripts\\20_rag_control_tower_export.py"

With print:
-----------
python ".\\08_scripts\\20_rag_control_tower_export.py" --print

Custom output:
--------------
python ".\\08_scripts\\20_rag_control_tower_export.py" --output ".\\05_reports\\TITAN_RAG_CONTROL_TOWER.xlsx"

Dependencies
------------
pip install openpyxl
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import traceback
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable


SCRIPT_NAME = "20_rag_control_tower_export.py"
SCRIPT_VERSION = "v2.0_AUDIT_LOCKED"

DEFAULT_BASE_DIR = Path(r"C:\Users\Korisnik\Desktop\TITAN_FULL_RAG")

REPORTS_DIR = Path("05_reports")
LOGS_DIR = Path("06_logs")

OUTPUT_XLSX = Path("05_reports") / "TITAN_RAG_CONTROL_TOWER.xlsx"
OUTPUT_SUMMARY_JSON = Path("05_reports") / "rag_control_tower_export_summary.json"
OUTPUT_SUMMARY_MD = Path("05_reports") / "rag_control_tower_export_summary.md"

AUDIT_LOG = Path("06_logs") / "rag_control_tower_export_audit.jsonl"
ERROR_LOG = Path("06_logs") / "rag_control_tower_export_errors.jsonl"

MAX_EXCEL_CELL_CHARS = 32000
MAX_ROWS_PER_SHEET = 50000

REPORT_JSON_FILES = {
    "night_run": Path("05_reports") / "night_run_summary.json",
    "morning_briefing": Path("05_reports") / "morning_briefing.json",
    "incremental": Path("05_reports") / "incremental_refresh_summary.json",
    "evidence_pack": Path("05_reports") / "latest_evidence_pack.json",
    "answer": Path("05_reports") / "latest_rag_answer.json",
    "audit_record": Path("05_reports") / "latest_audit_record.json",
    "evidence_report": Path("05_reports") / "evidence_pack_report.json",
    "citation": Path("05_reports") / "citation_verification_report.json",
    "conflict": Path("05_reports") / "conflict_report.json",
    "ssot": Path("05_reports") / "ssot_candidates_report.json",
    "financial": Path("05_reports") / "financial_signals_report.json",
    "risk": Path("05_reports") / "risk_signals_report.json",
    "priority": Path("05_reports") / "document_priority_rank.json",
    "embedding": Path("05_reports") / "embedding_summary.json",
    "chunking": Path("05_reports") / "chunking_summary.json",
    "extraction": Path("05_reports") / "extraction_summary.json",
    "safety": Path("05_reports") / "safety_filter_summary.json",
    "discovery": Path("05_reports") / "discovery_summary.json",
}

REPORT_CSV_FILES = {
    "document_priority_rank": Path("05_reports") / "document_priority_rank.csv",
    "risk_signals": Path("05_reports") / "risk_signals.csv",
    "financial_signals": Path("05_reports") / "financial_signals.csv",
    "ssot_candidates": Path("05_reports") / "ssot_candidates.csv",
    "conflict_report": Path("05_reports") / "conflict_report.csv",
    "citation_verification": Path("05_reports") / "citation_verification_report.csv",
    "evidence_pack_report": Path("05_reports") / "evidence_pack_report.csv",
    "evidence_sources": Path("05_reports") / "evidence_sources.csv",
    "morning_tasks": Path("05_reports") / "morning_briefing_tasks.csv",
    "night_run_steps": Path("05_reports") / "night_run_steps.csv",
    "incremental_file_index": Path("01_index") / "incremental_file_index.csv",
}

LOG_JSONL_FILES = {
    "night_run_log": Path("06_logs") / "night_run_orchestrator_audit.jsonl",
    "control_tower_log": Path("06_logs") / "rag_control_tower_export_audit.jsonl",
    "final_audit_log": Path("06_logs") / "final_audit_log.jsonl",
    "risk_log": Path("06_logs") / "risk_signal_extractor_audit.jsonl",
    "financial_log": Path("06_logs") / "financial_signal_extractor_audit.jsonl",
    "ssot_log": Path("06_logs") / "ssot_candidate_extractor_audit.jsonl",
    "conflict_log": Path("06_logs") / "conflict_detector_audit.jsonl",
    "citation_log": Path("06_logs") / "source_citation_verifier_audit.jsonl",
    "embedding_errors": Path("06_logs") / "embedding_errors.jsonl",
    "extraction_errors": Path("06_logs") / "extraction_errors.jsonl",
    "chunking_errors": Path("06_logs") / "chunking_errors.jsonl",
    "rag_search_errors": Path("06_logs") / "rag_search_errors.jsonl",
}

CORE_STATUS_FIELDS = [
    "source",
    "available",
    "status",
    "institutional_status",
    "audit_status",
    "verification_status",
    "report_status",
    "created_at",
    "updated_at",
    "total_records",
    "total_signals",
    "total_risks",
    "critical_count",
    "high_count",
    "total_conflicts",
    "total_candidates",
    "ranked_documents",
    "p0_count",
    "p1_count",
    "evidence_count",
    "source_file_count",
    "embedded_count",
    "error_count",
    "sha256",
    "path",
]


@dataclass(frozen=True)
class ExportPaths:
    base_dir: Path
    output_xlsx: Path
    summary_json: Path
    summary_md: Path
    audit_log: Path
    error_log: Path


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def normalize_inline(value: Any) -> str:
    return " ".join(str(value or "").replace("\x00", " ").split())


def safe_sheet_name(name: str) -> str:
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", "_", str(name))
    cleaned = cleaned.strip() or "Sheet"
    return cleaned[:31]


def safe_cell(value: Any) -> Any:
    if value is None:
        return None

    if isinstance(value, (int, float, bool)):
        return value

    if isinstance(value, (dict, list)):
        text = json.dumps(value, ensure_ascii=False)
    else:
        text = str(value)

    text = text.replace("\x00", "")
    if len(text) > MAX_EXCEL_CELL_CHARS:
        return text[:MAX_EXCEL_CELL_CHARS - 30] + " ...[TRUNCATED_FOR_EXCEL]"
    return text


def safe_float(value: Any, default: float = 0.0) -> float:
    try:
        if value is None or value == "":
            return default
        return float(value)
    except Exception:
        return default


def sha256_json(payload: dict[str, Any]) -> str:
    canonical = json.dumps(payload, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def sha256_file(path: Path) -> str | None:
    if not path.exists() or not path.is_file():
        return None
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def load_json_if_exists(path: Path) -> dict[str, Any] | None:
    if not path.exists():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        return payload if isinstance(payload, dict) else None
    except Exception:
        return None


def read_csv_rows(path: Path, max_rows: int = MAX_ROWS_PER_SHEET) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    rows: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8-sig", newline="", errors="ignore") as f:
        reader = csv.DictReader(f)
        for idx, row in enumerate(reader):
            if idx >= max_rows:
                break
            rows.append(dict(row))
    return rows


def read_jsonl_rows(path: Path, max_rows: int = 2000) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    rows: list[dict[str, Any]] = []
    lines = path.read_text(encoding="utf-8", errors="ignore").splitlines()
    for line in lines[-max_rows:]:
        raw = line.strip()
        if not raw:
            continue
        try:
            payload = json.loads(raw)
            if isinstance(payload, dict):
                rows.append(payload)
            else:
                rows.append({"raw": raw[:1000], "parse_status": "NON_OBJECT"})
        except Exception as exc:
            rows.append({"raw": raw[:1000], "parse_status": "INVALID_JSON", "error": str(exc)})
    return rows


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def append_jsonl(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(payload, ensure_ascii=False) + "\n")


def resolve_paths(base_dir: Path, output_arg: str | None) -> ExportPaths:
    return ExportPaths(
        base_dir=base_dir,
        output_xlsx=Path(output_arg) if output_arg else base_dir / OUTPUT_XLSX,
        summary_json=base_dir / OUTPUT_SUMMARY_JSON,
        summary_md=base_dir / OUTPUT_SUMMARY_MD,
        audit_log=base_dir / AUDIT_LOG,
        error_log=base_dir / ERROR_LOG,
    )


def import_openpyxl():
    try:
        import openpyxl
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        return openpyxl, Workbook, Font, PatternFill, Alignment, Border, Side, get_column_letter
    except ImportError as exc:
        raise ImportError("Missing dependency. Install with: pip install openpyxl") from exc


def flatten_for_status(source: str, path: Path, payload: dict[str, Any] | None) -> dict[str, Any]:
    if payload is None:
        return {
            "source": source,
            "available": False,
            "path": str(path),
            "sha256": None,
        }

    summary = payload.get("summary") if isinstance(payload.get("summary"), dict) else {}
    confidence = payload.get("confidence") if isinstance(payload.get("confidence"), dict) else {}

    row = {
        "source": source,
        "available": True,
        "path": str(path),
        "sha256": (
            payload.get("night_run_summary_sha256")
            or payload.get("briefing_sha256")
            or payload.get("audit_record_sha256")
            or payload.get("evidence_pack_sha256")
            or payload.get("answer_sha256")
            or payload.get("citation_verification_sha256")
            or payload.get("conflict_report_sha256")
            or payload.get("ssot_candidates_report_sha256")
            or payload.get("financial_signals_report_sha256")
            or payload.get("risk_signals_report_sha256")
            or payload.get("document_priority_rank_sha256")
            or payload.get("evidence_pack_report_sha256")
            or sha256_json(payload)
        ),
        "created_at": payload.get("created_at"),
        "updated_at": payload.get("updated_at"),
        "status": payload.get("status") or payload.get("final_status") or payload.get("overall_status"),
        "institutional_status": summary.get("institutional_status"),
        "audit_status": payload.get("audit_status"),
        "verification_status": payload.get("verification_status"),
        "report_status": summary.get("report_status"),
        "total_records": payload.get("total_records") or summary.get("total_records"),
        "total_signals": summary.get("total_signals"),
        "total_risks": summary.get("total_risks"),
        "critical_count": summary.get("critical_count"),
        "high_count": summary.get("high_count"),
        "total_conflicts": summary.get("total_conflicts"),
        "total_candidates": summary.get("total_candidates"),
        "ranked_documents": summary.get("ranked_documents"),
        "p0_count": summary.get("p0_count"),
        "p1_count": summary.get("p1_count"),
        "evidence_count": payload.get("evidence_count") or summary.get("evidence_count"),
        "source_file_count": payload.get("source_file_count") or summary.get("source_file_count"),
        "embedded_count": payload.get("embedded_count") or summary.get("embedded_count"),
        "error_count": payload.get("error_count") or summary.get("error_count"),
        "confidence_score": confidence.get("confidence_score"),
        "confidence_label": confidence.get("confidence_label"),
        "query": payload.get("query"),
    }

    return row


def load_all_reports(base_dir: Path) -> dict[str, dict[str, Any] | None]:
    return {
        name: load_json_if_exists(base_dir / rel_path)
        for name, rel_path in REPORT_JSON_FILES.items()
    }


def make_status_rows(base_dir: Path, reports: dict[str, dict[str, Any] | None]) -> list[dict[str, Any]]:
    rows = []
    for name, rel_path in REPORT_JSON_FILES.items():
        rows.append(flatten_for_status(name, base_dir / rel_path, reports.get(name)))
    return rows


def make_dashboard_rows(status_rows: list[dict[str, Any]], reports: dict[str, dict[str, Any] | None]) -> list[dict[str, Any]]:
    def find(source: str, field: str, default: Any = None) -> Any:
        for row in status_rows:
            if row.get("source") == source:
                return row.get(field, default)
        return default

    return [
        {"Metric": "Night Run Status", "Value": find("night_run", "status"), "Interpretation": "Critical if FAILED_CRITICAL_STEP."},
        {"Metric": "Morning Briefing Status", "Value": find("morning_briefing", "status"), "Interpretation": "GREEN/AMBER/RED/GRAY operational view."},
        {"Metric": "Latest Audit Status", "Value": find("audit_record", "audit_status"), "Interpretation": "AUDIT_PASS preferred for institutional use."},
        {"Metric": "Citation Verification", "Value": find("citation", "verification_status"), "Interpretation": "Citation verification must pass before lender/board use."},
        {"Metric": "Conflict Status", "Value": find("conflict", "institutional_status"), "Interpretation": "High conflict requires manual SSOT review."},
        {"Metric": "Risk Status", "Value": find("risk", "institutional_status"), "Interpretation": "Critical/high risks require review."},
        {"Metric": "Critical Risks", "Value": find("risk", "critical_count", 0), "Interpretation": "Blocks institutional use if > 0."},
        {"Metric": "High Risks", "Value": find("risk", "high_count", 0), "Interpretation": "Requires review if > 0."},
        {"Metric": "Financial Signals", "Value": find("financial", "total_signals", 0), "Interpretation": "Candidate financial input universe."},
        {"Metric": "SSOT Candidates", "Value": find("ssot", "total_candidates", 0), "Interpretation": "Candidates only; not final locked truth."},
        {"Metric": "P0 Documents", "Value": find("priority", "p0_count", 0), "Interpretation": "Immediate review queue."},
        {"Metric": "P1 Documents", "Value": find("priority", "p1_count", 0), "Interpretation": "High priority review queue."},
        {"Metric": "Embedded Chunks", "Value": find("embedding", "embedded_count", 0), "Interpretation": "Vector index population."},
        {"Metric": "Evidence Count", "Value": find("evidence_pack", "evidence_count", 0), "Interpretation": "Current retrieval evidence items."},
    ]


def rows_from_report_collection(report: dict[str, Any] | None, key: str) -> list[dict[str, Any]]:
    if not report:
        return []
    rows = report.get(key, [])
    if isinstance(rows, list):
        return [x for x in rows if isinstance(x, dict)]
    return []


def flatten_dict_row(row: dict[str, Any]) -> dict[str, Any]:
    out = {}
    for key, value in row.items():
        if isinstance(value, (dict, list)):
            out[key] = json.dumps(value, ensure_ascii=False)
        else:
            out[key] = value
    return out


def append_rows(ws: Any, rows: list[dict[str, Any]], header_fill: Any, header_font: Any, thin_border: Any, max_rows: int = MAX_ROWS_PER_SHEET) -> None:
    if not rows:
        ws.append(["No data"])
        return

    # Union headers, stable first-row order.
    headers: list[str] = []
    for row in rows:
        for key in row.keys():
            if key not in headers:
                headers.append(key)

    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    for row in rows[:max_rows]:
        ws.append([safe_cell(row.get(h)) for h in headers])

    # Basic filters/freeze.
    ws.freeze_panes = "A2"
    try:
        ws.auto_filter.ref = ws.dimensions
    except Exception:
        pass


def autosize(ws: Any, get_column_letter: Any, max_width: int = 55) -> None:
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = 8
        for cell in column_cells[:200]:
            value = cell.value
            if value is None:
                continue
            max_len = max(max_len, min(len(str(value)), max_width))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, max_width)


def style_sheet(ws: Any, Font: Any, PatternFill: Any, Alignment: Any, Border: Any, Side: Any, get_column_letter: Any) -> None:
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    autosize(ws, get_column_letter)


def create_workbook(paths: ExportPaths, reports: dict[str, dict[str, Any] | None], args: argparse.Namespace) -> dict[str, Any]:
    openpyxl, Workbook, Font, PatternFill, Alignment, Border, Side, get_column_letter = import_openpyxl()

    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    status_rows = make_status_rows(paths.base_dir, reports)
    dashboard_rows = make_dashboard_rows(status_rows, reports)

    sheet_map: list[tuple[str, list[dict[str, Any]]]] = [
        ("00_DASHBOARD", dashboard_rows),
        ("01_PIPELINE_STATUS", status_rows),
        ("02_NIGHT_STEPS", read_csv_rows(paths.base_dir / REPORT_CSV_FILES["night_run_steps"])),
        ("03_MORNING_TASKS", read_csv_rows(paths.base_dir / REPORT_CSV_FILES["morning_tasks"])),
        ("04_DOCUMENT_PRIORITY", read_csv_rows(paths.base_dir / REPORT_CSV_FILES["document_priority_rank"])),
        ("05_RISK_SIGNALS", read_csv_rows(paths.base_dir / REPORT_CSV_FILES["risk_signals"])),
        ("06_FINANCIAL_SIGNALS", read_csv_rows(paths.base_dir / REPORT_CSV_FILES["financial_signals"])),
        ("07_SSOT_CANDIDATES", read_csv_rows(paths.base_dir / REPORT_CSV_FILES["ssot_candidates"])),
        ("08_CONFLICTS", read_csv_rows(paths.base_dir / REPORT_CSV_FILES["conflict_report"])),
        ("09_CITATIONS", read_csv_rows(paths.base_dir / REPORT_CSV_FILES["citation_verification"])),
        ("10_EVIDENCE", read_csv_rows(paths.base_dir / REPORT_CSV_FILES["evidence_pack_report"])),
        ("11_EVIDENCE_SOURCES", read_csv_rows(paths.base_dir / REPORT_CSV_FILES["evidence_sources"])),
        ("12_INCREMENTAL", read_csv_rows(paths.base_dir / REPORT_CSV_FILES["incremental_file_index"])),
        ("13_LOG_TAIL", make_log_tail_rows(paths.base_dir, args.max_log_rows)),
    ]

    # Fallback rows from JSON reports when CSVs are not yet available.
    if not sheet_map[4][1]:
        sheet_map[4] = ("04_DOCUMENT_PRIORITY", [flatten_dict_row(x) for x in rows_from_report_collection(reports.get("priority"), "documents")])
    if not sheet_map[5][1]:
        sheet_map[5] = ("05_RISK_SIGNALS", [flatten_dict_row(x) for x in rows_from_report_collection(reports.get("risk"), "signals")])
    if not sheet_map[6][1]:
        sheet_map[6] = ("06_FINANCIAL_SIGNALS", [flatten_dict_row(x) for x in rows_from_report_collection(reports.get("financial"), "signals")])
    if not sheet_map[7][1]:
        sheet_map[7] = ("07_SSOT_CANDIDATES", [flatten_dict_row(x) for x in rows_from_report_collection(reports.get("ssot"), "candidates")])
    if not sheet_map[8][1]:
        sheet_map[8] = ("08_CONFLICTS", [flatten_dict_row(x) for x in rows_from_report_collection(reports.get("conflict"), "conflicts")])
    if not sheet_map[10][1]:
        sheet_map[10] = ("10_EVIDENCE", [flatten_dict_row(x) for x in rows_from_report_collection(reports.get("evidence_pack"), "evidence")])

    created_sheets = []

    for sheet_name, rows in sheet_map:
        ws = wb.create_sheet(safe_sheet_name(sheet_name))
        append_rows(ws, rows, header_fill, header_font, thin_border, max_rows=args.max_rows_per_sheet)
        style_sheet(ws, Font, PatternFill, Alignment, Border, Side, get_column_letter)
        created_sheets.append({
            "sheet_name": ws.title,
            "row_count": len(rows),
            "excel_rows_written": min(len(rows), args.max_rows_per_sheet),
        })

    # Manifest sheet last.
    manifest_rows = make_manifest_rows(paths.base_dir)
    ws = wb.create_sheet("99_MANIFEST")
    append_rows(ws, manifest_rows, header_fill, header_font, thin_border, max_rows=args.max_rows_per_sheet)
    style_sheet(ws, Font, PatternFill, Alignment, Border, Side, get_column_letter)
    created_sheets.append({
        "sheet_name": ws.title,
        "row_count": len(manifest_rows),
        "excel_rows_written": len(manifest_rows),
    })

    # Workbook properties.
    wb.properties.title = "TITAN RAG Control Tower"
    wb.properties.creator = "TITAN LOCAL EVIDENCE RAG ENGINE"
    wb.properties.subject = "Audit-grade RAG pipeline control workbook"
    wb.properties.keywords = "TITAN,RAG,SSOT,Audit,Evidence,Control Tower"
    wb.properties.created = datetime.now(timezone.utc).replace(tzinfo=None)

    paths.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(paths.output_xlsx)

    return {
        "workbook_path": str(paths.output_xlsx),
        "workbook_sha256": sha256_file(paths.output_xlsx),
        "sheets": created_sheets,
    }


def make_log_tail_rows(base_dir: Path, max_rows: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for name, rel_path in LOG_JSONL_FILES.items():
        path = base_dir / rel_path
        for item in read_jsonl_rows(path, max_rows=max_rows):
            row = flatten_dict_row(item)
            row["log_source"] = name
            row["log_path"] = str(path)
            rows.append(row)
    return rows[-max_rows:]


def make_manifest_rows(base_dir: Path) -> list[dict[str, Any]]:
    rows = []

    for name, rel_path in {**REPORT_JSON_FILES, **REPORT_CSV_FILES, **LOG_JSONL_FILES}.items():
        path = base_dir / rel_path
        rows.append({
            "name": name,
            "path": str(path),
            "exists": path.exists(),
            "file_size_bytes": path.stat().st_size if path.exists() and path.is_file() else None,
            "modified_time_utc": datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc).isoformat(timespec="seconds") if path.exists() and path.is_file() else None,
            "sha256": sha256_file(path),
            "category": "json" if rel_path.suffix.lower() == ".json" else "csv" if rel_path.suffix.lower() == ".csv" else "jsonl",
        })

    return rows


def build_summary(paths: ExportPaths, workbook_info: dict[str, Any], reports: dict[str, dict[str, Any] | None], args: argparse.Namespace) -> dict[str, Any]:
    status_rows = make_status_rows(paths.base_dir, reports)

    available_reports = sum(1 for row in status_rows if row.get("available"))
    missing_reports = [row.get("source") for row in status_rows if not row.get("available")]

    dashboard = make_dashboard_rows(status_rows, reports)

    summary = {
        "created_at": utc_now_iso(),
        "script": SCRIPT_NAME,
        "script_version": SCRIPT_VERSION,
        "policy": "RAG_CONTROL_TOWER_EXPORT_AUDIT_LOCKED",
        "base_dir": str(paths.base_dir),
        "output_xlsx": str(paths.output_xlsx),
        "output_xlsx_sha256": workbook_info.get("workbook_sha256"),
        "available_report_count": available_reports,
        "missing_reports": missing_reports,
        "sheet_count": len(workbook_info.get("sheets", [])),
        "sheets": workbook_info.get("sheets", []),
        "dashboard": dashboard,
        "max_rows_per_sheet": args.max_rows_per_sheet,
        "max_log_rows": args.max_log_rows,
        "governance_rule": {
            "export_consolidates_existing_reports_only": True,
            "source_documents_not_modified": True,
            "control_tower_is_review_interface_not_truth_source": True,
            "json_reports_remain_authoritative": True,
            "audit_precedes_decision": True,
        },
    }

    summary["rag_control_tower_export_sha256"] = sha256_json(summary)
    return summary


def md_escape(value: Any) -> str:
    return normalize_inline(value).replace("|", "\\|")


def dashboard_md(rows: list[dict[str, Any]]) -> str:
    lines = [
        "| Metric | Value | Interpretation |",
        "|---|---|---|",
    ]
    for row in rows:
        lines.append(
            f"| {md_escape(row.get('Metric'))} | {md_escape(row.get('Value'))} | {md_escape(row.get('Interpretation'))} |"
        )
    return "\n".join(lines)


def sheets_md(rows: list[dict[str, Any]]) -> str:
    if not rows:
        return "No sheets."
    lines = [
        "| Sheet | Source Rows | Excel Rows Written |",
        "|---|---:|---:|",
    ]
    for row in rows:
        lines.append(
            f"| {md_escape(row.get('sheet_name'))} | {row.get('row_count')} | {row.get('excel_rows_written')} |"
        )
    return "\n".join(lines)


def summary_to_markdown(summary: dict[str, Any]) -> str:
    return f"""# TITAN RAG Control Tower Export Summary

## 1. Export Identity

| Field | Value |
|---|---|
| Created At | {summary.get("created_at")} |
| Workbook | `{summary.get("output_xlsx")}` |
| Workbook SHA-256 | `{summary.get("output_xlsx_sha256")}` |
| Available Report Count | {summary.get("available_report_count")} |
| Missing Reports | {len(summary.get("missing_reports", []))} |
| Sheet Count | {summary.get("sheet_count")} |
| Export SHA-256 | `{summary.get("rag_control_tower_export_sha256")}` |

---

## 2. Dashboard Metrics

{dashboard_md(summary.get("dashboard", []))}

---

## 3. Sheets

{sheets_md(summary.get("sheets", []))}

---

## 4. Missing Reports

```json
{json.dumps(summary.get("missing_reports", []), indent=2, ensure_ascii=False)}
```

---

## 5. Governance Rule

```text
Control Tower consolidates existing reports only.
It does not modify source documents.
It is a review interface, not the authoritative truth source.
JSON reports and original source files remain authoritative.
Audit precedes decision.
```
"""


def print_summary(summary: dict[str, Any], paths: ExportPaths) -> None:
    print("=" * 100)
    print("TITAN RAG CONTROL TOWER EXPORT v2.0 AUDIT LOCKED")
    print("=" * 100)
    print(f"Workbook:                 {summary.get('output_xlsx')}")
    print(f"Workbook SHA-256:          {summary.get('output_xlsx_sha256')}")
    print(f"Available reports:         {summary.get('available_report_count')}")
    print(f"Missing reports:           {len(summary.get('missing_reports', []))}")
    print(f"Sheet count:               {summary.get('sheet_count')}")
    print("-" * 100)
    print(f"Summary JSON:              {paths.summary_json}")
    print(f"Summary Markdown:          {paths.summary_md}")
    print(f"Audit log:                 {paths.audit_log}")
    print("=" * 100)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="TITAN RAG Step 20 — Control Tower export.")

    parser.add_argument("--base-dir", default=str(DEFAULT_BASE_DIR), help="Base TITAN_FULL_RAG directory.")
    parser.add_argument("--output", default=None, help="Optional XLSX output path.")
    parser.add_argument("--max-rows-per-sheet", type=int, default=MAX_ROWS_PER_SHEET, help="Maximum rows per Excel sheet.")
    parser.add_argument("--max-log-rows", type=int, default=2000, help="Maximum log tail rows included.")
    parser.add_argument("--print", action="store_true", help="Print Markdown summary to console.")

    return parser.parse_args()


def main() -> None:
    args = parse_args()
    base_dir = Path(args.base_dir)
    paths = resolve_paths(base_dir, args.output)

    try:
        if not base_dir.exists():
            raise FileNotFoundError(f"Base directory does not exist: {base_dir}")

        reports = load_all_reports(base_dir)

        workbook_info = create_workbook(paths, reports, args)
        summary = build_summary(paths, workbook_info, reports, args)
        markdown = summary_to_markdown(summary)

        write_json(paths.summary_json, summary)
        write_text(paths.summary_md, markdown)

        append_jsonl(paths.audit_log, {
            "timestamp": utc_now_iso(),
            "event": "RAG_CONTROL_TOWER_EXPORT_COMPLETED",
            "script": SCRIPT_NAME,
            "script_version": SCRIPT_VERSION,
            "output_xlsx": str(paths.output_xlsx),
            "workbook_sha256": summary.get("output_xlsx_sha256"),
            "summary_sha256": summary.get("rag_control_tower_export_sha256"),
            "sheet_count": summary.get("sheet_count"),
            "available_report_count": summary.get("available_report_count"),
            "missing_reports": summary.get("missing_reports"),
        })

        print_summary(summary, paths)

        if args.print:
            print("\n" + markdown)

    except Exception as exc:
        append_jsonl(paths.error_log, {
            "timestamp": utc_now_iso(),
            "event": "RAG_CONTROL_TOWER_EXPORT_FAILED",
            "script": SCRIPT_NAME,
            "script_version": SCRIPT_VERSION,
            "base_dir": str(base_dir),
            "output_xlsx": str(paths.output_xlsx),
            "error": str(exc),
            "traceback": traceback.format_exc(),
        })

        print("=" * 100)
        print("TITAN RAG CONTROL TOWER EXPORT FAILED")
        print("=" * 100)
        print(f"Error: {exc}")
        print(f"Error log: {paths.error_log}")
        print("=" * 100)
        raise


if __name__ == "__main__":
    main()
