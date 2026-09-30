# ============================================================
# TITAN_KERNEL: 18_filefinding_evidence_weight_calculator.py
# PURPOSE: Calculate evidence/signal weight safely in FileFinding Register
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_BLOCK = 1
EXIT_FILE_ERROR = 2
EXIT_SCHEMA_ERROR = 3

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

SHEET = "FileFinding Register"

REQUIRED_COLUMNS = [
    "FileFinding_ID",
    "Evidence_ID",
    "Signal_ID",
    "File_Path",
    "SHA256",
    "Weight",
    "Status",
    "Reviewer",
    "Final_Use_Allowed"
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def norm(value):
    return "" if value is None else str(value).strip()

def headers(ws):
    return {norm(c.value): i for i, c in enumerate(ws[1], start=1) if norm(c.value)}

def emit_jsonl(path, event):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "a", encoding="utf-8") as f:
        f.write(json.dumps(event, ensure_ascii=False) + "\n")

def make_event(status, check_id, message, severity="INFO", extra=None):
    event = {
        "timestamp": now_iso(),
        "component": "FILEFINDING_EVIDENCE_WEIGHT_CALCULATOR",
        "version": "1.0",
        "status": status,
        "check_id": check_id,
        "severity": severity,
        "message": message,
        **CANON,
    }
    if extra:
        event.update(extra)
    return event

def calculate_weight(row):
    fid = row.get("FileFinding_ID", "")
    evid = row.get("Evidence_ID", "")
    signal = row.get("Signal_ID", "")
    file_path = row.get("File_Path", "")
    sha = row.get("SHA256", "")
    reviewer = row.get("Reviewer", "")
    status = row.get("Status", "").upper()

    reasons = []

    if not fid:
        reasons.append("Missing FileFinding_ID")
    if not evid:
        reasons.append("Missing Evidence_ID")
    if not signal:
        reasons.append("Missing Signal_ID")
    if not file_path:
        reasons.append("Missing File_Path")
    if not sha:
        reasons.append("Missing SHA256")
    if not reviewer:
        reasons.append("Missing Reviewer")
    if status not in {"REVIEWED", "VALIDATED", "APPROVED_REFERENCE"}:
        reasons.append("Status not reviewer-validated")

    if reasons:
        return 0, "REVIEW_REQUIRED", reasons

    return 1, "REVIEW_REQUIRED_EVIDENCE_SIGNAL", ["Mapped and review-ready, but not approved for final use"]

def main():
    parser = argparse.ArgumentParser(description="Calculate safe FileFinding evidence weights")
    parser.add_argument("--excel", required=True)
    parser.add_argument("--report", required=True)
    parser.add_argument("--dry-run", action="store_true", help="Do not modify workbook")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        e = make_event("BLOCK", "WEIGHT-001", "Excel file missing", "CRITICAL", {"excel": args.excel})
        emit_jsonl(args.report, e)
        print(json.dumps(e, ensure_ascii=False) if args.json else e["message"])
        return EXIT_FILE_ERROR

    try:
        wb = openpyxl.load_workbook(args.excel)
    except Exception as exc:
        e = make_event("BLOCK", "WEIGHT-002", f"Excel unreadable: {exc}", "CRITICAL")
        emit_jsonl(args.report, e)
        print(json.dumps(e, ensure_ascii=False) if args.json else e["message"])
        return EXIT_FILE_ERROR

    if SHEET not in wb.sheetnames:
        e = make_event("BLOCK", "WEIGHT-003", "FileFinding Register sheet missing", "CRITICAL")
        emit_jsonl(args.report, e)
        print(json.dumps(e, ensure_ascii=False) if args.json else e["message"])
        return EXIT_SCHEMA_ERROR

    ws = wb[SHEET]
    h = headers(ws)
    missing = [c for c in REQUIRED_COLUMNS if c not in h]

    if missing:
        e = make_event("BLOCK", "WEIGHT-004", "Required columns missing", "CRITICAL", {"missing_columns": missing})
        emit_jsonl(args.report, e)
        print(json.dumps(e, ensure_ascii=False) if args.json else e["message"])
        return EXIT_SCHEMA_ERROR

    checked = 0
    zero_weight = 0
    mapped_weight = 0
    row_results = []

    for r in range(2, ws.max_row + 1):
        row = {k: norm(ws.cell(r, h[k]).value) for k in h.keys()}
        if not any(row.values()):
            continue

        checked += 1
        weight, treatment, reasons = calculate_weight(row)

        if weight == 0:
            zero_weight += 1
        else:
            mapped_weight += 1

        if not args.dry_run:
            ws.cell(r, h["Weight"]).value = weight
            ws.cell(r, h["Status"]).value = treatment
            ws.cell(r, h["Final_Use_Allowed"]).value = "NO"

        row_results.append({
            "row": r,
            "FileFinding_ID": row.get("FileFinding_ID", ""),
            "Evidence_ID": row.get("Evidence_ID", ""),
            "Signal_ID": row.get("Signal_ID", ""),
            "Weight": weight,
            "Status": treatment,
            "Reasons": reasons,
            "Final_Use_Allowed": "NO"
        })

    if not args.dry_run:
        wb.save(args.excel)

    events = [
        make_event(
            "REVIEW_REQUIRED",
            "WEIGHT-005",
            "FileFinding evidence weights calculated",
            "Weights were calculated using non-approval safety rules",
            "HIGH",
            {
                "checked_rows": checked,
                "zero_weight_rows": zero_weight,
                "mapped_weight_rows": mapped_weight,
                "dry_run": args.dry_run,
                "row_results": row_results
            }
        ),
        make_event(
            "BLOCK",
            "SUMMARY",
            "Weight calculation complete but final use remains blocked",
            "SYSTEM RED — STEP102 LOCKED — NO FINAL USE",
            "CRITICAL",
            {"FINAL_USE_ALLOWED": "NO", "EVIDENCE_APPROVAL_ALLOWED": "NO"}
        )
    ]

    for e in events:
        emit_jsonl(args.report, e)

    if args.json:
        print(json.dumps({"events": events}, ensure_ascii=False, indent=2))
    else:
        print(f"✅ Checked rows: {checked}")
        print(f"Weight 0 rows: {zero_weight}")
        print(f"Mapped weight rows: {mapped_weight}")
        print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")

    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
