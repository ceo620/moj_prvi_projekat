# ============================================================
# TITAN_KERNEL: 78_dissent_objection_register.py
# PURPOSE: Create dissent/objection register for reviewers
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

HEADERS = [
    "Objection_ID",
    "Reviewer",
    "Date",
    "Objected_Item",
    "Objection_Type",
    "Description",
    "Severity",
    "Linked_Evidence_ID",
    "Linked_FileFinding_ID",
    "Requested_Action",
    "Resolution_Status",
    "Can_Approve_Evidence",
    "Can_Close_Gate",
    "Can_Unlock_STEP102",
    "Final_Use_Allowed",
    "Notes"
]

DEFAULT_ROWS = [
    {
        "Objection_ID": "OBJ-00001",
        "Reviewer": "",
        "Date": "",
        "Objected_Item": "NO_RELEASE_CLOSEOUT",
        "Objection_Type": "REVIEW_REQUIRED",
        "Description": "",
        "Severity": "HIGH",
        "Linked_Evidence_ID": "",
        "Linked_FileFinding_ID": "",
        "Requested_Action": "DOCUMENT_ONLY",
        "Resolution_Status": "OPEN",
        "Can_Approve_Evidence": "NO",
        "Can_Close_Gate": "NO",
        "Can_Unlock_STEP102": "NO",
        "Final_Use_Allowed": "NO",
        "Notes": "Template row only; does not approve or close anything."
    }
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def main():
    parser = argparse.ArgumentParser(description="Build TITAN dissent/objection register")
    parser.add_argument("--out-xlsx", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    payload = {
        "timestamp": now_iso(),
        "component": "DISSENT_OBJECTION_REGISTER",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "records": DEFAULT_ROWS,
        "decision": "Objection register captures dissent only. It does not authorize changes.",
        **CANON
    }

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dissent Objections"
        ws.append(HEADERS)
        for row in DEFAULT_ROWS:
            ws.append([row.get(h, "") for h in HEADERS])

        ws2 = wb.create_sheet("Canonical Locks")
        ws2.append(["Control", "Value"])
        for k, v in CANON.items():
            ws2.append([k, v])

        ws3 = wb.create_sheet("Instructions")
        ws3.append(["Instruction_ID", "Instruction"])
        ws3.append(["INS-001", "Objections do not approve evidence."])
        ws3.append(["INS-002", "Objections do not close gates."])
        ws3.append(["INS-003", "Objections do not unlock STEP102."])
        ws3.append(["INS-004", "SYSTEM RED remains active."])

        Path(args.out_xlsx).parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.out_xlsx)

        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else f"✅ Objection register: {args.out_xlsx}")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
