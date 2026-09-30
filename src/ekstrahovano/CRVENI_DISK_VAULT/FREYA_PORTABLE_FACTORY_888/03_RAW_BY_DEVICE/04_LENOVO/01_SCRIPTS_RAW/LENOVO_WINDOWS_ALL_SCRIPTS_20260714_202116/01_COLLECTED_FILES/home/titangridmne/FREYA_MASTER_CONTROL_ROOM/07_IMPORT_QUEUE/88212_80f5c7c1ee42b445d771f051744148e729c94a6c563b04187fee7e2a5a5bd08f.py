# ============================================================
# TITAN_KERNEL: 67_policy_exception_verifier.py
# PURPOSE: Verify policy exception register cannot override canon
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_BLOCK = 1
EXIT_FILE_ERROR = 2
EXIT_SCHEMA_ERROR = 3

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

SHEET = "Policy Exceptions"
DANGEROUS_YES_COLUMNS = [
    "Approved",
    "Can_Change_SYSTEM_RED",
    "Can_Unlock_STEP102",
    "Can_Close_Gate",
    "Can_Approve_Evidence",
    "Can_Write_SSoT",
    "Final_Use_Allowed"
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def norm(value):
    return "" if value is None else str(value).strip()

def headers(ws):
    return {norm(c.value): i for i, c in enumerate(ws[1], start=1) if norm(c.value)}

def main():
    parser = argparse.ArgumentParser(description="Verify TITAN policy exceptions")
    parser.add_argument("--exceptions-xlsx", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    if not os.path.exists(args.exceptions_xlsx):
        result = {
            "timestamp": now_iso(),
            "component": "POLICY_EXCEPTION_VERIFIER",
            "status": "BLOCK",
            "message": "Policy exception workbook missing",
            **CANON
        }
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else result["message"])
        return EXIT_FILE_ERROR

    wb = openpyxl.load_workbook(args.exceptions_xlsx, data_only=True)
    if SHEET not in wb.sheetnames:
        result = {
            "timestamp": now_iso(),
            "component": "POLICY_EXCEPTION_VERIFIER",
            "status": "BLOCK",
            "message": f"Missing sheet: {SHEET}",
            **CANON
        }
        Path(args.out_json).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        return EXIT_SCHEMA_ERROR

    ws = wb[SHEET]
    h = headers(ws)
    missing = [c for c in DANGEROUS_YES_COLUMNS if c not in h]
    violations = []

    for r in range(2, ws.max_row + 1):
        empty = True
        for c in range(1, ws.max_column + 1):
            if norm(ws.cell(r, c).value):
                empty = False
                break
        if empty:
            continue

        for col in DANGEROUS_YES_COLUMNS:
            if col in h and norm(ws.cell(r, h[col]).value).upper() == "YES":
                violations.append({
                    "row": r,
                    "column": col,
                    "value": "YES",
                    "violation": "Policy exception attempted forbidden approval/override"
                })

    blocked = bool(missing or violations)

    result = {
        "timestamp": now_iso(),
        "component": "POLICY_EXCEPTION_VERIFIER",
        "version": "1.0",
        "status": "BLOCK" if blocked else "REVIEW_REQUIRED",
        "missing_columns": missing,
        "violations": violations,
        "decision": "Any attempted exception approval is blocked. Canon remains unchanged.",
        **CANON
    }

    Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
    Path(args.out_json).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else ("⛔ Policy exception violation" if blocked else "✅ Policy exceptions verified as non-overriding"))
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_BLOCK if blocked else EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
