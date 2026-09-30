import os
import json
from datetime import datetime
try:
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.shared import OxmlElement, qn
except ImportError:
    print("[-] GREŠKA: Biblioteka python-docx nije instalirana. Pokreni: pip install python-docx")
    exit()

class Tier1DocxGenerator:
    def __init__(self):
        # Koristimo sigurnu Termux putanju
        self.titan_root = os.path.expanduser("~/TITAN_KERNEL")
        self.reports_dir = os.path.join(self.titan_root, "REPORTS")
        self.manifest_path = os.path.join(self.reports_dir, "vdr_builder_manifest_v9_4.json")
        self.docx_output_path = os.path.join(self.reports_dir, "TITAN_1_Institutional_VDR_Index.docx")

    def set_cell_background(self, cell, fill_color):
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), fill_color)
        shading.set(qn('w:val'), 'clear')
        tc_pr.append(shading)

    def generate_document(self):
        if not os.path.exists(self.manifest_path):
            print("[-] GREŠKA: VDR Manifest nije pronađen.")
            return

        with open(self.manifest_path, "r", encoding="utf-8") as f:
            vdr_data = json.load(f)

        doc = Document()
        
        for section in doc.sections:
            section.top_margin = Cm(2.54)
            section.bottom_margin = Cm(2.54)
            section.left_margin = Cm(2.54)
            section.right_margin = Cm(2.54)

        style = doc.styles['Normal']
        font = style.font
        font.name = 'Arial'
        font.size = Pt(10)
        font.color.rgb = RGBColor(33, 33, 33)

        doc.add_paragraph("\n\n\n\n\n")
        title = doc.add_heading('TITAN 1 INTEGRATED INDUSTRIAL ECOSYSTEM', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.runs[0].font.color.rgb = RGBColor(0, 32, 96)
        title.runs[0].font.size = Pt(18)
        
        subtitle = doc.add_paragraph('Virtual Data Room (VDR) – Master Audit Index')
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].bold = True
        subtitle.runs[0].font.size = Pt(12)
        
        doc.add_paragraph("\n\n")
        info_para = doc.add_paragraph()
        info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        info_para.add_run("Date: ").bold = True
        info_para.add_run(f"{datetime.now().strftime('%d %B %Y')}\n")
        info_para.add_run("Baseline: ").bold = True
        info_para.add_run(f"Total Project Cost €27.8M (Audited)\n")
        info_para.add_run("Classification: ").bold = True
        info_para.add_run("Strictly Confidential\n")
        info_para.add_run("Target Lenders: ").bold = True
        info_para.add_run("EIB, EBRD, IFC")
        
        doc.add_page_break()

        doc.add_heading('1. Executive Statement & Scope', level=1).runs[0].font.color.rgb = RGBColor(0, 32, 96)
        exec_text = (
            "This document serves as the canonical Virtual Data Room (VDR) index for the TITAN 1 Project, "
            "structured explicitly to align with the stringent due diligence frameworks of Tier-1 international "
            "financial institutions. The documentation architecture maps the project's strategic allocation "
            "within the KAP industrial zone, alongside preliminary utility approvals (12.5 MW grid allocation) "
            "and ESG compliance trajectories. The financial baseline is rigorously anchored at a Total Project "
            "Cost of €27.8M, optimizing the debt-to-equity leverage matrix for institutional evaluation."
        )
        para = doc.add_paragraph(exec_text)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        doc.add_heading('2. Master Document Index', level=1).runs[0].font.color.rgb = RGBColor(0, 32, 96)
        
        for stub_name, documents in vdr_data['vdr_index'].items():
            if not documents: continue 
                
            stub_heading = doc.add_heading(stub_name.replace("_", " "), level=2)
            stub_heading.runs[0].font.size = Pt(11)
            stub_heading.runs[0].font.color.rgb = RGBColor(80, 80, 80)
            
            table = doc.add_table(rows=1, cols=4)
            table.autofit = True
            
            hdr_cells = table.rows[0].cells
            headers = ['Document Identification', 'Maturity Level', 'Lender Readiness Status', 'Flag']
            for i, text in enumerate(headers):
                hdr_cells[i].text = text
                hdr_cells[i].paragraphs[0].runs[0].bold = True
                hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                self.set_cell_background(hdr_cells[i], "002060")

            for d in documents:
                row_cells = table.add_row().cells
                row_cells[0].text = str(d.get('document_name', 'N/A'))
                row_cells[0].paragraphs[0].runs[0].font.size = Pt(9)
                
                row_cells[1].text = str(d.get('evidence_maturity_level', 'N/A')).replace("_", " ")
                row_cells[1].paragraphs[0].runs[0].font.size = Pt(9)
                
                status = str(d.get('lender_readiness_status', 'PENDING')).replace("_", " ")
                row_cells[2].text = status
                row_cells[2].paragraphs[0].runs[0].font.size = Pt(9)
                
                flag = str(d.get('red_flag', 'YES'))
                row_cells[3].text = flag
                row_cells[3].paragraphs[0].runs[0].font.size = Pt(9)
                row_cells[3].paragraphs[0].runs[0].bold = True
                if flag == "YES":
                    row_cells[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(192, 0, 0)
                else:
                    row_cells[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 128, 0)

            doc.add_paragraph("\n")

        doc.save(self.docx_output_path)
        print(f"[✓] TIER-1 DOCX GENERISAN U LOKALNOJ MEMORIJI!")

if __name__ == "__main__":
    generator = Tier1DocxGenerator()
    generator.generate_document()
