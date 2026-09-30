import json, os
from datetime import datetime
from docx import Document
from docx.shared import Pt

with open("00_MASTER_TOKEN_AUTHORITY/MASTER_TOKEN_REGISTER_v4.0.json") as f:
    master = json.load(f)

print("🔒 CTOS v4.0 — Controlled Generation (ANEKS 1)")

doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

doc.add_heading('ANEKS 1', 0).alignment = 1
doc.add_heading('REVIZIJA CAPEX OKVIRA ZA PROJEKAT TITAN 1', 1).alignment = 1
doc.add_paragraph(f'Datum: 2026-05-10\n')

text = f"""Ovaj Aneks 1 odnosi se isključivo na projekat TITAN 1 – Fabrika transformatorskih kazana (tankova).

Nosilac projekta:
{master['canonical_identity']['company']}

Lokacija projekta:
{master['canonical_identity']['location']}, Crna Gora

Ažurirani CAPEX okvir za TITAN 1:
{master['canonical_capex']['label']}

Ovaj Aneks predstavlja prateće objašnjenje uz postojeću projektnu dokumentaciju i koristi se isključivo za potrebe institucionalne komunikacije."""

for paragraph in text.split('\n'):
    if paragraph.strip():
        doc.add_paragraph(paragraph)

doc.add_paragraph('\nSa poštovanjem,')
doc.add_paragraph('DANIJELA KESKIN')
doc.add_paragraph('Director')
doc.add_paragraph('Ars Metal Industries D.O.O. Podgorica')

filename = "ANEKS_1_REVIZIJA_CAPEX_TITAN1_v4_2026-05-10.docx"
doc.save(filename)

print(f"\n✅ ANEKS 1 v4 generisan: {filename}")
print("   Generisan iz MASTER_TOKEN_REGISTER_v4.0")
