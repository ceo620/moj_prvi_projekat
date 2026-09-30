# ============================================================
# TITAN_KERNEL: 17_evidence_link_validator.py
# PURPOSE: Validate Evidence Gap linked FileFinding IDs exist
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_BLOCK = 1
EXIT_FILE_ERROR = 2
EXIT_SCHEMA_ERROR = 3

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

FILEFINDING_SHEET = "FileFinding Register"
GAP_SHEET = "Evidence Gap Register"

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def norm(value):
    return "" if value is None else str(value).strip()

def emit_jsonl(path, event):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "a", encoding="utf-8") as f:
        f.write(json.dumps(event, ensure_ascii=False) + "\n")

def make_event(status, check_id, message, severity="INFO", extra=None):
    event = {
        "timestamp": now_iso(),
        "component": "EVIDENCE_LINK_VALIDATOR",
        "version": "1.0",
        "status": status,
        "check_id": check_id,
        "severity": severity,
        "message": message,
        **CANON,
    }
    if extra:
        event.update(extra)
    return event

def headers(ws):
    return {norm(c.value): i for i, c in enumerate(ws[1], start=1) if norm(c.value)}

def main():
    parser = argparse.ArgumentParser(description="Validate FileFinding links in Evidence Gap Register")
    parser.add_argument("--excel", required=True)
    parser.add_argument("--report", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        e = make_event("BLOCK", "LINK-001", "Excel file missing", "CRITICAL", {"excel": args.excel})
        emit_jsonl(args.report, e)
        print(json.dumps(e, ensure_ascii=False) if args.json else e["message"])
        return EXIT_FILE_ERROR

    try:
        wb = openpyxl.load_workbook(args.excel, data_only=True)
    except Exception as exc:
        e = make_event("BLOCK", "LINK-002", f"Excel unreadable: {exc}", "CRITICAL")
        emit_jsonl(args.report, e)
        print(json.dumps(e, ensure_ascii=False) if args.json else e["message"])
        return EXIT_FILE_ERROR

    missing_sheets = [s for s in [FILEFINDING_SHEET, GAP_SHEET] if s not in wb.sheetnames]
    if missing_sheets:
        e = make_event("BLOCK", "LINK-003", "Required sheet missing", "CRITICAL", {"missing_sheets": missing_sheets})
        emit_jsonl(args.report, e)
        print(json.dumps(e, ensure_ascii=False) if args.json else e["message"])
        return EXIT_SCHEMA_ERROR

    ff_ws = wb[FILEFINDING_SHEET]
    gap_ws = wb[GAP_SHEET]
    ff_h = headers(ff_ws)
    gap_h = headers(gap_ws)

    required_ff = ["FileFinding_ID"]
    required_gap = ["Gap_ID", "Linked_FileFinding_ID", "Status", "Blocking_STEP"]
    missing_cols = [c for c in required_ff if c not in ff_h] + [c for c in required_gap if c not in gap_h]

    if missing_cols:
        e = make_event("BLOCK", "LINK-004", "Required columns missing", "CRITICAL", {"missing_columns": missing_cols})
        emit_jsonl(args.report, e)
        print(json.dumps(e, ensure_ascii=False) if args.json else e["message"])
        return EXIT_SCHEMA_ERROR

    filefinding_ids = set()
    for r in range(2, ff_ws.max_row + 1):
        fid = norm(ff_ws.cell(r, ff_h["FileFinding_ID"]).value)
        if fid:
            filefinding_ids.add(fid)

    violations = []
    checked_gaps = 0
    open_gaps = 0
    linked_gaps = 0

    closed = {"CLOSED", "APPROVED", "RESOLVED", "VALIDATED"}

    for r in range(2, gap_ws.max_row + 1):
        gap_id = norm(gap_ws.cell(r, gap_h["Gap_ID"]).value)
        linked = norm(gap_ws.cell(r, gap_h["Linked_FileFinding_ID"]).value)
        status = norm(gap_ws.cell(r, gap_h["Status"]).value).upper()
        blocking_step = norm(gap_ws.cell(r, gap_h["Blocking_STEP"]).value)

        if not gap_id and not linked and not status:
            continue

        checked_gaps += 1

        if status not in closed:
            open_gaps += 1

        if linked:
            linked_gaps += 1
            ids = [x.strip() for x in linked.replace(";", ",").split(",") if x.strip()]
            for fid in ids:
                if fid not in filefinding_ids:
                    violations.append({
                        "row": r,
                        "Gap_ID": gap_id,
                        "Linked_FileFinding_ID": fid,
                        "violation": "Linked_FileFinding_ID does not exist in FileFinding Register"
                    })

    blocked = bool(violations)

    events = [
        make_event(
            "PASS" if not violations else "BLOCK",
            "LINK-005",
            "Evidence gap FileFinding links validated",
            "Linked_FileFinding_ID values checked against FileFinding Register",
            "CRITICAL",
            {
                "filefinding_ids_found": len(filefinding_ids),
                "checked_gaps": checked_gaps,
                "linked_gaps": linked_gaps,
                "open_gaps": open_gaps,
                "violations": violations
            }
        ),
        make_event(
            "BLOCK" if open_gaps > 0 else "REVIEW_REQUIRED",
            "LINK-006",
            "Open evidence gaps keep STEP102 locked",
            "STEP102 remains LOCKED / NOT ACCEPTED while open gaps exist",
            "CRITICAL",
            {"open_gaps": open_gaps, "STEP102": "LOCKED / NOT ACCEPTED", "FINAL_USE_ALLOWED": "NO"}
        ),
        make_event(
            "BLOCK" if blocked or open_gaps > 0 else "REVIEW_REQUIRED",
            "SUMMARY",
            "Evidence link validation complete",
            "SYSTEM RED — STEP102 LOCKED — NO FINAL USE",
            "CRITICAL",
            {"blocked": blocked or open_gaps > 0}
        )
    ]

    for e in events:
        emit_jsonl(args.report, e)

    if args.json:
        print(json.dumps({"events": events}, ensure_ascii=False, indent=2))
    else:
        for e in events:
            print(f"{e['status']} | {e['check_id']} | {e['message']}")
        print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")

    return EXIT_BLOCK if blocked or open_gaps > 0 else EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
