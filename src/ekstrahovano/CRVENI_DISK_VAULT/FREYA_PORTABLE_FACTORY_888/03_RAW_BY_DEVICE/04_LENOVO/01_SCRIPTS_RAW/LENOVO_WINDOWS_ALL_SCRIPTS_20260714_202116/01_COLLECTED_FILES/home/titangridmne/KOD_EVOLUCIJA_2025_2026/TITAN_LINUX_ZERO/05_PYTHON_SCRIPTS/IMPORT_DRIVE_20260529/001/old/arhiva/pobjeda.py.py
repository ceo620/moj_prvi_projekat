import os
from docx import Document

# Kompjuter sam pronalazi putanju do tvog Desktopa
desktop = os.path.join(os.path.expanduser("~"), "Desktop")
folder = os.path.join(desktop, "TITAN_309_FINAL")

if not os.path.exists(folder):
    os.makedirs(folder)

print("🚀 TITAN MASINA KREĆE... Štampam 309 dokumenata direktno na Desktop...")

for i in range(1, 310):
    doc = Document()
    doc.add_heading(f'TITAN GRID - STRATEŠKI SCENARIO {i:03d}', 0)
    doc.add_paragraph(f"Scenario ID: TG-2026-{i:03d}")
    doc.add_paragraph("Analiza investicije i operativne efikasnosti - Tuzi, Montenegro.")
    doc.save(os.path.join(folder, f"TITAN_IZVESTAJ_{i:03d}.docx"))
    
    if i % 100 == 0:
        print(f"✅ Obrađeno: {i}...")

print(f"\n🏁 TRIJUMF! Onur, folder 'TITAN_309_FINAL' te čeka na Desktopu.")
input("Pritisni Enter da zatvoriš ovaj prozor...")