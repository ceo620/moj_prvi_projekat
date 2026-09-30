import json, os
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

with open("00_MASTER_TOKEN_AUTHORITY/MASTER_TOKEN_REGISTER.json") as f:
    master = json.load(f)

def create_eu_branded_doc(title, filename):
    doc = Document()
    # EU stil
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)

    # Header (EU plava)
    header = doc.sections[0].header
    h = header.paragraphs[0]
    h.text = f"{master['canonical_company']}  |  TITAN 1"
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.style.font.color.rgb = RGBColor(0, 51, 153)  # EU plava

    # Naslov
    doc.add_heading(title, 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f'Datum: 2026-05-10\n').alignment = WD_ALIGN_PARAGRAPH.CENTER

    return doc

# 1. SIGNALNI MEMORANDUM
doc1 = create_eu_branded_doc('SIGNALNI MEMORANDUM', 'SIGNALNI_MEMORANDUM_TITAN1_EU_BRAND_2026-05-10.docx')
p = doc1.add_paragraph("""Poštovani,

Ovaj memorandum se odnosi isključivo na projekat TITAN 1 – Fabrika transformatorskih kazana (tankova).

Nosilac projekta: 
""" + master['canonical_company'] + """

Lokacija: 
""" + master['canonical_location'] + """, Crna Gora

Projekat TITAN 1 predviđa investiciju od približno """ + master['canonical_capex_label'] + """ i ima za cilj jačanje metaloprerađivačke industrije, otvaranje novih radnih mjesta te doprinos energetskom lancu i industrijskoj tranziciji u Crnoj Gori.

U cilju dalje institucionalne koordinacije molimo za informaciju o raspoloživom kapacitetu i statusu elektrodistributivne infrastrukture, proceduru i tehničke uslove priključenja te mogućnosti institucionalnog usmjeravanja i podrške u okviru važećih propisa.

Sa poštovanjem,""")
doc1.add_paragraph('DANIJELA KESKIN')
doc1.add_paragraph('Director')
doc1.add_paragraph('Ars Metal Industries D.O.O. Podgorica')
doc1.save('SIGNALNI_MEMORANDUM_TITAN1_EU_BRAND_2026-05-10.docx')

# 2. ANEKS 1
doc2 = create_eu_branded_doc('ANEKS 1 – REVIZIJA CAPEX OKVIRA ZA PROJEKAT TITAN 1', 'ANEKS_1_REVIZIJA_CAPEX_TITAN1_EU_BRAND_2026-05-10.docx')
doc2.add_paragraph(f"""Ovaj Aneks 1 odnosi se isključivo na projekat TITAN 1.

Nosilac projekta: 
{master['canonical_company']}

Lokacija projekta: 
{master['canonical_location']}, Crna Gora

Ažurirani CAPEX okvir za TITAN 1: 
{master['canonical_capex_label']}

Ovaj Aneks predstavlja prateće objašnjenje uz postojeću projektnu dokumentaciju i koristi se isključivo za potrebe institucionalne komunikacije i preliminarnog usklađivanja.""")
doc2.add_paragraph('\nSa poštovanjem,')
doc2.add_paragraph('DANIJELA KESKIN')
doc2.add_paragraph('Director')
doc2.add_paragraph('Ars Metal Industries D.O.O. Podgorica')
doc2.save('ANEKS_1_REVIZIJA_CAPEX_TITAN1_EU_BRAND_2026-05-10.docx')

print("✅ EU BRANDIRANI DOKUMENTI KREIRANI")
print("   • SIGNALNI_MEMORANDUM_TITAN1_EU_BRAND_2026-05-10.docx")
print("   • ANEKS_1_REVIZIJA_CAPEX_TITAN1_EU_BRAND_2026-05-10.docx")
print("   Oba dokumenta su sada u EU institucionalnom stilu i spremna za tablet.")
