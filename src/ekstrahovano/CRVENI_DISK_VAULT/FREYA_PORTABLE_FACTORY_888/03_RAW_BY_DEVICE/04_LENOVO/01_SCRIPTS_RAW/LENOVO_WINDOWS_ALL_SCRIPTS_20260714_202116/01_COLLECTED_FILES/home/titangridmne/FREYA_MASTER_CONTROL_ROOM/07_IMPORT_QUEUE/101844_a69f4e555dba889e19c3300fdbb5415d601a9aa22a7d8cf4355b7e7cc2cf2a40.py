# ============================================================
# TITAN_KERNEL: 47_control_tower_dashboard_exporter.py
# PURPOSE: Export readable dashboard from Control Tower JSONL reports
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_OK = 0
EXIT_FILE_ERROR = 1
EXIT_WRITE_ERROR = 2
EXIT_DEPENDENCY_ERROR = 5

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def read_jsonl(path):
    events = []
    if not os.path.exists(path):
        return events, f"Missing report: {path}"

    with open(path, "r", encoding="utf-8") as f:
        for idx, line in enumerate(f, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                row = json.loads(line)
                row["_source_file"] = path
                row["_line"] = idx
                events.append(row)
            except json.JSONDecodeError:
                events.append({
                    "timestamp": now_iso(),
                    "status": "INVALID_JSON",
                    "message": line,
                    "_source_file": path,
                    "_line": idx
                })
    return events, None

def flatten_event(event):
    return {
        "timestamp": event.get("timestamp", ""),
        "component": event.get("component", event.get("tool", "")),
        "status": event.get("status", ""),
        "severity": event.get("severity", ""),
        "rule_or_check_id": event.get("rule_id", event.get("check_id", "")),
        "name": event.get("rule_name", event.get("check_name", "")),
        "message": event.get("message", ""),
        "source_file": event.get("_source_file", ""),
        "line": event.get("_line", ""),
        "final_use_allowed": event.get("FINAL_USE_ALLOWED", event.get("final_use_allowed", CANON["FINAL_USE_ALLOWED"])),
    }

def write_xlsx(events, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Dashboard"
    status_counts = Counter(e.get("status", "UNKNOWN") for e in events)
    severity_counts = Counter(e.get("severity", "UNKNOWN") for e in events)

    ws.append(["Metric", "Value"])
    ws.append(["Generated_At", now_iso()])
    ws.append(["Events_Total", len(events)])
    ws.append(["Status_Counts", json.dumps(dict(status_counts), ensure_ascii=False)])
    ws.append(["Severity_Counts", json.dumps(dict(severity_counts), ensure_ascii=False)])
    for k, v in CANON.items():
        ws.append([k, v])

    ws2 = wb.create_sheet("Events")
    headers = ["timestamp", "component", "status", "severity", "rule_or_check_id", "name", "message", "source_file", "line", "final_use_allowed"]
    ws2.append(headers)
    for e in events:
        flat = flatten_event(e)
        ws2.append([flat.get(h, "") for h in headers])

    ws3 = wb.create_sheet("Blocking Events")
    ws3.append(headers)
    for e in events:
        if e.get("status") == "BLOCK":
            flat = flatten_event(e)
            ws3.append([flat.get(h, "") for h in headers])

    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def write_html(events, out_html):
    status_counts = Counter(e.get("status", "UNKNOWN") for e in events)
    block_events = [flatten_event(e) for e in events if e.get("status") == "BLOCK"]

    rows = []
    for e in block_events:
        rows.append(
            "<tr>"
            f"<td>{e['timestamp']}</td>"
            f"<td>{e['component']}</td>"
            f"<td>{e['rule_or_check_id']}</td>"
            f"<td>{e['severity']}</td>"
            f"<td>{e['message']}</td>"
            "</tr>"
        )

    html = f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>TITAN Control Tower Dashboard</title>
<style>
body {{ font-family: Arial, sans-serif; margin: 24px; }}
h1 {{ color: #8B0000; }}
table {{ border-collapse: collapse; width: 100%; margin-top: 12px; }}
td, th {{ border: 1px solid #ccc; padding: 8px; text-align: left; }}
.status {{ font-weight: bold; color: #8B0000; }}
</style>
</head>
<body>
<h1>TITAN Control Tower Dashboard</h1>
<p class="status">SYSTEM RED — STEP102 LOCKED — NO FINAL USE</p>
<h2>Summary</h2>
<table>
<tr><th>Metric</th><th>Value</th></tr>
<tr><td>Generated At</td><td>{now_iso()}</td></tr>
<tr><td>Events Total</td><td>{len(events)}</td></tr>
<tr><td>Status Counts</td><td>{json.dumps(dict(status_counts), ensure_ascii=False)}</td></tr>
<tr><td>FINAL_USE_ALLOWED</td><td>NO</td></tr>
</table>
<h2>Blocking Events</h2>
<table>
<tr><th>Timestamp</th><th>Component</th><th>Rule/Check</th><th>Severity</th><th>Message</th></tr>
{''.join(rows)}
</table>
</body>
</html>"""

    Path(out_html).parent.mkdir(parents=True, exist_ok=True)
    Path(out_html).write_text(html, encoding="utf-8")

def main():
    parser = argparse.ArgumentParser(description="Export Control Tower dashboard")
    parser.add_argument("--reports", nargs="+", required=True, help="One or more JSONL report paths")
    parser.add_argument("--out-xlsx", required=True)
    parser.add_argument("--out-html")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    all_events = []
    missing = []

    for report in args.reports:
        events, error = read_jsonl(report)
        all_events.extend(events)
        if error:
            missing.append(error)

    try:
        write_xlsx(all_events, args.out_xlsx)
        if args.out_html:
            write_html(all_events, args.out_html)
    except Exception as exc:
        status = {"timestamp": now_iso(), "component": "CONTROL_TOWER_DASHBOARD_EXPORTER", "status": "BLOCK", "message": str(exc), **CANON}
        print(json.dumps(status, ensure_ascii=False) if args.json else status["message"])
        return EXIT_WRITE_ERROR

    status = {
        "timestamp": now_iso(),
        "component": "CONTROL_TOWER_DASHBOARD_EXPORTER",
        "status": "REVIEW_REQUIRED",
        "message": "Dashboard exported; not approval",
        "events_total": len(all_events),
        "missing_reports": missing,
        "out_xlsx": args.out_xlsx,
        "out_html": args.out_html or "N/A",
        **CANON,
    }
    print(json.dumps(status, ensure_ascii=False) if args.json else "✅ Dashboard exported")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
