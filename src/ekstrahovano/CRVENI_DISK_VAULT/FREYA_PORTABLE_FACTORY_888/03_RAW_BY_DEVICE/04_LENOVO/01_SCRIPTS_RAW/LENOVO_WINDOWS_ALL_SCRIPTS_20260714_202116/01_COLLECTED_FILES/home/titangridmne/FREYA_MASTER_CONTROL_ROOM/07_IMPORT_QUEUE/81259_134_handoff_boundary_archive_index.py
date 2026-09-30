# ============================================================
# TITAN_KERNEL: 134_handoff_boundary_archive_index.py
# PURPOSE: Index handoff-denial and review-only boundary archives/artifacts
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import hashlib
import json
import zipfile
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_OK = 0
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def inspect_zip(path):
    try:
        with zipfile.ZipFile(path, "r") as z:
            bad = z.testzip()
            names = z.namelist()
        return bad is None, bad or "", names
    except Exception as exc:
        return False, str(exc), []

def classify(path):
    n = path.name.lower()
    if "handoff_denial" in n:
        return "HANDOFF_DENIAL_ARCHIVE"
    if "handoff_boundary" in n or "review_only_handoff" in n:
        return "REVIEW_ONLY_HANDOFF_BOUNDARY"
    if "no_operational_handoff" in n:
        return "NO_OPERATIONAL_HANDOFF"
    return "HANDOFF_RELATED_ARTIFACT"

def collect(roots):
    rows = []
    contents = []
    for root in roots:
        p = Path(root)
        if not p.exists():
            continue
        iterator = [p] if p.is_file() else p.rglob("*")
        for item in iterator:
            if not item.is_file():
                continue
            lname = item.name.lower()
            if not any(k in lname for k in ["handoff", "operational"]):
                continue
            if item.suffix.lower() not in {".json", ".md", ".txt", ".xlsx", ".zip"}:
                continue

            valid_zip = "N/A"
            bad = ""
            item_count = 0
            if item.suffix.lower() == ".zip":
                valid_zip, bad, names = inspect_zip(item)
                item_count = len(names)
                for name in names:
                    contents.append({
                        "Archive_Name": item.name,
                        "Archive_Path": str(item),
                        "Item_Name": name,
                        "Final_Use_Allowed": "NO",
                        "Handoff_Allowed": "NO"
                    })

            rows.append({
                "Handoff_Index_ID": f"HIDX-{len(rows)+1:05d}",
                "Artifact_Name": item.name,
                "Artifact_Type": classify(item),
                "Path": str(item),
                "SHA256": sha256_file(item),
                "Size_Bytes": item.stat().st_size,
                "Modified_UTC": datetime.utcfromtimestamp(item.stat().st_mtime).isoformat(timespec="seconds"),
                "Valid_Zip": valid_zip,
                "Bad_Member": bad,
                "Item_Count": item_count,
                "Handoff_Allowed": "NO",
                "Canonical_Effect": "NONE",
                "Final_Use_Allowed": "NO",
            })
    return rows, contents

def write_xlsx(payload, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Handoff Boundary Index"
    headers = ["Handoff_Index_ID", "Artifact_Name", "Artifact_Type", "Path", "SHA256", "Size_Bytes", "Modified_UTC", "Valid_Zip", "Bad_Member", "Item_Count", "Handoff_Allowed", "Canonical_Effect", "Final_Use_Allowed"]
    ws.append(headers)
    for row in payload["records"]:
        ws.append([row.get(h, "") for h in headers])

    ws2 = wb.create_sheet("Archive Contents")
    ws2.append(["Archive_Name", "Archive_Path", "Item_Name", "Handoff_Allowed", "Final_Use_Allowed"])
    for row in payload["contents"]:
        ws2.append([row.get("Archive_Name"), row.get("Archive_Path"), row.get("Item_Name"), row.get("Handoff_Allowed"), row.get("Final_Use_Allowed")])

    ws3 = wb.create_sheet("Canonical Locks")
    ws3.append(["Control", "Value"])
    for k, v in CANON.items():
        ws3.append([k, v])

    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def main():
    parser = argparse.ArgumentParser(description="Build handoff boundary archive index")
    parser.add_argument("--roots", nargs="+", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--out-xlsx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    records, contents = collect(args.roots)

    payload = {
        "timestamp": now_iso(),
        "component": "HANDOFF_BOUNDARY_ARCHIVE_INDEX",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "record_count": len(records),
        "content_count": len(contents),
        "records": records,
        "contents": contents,
        "decision": "Handoff boundary index is audit inventory only. It grants no handoff or final use.",
        **CANON
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.out_xlsx:
            write_xlsx(payload, args.out_xlsx)
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else f"✅ Handoff boundary index built: {len(records)} records")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    raise SystemExit(main())
