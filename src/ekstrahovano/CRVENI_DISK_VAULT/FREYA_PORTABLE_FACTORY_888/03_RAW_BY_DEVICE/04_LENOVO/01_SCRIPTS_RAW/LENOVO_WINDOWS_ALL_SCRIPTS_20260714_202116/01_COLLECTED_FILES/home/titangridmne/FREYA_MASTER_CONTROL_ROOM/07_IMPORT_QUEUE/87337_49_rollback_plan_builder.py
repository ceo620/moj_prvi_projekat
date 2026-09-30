# ============================================================
# TITAN_KERNEL: 49_rollback_plan_builder.py
# PURPOSE: Build rollback plan from backups, manifests and current state
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_OK = 0
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def collect_backups(backup_root):
    root = Path(backup_root)
    rows = []
    if not root.exists():
        return rows
    for p in sorted(root.rglob("*")):
        if p.is_file():
            rows.append({
                "Backup_Item": p.name,
                "Path": str(p),
                "Size_Bytes": p.stat().st_size,
                "Modified_UTC": datetime.utcfromtimestamp(p.stat().st_mtime).isoformat(timespec="seconds"),
            })
    return rows

def write_xlsx(payload, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rollback Plan"
    headers = ["Step_ID", "Rollback_Action", "Source", "Target", "Risk", "Requires_Manual_Approval", "Auto_Execute", "Final_Use_Allowed"]
    ws.append(headers)
    for row in payload["rollback_steps"]:
        ws.append([row.get(h, "") for h in headers])

    ws2 = wb.create_sheet("Backup Inventory")
    b_headers = ["Backup_Item", "Path", "Size_Bytes", "Modified_UTC"]
    ws2.append(b_headers)
    for row in payload["backup_inventory"]:
        ws2.append([row.get(h, "") for h in b_headers])

    ws3 = wb.create_sheet("Canonical Locks")
    ws3.append(["Control", "Value"])
    for k, v in CANON.items():
        ws3.append([k, v])
    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def main():
    parser = argparse.ArgumentParser(description="Build TITAN rollback plan")
    parser.add_argument("--backup-root", required=True)
    parser.add_argument("--root", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--out-xlsx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    backups = collect_backups(args.backup_root)

    rollback_steps = [
        {
            "Step_ID": "RB-001",
            "Rollback_Action": "Stop all TITAN script execution",
            "Source": "Operator action",
            "Target": args.root,
            "Risk": "LOW",
            "Requires_Manual_Approval": "YES",
            "Auto_Execute": "NO",
            "Final_Use_Allowed": "NO",
        },
        {
            "Step_ID": "RB-002",
            "Rollback_Action": "Create current-state backup before rollback",
            "Source": args.root,
            "Target": args.backup_root,
            "Risk": "MEDIUM",
            "Requires_Manual_Approval": "YES",
            "Auto_Execute": "NO",
            "Final_Use_Allowed": "NO",
        },
        {
            "Step_ID": "RB-003",
            "Rollback_Action": "Select reviewed backup manifest",
            "Source": args.backup_root,
            "Target": "Manual reviewer",
            "Risk": "HIGH",
            "Requires_Manual_Approval": "YES",
            "Auto_Execute": "NO",
            "Final_Use_Allowed": "NO",
        },
        {
            "Step_ID": "RB-004",
            "Rollback_Action": "Restore selected files manually after SHA-256 comparison",
            "Source": "Selected backup",
            "Target": args.root,
            "Risk": "HIGH",
            "Requires_Manual_Approval": "YES",
            "Auto_Execute": "NO",
            "Final_Use_Allowed": "NO",
        },
        {
            "Step_ID": "RB-005",
            "Rollback_Action": "Rerun self-test, integrity checker and Control Tower validator",
            "Source": "Restored system",
            "Target": "REPORTS",
            "Risk": "MEDIUM",
            "Requires_Manual_Approval": "YES",
            "Auto_Execute": "NO",
            "Final_Use_Allowed": "NO",
        },
    ]

    payload = {
        "timestamp": now_iso(),
        "component": "ROLLBACK_PLAN_BUILDER",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "backup_root": args.backup_root,
        "root": args.root,
        "backup_count": len(backups),
        "backup_inventory": backups,
        "rollback_steps": rollback_steps,
        "decision": "Rollback plan is non-executing. No automatic restore is performed.",
        **CANON
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.out_xlsx:
            write_xlsx(payload, args.out_xlsx)
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else f"✅ Rollback plan: {args.out_json}")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
