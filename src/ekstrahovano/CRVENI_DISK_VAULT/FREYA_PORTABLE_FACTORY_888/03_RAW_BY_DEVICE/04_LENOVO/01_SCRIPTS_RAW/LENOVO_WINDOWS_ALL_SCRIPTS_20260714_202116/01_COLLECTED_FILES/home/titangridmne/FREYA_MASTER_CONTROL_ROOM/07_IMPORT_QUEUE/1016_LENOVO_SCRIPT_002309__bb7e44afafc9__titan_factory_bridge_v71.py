import json
from pathlib import Path

import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# =========================================================
# TITAN v7.1 FACTORY BRIDGE ADD-ON
# Non-destructive workbook enhancer
# Reads existing v7 lender workbook and appends factory sheets
# =========================================================

WORKBOOK_NAME = "TITAN_Central_Brain_Finance_Master_v7_Lender.xlsx"

FACTORY_CONFIG = {
    "Factory 1": {
        "capacity_units": 1200,
        "capex_share": 0.36,
        "fixed_opex_share": 0.36,
        "utilization": {
            2027: 0.35, 2028: 0.70, 2029: 0.90, 2030: 0.95, 2031: 0.95,
            2032: 0.95, 2033: 0.95, 2034: 0.95, 2035: 0.95, 2036: 0.95
        },
    },
    "Factory 2": {
        "capacity_units": 1200,
        "capex_share": 0.33,
        "fixed_opex_share": 0.33,
        "utilization": {
            2027: 0.00, 2028: 0.25, 2029: 0.80, 2030: 0.92, 2031: 0.95,
            2032: 0.95, 2033: 0.95, 2034: 0.95, 2035: 0.95, 2036: 0.95
        },
    },
    "Factory 3": {
        "capacity_units": 1200,
        "capex_share": 0.31,
        "fixed_opex_share": 0.31,
        "utilization": {
            2027: 0.00, 2028: 0.00, 2029: 0.60, 2030: 0.90, 2031: 0.95,
            2032: 0.95, 2033: 0.95, 2034: 0.95, 2035: 0.95, 2036: 0.95
        },
    },
}

REQUIRED_COLUMNS = {
    "Year",
    "Units_Sold",
    "ASP_EUR_per_Unit",
    "Revenue",
    "COGS",
    "Fixed_OPEX",
    "Variable_OPEX",
    "EBITDA",
    "Tax",
    "Working_Capital_Required",
    "Working_Capital_Change",
    "Growth_CAPEX",
    "Maintenance_CAPEX",
    "Total_CAPEX",
    "CFADS_Pre_Debt",
}

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def read_assumptions(xlsx_path: str) -> dict:
    try:
        df = pd.read_excel(xlsx_path, sheet_name="Assumptions")
    except Exception:
        return {}

    out = {}
    for _, row in df.iterrows():
        key = str(row.iloc[0])
        val = row.iloc[1]
        if isinstance(val, str):
            try:
                out[key] = json.loads(val)
                continue
            except Exception:
                pass
        out[key] = val
    return out


def read_project_metadata(xlsx_path: str) -> dict:
    try:
        df = pd.read_excel(xlsx_path, sheet_name="Project_Metadata")
    except Exception:
        return {}
    return {str(r["Field"]): r["Value"] for _, r in df.iterrows()}


def validate_operating_model(df: pd.DataFrame) -> None:
    missing = REQUIRED_COLUMNS - set(df.columns)
    if missing:
        raise ValueError(f"Operating_Model missing required columns: {sorted(missing)}")


def build_factory_model(op_model: pd.DataFrame, assumptions: dict, project_meta: dict, factory_name: str, cfg: dict) -> pd.DataFrame:
    rows = []
    prev_wc = 0.0

    base_capex_eur = float(project_meta.get("Base_CAPEX_EUR_m", 43.56)) * 1_000_000
    inflation_opex = float(assumptions.get("Inflation_Opex", 0.025))
    fixed_opex_base = float(assumptions.get("Fixed_OPEX_Base", 2_400_000))
    maint_capex_pct = float(assumptions.get("Maintenance_CAPEX_pct_of_Revenue", 0.015))
    tax_rate = float(assumptions.get("Tax_Rate", 0.09))
    capex_sched = assumptions.get("CAPEX_Schedule_pct", {})

    for _, r in op_model.iterrows():
        year = int(r["Year"])
        util = float(cfg["utilization"].get(year, 0.0))
        units = cfg["capacity_units"] * util
        asp = float(r["ASP_EUR_per_Unit"])
        revenue = units * asp

        group_units = float(r["Units_Sold"]) if float(r["Units_Sold"]) else 0.0
        cogs_per_unit = (float(r["COGS"]) / group_units) if group_units > 0 else 0.0
        var_opex_per_unit = (float(r["Variable_OPEX"]) / group_units) if group_units > 0 else 0.0

        cogs = units * cogs_per_unit
        gross_profit = revenue - cogs

        year_index = year - int(op_model["Year"].min())
        fixed_opex = (fixed_opex_base * cfg["fixed_opex_share"]) * ((1 + inflation_opex) ** year_index)
        variable_opex = units * var_opex_per_unit
        opex_total = fixed_opex + variable_opex
        ebitda = gross_profit - opex_total

        wc_required = revenue * (float(r["Working_Capital_Required"]) / float(r["Revenue"])) if float(r["Revenue"]) else 0.0
        wc_change = wc_required - prev_wc
        prev_wc = wc_required

        growth_capex = base_capex_eur * cfg["capex_share"] * float(capex_sched.get(str(year), capex_sched.get(year, 0.0)))
        maintenance_capex = revenue * maint_capex_pct if year >= 2030 else 0.0
        total_capex = growth_capex + maintenance_capex

        tax = max(ebitda, 0.0) * tax_rate
        cfads_pre_debt = ebitda - tax - wc_change

        rows.append({
            "Year": year,
            "Factory": factory_name,
            "Capacity_Units": round(cfg["capacity_units"], 2),
            "Utilization": round(util, 4),
            "Units_Sold": round(units, 2),
            "ASP_EUR_per_Unit": round(asp, 2),
            "Revenue": round(revenue, 2),
            "COGS": round(cogs, 2),
            "Gross_Profit": round(gross_profit, 2),
            "Fixed_OPEX": round(fixed_opex, 2),
            "Variable_OPEX": round(variable_opex, 2),
            "OPEX_Total": round(opex_total, 2),
            "EBITDA": round(ebitda, 2),
            "Tax": round(tax, 2),
            "Working_Capital_Required": round(wc_required, 2),
            "Working_Capital_Change": round(wc_change, 2),
            "Growth_CAPEX": round(growth_capex, 2),
            "Maintenance_CAPEX": round(maintenance_capex, 2),
            "Total_CAPEX": round(total_capex, 2),
            "CFADS_Pre_Debt": round(cfads_pre_debt, 2),
            "Source_Tag": "Factory bridge additive v7.1",
        })

    return pd.DataFrame(rows)


def build_factory_summary(factory_dfs):
    all_df = pd.concat(factory_dfs, ignore_index=True)
    return all_df.groupby("Year", as_index=False).agg({
        "Units_Sold": "sum",
        "Revenue": "sum",
        "COGS": "sum",
        "Gross_Profit": "sum",
        "Fixed_OPEX": "sum",
        "Variable_OPEX": "sum",
        "OPEX_Total": "sum",
        "EBITDA": "sum",
        "Tax": "sum",
        "Working_Capital_Required": "sum",
        "Working_Capital_Change": "sum",
        "Growth_CAPEX": "sum",
        "Maintenance_CAPEX": "sum",
        "Total_CAPEX": "sum",
        "CFADS_Pre_Debt": "sum",
    })


def build_bridge(op_model: pd.DataFrame, f1: pd.DataFrame, f2: pd.DataFrame, f3: pd.DataFrame) -> pd.DataFrame:
    def slim(df: pd.DataFrame, suffix: str) -> pd.DataFrame:
        return df[["Year", "Units_Sold", "Revenue", "EBITDA", "Total_CAPEX"]].rename(columns={
            "Units_Sold": f"Units_{suffix}",
            "Revenue": f"Revenue_{suffix}",
            "EBITDA": f"EBITDA_{suffix}",
            "Total_CAPEX": f"CAPEX_{suffix}",
        })

    bridge = op_model[["Year", "Units_Sold", "Revenue", "EBITDA", "Total_CAPEX"]].rename(columns={
        "Units_Sold": "Units_Operating_Model",
        "Revenue": "Revenue_Operating_Model",
        "EBITDA": "EBITDA_Operating_Model",
        "Total_CAPEX": "CAPEX_Operating_Model",
    })

    bridge = bridge.merge(slim(f1, "F1"), on="Year", how="left")
    bridge = bridge.merge(slim(f2, "F2"), on="Year", how="left")
    bridge = bridge.merge(slim(f3, "F3"), on="Year", how="left")

    bridge["Units_Total_Factory"] = bridge["Units_F1"] + bridge["Units_F2"] + bridge["Units_F3"]
    bridge["Revenue_Total_Factory"] = bridge["Revenue_F1"] + bridge["Revenue_F2"] + bridge["Revenue_F3"]
    bridge["EBITDA_Total_Factory"] = bridge["EBITDA_F1"] + bridge["EBITDA_F2"] + bridge["EBITDA_F3"]
    bridge["CAPEX_Total_Factory"] = bridge["CAPEX_F1"] + bridge["CAPEX_F2"] + bridge["CAPEX_F3"]

    bridge["Units_Variance"] = bridge["Units_Total_Factory"] - bridge["Units_Operating_Model"]
    bridge["Revenue_Variance"] = bridge["Revenue_Total_Factory"] - bridge["Revenue_Operating_Model"]
    bridge["EBITDA_Variance"] = bridge["EBITDA_Total_Factory"] - bridge["EBITDA_Operating_Model"]
    bridge["CAPEX_Variance"] = bridge["CAPEX_Total_Factory"] - bridge["CAPEX_Operating_Model"]

    tol = 1.0
    bridge["Bridge_Status"] = np.where(
        (bridge["Units_Variance"].abs() <= tol) &
        (bridge["Revenue_Variance"].abs() <= tol) &
        (bridge["EBITDA_Variance"].abs() <= tol) &
        (bridge["CAPEX_Variance"].abs() <= tol),
        "MATCH",
        "REVIEW",
    )

    ordered_cols = [
        "Year",
        "Revenue_F1", "Revenue_F2", "Revenue_F3", "Revenue_Total_Factory", "Revenue_Operating_Model", "Revenue_Variance",
        "EBITDA_F1", "EBITDA_F2", "EBITDA_F3", "EBITDA_Total_Factory", "EBITDA_Operating_Model", "EBITDA_Variance",
        "Units_F1", "Units_F2", "Units_F3", "Units_Total_Factory", "Units_Operating_Model", "Units_Variance",
        "CAPEX_F1", "CAPEX_F2", "CAPEX_F3", "CAPEX_Total_Factory", "CAPEX_Operating_Model", "CAPEX_Variance",
        "Bridge_Status"
    ]

    return bridge[ordered_cols]


def format_new_sheets(xlsx_path: str, sheet_names):
    wb = load_workbook(xlsx_path)
    for name in sheet_names:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 40)
        ws.freeze_panes = "A2"
    wb.save(xlsx_path)


def append_factory_bridge(workbook_path: str) -> str:
    xlsx_path = Path(workbook_path)
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Workbook not found: {xlsx_path}")

    op_model = pd.read_excel(xlsx_path, sheet_name="Operating_Model")
    validate_operating_model(op_model)

    assumptions = read_assumptions(str(xlsx_path))
    project_meta = read_project_metadata(str(xlsx_path))

    f1 = build_factory_model(op_model, assumptions, project_meta, "Factory 1", FACTORY_CONFIG["Factory 1"])
    f2 = build_factory_model(op_model, assumptions, project_meta, "Factory 2", FACTORY_CONFIG["Factory 2"])
    f3 = build_factory_model(op_model, assumptions, project_meta, "Factory 3", FACTORY_CONFIG["Factory 3"])

    summary = build_factory_summary([f1, f2, f3])
    bridge = build_bridge(op_model, f1, f2, f3)

    with pd.ExcelWriter(xlsx_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        f1.to_excel(writer, sheet_name="Factory_1_Model", index=False)
        f2.to_excel(writer, sheet_name="Factory_2_Model", index=False)
        f3.to_excel(writer, sheet_name="Factory_3_Model", index=False)
        summary.to_excel(writer, sheet_name="Factory_Summary", index=False)
        bridge.to_excel(writer, sheet_name="Factory_Bridge_to_Operating_Model", index=False)

    format_new_sheets(str(xlsx_path), [
        "Factory_1_Model",
        "Factory_2_Model",
        "Factory_3_Model",
        "Factory_Summary",
        "Factory_Bridge_to_Operating_Model",
    ])

    return str(xlsx_path)


if __name__ == "__main__":
    append_factory_bridge(WORKBOOK_NAME)
    print("✅ Factory bridge appended successfully.")