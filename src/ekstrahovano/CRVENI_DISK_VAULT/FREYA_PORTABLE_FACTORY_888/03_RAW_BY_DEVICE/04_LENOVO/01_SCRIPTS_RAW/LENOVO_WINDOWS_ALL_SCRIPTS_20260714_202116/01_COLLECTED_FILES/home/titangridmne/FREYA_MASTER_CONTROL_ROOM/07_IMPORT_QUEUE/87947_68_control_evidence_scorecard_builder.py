# ============================================================
# TITAN_KERNEL: 68_control_evidence_scorecard_builder.py
# PURPOSE: Build non-approval control/evidence scorecard
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_FILE_ERROR = 2
EXIT_SCHEMA_ERROR = 3
EXIT_WRITE_ERROR = 4

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "RISK_BASELINE": 890,
    "P0_GATES": "10/10 BLOCKED",
    "EVIDENCE_GAPS_REMAINING": 6,
    "EVIDENCE_APPROVED": 0,
    "GATES_CLOSED": 0,
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def norm(value):
    return "" if value is None else str(value).strip()

def headers(ws):
    return {norm(c.value): i for i, c in enumerate(ws[1], start=1) if norm(c.value)}

def safe_float(value):
    try:
        return float(str(value).strip())
    except Exception:
        return 0.0

def main():
    parser = argparse.ArgumentParser(description="Build TITAN non-approval control evidence scorecard")
    parser.add_argument("--excel", required=True, help="FileFinding/EvidenceGap workbook")
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--out-xlsx", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print(f"❌ Excel ne postoji: {args.excel}")
        return EXIT_FILE_ERROR

    wb_in = openpyxl.load_workbook(args.excel, data_only=True)
    required_sheets = ["FileFinding Register", "Evidence Gap Register"]
    missing_sheets = [s for s in required_sheets if s not in wb_in.sheetnames]
    if missing_sheets:
        print(f"❌ Nedostaju sheetovi: {missing_sheets}")
        return EXIT_SCHEMA_ERROR

    ff = wb_in["FileFinding Register"]
    gaps = wb_in["Evidence Gap Register"]
    ff_h = headers(ff)
    gap_h = headers(gaps)

    ff_required = ["FileFinding_ID", "Evidence_ID", "Signal_ID", "SHA256", "Weight", "Status", "Reviewer"]
    gap_required = ["Gap_ID", "Linked_FileFinding_ID", "Status", "Blocking_STEP"]

    missing_cols = [c for c in ff_required if c not in ff_h] + [c for c in gap_required if c not in gap_h]
    if missing_cols:
        print(f"❌ Nedostaju kolone: {missing_cols}")
        return EXIT_SCHEMA_ERROR

    findings_total = 0
    findings_weight_positive = 0
    findings_reviewed = 0
    findings_complete = 0

    for r in range(2, ff.max_row + 1):
        fid = norm(ff.cell(r, ff_h["FileFinding_ID"]).value)
        if not fid:
            continue
        findings_total += 1
        evidence_id = norm(ff.cell(r, ff_h["Evidence_ID"]).value)
        signal_id = norm(ff.cell(r, ff_h["Signal_ID"]).value)
        sha = norm(ff.cell(r, ff_h["SHA256"]).value)
        weight = safe_float(ff.cell(r, ff_h["Weight"]).value)
        status = norm(ff.cell(r, ff_h["Status"]).value).upper()
        reviewer = norm(ff.cell(r, ff_h["Reviewer"]).value)

        if weight > 0:
            findings_weight_positive += 1
        if reviewer:
            findings_reviewed += 1
        if evidence_id and signal_id and sha and weight > 0 and reviewer and status in {"REVIEWED", "VALIDATED", "APPROVED_REFERENCE", "REVIEW_REQUIRED_EVIDENCE_SIGNAL"}:
            findings_complete += 1

    gaps_total = 0
    gaps_open = 0
    gaps_linked = 0
    closed = {"CLOSED", "APPROVED", "RESOLVED", "VALIDATED"}

    for r in range(2, gaps.max_row + 1):
        gap_id = norm(gaps.cell(r, gap_h["Gap_ID"]).value)
        if not gap_id:
            continue
        gaps_total += 1
        status = norm(gaps.cell(r, gap_h["Status"]).value).upper()
        linked = norm(gaps.cell(r, gap_h["Linked_FileFinding_ID"]).value)
        if status not in closed:
            gaps_open += 1
        if linked:
            gaps_linked += 1

    scores = {
        "FileFinding_Completeness_Percent": round((findings_complete / findings_total * 100), 2) if findings_total else 0,
        "Evidence_Linkage_Percent": round((gaps_linked / gaps_total * 100), 2) if gaps_total else 0,
        "Open_Gap_Count": gaps_open,
        "Positive_Weight_Findings": findings_weight_positive,
        "Reviewed_Findings": findings_reviewed,
    }

    result = {
        "timestamp": now_iso(),
        "component": "CONTROL_EVIDENCE_SCORECARD_BUILDER",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "findings_total": findings_total,
        "findings_complete": findings_complete,
        "gaps_total": gaps_total,
        "gaps_open": gaps_open,
        "scores": scores,
        "decision": "Scorecard is informational only. It cannot approve evidence or close gaps.",
        **CANON,
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Scorecard"
        ws.append(["Metric", "Value"])
        ws.append(["Generated_At", result["timestamp"]])
        ws.append(["Findings_Total", findings_total])
        ws.append(["Findings_Complete", findings_complete])
        ws.append(["Gaps_Total", gaps_total])
        ws.append(["Gaps_Open", gaps_open])
        for k, v in scores.items():
            ws.append([k, v])
        for k, v in CANON.items():
            ws.append([k, v])

        ws2 = wb.create_sheet("Interpretation")
        ws2.append(["Rule", "Meaning"])
        ws2.append(["Scorecard is non-approval", "Even 100% score cannot close a gate automatically."])
        ws2.append(["Open gaps block STEP102", "Any open gap keeps STEP102 LOCKED / NOT ACCEPTED."])
        ws2.append(["Final use remains NO", "This workbook cannot change final-use status."])

        Path(args.out_xlsx).parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.out_xlsx)

    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(result, ensure_ascii=False, indent=2) if args.json else "✅ Scorecard built")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
