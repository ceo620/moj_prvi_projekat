import json, os
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

with open("00_MASTER_TOKEN_AUTHORITY/MASTER_TOKEN_REGISTER.json") as f:
    master = json.load(f)

print("🔒 CTOS — Institutional Tone Revision (ANEKS 1)")

doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

doc.add_heading('ANEKS 1', 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading('REVIZIJA CAPEX OKVIRA ZA PROJEKAT TITAN 1', 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph(f'Datum: 2026-05-10\n').alignment = WD_ALIGN_PARAGRAPH.CENTER

text = f"""Poštovani,

Ovaj Aneks 1 odnosi se isključivo na projekat **TITAN 1 – Fabrika transformatorskih kazana (tankova)**, koji predstavlja aktivni investicioni prioritet kompanije Ars Metal Industries D.O.O. Podgorica.

Svrha ovog aneksa je da se nadležnim institucijama i partnerima pruži ažurirano objašnjenje CAPEX okvira projekta TITAN 1, nastalo nakon detaljnije interne revizije tehničkih i troškovnih procjena.

Originalni Elaborat je izrađen u ranijoj fazi razvoja projekta i sadržavao je preliminarne procjene koje su bile podložne daljoj optimizaciji. Nakon preciznijeg definisanja proizvodnog procesa, ažuriranja tehničkih parametara i optimizacije troškovnog modela, CAPEX okvir za TITAN 1 je revidiran na **{master['canonical_capex_label']}**.

Ova revizija omogućava preciznije planiranje i usklađivanje projektne dokumentacije sa trenutno raspoloživim podacima i najnovijim internim procjenama.

Ovaj Aneks služi kao prateće objašnjenje uz postojeću projektnu dokumentaciju i koristi se isključivo u svrhu institucionalne komunikacije, preliminarnog usklađivanja i daljih procedura vezanih za projekat TITAN 1.

Sa poštovanjem,"""

for paragraph in text.split('\n'):
    if paragraph.strip():
        doc.add_paragraph(paragraph)

doc.add_paragraph('DANIJELA KESKIN')
doc.add_paragraph('Director')
doc.add_paragraph('Ars Metal Industries D.O.O. Podgorica')

filename = "ANEKS_1_REVIZIJA_CAPEX_TITAN1_INSTITUTIONAL_FINAL_2026-05-10.docx"
doc.save(filename)

print(f"\n✅ INSTITUTIONALNA VERZIJA ANEKS 1 generisana: {filename}")
print("   Ton revidiran – formalniji, diplomatskiji i institucionalniji")
print("   Spreman za manual review, human signoff i slanje institucijama.")
