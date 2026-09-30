# ============================================================
# TITAN_KERNEL: 59_all_packages_checksum_ledger.py
# PURPOSE: Build checksum ledger for all TITAN ZIP packages
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import hashlib
import json
import os
import sys
import zipfile
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None

EXIT_OK = 0
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def inspect_zip(path):
    try:
        with zipfile.ZipFile(path, "r") as z:
            bad = z.testzip()
            names = z.namelist()
        return bad is None, bad, names
    except Exception as exc:
        return False, str(exc), []

def collect_packages(root, recursive):
    root = Path(root)
    if root.is_file() and root.suffix.lower() == ".zip":
        return [root]
    if not root.exists():
        return []
    return sorted(root.rglob("*.zip") if recursive else root.glob("*.zip"))

def write_xlsx(payload, out_xlsx):
    if openpyxl is None:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Checksum Ledger"
    headers = ["Ledger_ID", "Package_Name", "Path", "SHA256", "Valid_Zip", "Bad_Member", "Item_Count", "Size_Bytes", "Modified_UTC", "Review_Status", "Final_Use_Allowed"]
    ws.append(headers)
    for row in payload["packages"]:
        ws.append([row.get(h, "") for h in headers])

    ws2 = wb.create_sheet("Package Contents")
    ws2.append(["Ledger_ID", "Package_Name", "Item_Name"])
    for row in payload["packages"]:
        for item in row.get("Items", []):
            ws2.append([row["Ledger_ID"], row["Package_Name"], item])

    ws3 = wb.create_sheet("Canonical Locks")
    ws3.append(["Control", "Value"])
    for k, v in CANON.items():
        ws3.append([k, v])

    Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)

def main():
    parser = argparse.ArgumentParser(description="TITAN all packages checksum ledger")
    parser.add_argument("--packages-root", required=True)
    parser.add_argument("--recursive", action="store_true")
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--out-xlsx")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    rows = []
    for p in collect_packages(args.packages_root, args.recursive):
        valid, bad, items = inspect_zip(p)
        rows.append({
            "Ledger_ID": f"LEDGER-{len(rows)+1:05d}",
            "Package_Name": p.name,
            "Path": str(p),
            "SHA256": sha256_file(p),
            "Valid_Zip": valid,
            "Bad_Member": bad or "",
            "Item_Count": len(items),
            "Items": items,
            "Size_Bytes": p.stat().st_size,
            "Modified_UTC": datetime.utcfromtimestamp(p.stat().st_mtime).isoformat(timespec="seconds"),
            "Review_Status": "REVIEW_REQUIRED",
            "Final_Use_Allowed": "NO"
        })

    payload = {
        "timestamp": now_iso(),
        "component": "ALL_PACKAGES_CHECKSUM_LEDGER",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "package_count": len(rows),
        "packages": rows,
        "decision": "Checksum ledger is audit-only. Package validity does not approve final use.",
        **CANON
    }

    try:
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if args.out_xlsx:
            write_xlsx(payload, args.out_xlsx)
    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else f"✅ Checksum ledger built: {len(rows)} packages")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
