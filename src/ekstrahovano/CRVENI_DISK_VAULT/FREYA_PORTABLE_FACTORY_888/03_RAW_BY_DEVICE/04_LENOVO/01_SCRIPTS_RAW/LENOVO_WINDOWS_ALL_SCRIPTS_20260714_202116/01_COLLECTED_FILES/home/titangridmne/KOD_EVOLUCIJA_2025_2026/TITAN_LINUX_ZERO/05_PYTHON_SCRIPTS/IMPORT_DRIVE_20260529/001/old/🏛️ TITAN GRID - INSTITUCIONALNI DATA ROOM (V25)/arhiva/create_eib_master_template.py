import qrcode
import os
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ====================== POMOĆNE XML FUNKCIJE (DIZAJN) ======================
def set_cell_margins(cell, top=50, start=100, bottom=50, end=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('left', start), ('bottom', bottom), ('right', end)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def add_page_number(run):
    for fld_type in ['begin', 'separate', 'end']:
        fldChar = OxmlElement('w:fldChar')
        fldChar.set(qn('w:fldCharType'), fld_type)
        if fld_type == 'begin':
            run._r.append(fldChar)
            instrText = OxmlElement('w:instrText')
            instrText.set(qn('xml:space'), 'preserve')
            instrText.text = " PAGE "
            run._r.append(instrText)
        else:
            run._r.append(fldChar)

# ====================== INICIJALIZACIJA ======================
doc = Document()

# Paleta boja
BRAND_BLUE_HEX = '0F2841'
BRAND_BLUE = RGBColor(15, 40, 65)
DARK_GRAY = RGBColor(64, 64, 64)
LIGHT_GRAY_HEX = 'F5F7FA'
CONFIDENTIAL_RED = RGBColor(192, 0, 0)
WHITE = RGBColor(255, 255, 255)

# ====================== GLOBALNI STILOVI ======================
styles = doc.styles

# Heading 1 (Glavni naslovi sekcija)
h1_style = styles['Heading 1']
h1_font = h1_style.font
h1_font.name = 'Calibri Light'
h1_font.size = Pt(16)
h1_font.color.rgb = BRAND_BLUE
h1_font.bold = True
h1_style.paragraph_format.space_before = Pt(18)
h1_style.paragraph_format.space_after = Pt(6)

# Normal tekst
normal_style = styles['Normal']
normal_font = normal_style.font
normal_font.name = 'Arial'
normal_font.size = Pt(10.5)
normal_font.color.rgb = DARK_GRAY
normal_style.paragraph_format.space_after = Pt(10)
normal_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
normal_style.paragraph_format.line_spacing = 1.15

# ====================== NASLOVNA STRANA ======================
section = doc.sections[0]
for margin in ['left_margin', 'right_margin', 'top_margin', 'bottom_margin']:
    setattr(section, margin, Inches(1.25))

doc.add_paragraph("\n\n\n\n")

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_title = p_title.add_run("TITAN GRID")
run_title.font.name = 'Calibri Light'
run_title.font.size = Pt(56)
run_title.font.bold = True
run_title.font.color.rgb = BRAND_BLUE

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_sub = p_sub.add_run("ARS METAL INDUSTRIES D.O.O.")
run_sub.font.name = 'Arial'
run_sub.font.size = Pt(12)
run_sub.font.color.rgb = DARK_GRAY
run_sub.font.bold = True
p_sub.paragraph_format.space_after = Pt(40)

p_line = doc.add_paragraph()
p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_line.add_run("________________________________________").font.color.rgb = BRAND_BLUE

p_type = doc.add_paragraph()
p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_type.paragraph_format.space_before = Pt(20)
run_type = p_type.add_run("EXECUTIVE INVESTMENT TEASER")
run_type.font.name = 'Calibri Light'
run_type.font.size = Pt(18)
run_type.font.color.rgb = BRAND_BLUE

p_meta = doc.add_paragraph()
p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_meta.paragraph_format.space_before = Pt(10)
p_meta.add_run(f"Version V25 | {datetime.now().strftime('%d.%m.%Y.')}").font.size = Pt(10)

doc.add_paragraph("\n\n")

# QR Kod
qr = qrcode.QRCode(box_size=6, border=1)
qr.add_data("https://titan-grid.dataroom.com")
qr.make(fit=True)
img = qr.make_image(fill_color="#0F2841", back_color="white")
img.save("qr_teaser.png")

p_qr = doc.add_paragraph()
p_qr.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_qr.add_run().add_picture("qr_teaser.png", width=Inches(1.2))
p_qr_text = doc.add_paragraph("Scan for secure Data Room access")
p_qr_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_qr_text.runs[0].font.size = Pt(8)
os.remove("qr_teaser.png")

doc.add_page_break()

# ====================== KONTROLA & DISCLAIMER ======================
doc.add_paragraph("DOCUMENT CONTROL", style="Heading 1")
table = doc.add_table(rows=5, cols=2)
table.style = 'Table Grid'
control_data = [
    ("Document Reference", "TG-TEASER-V25"),
    ("Date of Issue", datetime.now().strftime("%d.%m.%Y.")),
    ("Prepared by", "Onur - Project CFO"),
    ("Reviewed by", "Danijela Đurović Keskin - Project Manager"),
    ("Classification", "STRICTLY CONFIDENTIAL")
]
for i, (col1, col2) in enumerate(control_data):
    cell1, cell2 = table.cell(i, 0), table.cell(i, 1)
    cell1.text, cell2.text = col1, col2
    set_cell_background(cell1, LIGHT_GRAY_HEX)
    cell1.paragraphs[0].runs[0].font.bold = True
    cell1.paragraphs[0].runs[0].font.color.rgb = BRAND_BLUE
    set_cell_margins(cell1); set_cell_margins(cell2)

doc.add_paragraph("\n")
p_disc = doc.add_paragraph("CONFIDENTIALITY & DISCLAIMER")
p_disc.runs[0].font.name = 'Calibri Light'
p_disc.runs[0].font.size = Pt(12)
p_disc.runs[0].font.bold = True
p_disc.runs[0].font.color.rgb = CONFIDENTIAL_RED

doc.add_paragraph(
    "This document contains strictly confidential and proprietary information of Titan Grid and ARS Metal Industries D.O.O. Montenegro. "
    "It is intended solely for the named recipient(s) for the purpose of evaluating potential project finance structuring. "
    "Unauthorized use, reproduction, or distribution is strictly prohibited.", style="Normal"
).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

doc.add_page_break()

# ====================== TEASER SADRŽAJ (BODY TEXT) ======================
def add_bullet(doc, bold_text, normal_text):
    """Pomoćna funkcija za kreiranje formatiranih bullet pointova"""
    p = doc.add_paragraph(style='List Bullet')
    run_bold = p.add_run(bold_text)
    run_bold.font.bold = True
    p.add_run(normal_text)

# Sekcija 1
doc.add_paragraph("1. PROJECT OVERVIEW & THE OPPORTUNITY", style="Heading 1")
p1 = doc.add_paragraph(style="Normal")
p1.add_run("ARS Metal Industries D.O.O.").font.bold = True
p1.add_run(" is initiating the development of ")
p1.add_run("TITAN GRID").font.bold = True
p1.add_run(", a state-of-the-art, greenfield manufacturing facility located in Tuzi, Montenegro. The project encompasses the construction and operation of a high-tech \"Smart Factory\" dedicated to the production of premium power transformer tanks and advanced cooling systems.\n\nDesigned to meet the stringent quality standards of the European Union, TITAN GRID leverages Industry 4.0 automation, optimized CAPEX deployment, and a highly strategic geographic location to disrupt the regional supply chain for high-voltage energy infrastructure.\n\nWith Europe undergoing a massive energy transition and grid modernization, the demand for transformer components is at an all-time high, while existing production capacities remain heavily constrained. TITAN GRID is positioned to capture this supply-demand gap, offering rapid scalability, cost-efficient production, and uncompromising premium quality.")

# Sekcija 2
doc.add_paragraph("2. STRATEGIC LOCATION: TUZI, MONTENEGRO", style="Heading 1")
doc.add_paragraph("The facility is strategically situated in Tuzi, Montenegro, offering distinct geopolitical and logistical advantages:", style="Normal")
add_bullet(doc, "Logistical Hub: ", "Immediate proximity to major transport corridors and the Port of Bar, enabling seamless export to the EU, Middle East, and North African markets.")
add_bullet(doc, "Cost Competitiveness: ", "Access to a highly competitive energy tariff and a favorable corporate tax environment (9-15%), driving superior EBITDA margins compared to EU-based competitors.")
add_bullet(doc, "Nearshoring Advantage: ", "For European OEMs, TITAN GRID represents a secure, nearshore supply chain alternative, mitigating global shipping risks and geopolitical instability.")

# Sekcija 3
doc.add_paragraph("3. TECHNICAL EXCELLENCE & PRODUCT PORTFOLIO", style="Heading 1")
doc.add_paragraph("TITAN GRID will operate a ~10,000 m² modern production facility, utilizing advanced CNC machinery, robotic welding stations, and automated anti-corrosion coating lines. The core product lines include:", style="Normal")
add_bullet(doc, "High-Voltage Transformer Tanks: ", "Custom-engineered, hermetically sealed steel tanks built to withstand extreme mechanical and environmental stress.")
add_bullet(doc, "Advanced Cooling Systems: ", "High-efficiency radiator banks and cooling modules essential for thermal regulation in mega-transformers.")
add_bullet(doc, "Smart Manufacturing: ", "Full integration of a \"Central Brain\" ERP/MES system for real-time tracking of materials, energy consumption, and quality control.")

# Sekcija 4
doc.add_paragraph("4. ESG & EU TAXONOMY ALIGNMENT", style="Heading 1")
doc.add_paragraph("The TITAN GRID project is fully aligned with the European Investment Bank (EIB) Climate Action goals and the EU Taxonomy for Sustainable Activities.", style="Normal")
add_bullet(doc, "Energy Efficiency: ", "The facility design incorporates rooftop solar PV installations to offset operational energy consumption.")
add_bullet(doc, "Resource Optimization: ", "Implementing closed-loop water cooling systems and zero-waste steel optimization algorithms.")
add_bullet(doc, "Green Supply Chain: ", "Enabling the broader European transition to renewable energy by supplying critical bottleneck components for grid upgrades.")

doc.add_page_break()

# Sekcija 5 (Tabela umesto buleta za bolji vizuelni efekat)
doc.add_paragraph("5. FINANCIAL HIGHLIGHTS & CAPITAL STRUCTURE", style="Heading 1")
doc.add_paragraph("The financial model for TITAN GRID has been rigorously stress-tested. The project offers a highly attractive risk-reward profile, characterized by strong cash flow generation and rapid deleveraging capabilities.", style="Normal")

fin_table = doc.add_table(rows=6, cols=2)
fin_table.style = 'Table Grid'

# Zaglavlje tabele
hdr_cells = fin_table.rows[0].cells
hdr_cells[0].text, hdr_cells[1].text = "METRIC", "PROJECTED VALUE"
for cell in hdr_cells:
    set_cell_background(cell, BRAND_BLUE_HEX)
    set_cell_margins(cell, top=80, bottom=80)
    run = cell.paragraphs[0].runs[0]
    run.font.bold = True
    run.font.color.rgb = WHITE
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Podaci
fin_data = [
    ("Estimated Total CAPEX", "€ [XX.X] Million"),
    ("Target Internal Rate of Return (IRR)", "[XX.X]%"),
    ("Stabilized EBITDA Margin", "[XX.X]%"),
    ("Payback Period", "[X.X] Years"),
    ("Debt Service Coverage Ratio (DSCR)", "Min [X.XX]x")
]
for i, (k, v) in enumerate(fin_data, start=1):
    row_cells = fin_table.rows[i].cells
    row_cells[0].text, row_cells[1].text = k, v
    row_cells[0].paragraphs[0].runs[0].font.bold = True
    row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_cell_margins(row_cells[0]); set_cell_margins(row_cells[1])
    if i % 2 != 0:
        set_cell_background(row_cells[0], LIGHT_GRAY_HEX)
        set_cell_background(row_cells[1], LIGHT_GRAY_HEX)

doc.add_paragraph("\n")

# Sekcija 6
doc.add_paragraph("6. THE ASK & NEXT STEPS", style="Heading 1")
doc.add_paragraph("ARS Metal Industries is currently seeking a strategic partnership and institutional financing to execute the TITAN GRID vision. The immediate use of proceeds will be directed toward:", style="Normal")

p_num1 = doc.add_paragraph(style='List Number')
p_num1.add_run("Finalization of land acquisition and civil works in Tuzi.")
p_num2 = doc.add_paragraph(style='List Number')
p_num2.add_run("Procurement of proprietary manufacturing technology and robotic assembly lines.")
p_num3 = doc.add_paragraph(style='List Number')
p_num3.add_run("Working capital allocation for the initial 18 months of operations.")

p_close = doc.add_paragraph("\nAccess to the comprehensive ", style="Normal")
p_close.add_run("Titan Grid Data Room").font.bold = True
p_close.add_run("—including the Master Financial Model, detailed CAPEX breakdowns, and technical blueprints—will be granted to qualified partners following the execution of an NDA.")

# ====================== ELEGANTAN FOOTER ======================
footer = section.footer
p_footer_line = footer.paragraphs[0]
p_footer_line.add_run("________________________________________________________________________").font.color.rgb = RGBColor(200, 200, 200)

f_table = footer.add_table(rows=1, cols=3, width=Inches(8))
f_table.autofit = True

cell_left = f_table.cell(0, 0)
run_left = cell_left.paragraphs[0].add_run("STRICTLY CONFIDENTIAL")
run_left.font.size, run_left.font.bold, run_left.font.color.rgb = Pt(8), True, CONFIDENTIAL_RED

cell_center = f_table.cell(0, 1)
cell_center.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
run_c = cell_center.paragraphs[0].add_run("TITAN GRID | ARS Metal Industries")
run_c.font.size, run_c.font.color.rgb = Pt(8), DARK_GRAY

cell_right = f_table.cell(0, 2)
cell_right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
run_right = cell_right.paragraphs[0].add_run("Page ")
run_right.font.size, run_right.font.color.rgb = Pt(8), DARK_GRAY
add_page_number(cell_right.paragraphs[0].add_run())

for row in f_table.rows:
    for cell in row.cells:
        tcBorders = OxmlElement('w:tcBorders')
        for border_name in ['top', 'left', 'bottom', 'right']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'nil')
            tcBorders.append(border)
        cell._tc.get_or_add_tcPr().append(tcBorders)

# ====================== ČUVANJE DOKUMENTA ======================
filename = "02_Executive_Teaser_Titan_Grid.docx"
doc.save(filename)

print(f"✅ EXECUTIVE TEASER USPEŠNO GENERISAN: {filename}")