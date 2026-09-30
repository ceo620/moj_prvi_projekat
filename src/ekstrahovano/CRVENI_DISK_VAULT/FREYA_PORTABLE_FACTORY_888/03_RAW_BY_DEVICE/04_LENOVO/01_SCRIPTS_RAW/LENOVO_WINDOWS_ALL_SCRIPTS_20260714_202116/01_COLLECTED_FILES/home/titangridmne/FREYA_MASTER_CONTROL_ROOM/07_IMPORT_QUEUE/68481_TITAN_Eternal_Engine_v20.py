import os, re, pandas as pd, warnings
from datetime import datetime
from docx import Document
from pdf2image import convert_from_path
import pytesseract
from concurrent.futures import ProcessPoolExecutor
warnings.filterwarnings('ignore')

RAW = r'C:\Users\Lenovo\Desktop\CTITAN_CENTRAL_BRAINRaw_Data'
ROOM = r'C:\Users\Lenovo\Desktop\TITAN_GRID_Data_Room_FINAL_April2026'
OUT = os.path.join(ROOM, '00_MASTER', 'TITAN_ETERNAL_ENGINE_v20.xlsx')

for sub in ['00_MASTER', '01_FINANCE', '02_LEGAL', '03_STRATEGY', '04_TECHNICAL', '08_RISK']:
    os.makedirs(os.path.join(ROOM, sub), exist_ok=True)

def clean_text(text):
    if not text: return ''
    return re.sub(r'\s+', ' ', text).strip()

def advanced_ocr(pdf_path):
    try:
        images = convert_from_path(pdf_path, dpi=400)
        return clean_text('\n'.join([pytesseract.image_to_string(img, lang='eng+srp+hrv') for img in images]))
    except: return ''

def extract_docx(fpath):
    try:
        doc = Document(fpath)
        return clean_text('\n'.join([p.text for p in doc.paragraphs]))
    except: return ''

def extract_financials(text):
    pattern = r'(\d{1,3}(?:\.\d{3})*,\d{1,2}|\d{1,3}(?:\.\d{3})*)'
    amounts = re.findall(pattern, text)
    return [float(a.replace('.', '').replace(',', '.')) for a in amounts if len(a) > 2]

def process_file(fpath):
    name = os.path.basename(fpath)
    ext = os.path.splitext(fpath)[1].lower()
    content = ''

    try:
        if ext == '.docx':
            content = extract_docx(fpath)
        elif ext == '.pdf':
            content = advanced_ocr(fpath)
        elif ext in ['.txt', '.csv']:
            with open(fpath, 'r', encoding='utf-8', errors='ignore') as f:
                content = clean_text(f.read())
    except: return None

    if len(content) < 30: return None

    vals = extract_financials(content)
    total = sum([v for v in vals if v > 100])
    risk = 'HIGH' if any(k in content.upper() for k in ['DELAY','RIZIK','PENAL','DEFAULT','OVERRUN']) else 'MEDIUM'

    # Auto-generisanje ID dokumenta
    if total > 5000 or 'CAPEX' in content.upper() or risk == 'HIGH':
        try:
            cat_folder = '01_FINANCE' if total > 0 else '04_TECHNICAL'
            doc = Document()
            doc.add_heading(f'ID 4.2 — Quantum Analysis (Auto v20.0)', 0)
            doc.add_paragraph(f'Fajl: {name}')
            doc.add_paragraph(f'Vrijednost: {total:,.2f} EUR | Rizik: {risk}')
            doc.add_paragraph(content[:3000])
            doc.save(os.path.join(ROOM, cat_folder, f'ID_4.2_Auto_{name[:40]}.docx'))
        except: pass

    return {
        'Data_ID': f'TITAN-V20-{datetime.now().strftime("%f")[:6]}',
        'File_Name': name,
        'Category': 'FINANCE' if total > 0 else 'TECHNICAL',
        'Risk_Level': risk,
        'Net_Value_EUR': round(total, 2),
        'VAT_21_EUR': round(total * 0.21, 2),
        'Gross_Total_EUR': round(total * 1.21, 2),
        'Text_Insight': content[:2500].strip(),
        'Words_Count': len(content.split()),
        'Last_Modified': datetime.fromtimestamp(os.path.getmtime(fpath)).strftime('%Y-%m-%d')
    }

if __name__ == '__main__':
    print('🚀 TITAN v20.0 — ETERNAL ENGINE POKRENUT')
    paths = [os.path.join(r, f) for r, d, files in os.walk(RAW) for f in files]
    print(f'Pronađeno {len(paths)} fajlova...')

    with ProcessPoolExecutor() as executor:
        results = [r for r in list(executor.map(process_file, paths)) if r]

    df = pd.DataFrame(results)
    df = df.sort_values(by='Net_Value_EUR', ascending=False)
    df.to_excel(OUT, index=False)

    print(f'✅ POBJEDA! Generisano {len(df)} redova sa najjačom analizom.')
    print(f'📂 Master fajl: {OUT}')
    print('Data Room je automatski popunjen novim ID dokumentima.')
