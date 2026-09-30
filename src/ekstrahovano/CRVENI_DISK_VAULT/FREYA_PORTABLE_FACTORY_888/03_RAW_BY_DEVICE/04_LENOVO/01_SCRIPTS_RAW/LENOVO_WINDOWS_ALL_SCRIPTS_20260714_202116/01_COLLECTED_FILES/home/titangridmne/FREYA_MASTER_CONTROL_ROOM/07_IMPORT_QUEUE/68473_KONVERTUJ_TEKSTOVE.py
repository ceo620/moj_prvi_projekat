import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# DEFINISANJE PUTANJA
SOURCE_DIR = Path(r"C:\TITAN_GRID_DATA_ROOM\TEKSTOVI")
DEST_DIR = Path(r"C:\Users\Lenovo\Desktop\TITAN_GRID_FINANCING_PACKAGE_2026\00_MASTER_INDEX_CONTROL")
DEST_DIR.mkdir(parents=True, exist_ok=True)

def batch_convert():
    print(f"🚀 Pokrećem masovnu konverziju iz: {SOURCE_DIR}")
    
    # Prolazimo kroz svaki tekstualni fajl
    for txt_file in SOURCE_DIR.glob("*.txt"):
        with open(txt_file, 'r', encoding='utf-8') as f:
            sadrzaj = f.read()

        doc = Document()
        
        # Postavljanje vizuelnog autoriteta (Navy Blue & Segoe UI)
        style = doc.styles['Normal']
        style.font.name = 'Segoe UI'
        style.font.size = Pt(11)

        # Naslov dokumenta
        heading = doc.add_heading(f'TITAN GRID - DOKUMENT {txt_file.stem}', 0)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0, 32, 96)

        # Ubacivanje tvog teksta
        doc.add_paragraph(sadrzaj)
        
        # Audit Trace Footer (CFO Authority)
        footer = doc.sections[0].footer.paragraphs[0]
        footer.text = f"Electronic Trace: Danijela Keskin, CFO | TITAN V25 | Board Ready 23.04.2026"
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Snimanje DOCX fajla
        save_path = DEST_DIR / f"TITAN_DOC_{txt_file.stem}.docx"
        doc.save(str(save_path))
        print(f"✅ Kreiran Word: {save_path.name}")

if __name__ == "__main__":
    batch_convert()