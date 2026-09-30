import json, os
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

with open("00_MASTER_TOKEN_AUTHORITY/MASTER_TOKEN_REGISTER.json") as f:
    master = json.load(f)

print("🔒 CTOS — Improved & Professional ANEKS 1 Generation")

doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

# Naslov
doc.add_heading('ANEKS 1', 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading('REVIZIJA CAPEX OKVIRA ZA PROJEKAT TITAN 1', 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph(f'Datum: 2026-05-10\n').alignment = WD_ALIGN_PARAGRAPH.CENTER

text = f"""Poštovani,

Ovaj Aneks 1 odnosi se isključivo na projekat **TITAN 1 – Fabrika transformatorskih kazana (tankova)**.

Nosilac projekta:  
{master['canonical_company']}

Lokacija projekta:  
{master['canonical_location']}, Crna Gora

2. SVRHA REVIZIJE

Svrha ovog Aneksa je usklađivanje CAPEX okvira projekta TITAN 1 sa ažuriranim internim inženjerskim, tehničkim i troškovnim procjenama nastalim nakon izrade originalnog Elaborata.

Originalni Elaborat je izrađen u ranijoj razvojnoj fazi projekta i sadržao je preliminarne finansijske procjene koje su bile podložne daljoj reviziji i usklađivanju na osnovu preciznijih tehničkih parametara i ponuda dobavljača.

3. AŽURIRANI CAPEX OKVIR

Ažurirani CAPEX okvir za TITAN 1 iznosi **{master['canonical_capex_label']}**.

Revizija je izvršena na osnovu:
• preciznijeg definisanja proizvodnog procesa,
• ažuriranih tehničkih parametara,
• novih ponuda dobavljača,
• optimizacije planirane proizvodne konfiguracije,
• detaljnijeg troškovnog modeliranja specifičnog za proizvodnju transformatorskih kazana.

4. STATUS DOKUMENTA

Ovaj Aneks predstavlja prateće objašnjenje uz postojeću projektnu dokumentaciju i koristi se isključivo za potrebe institucionalne komunikacije, preliminarnog usklađivanja dokumentacije i daljih procedura vezanih za projekat TITAN 1.

Originalni Elaborat ostaje važeći u dijelovima koji nisu predmet ove CAPEX revizije. Formalna integracija ažuriranja u buduću projektnu dokumentaciju biće izvršena u skladu sa važećim procedurama i zahtjevima nadležnih institucija."""

for paragraph in text.split('\n'):
    if paragraph.strip():
        doc.add_paragraph(paragraph)

doc.add_paragraph('\nSa poštovanjem,')
doc.add_paragraph('DANIJELA KESKIN')
doc.add_paragraph('Director')
doc.add_paragraph('Ars Metal Industries D.O.O. Podgorica')

# EU Footer
section = doc.sections[0]
footer = section.footer
p = footer.paragraphs[0]
p.text = "Ars Metal Industries D.O.O. Podgorica  |  TITAN 1"
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.style.font.color.rgb = RGBColor(0, 51, 153)

filename = "ANEKS_1_REVIZIJA_CAPEX_TITAN1_IMPROVED_2026-05-10.docx"
doc.save(filename)

print(f"\n✅ POBOLJŠANA VERZIJA ANEKS 1 generisana: {filename}")
print("   • Dodani svi ključni pillari i narativni standardi")
print("   • Strogo TITAN 1 scope")
print("   • Spreman za manual review i human signoff")
