from docx import Document
import pandas as pd

BASE = r"C:\Users\Lenovo\Desktop\DATA_ROOM_FINAL"
MODEL_FILE = BASE + r"\OUTPUT\TITAN_v7_INVESTMENT.xlsx"
OUTPUT_DOC = BASE + r"\OUTPUT\TITAN_INVESTMENT_MEMO.docx"

# LOAD MODEL
df_model = pd.read_excel(MODEL_FILE, sheet_name="MODEL")
df_kpi = pd.read_excel(MODEL_FILE, sheet_name="KPI")

# CREATE DOCUMENT
doc = Document()

# =========================
# TITLE
# =========================
doc.add_heading('TITAN GRID INVESTMENT MEMORANDUM', 0)

# =========================
# EXECUTIVE SUMMARY
# =========================
doc.add_heading('1. Executive Summary', 1)

irr = float(df_kpi[df_kpi["Metric"]=="IRR"]["Value"])
npv = float(df_kpi[df_kpi["Metric"]=="NPV"]["Value"])
dscr = float(df_kpi[df_kpi["Metric"]=="Min DSCR"]["Value"])

doc.add_paragraph(
f"""
TITAN Grid predstavlja industrijski projekat proizvodnje transformatora i komponenti
u okviru nearshoring strategije za EU tržište.

Ključni finansijski indikatori:
- IRR: {irr:.2%}
- NPV: {npv:,.0f} EUR
- Minimalni DSCR: {dscr:.2f}

Projekat pokazuje snažnu finansijsku održivost i sposobnost servisiranja duga.
"""
)

# =========================
# PROJECT OVERVIEW
# =========================
doc.add_heading('2. Project Overview', 1)
doc.add_paragraph(
"""
TITAN projekat uključuje razvoj proizvodnog kapaciteta za transformatorske tankove
i komponente sa fokusom na EU tržište.

Lokacija: Montenegro (Tuzi industrial zone)
Kapacitet: 10–400 MVA segment
"""
)

# =========================
# FINANCIAL SUMMARY
# =========================
doc.add_heading('3. Financial Summary', 1)

doc.add_paragraph("Model Summary:")

for index, row in df_model.head(5).iterrows():
    doc.add_paragraph(str(row.to_dict()))

# =========================
# RISK SECTION
# =========================
doc.add_heading('4. Risk Analysis', 1)

doc.add_paragraph(
"""
Glavni rizici:
- Energy cost volatility
- Regulatory (excise / CBAM)
- Production ramp-up

Mitigation:
- Solar integration
- Pass-through pricing
- Modular expansion
"""
)

# =========================
# INVESTMENT CASE
# =========================
doc.add_heading('5. Investment Case', 1)

doc.add_paragraph(
"""
Projekat generiše:
- Stabilan cash flow
- DSCR iznad lender threshold
- Visok IRR

Predstavlja atraktivnu investicionu priliku za EIB/EBRD finansiranje.
"""
)

# SAVE
doc.save(OUTPUT_DOC)

print("✅ INVESTMENT MEMO CREATED:", OUTPUT_DOC)