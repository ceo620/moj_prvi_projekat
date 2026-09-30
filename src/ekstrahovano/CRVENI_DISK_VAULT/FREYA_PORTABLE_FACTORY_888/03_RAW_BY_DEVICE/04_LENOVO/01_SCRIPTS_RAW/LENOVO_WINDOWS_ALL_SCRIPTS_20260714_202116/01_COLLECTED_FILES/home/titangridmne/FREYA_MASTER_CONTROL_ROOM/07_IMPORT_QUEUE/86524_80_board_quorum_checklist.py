# ============================================================
# TITAN_KERNEL: 80_board_quorum_checklist.py
# PURPOSE: Create audit board quorum checklist template
# VERSION: v1.0
# STATUS: SYSTEM RED — STEP102 LOCKED — NO FINAL USE
# ============================================================

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("❌ Nedostaje openpyxl. Instaliraj: pip install openpyxl")
    sys.exit(5)

EXIT_OK = 0
EXIT_WRITE_ERROR = 2

CANON = {
    "SYSTEM_STATUS": "SYSTEM RED",
    "STEP102": "LOCKED / NOT ACCEPTED",
    "FINAL_USE_ALLOWED": "NO",
    "SSOT_WRITE_ALLOWED": "NO",
    "EVIDENCE_APPROVAL_ALLOWED": "NO",
    "GATE_CLOSURE_ALLOWED": "NO",
}

HEADERS = [
    "Quorum_ID",
    "Reviewer_Name",
    "Reviewer_Role",
    "Present",
    "Conflict_Disclosed",
    "Can_Approve_Evidence",
    "Can_Close_Gate",
    "Can_Unlock_STEP102",
    "Can_Write_SSoT",
    "Final_Use_Allowed",
    "Notes"
]

DEFAULT_ROWS = [
    {
        "Quorum_ID": "QRM-00001",
        "Reviewer_Name": "",
        "Reviewer_Role": "CHAIR",
        "Present": "NO",
        "Conflict_Disclosed": "REVIEW_REQUIRED",
        "Can_Approve_Evidence": "NO",
        "Can_Close_Gate": "NO",
        "Can_Unlock_STEP102": "NO",
        "Can_Write_SSoT": "NO",
        "Final_Use_Allowed": "NO",
        "Notes": "Template row only; no authority granted."
    },
    {
        "Quorum_ID": "QRM-00002",
        "Reviewer_Name": "",
        "Reviewer_Role": "TECHNICAL_REVIEWER",
        "Present": "NO",
        "Conflict_Disclosed": "REVIEW_REQUIRED",
        "Can_Approve_Evidence": "NO",
        "Can_Close_Gate": "NO",
        "Can_Unlock_STEP102": "NO",
        "Can_Write_SSoT": "NO",
        "Final_Use_Allowed": "NO",
        "Notes": "Template row only; no authority granted."
    },
    {
        "Quorum_ID": "QRM-00003",
        "Reviewer_Name": "",
        "Reviewer_Role": "CONTROL_REVIEWER",
        "Present": "NO",
        "Conflict_Disclosed": "REVIEW_REQUIRED",
        "Can_Approve_Evidence": "NO",
        "Can_Close_Gate": "NO",
        "Can_Unlock_STEP102": "NO",
        "Can_Write_SSoT": "NO",
        "Final_Use_Allowed": "NO",
        "Notes": "Template row only; no authority granted."
    }
]

def now_iso():
    return datetime.now().isoformat(timespec="seconds")

def main():
    parser = argparse.ArgumentParser(description="Build audit board quorum checklist")
    parser.add_argument("--out-xlsx", required=True)
    parser.add_argument("--out-json", required=True)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    payload = {
        "timestamp": now_iso(),
        "component": "BOARD_QUORUM_CHECKLIST",
        "version": "1.0",
        "status": "REVIEW_REQUIRED",
        "records": DEFAULT_ROWS,
        "decision": "Quorum checklist is attendance/review material only. It grants no approval authority.",
        **CANON
    }

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Board Quorum"
        ws.append(HEADERS)
        for row in DEFAULT_ROWS:
            ws.append([row.get(h, "") for h in HEADERS])

        ws2 = wb.create_sheet("Canonical Locks")
        ws2.append(["Control", "Value"])
        for k, v in CANON.items():
            ws2.append([k, v])

        ws3 = wb.create_sheet("Instructions")
        ws3.append(["Instruction_ID", "Instruction"])
        ws3.append(["INS-001", "Quorum presence does not approve evidence."])
        ws3.append(["INS-002", "Quorum presence does not close gates."])
        ws3.append(["INS-003", "Quorum presence does not unlock STEP102."])
        ws3.append(["INS-004", "Final use remains NO."])

        Path(args.out_xlsx).parent.mkdir(parents=True, exist_ok=True)
        wb.save(args.out_xlsx)
        Path(args.out_json).parent.mkdir(parents=True, exist_ok=True)
        Path(args.out_json).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    except Exception as exc:
        print(f"❌ Write error: {exc}")
        return EXIT_WRITE_ERROR

    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else f"✅ Board quorum checklist: {args.out_xlsx}")
    print("SYSTEM RED — STEP102 LOCKED — NO FINAL USE")
    return EXIT_OK

if __name__ == "__main__":
    sys.exit(main())
