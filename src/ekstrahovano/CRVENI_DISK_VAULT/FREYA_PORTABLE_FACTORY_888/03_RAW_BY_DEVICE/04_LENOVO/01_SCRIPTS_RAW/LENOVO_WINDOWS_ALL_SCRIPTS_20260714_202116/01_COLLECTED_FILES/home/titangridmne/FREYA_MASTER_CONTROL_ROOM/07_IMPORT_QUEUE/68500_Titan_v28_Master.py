import os, re, pandas as pd, warnings
from datetime import datetime
try:
    from docx import Document
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
except:
    pass

warnings.filterwarnings('ignore')

RAW = r'C:\Users\Lenovo\Desktop\CTITAN_CENTRAL_BRAINRaw_Data'
OUT = os.path.join(os.environ['USERPROFILE'], 'Desktop', '00_MAJKA_TITANA_v28_CENTRALNI_MOZAK.xlsx')

def clean_name(name):
    """Uklanja sve karaktere koji blokiraju Excel"""
    return re.sub(r'[\\/*?:\[\]]', '', name)[:30]

def get_pilar(text):
    t = text.upper()
    if any(k in t for k in ["EUR", "CAPEX", "BUDGET"]): return "01_FINANCE"
    if any(k in t for k in ["SCADA", "ABB", "TECHNICAL"]): return "02_TECHNICAL"
    if any(k in t for k in ["HAMZA", "STRATEGY"]): return "03_STRATEGY"
    if any(k in t for k in ["RIZIK", "DELAY", "PENAL"]): return "04_RISK"
    if any(k in t for k in ["ESG", "CBAM", "GREEN"]): return "05_ESG"
    return "06_GENERAL"

def process_v28(fpath):
    name = os.path.basename(fpath)
    if name.startswith('~$') or name.startswith('.'): return None
    ext = os.path.splitext(fpath)[1].lower()
    content = ""
    try:
        if ext == '.docx':
            doc = Document(fpath)
            content = ' '.join([p.text for p in doc.paragraphs])
        else:
            with open(fpath, 'r', encoding='utf-8', errors='ignore') as f: content = f.read()
    except: return None
    if len(content) < 5: return None
    
    vals = [float(a.replace('.', '').replace(',', '.')) for a in re.findall(r'\d{1,3}(?:\.\d{3})*,\d{2}', content)]
    net_val = sum(vals)
    
    return {
        'PILAR': get_pilar(content),
        'DOKUMENT': name,
        'NETO_EUR': round(net_val, 2),
        'PDV_21': round(net_val * 0.21, 2),
        'BRUTO': round(net_val * 1.21, 2),
        'STRATESKI_AKTERI': ', '.join([n for n in ['HAMZA', 'DARKO', 'ABB', 'MAREL', 'ARS METAL'] if n in content.upper()]),
        'SKALABILNOST': '95%' if 'SCADA' in content.upper() else '65%',
        'SIROVA_MATERIJA': content[:3500].strip().replace('\n', ' '),
        'CFO_FLAG': '🚨 HITNO' if net_val > 50000 or 'DELAY' in content.upper() else 'OK',
        'VLASNIK': 'ONUR (100% MAREL)',
        'DATUM': '2026-04-04'
    }

if __name__ == '__main__':
    paths = [os.path.join(r, f) for r, d, files in os.walk(RAW) for f in files if not f.startswith('~$')]
    print(f"🚀 POKRECEM FINALNU HARMONIZACIJU v28: Analiziram {len(paths)} dokumenata...")
    results = [r for r in [process_v28(p) for p in paths] if r]
    
    if results:
        df = pd.DataFrame(results)
        with pd.ExcelWriter(OUT, engine='openpyxl') as writer:
            for pilar in sorted(df['PILAR'].unique()):
                sh_name = clean_name(pilar)
                df[df['PILAR'] == pilar].to_excel(writer, sheet_name=sh_name, index=False, startrow=1)
                
                ws = writer.sheets[sh_name]
                navy_fill = PatternFill(start_color='000080', end_color='000080', fill_type='solid')
                white_font = Font(color='FFFFFF', bold=True)
                
                for cell in ws[2]:
                    cell.fill = navy_fill
                    cell.font = white_font
                    cell.alignment = Alignment(horizontal='center')
                
                ws.merge_cells(f'A1:{get_column_letter(ws.max_column)}1')
                ws['A1'] = f"MAJKA TITANA v28 - CENTRALNI MOZAK - PROPERTY OF ONUR"
                ws['A1'].font = Font(bold=True, size=18, color='000080')
                ws['A1'].alignment = Alignment(horizontal='center')
                
                for col in ws.columns:
                    ws.column_dimensions[col[0].column_letter].width = 30
                
                ws.freeze_panes = "A3"
        
        print(f"💎 TRIJUMF! v28 JE STVORENA NA DESKTOPU!")
        os.startfile(OUT)
