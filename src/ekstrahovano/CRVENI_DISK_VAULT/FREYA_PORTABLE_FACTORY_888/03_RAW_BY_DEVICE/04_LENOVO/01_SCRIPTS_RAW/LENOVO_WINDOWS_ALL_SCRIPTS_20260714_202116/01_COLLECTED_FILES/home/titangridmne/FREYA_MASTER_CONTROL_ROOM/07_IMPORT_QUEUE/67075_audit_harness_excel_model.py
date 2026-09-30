from pathlib import Path
import zipfile
from openpyxl import load_workbook

EXPECTED_SHEETS = [
    "INDEX / NAVIGATION",
    "MODEL CONTROL PANEL",
    "ERROR CHECKS",
    "ASSUMPTIONS",
    "MACRO DRIVERS",
    "SCENARIO ENGINE",
    "PROJECT TIMELINE",
    "PRODUCTION MODEL",
    "PRODUCT MIX",
    "REVENUE MODEL",
    "RAW MATERIAL MODEL",
    "ENERGY MODEL",
    "PAYROLL MODEL",
    "OPEX SUMMARY",
    "CAPEX STRUCTURE",
    "CAPEX PHASING",
    "DEPRECIATION MODEL",
    "P&L STATEMENT",
    "CASH FLOW STATEMENT",
    "BALANCE SHEET",
    "WORKING CAPITAL MODEL",
    "FUNDING STRUCTURE",
    "DEBT MODEL",
    "DSCR MODEL",
    "VALUATION",
]

EXPECTED_NAMES = [
    "Base_Year",
    "Full_Capacity",
    "Util_Y1",
    "ASP_Standard",
    "ASP_Premium",
    "Premium_Share",
    "COGS_per_t",
    "Grant_Pct",
    "Contingency_Pct",
    "IDC_Pct",
    "Debt_Pct",
    "Equity_Pct",
    "Interest_Rate",
    "Tenor_Years",
    "Tax_Rate",
    "Discount_Rate",
]

def assert_zip_integrity(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as zf:
        bad = zf.testzip()
        if bad is not None:
            raise AssertionError(f"Corrupted zip member: {bad}")

def assert_workbook_structure(path: Path) -> None:
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == EXPECTED_SHEETS, f"Unexpected sheet order: {wb.sheetnames}"
    defined_names = list(wb.defined_names.keys())
    for name in EXPECTED_NAMES:
        assert name in defined_names, f"Missing named range: {name}"

    # basic link / formula sanity
    ws_index = wb["INDEX / NAVIGATION"]
    assert ws_index["D4"].hyperlink is not None, "Missing first hyperlink"

    ws_err = wb["ERROR CHECKS"]
    assert str(ws_err["D4"].value).startswith("=IF("), "Unexpected status formula"

    ws_ass = wb["ASSUMPTIONS"]
    assert ws_ass["B4"].value == 2025, "Unexpected base year"

    ws_val = wb["VALUATION"]
    assert isinstance(ws_val["B5"].value, str) and ws_val["B5"].value.startswith("="), "NPV formula missing"

def run(path_str: str) -> None:
    path = Path(path_str)
    assert path.exists(), f"File not found: {path}"
    assert path.suffix.lower() == ".xlsx", "Expected .xlsx file"
    assert_zip_integrity(path)
    assert_workbook_structure(path)
    print("PASS - workbook integrity and structure checks completed successfully.")

if __name__ == "__main__":
    import sys
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python audit_harness.py <workbook.xlsx>")
    run(sys.argv[1])
