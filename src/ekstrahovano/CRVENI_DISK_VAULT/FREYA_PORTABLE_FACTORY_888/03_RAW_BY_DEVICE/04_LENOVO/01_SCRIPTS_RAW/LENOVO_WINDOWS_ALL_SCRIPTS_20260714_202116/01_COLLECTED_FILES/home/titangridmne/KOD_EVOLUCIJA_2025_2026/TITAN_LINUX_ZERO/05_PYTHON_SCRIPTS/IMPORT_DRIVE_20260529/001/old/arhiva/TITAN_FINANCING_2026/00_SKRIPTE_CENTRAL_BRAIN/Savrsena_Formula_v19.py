import os, re, pandas as pd, warnings
from datetime import datetime
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

warnings.filterwarnings('ignore')

RAW = r'C:\Users\Lenovo\Desktop\CTITAN_CENTRAL_BRAINRaw_Data'
ROOM = r'C:\Users\Lenovo\Desktop\TITAN_GRID_Data_Room_FINAL_April2026'
OUT = os.path.join(ROOM, '00_MASTER', 'TITAN_FORMULA_SAVRSENA.xlsx')

def process_fast(fpath):
    name = os.path.basename(fpath)
    ext = os.path.splitext(fpath)[1].lower()
    content = ""
    try:
        if ext == '.docx':
            content = ' '.join([p.text for p in Document(fpath).paragraphs])
        else:
            with open(fpath, 'r', encoding='utf-8', errors='ignore') as f: content = f.read()
    except: return None
    
    cat = "04_TECHNICAL"
    if any(k in content.upper() for k in ["EUR", "CAPEX", "BUDGET", "COST"]): cat = "01_FINANCE"
    elif any(k in content.upper() for k in ["HAMZA", "INVESTOR", "STRATEGY"]): cat = "02_STRATEGY"
    elif any(k in content.upper() for k in ["RIZIK", "DELAY", "PENAL"]): cat = "03_RISK"

    vals = [float(a.replace('.', '').replace(',', '.')) for a in re.findall(r'\d{1,3}(?:\.\d{3})*,\d{2}', content)]
    
    return {
        'Category': cat,
        'Dokument': name,
        'Vrijednost_EUR': round(sum(vals), 2),
        'Akteri': ', '.join([n for n in ['HAMZA', 'DARKO', 'ABB', 'MAREL'] if n in content.upper()]),
        'Sustina': content[:1500].strip().replace('\n', ' '),
        'Datum': '2026-04-04'
    }

if __name__ == '__main__':
    if not os.path.exists(os.path.dirname(OUT)): os.makedirs(os.path.dirname(OUT))
    paths = [os.path.join(r, f) for r, d, files in os.walk(RAW) for f in files]
    results = [res for res in [process_fast(p) for p in paths] if res]
    
    if results:
        df = pd.DataFrame(results)
        with pd.ExcelWriter(OUT, engine='openpyxl') as writer:
            for category in sorted(df['Category'].unique()):
                temp_df = df[df['Category'] == category]
                temp_df.to_excel(writer, sheet_name=category, index=False)
                
                # BRENDIRANJE: Dodavanje "Onur Logo" zaglavlja
                ws = writer.sheets[category]
                ws.insert_rows(1)
                ws['A1'] = "ARS METAL INDUSTRIES - PROPERTY OF ONUR (CFO)"
                ws['A1'].font = Font(bold=True, size=14, color="000080")
                ws['A1'].alignment = Alignment(horizontal="center")
        
        print(f"💎 FORMULA JE U TREZORU! {OUT}")
        os.startfile(OUT)
