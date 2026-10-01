import os, re, pandas as pd, warnings
from datetime import datetime
try:
    from docx import Document
    from openpyxl.styles import Font, Alignment, PatternFill
except:
    pass

warnings.filterwarnings('ignore')

RAW = r'C:\Users\Lenovo\Desktop\CTITAN_CENTRAL_BRAINRaw_Data'
OUT = os.path.join(os.environ['USERPROFILE'], 'Desktop', '00_MAJKA_TITANA_v30_TRIJUMF.xlsx')

def clean_sheet_name(name):
    """Agresivno čišćenje za Excel tabove"""
    clean = re.sub(r'[\\/*?:\[\]]', '', name)
    return clean[:25] # Maksimalno 25 karaktera za sigurnost

def get_pilar(text):
    t = text.upper()
    if any(k in t for k in ["EUR", "CAPEX", "BUDGET", "FAKTURA"]): return "01_FINANCE"
    if any(k in t for k in ["SCADA", "ABB", "TECHNICAL"]): return "02_TECHNICAL"
    if any(k in t for k in ["HAMZA", "STRATEGY"]): return "03_STRATEGY"
    if any(k in t for k in ["RIZIK", "RISK", "PENAL"]): return "04_RISK"
    if any(k in t for k in ["ESG", "CBAM", "GREEN"]): return "05_ESG"
    return "06_GENERAL"

def process_v30(fpath):
    fname = os.path.basename(fpath)
    if fname.startswith('~$'): return None
    ext = os.path.splitext(fpath)[1].lower()
    content = ""
    try:
        if ext == '.docx':
            doc = Document(fpath)
            content = ' '.join([p.text for p in doc.paragraphs])
        elif ext in ['.txt', '.csv']:
            with open(fpath, 'r', encoding='utf-8', errors='ignore') as f: content = f.read()
    except: return None
    
    if len(content) < 10: return None
    
    # Detekcija novca
    vals = [float(a.replace('.', '').replace(',', '.')) for a in re.findall(r'\d{1,3}(?:\.\d{3})*,\d{2}', content)]
    neto = sum(vals)
    
    return {
        'PILAR': get_pilar(content),
        'DOKUMENT': fname,
        'IZNOS_NETO': round(neto, 2),
        'PDV_21': round(neto * 0.21, 2),
        'BRUTO': round(neto * 1.21, 2),
        'SENTINEL': 'Onur (CFO)' if neto > 0 else 'Darko D.',
        'UTICAJ': 'VISOK' if neto > 50000 else 'OPERATIVNI',
        'STATUS': 'Verifikovano v30',
        'TEKST_DNA': content[:2000].replace('\n', ' ')
    }

if __name__ == '__main__':
    print("🚀 POKRETANJE MAJKE TITANA v30: Harmonizacija carstva u toku...")
    paths = [os.path.join(r, f) for r, d, files in os.walk(RAW) for f in files if not f.startswith('~$')]
    data = [r for r in [process_v30(p) for p in paths] if r]
    
    if data:
        df = pd.DataFrame(data)
        with pd.ExcelWriter(OUT, engine='openpyxl') as writer:
            for pilar in sorted(df['PILAR'].unique()):
                s_name = clean_sheet_name(pilar)
                df[df['PILAR'] == pilar].to_excel(writer, sheet_name=s_name, index=False, startrow=1)
                
                ws = writer.sheets[s_name]
                fill = PatternFill(start_color='000080', end_color='000080', fill_type='solid')
                font = Font(color='FFFFFF', bold=True)
                
                for cell in ws[2]:
                    cell.fill = fill
                    cell.font = font
                
                ws['A1'] = f"TITAN GRID v30 - CENTRALNI MOZAK - VLASNIK: ONUR"
                ws['A1'].font = Font(size=16, bold=True, color='000080')
        
        print(f"💎 TRIJUMF! Fajl je na Desktopu pod nazivom: 00_MAJKA_TITANA_v30_TRIJUMF.xlsx")
        os.startfile(OUT)
