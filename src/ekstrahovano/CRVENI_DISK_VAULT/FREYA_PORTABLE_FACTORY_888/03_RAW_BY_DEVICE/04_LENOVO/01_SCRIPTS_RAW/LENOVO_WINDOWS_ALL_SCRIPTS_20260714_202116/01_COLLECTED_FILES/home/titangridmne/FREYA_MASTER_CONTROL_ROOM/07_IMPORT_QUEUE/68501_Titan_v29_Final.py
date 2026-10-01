import os, re, pandas as pd, warnings
from datetime import datetime
try:
    from docx import Document
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
except:
    pass

warnings.filterwarnings('ignore')

RAW = r'C:\Users\Lenovo\Desktop\CTITAN_CENTRAL_BRAINRaw_Data'
OUT = os.path.join(os.environ['USERPROFILE'], 'Desktop', '00_MAJKA_TITANA_v29_MASTER.xlsx')

def safe_sheet_name(name):
    """Pretvara bilo koje ime u siguran Excel tab"""
    clean = re.sub(r'[\\/*?:\[\]]', '_', name)
    return clean[:30] if clean else "SHEET_UNAMED"

def get_pilar(text):
    t = text.upper()
    if any(k in t for k in ["EUR", "CAPEX", "COST", "FAKTURA"]): return "01_FINANCE"
    if any(k in t for k in ["SCADA", "ABB", "TECHNICAL", "TEHNIKA"]): return "02_TECHNICAL"
    if any(k in t for k in ["HAMZA", "STRATEGY", "VIZIJA"]): return "03_STRATEGY"
    if any(k in t for k in ["RIZIK", "DELAY", "PENAL", "RISK"]): return "04_RISK"
    if any(k in t for k in ["ESG", "CBAM", "GREEN", "ZELENO"]): return "05_ESG"
    return "06_GENERAL"

def process_v29(fpath):
    name = os.path.basename(fpath)
    if name.startswith('~$') or name.startswith('.'): return None
    ext = os.path.splitext(fpath)[1].lower()
    content = ""
    try:
        if ext == '.docx':
            doc = Document(fpath)
            content = ' '.join([p.text for p in doc.paragraphs])
        else:
            with open(fpath, 'r', encoding='utf-8', errors='ignore') as f: content = f.read()
    except: return None
    if len(content) < 5: return None
    
    # Napredna detekcija cifara za CFO
    vals = [float(a.replace('.', '').replace(',', '.')) for a in re.findall(r'\d{1,3}(?:\.\d{3})*,\d{2}', content)]
    net_val = sum(vals)
    
    return {
        'PILAR': get_pilar(content),
        'DOKUMENT_NAZIV': name,
        'IZNOS_NETO': round(net_val, 2),
        'PDV_21': round(net_val * 0.21, 2),
        'BRUTO_TOTAL': round(net_val * 1.21, 2),
        'KLJUCNI_AKTERI': ', '.join([n for n in ['HAMZA', 'DARKO', 'ABB', 'MAREL', 'ARS METAL'] if n in content.upper()]),
        'SKALABILNOST': '95%' if 'SCADA' in content.upper() or 'MODULAR' in content.upper() else '65%',
        'SIROVA_MATERIJA_ZA_DOKUMENTE': content[:3000].strip().replace('\n', ' '),
        'CFO_ALARM': '🚨 HITNO' if net_val > 50000 or 'PENAL' in content.upper() else 'STABILNO',
        'VLASNISTVO': 'ONUR (100% MAREL)',
        'DATUM_OBRADE': '2026-04-04'
    }

if __name__ == '__main__':
    paths = [os.path.join(r, f) for r, d, files in os.walk(RAW) for f in files if not f.startswith('~$')]
    print(f"🚀 POKRECEM v29 - TOTALNA HARMONIZACIJA ZA ONUR...")
    print(f"📂 Analiziram bazu od {len(paths)} dokumenata. Molim sacekajte dijamant...")
    
    results = [r for r in [process_v29(p) for p in paths] if r]
    
    if results:
        df = pd.DataFrame(results)
        with pd.ExcelWriter(OUT, engine='openpyxl') as writer:
            for pilar in sorted(df['PILAR'].unique()):
                s_name = safe_sheet_name(pilar)
                df[df['PILAR'] == pilar].to_excel(writer, sheet_name=s_name, index=False, startrow=1)
                
                ws = writer.sheets[s_name]
                navy = PatternFill(start_color='000080', end_color='000080', fill_type='solid')
                white = Font(color='FFFFFF', bold=True)
                
                # Zaglavlje teget plavo - ONUR STYLE
                for cell in ws[2]:
                    cell.fill = navy
                    cell.font = white
                    cell.alignment = Alignment(horizontal='center')
                
                # Master Titul
                ws.merge_cells(f'A1:{get_column_letter(ws.max_column)}1')
                ws['A1'] = f"TITAN GRID - CENTRALNI MOZAK v29 - VLASNIK: ONUR (CFO)"
                ws['A1'].font = Font(bold=True, size=20, color='000080')
                ws['A1'].alignment = Alignment(horizontal='center')
                
                ws.column_dimensions['B'].width = 40
                ws.column_dimensions['H'].width = 80
                ws.freeze_panes = "A3"
        
        print(f"💎 TRIJUMF! 'MAJKA TITANA v29' JE NA DESKTOPU!")
        os.startfile(OUT)
