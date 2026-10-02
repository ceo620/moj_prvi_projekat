#!/usr/bin/env python3

import sys
import hashlib
from pathlib import Path
from datetime import datetime, timezone

try:
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.colors import HexColor, black
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    )
except ImportError as exc:
    print("RESULT=HOLD")
    print(f"MISSING_PYTHON_MODULE={exc}")
    print("INSTALL_COMMAND=sudo apt-get update && sudo apt-get install -y python3-docx python3-reportlab")
    sys.exit(20)

PROTOCOL = "888"
ROOT = Path("/mnt/c/Users/ceo/OneDrive/Desktop/FIRMA_DOKUMENTACIJA")
RUN_TS = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
DARK_BLUE = RGBColor(18, 52, 86)
PDF_BLUE = HexColor("#123456")
PLACEHOLDER = "[HUMAN_INPUT_REQUIRED]"

COMPANIES = [
    {
        "folder": "01_ARSMETAL",
        "name": "ARS METAL INDUSTRIES DOO",
        "prefix": "ARS",
    },
    {
        "folder": "02_ADS",
        "name": "ADS METAL TURKEY",
        "prefix": "ADS",
    },
    {
        "folder": "03_ECO",
        "name": "ECO TRANSFORMER",
        "prefix": "ECO",
    },
    {
        "folder": "04_MAREL",
        "name": "MAREL DOO",
        "prefix": "MAREL",
    },
]

def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()

def safe_output(folder, base_name, extension):
    candidate = folder / f"{base_name}{extension}"
    if not candidate.exists():
        return candidate

    version = 2
    while True:
        candidate = folder / f"{base_name}_V{version:03d}{extension}"
        if not candidate.exists():
            return candidate
        version += 1

def set_cell_text(cell, text, bold=False, size=9, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def create_docx(company, output_path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(17)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    header = doc.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    header.columns[0].width = Mm(42)
    header.columns[1].width = Mm(128)

    set_cell_text(
        header.cell(0, 0),
        "[LOGO_PLACEHOLDER]",
        bold=True,
        size=9,
        color=DARK_BLUE,
    )

    company_data = (
        f"{company['name']}\n"
        f"Address: {PLACEHOLDER}\n"
        f"Registration number: {PLACEHOLDER}\n"
        f"Tax ID: {PLACEHOLDER}\n"
        f"Contact: {PLACEHOLDER}"
    )
    set_cell_text(header.cell(0, 1), company_data, bold=False, size=9)

    doc.add_paragraph()

    identity = doc.add_table(rows=3, cols=2)
    identity.alignment = WD_TABLE_ALIGNMENT.CENTER
    identity.style = "Table Grid"
    identity_data = [
        ("DOCUMENT NUMBER", f"{company['prefix']}-000001"),
        ("DATE / PLACE", f"{PLACEHOLDER} / {PLACEHOLDER}"),
        ("DOCUMENT STATUS", "TEMPLATE — HUMAN REVIEW REQUIRED"),
    ]
    for row_index, (label, value) in enumerate(identity_data):
        set_cell_text(identity.cell(row_index, 0), label, bold=True, size=9, color=DARK_BLUE)
        set_cell_text(identity.cell(row_index, 1), value, size=9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("OFFICIAL CORPORATE MEMORANDUM")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = DARK_BLUE

    content = doc.add_table(rows=3, cols=2)
    content.style = "Table Grid"
    content.alignment = WD_TABLE_ALIGNMENT.CENTER
    fields = [
        ("SUBJECT", PLACEHOLDER),
        ("RECIPIENT", PLACEHOLDER),
        ("REFERENCE", PLACEHOLDER),
    ]
    for row_index, (label, value) in enumerate(fields):
        set_cell_text(content.cell(row_index, 0), label, bold=True, color=DARK_BLUE)
        set_cell_text(content.cell(row_index, 1), value)

    doc.add_paragraph()
    body_label = doc.add_paragraph()
    run = body_label.add_run("BODY TEXT")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = DARK_BLUE

    body = doc.add_paragraph(
        f"{PLACEHOLDER}\n\n"
        "This area is reserved for the official memorandum content.\n\n"
        f"{PLACEHOLDER}"
    )
    body.paragraph_format.space_after = Pt(20)

    signature = doc.add_table(rows=4, cols=2)
    signature.style = "Table Grid"
    signature.alignment = WD_TABLE_ALIGNMENT.CENTER
    signature_data = [
        ("AUTHORIZED PERSON", PLACEHOLDER),
        ("TITLE", PLACEHOLDER),
        ("ELECTRONIC SIGNATURE", "[ELECTRONIC_SIGNATURE_AREA]"),
        ("STAMP", "[STAMP_AREA]"),
    ]
    for row_index, (label, value) in enumerate(signature_data):
        set_cell_text(signature.cell(row_index, 0), label, bold=True, color=DARK_BLUE)
        set_cell_text(signature.cell(row_index, 1), value)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = paragraph.add_run(
        f"{company['name']} | CONFIDENTIAL | "
        f"Document control: {company['prefix']}-000001 | Protocol 888"
    )
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(7)
    footer_run.font.color.rgb = DARK_BLUE

    doc.save(output_path)

def create_pdf(company, output_path):
    pdf = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=17 * mm,
        bottomMargin=17 * mm,
        title=f"{company['prefix']}-000001 Corporate Memorandum Template",
        author=company["name"],
    )

    styles = getSampleStyleSheet()
    normal = ParagraphStyle(
        "CorporateNormal",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=black,
    )
    small = ParagraphStyle(
        "CorporateSmall",
        parent=normal,
        fontSize=7,
        leading=9,
    )
    title = ParagraphStyle(
        "CorporateTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        alignment=TA_CENTER,
        textColor=PDF_BLUE,
        spaceAfter=8 * mm,
    )
    label = ParagraphStyle(
        "CorporateLabel",
        parent=normal,
        fontName="Helvetica-Bold",
        textColor=PDF_BLUE,
        alignment=TA_LEFT,
    )

    story = []

    header_data = [[
        Paragraph("<b>[LOGO_PLACEHOLDER]</b>", label),
        Paragraph(
            f"<b>{company['name']}</b><br/>"
            f"Address: {PLACEHOLDER}<br/>"
            f"Registration number: {PLACEHOLDER}<br/>"
            f"Tax ID: {PLACEHOLDER}<br/>"
            f"Contact: {PLACEHOLDER}",
            normal,
        ),
    ]]
    header = Table(header_data, colWidths=[42 * mm, 128 * mm])
    header.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.7, PDF_BLUE),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, PDF_BLUE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([header, Spacer(1, 6 * mm)])

    identity_data = [
        [Paragraph("DOCUMENT NUMBER", label), Paragraph(f"{company['prefix']}-000001", normal)],
        [Paragraph("DATE / PLACE", label), Paragraph(f"{PLACEHOLDER} / {PLACEHOLDER}", normal)],
        [Paragraph("DOCUMENT STATUS", label), Paragraph("TEMPLATE — HUMAN REVIEW REQUIRED", normal)],
    ]
    identity = Table(identity_data, colWidths=[48 * mm, 122 * mm])
    identity.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.6, black),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 0), (0, -1), HexColor("#EDF2F7")),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([identity, Spacer(1, 8 * mm)])
    story.append(Paragraph("OFFICIAL CORPORATE MEMORANDUM", title))

    content_data = [
        [Paragraph("SUBJECT", label), Paragraph(PLACEHOLDER, normal)],
        [Paragraph("RECIPIENT", label), Paragraph(PLACEHOLDER, normal)],
        [Paragraph("REFERENCE", label), Paragraph(PLACEHOLDER, normal)],
    ]
    content = Table(content_data, colWidths=[40 * mm, 130 * mm])
    content.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.6, black),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, black),
        ("BACKGROUND", (0, 0), (0, -1), HexColor("#EDF2F7")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([
        content,
        Spacer(1, 7 * mm),
        Paragraph("BODY TEXT", label),
        Spacer(1, 3 * mm),
        Paragraph(
            f"{PLACEHOLDER}<br/><br/>"
            "This area is reserved for the official memorandum content."
            f"<br/><br/>{PLACEHOLDER}",
            normal,
        ),
        Spacer(1, 14 * mm),
    ])

    signature_data = [
        [Paragraph("AUTHORIZED PERSON", label), Paragraph(PLACEHOLDER, normal)],
        [Paragraph("TITLE", label), Paragraph(PLACEHOLDER, normal)],
        [Paragraph("ELECTRONIC SIGNATURE", label), Paragraph("[ELECTRONIC_SIGNATURE_AREA]", normal)],
        [Paragraph("STAMP", label), Paragraph("[STAMP_AREA]", normal)],
    ]
    signature = Table(signature_data, colWidths=[55 * mm, 115 * mm], rowHeights=[10*mm, 10*mm, 16*mm, 16*mm])
    signature.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.6, black),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, black),
        ("BACKGROUND", (0, 0), (0, -1), HexColor("#EDF2F7")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))

    story.append(KeepTogether(signature))
    story.extend([
        Spacer(1, 8 * mm),
        Paragraph(
            f"{company['name']} | CONFIDENTIAL | "
            f"Document control: {company['prefix']}-000001 | Protocol 888",
            small,
        ),
    ])

    pdf.build(story)

ROOT.mkdir(parents=True, exist_ok=True)
all_outputs = []
master_rows = []

for company in COMPANIES:
    company_root = ROOT / company["folder"]
    control = company_root / "00_CONTROL"
    template_dir = company_root / "01_MEMORANDUM_TEMPLATE"
    signature_dir = company_root / "02_SIGNATURE_READY"
    archive_dir = company_root / "03_ARCHIVE"

    for directory in (control, template_dir, signature_dir, archive_dir):
        directory.mkdir(parents=True, exist_ok=True)

    base = f"{company['prefix']}-000001_CORPORATE_MEMORANDUM_TEMPLATE"
    docx_path = safe_output(template_dir, base, ".docx")
    pdf_path = safe_output(template_dir, base, ".pdf")

    create_docx(company, docx_path)
    create_pdf(company, pdf_path)

    company_outputs = [docx_path, pdf_path]
    all_outputs.extend(company_outputs)

    manifest_path = control / f"{company['prefix']}_MANIFEST_{RUN_TS}.tsv"
    with manifest_path.open("x", encoding="utf-8", newline="\n") as manifest:
        manifest.write("PROTOCOL\tCOMPANY\tTYPE\tFILE\tSHA256\tSTATUS\n")
        for path in company_outputs:
            digest = sha256(path)
            file_type = path.suffix.lstrip(".").upper()
            manifest.write(
                f"{PROTOCOL}\t{company['name']}\t{file_type}\t"
                f"{path}\t{digest}\tPASS\n"
            )
            master_rows.append(
                (company["name"], company["prefix"], file_type, str(path), digest)
            )

master_manifest = ROOT / f"MASTER_MANIFEST_{RUN_TS}.tsv"
with master_manifest.open("x", encoding="utf-8", newline="\n") as handle:
    handle.write("PROTOCOL\tCOMPANY\tPREFIX\tTYPE\tFILE\tSHA256\tSTATUS\n")
    for company_name, prefix, file_type, path, digest in master_rows:
        handle.write(
            f"{PROTOCOL}\t{company_name}\t{prefix}\t{file_type}\t"
            f"{path}\t{digest}\tPASS\n"
        )

hash_manifest = ROOT / f"SHA256SUMS_{RUN_TS}.txt"
with hash_manifest.open("x", encoding="utf-8", newline="\n") as handle:
    for path in all_outputs:
        handle.write(f"{sha256(path)}  {path}\n")

final_status = ROOT / f"FINAL_STATUS_{RUN_TS}.env"
with final_status.open("x", encoding="utf-8", newline="\n") as handle:
    handle.write(
        "PROTOCOL=888\n"
        "NODE=ASUS_DEBIAN_DOCUMENT_FACTORY\n"
        "MEMORANDUM_FACTORY=PASS\n"
        "COMPANIES_CREATED=4\n"
        "TEMPLATES_CREATED=8\n"
        "HASH_MANIFEST=CREATED\n"
        "TEMPLATE_READINESS=PASS\n"
        "MASTER_DATA_COMPLETENESS=HOLD\n"
        "PLACEHOLDERS_PRESENT=YES\n"
        "READY_FOR_HUMAN_REVIEW=YES\n"
        "READY_FOR_SIGNATURE=NO\n"
        "HUMAN_GATE=REQUIRED\n"
        "DELETE=NO\n"
        "MOVE=NO\n"
        "OVERWRITE=NO\n"
        f"OUTPUT_ROOT={ROOT}\n"
        f"RUN_TIMESTAMP={RUN_TS}\n"
    )

print("============================================================")
print("MAGNUS DEBIAN DOCUMENT FACTORY 888")
print("============================================================")
print("MEMORANDUM_FACTORY=PASS")
print("COMPANIES_CREATED=4")
print(f"TEMPLATES_CREATED={len(all_outputs)}")
print("HASH_MANIFEST=CREATED")
print("PLACEHOLDERS_PRESENT=YES")
print("READY_FOR_HUMAN_REVIEW=YES")
print("READY_FOR_SIGNATURE=NO")
print("HUMAN_GATE=REQUIRED")
print(f"OUTPUT_ROOT={ROOT}")
print(f"MASTER_MANIFEST={master_manifest}")
print(f"SHA256_MANIFEST={hash_manifest}")
print(f"FINAL_STATUS={final_status}")
print("")
print("GENERATED_FILES:")
for path in all_outputs:
    print(path)
print("============================================================")
print("STOP=HUMAN_GATE")
